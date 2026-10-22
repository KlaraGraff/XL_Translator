"""诊断 Word 里带公式（OMML）的段落为什么没按占位符放回原位（反馈 #19 临时用，修好后删）。

用法：python3 scripts/diag_word_math.py 原文.docx [双语或译文.docx]
只打印段落编号、结构标签和计数，不打印正文内容。
"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from core import word_inline_marks as wim  # noqa: E402


def _local(tag) -> str:
    if not isinstance(tag, str):
        return "<comment/pi>"
    ns, _, name = tag.rpartition("}")
    prefix = {
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main": "w",
        "{http://schemas.openxmlformats.org/officeDocument/2006/math": "m",
    }.get(ns, ns[-12:])
    return f"{prefix}:{name}"


def _reason(paragraph) -> str:
    element = paragraph._p
    odd_children = [
        _local(child.tag) for child in element
        if not isinstance(child.tag, str) or child.tag not in wim._PARAGRAPH_CHILDREN
    ]
    if odd_children:
        return "段落里有不支持的节点 " + ",".join(sorted(Counter(odd_children)))
    odd_run = Counter()
    for child in element:
        runs = [child] if child.tag == qn("w:r") else list(child.iterchildren(qn("w:r")))
        if child.tag == qn("w:hyperlink"):
            for nested in child:
                if nested.tag not in wim._HYPERLINK_CHILDREN:
                    odd_run["(链接里)" + _local(nested.tag)] += 1
        for run in runs:
            for node in run:
                if not isinstance(node.tag, str) or node.tag not in wim._RUN_CHILDREN | wim._SLOT_TAGS:
                    odd_run[_local(node.tag)] += 1
    if odd_run:
        return "run 里有不支持的节点 " + ",".join(f"{k}×{v}" for k, v in sorted(odd_run.items()))
    runs = wim._content_runs(element)
    items = wim._content_items(runs)
    plain_full = "".join(item.text for item in items)
    plain = plain_full.strip()
    if not plain:
        return "段里没有文字"
    if plain != (paragraph.text or "").strip():
        return "纯文本和 paragraph.text 对不上"
    if wim._TAG_RE.search(plain) or wim._SLOT_RE.search(plain):
        return "原文里本来就有像标记的字符"
    slots = sum(1 for item in items if item.is_slot)
    groups = len({(wim._look_key(item.rpr)) for item in items if not item.is_slot})
    return f"paragraph_marks=None（占位约 {slots} 个，上限 {wim.MAX_PLACEHOLDERS_PER_UNIT}；格式种类 {groups}）"


def _iter_paragraphs(doc):
    for index, paragraph in enumerate(doc.paragraphs):
        yield f"正文段 {index}", paragraph
    for t_index, table in enumerate(doc.tables):
        for r_index, row in enumerate(table.rows):
            for c_index, cell in enumerate(row.cells):
                for p_index, paragraph in enumerate(cell.paragraphs):
                    yield f"表{t_index} 行{r_index} 列{c_index} 段{p_index}", paragraph


def diagnose_source(path: str) -> None:
    doc = Document(path)
    total = ok = 0
    reasons = Counter()
    print(f"== 原文 {Path(path).name}")
    for where, paragraph in _iter_paragraphs(doc):
        maths = list(paragraph._p.iterchildren(*wim.MATH_TAGS))
        if not maths:
            continue
        nested = sum(1 for _ in paragraph._p.iter(*wim.MATH_TAGS)) - len(maths)
        if not (paragraph.text or "").strip():
            continue  # 只有公式的独立行，不翻译
        total += 1
        marks = wim.paragraph_marks(paragraph)
        if marks is not None:
            ok += 1
            print(f"  {where}: 公式 {len(maths)}，可放回原位（占位 {len(marks.placeholders)}）")
            continue
        reason = _reason(paragraph)
        reasons[reason.split("（")[0]] += 1
        extra = f"，另有 {nested} 个公式嵌在别的节点里" if nested else ""
        print(f"  {where}: 公式 {len(maths)}{extra}，放不回：{reason}")
    print(f"带公式的文字段 {total} 个，能放回 {ok} 个")
    for reason, count in reasons.most_common():
        print(f"  {count} × {reason}")


def diagnose_output(path: str) -> None:
    doc = Document(path)
    at_end = inside = 0
    print(f"== 译文 {Path(path).name}")
    for where, paragraph in _iter_paragraphs(doc):
        children = list(paragraph._p)
        math_positions = [i for i, child in enumerate(children) if child.tag in wim.MATH_TAGS]
        if not math_positions or not (paragraph.text or "").strip():
            continue
        text_positions = [
            i for i, child in enumerate(children)
            if any((node.text or "").strip() for node in child.iter(qn("w:t")))
        ]
        if text_positions and min(math_positions) > max(text_positions):
            at_end += 1
            print(f"  {where}: {len(math_positions)} 个公式全在文字后面")
        else:
            inside += 1
    print(f"公式在句中 {inside} 段，堆在段末 {at_end} 段")


if __name__ == "__main__":
    diagnose_source(sys.argv[1])
    if len(sys.argv) > 2:
        diagnose_output(sys.argv[2])
