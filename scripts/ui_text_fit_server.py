"""界面文字检查（scripts/ui_text_fit_check.mjs）用的后端：隔离的数据目录 + 样例文件 + 假模型 + API。

只给检查脚本用，不连任何真实服务、不碰用户数据：
- 数据目录、样例文件都在 --root 下（检查脚本给一个临时目录）。
- 假模型是本机 127.0.0.1 上一个兼容 OpenAI 的小服务，把每条原文改写成「法文」样子的长句，
  让英法界面里出现真实长度的文件名、任务名和结果。
- API 用固定令牌、放行 Vite 开发页（TRANSLATOR_DEV_ORIGIN），和 ui/src/dev-tauri-shim.ts 的用法一致。

启动后在标准输出打一行 JSON（端口、令牌、样例目录），检查脚本读到这一行再往下走。
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

# 故意取得长：文件表、任务列表、任务详情、反馈页里的「内容」要真被截断，才查得到悬停全文。
LONG_STEM = "Rapport_annuel_de_maintenance_des_installations_electriques_2026"
CJK_TEXTS = [
    "设备名称", "检修日期", "备注：请在每次检修后填写本表并交给项目经理签字确认",
    "电气设备年度维护报告", "消防系统", "合计",
]


def _fake_translation(text: str) -> str:
    # 法文比中文长一倍多：按字数放大，界面里的长度才接近真实情况。
    words = ["Traduction", "provisoire", "de", "l'élément", "d'origine", "pour", "vérification"]
    count = max(2, min(len(words), len(text)))
    return " ".join(words[:count])


def _reply_for(system: str, user: str) -> str:
    if "OK" in user and len(user) < 200 or "只回复 OK" in system:
        return "OK"
    try:
        payload = json.loads(user)
    except (TypeError, ValueError):
        payload = None
    if isinstance(payload, list):
        out = []
        for item in payload:
            if isinstance(item, str):
                out.append({"translation": _fake_translation(item), "source_lang": "zh"})
            elif isinstance(item, dict) and "id" in item:
                source = str(item.get("text") or item.get("source") or "")
                out.append({
                    "id": item["id"],
                    "action": "translate",
                    "translation": _fake_translation(source),
                    "suggested": _fake_translation(source),
                    "source_lang": "zh",
                    "note": "",
                })
            else:
                out.append({"translation": "Traduction", "source_lang": "zh"})
        return json.dumps(out, ensure_ascii=False)
    if "source_langs" in system or "source_langs" in user:
        return json.dumps({"source_langs": ["zh"]})
    return "[]"


class _FakeModel(BaseHTTPRequestHandler):
    def log_message(self, *_args) -> None:  # 不刷屏
        return

    def _send(self, body: dict) -> None:
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self) -> None:  # /v1/models
        self._send({"object": "list", "data": [{"id": "fake-model", "object": "model"}]})

    def do_POST(self) -> None:
        length = int(self.headers.get("Content-Length") or 0)
        try:
            body = json.loads(self.rfile.read(length) or b"{}")
        except ValueError:
            body = {}
        messages = body.get("messages") or []
        system = "\n".join(str(m.get("content")) for m in messages if m.get("role") == "system")
        user = "\n".join(
            m["content"] if isinstance(m.get("content"), str) else json.dumps(m.get("content"), ensure_ascii=False)
            for m in messages if m.get("role") == "user"
        )
        content = _reply_for(system, user)
        self._send({
            "id": "fake",
            "object": "chat.completion",
            "model": body.get("model") or "fake-model",
            "choices": [{"index": 0, "finish_reason": "stop", "message": {"role": "assistant", "content": content}}],
            "usage": {"prompt_tokens": 10, "completion_tokens": 10, "total_tokens": 20},
        })


def _write_fixtures(folder: Path) -> dict[str, str]:
    """每种类型一个文件夹，各放两个文件：一个长名字（真会被截断）、一个短名字。"""
    import openpyxl
    from docx import Document

    paths: dict[str, str] = {}
    excel = folder / "excel"
    excel.mkdir(parents=True, exist_ok=True)
    for stem in (LONG_STEM, "Budget"):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "维护记录"
        for row, text in enumerate(CJK_TEXTS, start=1):
            sheet.cell(row=row, column=1, value=text)
            sheet.cell(row=row, column=2, value=row * 3)
        book.save(excel / f"{stem}.xlsx")
    paths["excel"] = str(excel)

    word = folder / "word"
    word.mkdir(parents=True, exist_ok=True)
    for stem in (LONG_STEM, "Note"):
        doc = Document()
        for text in CJK_TEXTS:
            doc.add_paragraph(text)
        doc.save(str(word / f"{stem}.docx"))
    paths["word"] = str(word)
    return paths


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True, help="临时目录：数据目录和样例文件都放这里")
    parser.add_argument("--port", type=int, default=18765)
    parser.add_argument("--token", default="uitextfit")
    parser.add_argument("--origin", default="http://127.0.0.1:1420")
    args = parser.parse_args()

    root = Path(args.root).resolve()
    app_data = root / "appdata"
    app_data.mkdir(parents=True, exist_ok=True)
    os.environ["TRANSLATOR_APP_DATA_DIR"] = str(app_data)
    os.environ["TRANSLATOR_DEV_ORIGIN"] = args.origin
    for key in ("NO_PROXY", "no_proxy"):
        os.environ[key] = ",".join(filter(None, [os.environ.get(key, ""), "127.0.0.1", "localhost"]))
    fixtures = _write_fixtures(root / "samples")

    model = ThreadingHTTPServer(("127.0.0.1", 0), _FakeModel)
    threading.Thread(target=model.serve_forever, daemon=True).start()

    sys.path.insert(0, str(ROOT))
    import uvicorn
    from api.app import create_app

    print(json.dumps({
        "api_port": args.port,
        "token": args.token,
        "model_url": f"http://127.0.0.1:{model.server_address[1]}/v1",
        "samples": fixtures,
        "app_data": str(app_data),
    }, ensure_ascii=False), flush=True)
    uvicorn.run(create_app(auth_token=args.token), host="127.0.0.1", port=args.port, log_level="warning")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
