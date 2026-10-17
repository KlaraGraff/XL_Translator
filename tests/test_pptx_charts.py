"""Chart text in PPT files (``core/pptx_charts.py``) on ``tests/fixtures/pptx/graphics.pptx``.

Slides 1–3 of ``graphics.pptx`` (built by ``build_graphics.py``) hold a column
chart with a rich title, axis titles, two series and an Excel table over its
data sheet, a bar chart with two-level categories, and a pie chart.
"""

from __future__ import annotations

import io
import re
import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest
from lxml import etree

from core.pptx_charts import REASON_DATA_NOT_SYNCED, _parse_range, patch_workbook
from core.pptx_text import Translation, extract_units, scan_file, write_file

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "pptx"
GRAPHICS = FIXTURES / "graphics.pptx"
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

FRENCH = {
    "2026 年各季度产值（万元）": "Production par trimestre en 2026 (10 000 CNY)",
    "季度": "Trimestre",
    "产值": "Production",
    "计划产值": "Production prévue",
    "实际产值": "Production réelle",
    "第一季度": "1er trimestre",
    "第二季度": "2e trimestre",
    "第三季度": "3e trimestre",
    "第四季度": "4e trimestre",
    "在岗人数": "Effectif",
    "钢筋班": "Ferraillage",
    "木工班": "Coffrage",
    "北区": "Zone nord",
    "南区": "Zone sud",
    "费用": "Coûts",
    "人工费": "Main-d’œuvre",
    "材料费": "Matériaux",
    "机械费": "Engins",
}


def _chart_units(source: Path = GRAPHICS):
    return [unit for unit in extract_units(source) if "/charts/" in unit.id]


def _write(tmp_path: Path, *, bilingual: bool = False, source: Path = GRAPHICS):
    translations = {unit.id: Translation(FRENCH.get(unit.text, f"FR {unit.text}")) for unit in extract_units(source)}
    out = tmp_path / ("bilingual.pptx" if bilingual else "translated.pptx")
    return write_file(source, out, translations, bilingual=bilingual, target_lang="fr"), out


def _part(path: Path, name: str):
    with zipfile.ZipFile(path) as package:
        return etree.fromstring(package.read(name))


def _cached(path: Path, chart: str) -> list[str]:
    return [value.text for value in _part(path, f"ppt/charts/{chart}").iter(f"{{{NS['c']}}}v")]


def _workbook(path: Path, chart: str):
    with zipfile.ZipFile(path) as package:
        rels = etree.fromstring(package.read(f"ppt/charts/_rels/{chart}.rels"))
        target = "ppt/" + rels[0].get("Target").removeprefix("../")
        return openpyxl.load_workbook(io.BytesIO(package.read(target)))


def test_units_follow_the_slide_and_share_one_unit_per_cell() -> None:
    units = _chart_units()
    first = [(unit.text, unit.context.split(" · ")[-1], unit.kind) for unit in units if "chart1" in unit.id]
    assert first == [
        ("2026 年各季度产值（万元）", "图表标题", "heading"),
        ("季度", "坐标轴标题", "heading"),
        ("产值", "坐标轴标题", "heading"),
        ("计划产值", "系列名称", "cell"),
        ("第一季度", "类别", "cell"),
        ("第二季度", "类别", "cell"),
        ("第三季度", "类别", "cell"),
        ("第四季度", "类别", "cell"),
        ("实际产值", "系列名称", "cell"),
    ]
    # 两个系列共用一列类别：类别各只出现一次。
    texts = [unit.text for unit in extract_units(GRAPHICS)]
    assert texts.index("季度产值") < texts.index("2026 年各季度产值（万元）") < texts.index("各工区人数")
    # 多级类别里两个「钢筋班」在不同格子，各是一个单元；纯数字的「2025」不翻。
    second = [unit.text for unit in units if "chart2" in unit.id]
    assert second == ["在岗人数", "钢筋班", "木工班", "钢筋班", "北区", "南区"]
    assert len({unit.id for unit in units}) == len(units)


def test_scan_no_longer_reports_supported_charts() -> None:
    scan = scan_file(GRAPHICS)
    assert scan.unit_count == len(extract_units(GRAPHICS))
    assert scan.warnings == []


def test_translation_only_rewrites_cache_workbook_and_table(tmp_path: Path) -> None:
    result, out = _write(tmp_path)
    assert result.review == []
    values = _cached(out, "chart1.xml")
    assert values[:5] == ["Production prévue", "1er trimestre", "2e trimestre", "3e trimestre", "4e trimestre"]
    assert "Production réelle" in values
    titles = ["".join(p.itertext()) for p in _part(out, "ppt/charts/chart1.xml").iter(f"{{{NS['a']}}}p")]
    assert "Production par trimestre en 2026 (10 000 CNY)" in titles and "Trimestre" in titles
    sheet = _workbook(out, "chart1.xml").active
    assert [cell.value for cell in sheet[1]] == [" ", "Production prévue", "Production réelle"]
    assert [sheet.cell(row, 1).value for row in range(2, 6)] == ["1er trimestre", "2e trimestre", "3e trimestre", "4e trimestre"]
    assert sheet["B2"].value == 120
    table = next(iter(sheet.tables.values()))
    assert [column.name for column in table.tableColumns] == [" ", "Production prévue", "Production réelle"]


def test_numbers_and_references_stay(tmp_path: Path) -> None:
    _result, out = _write(tmp_path)
    for chart in ("chart1.xml", "chart2.xml", "chart3.xml"):
        before, after = _part(GRAPHICS, f"ppt/charts/{chart}"), _part(out, f"ppt/charts/{chart}")
        for tag in ("numCache", "f"):
            assert [etree.tostring(item) for item in before.iter(f"{{{NS['c']}}}{tag}")] == [
                etree.tostring(item) for item in after.iter(f"{{{NS['c']}}}{tag}")
            ]


def test_multi_level_categories_reach_the_right_cells(tmp_path: Path) -> None:
    _result, out = _write(tmp_path)
    sheet = _workbook(out, "chart2.xml").active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=1, max_row=5)]
    assert ["Zone nord", "Ferraillage"] == rows[1][:2]
    assert rows[2][1] == "Coffrage"
    assert ["Zone sud", "Ferraillage"] == rows[3][:2]
    assert str(rows[4][1]) == "2025"
    assert rows[0][2] == "Effectif"


def test_bilingual_writes_source_line_break_translation(tmp_path: Path) -> None:
    result, out = _write(tmp_path, bilingual=True)
    assert result.review == []
    values = _cached(out, "chart3.xml")
    assert "人工费\nMain-d’œuvre" in values
    assert _workbook(out, "chart3.xml").active["A2"].value == "人工费\nMain-d’œuvre"
    table = next(iter(_workbook(out, "chart1.xml").active.tables.values()))
    assert [column.name for column in table.tableColumns][1] == "计划产值\nProduction prévue"
    with zipfile.ZipFile(out) as package:
        rels = etree.fromstring(package.read("ppt/charts/_rels/chart1.xml.rels"))
        embedded = zipfile.ZipFile(io.BytesIO(package.read("ppt/" + rels[0].get("Target").removeprefix("../"))))
        assert b'name="\xe8\xae\xa1\xe5\x88\x92\xe4\xba\xa7\xe5\x80\xbc_x000a_Production pr' in embedded.read("xl/tables/table1.xml")
    # 标题：原段后接一段译文。
    title = _part(out, "ppt/charts/chart1.xml").find(".//c:title/c:tx/c:rich", NS)
    assert ["".join(p.itertext()) for p in title.findall("a:p", NS)] == [
        "2026 年各季度产值（万元）",
        "Production par trimestre en 2026 (10 000 CNY)",
    ]


def _rewrite(source: Path, target: Path, name: str, change) -> None:
    with zipfile.ZipFile(source) as package, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as out:
        for info in package.infolist():
            data = package.read(info.filename)
            if info.filename == name:
                data = change(data)
            out.writestr(info, data)


def test_newer_chart_types_are_counted_not_translated(tmp_path: Path) -> None:
    source = tmp_path / "chartex.pptx"
    _rewrite(
        GRAPHICS,
        source,
        "ppt/slides/slide3.xml",
        lambda data: data.replace(
            b"http://schemas.openxmlformats.org/drawingml/2006/chart",
            b"http://schemas.microsoft.com/office/drawing/2014/chartex",
        ),
    )
    assert scan_file(source).warnings == ["1 个图表中的文字未翻译"]
    assert not [unit for unit in extract_units(source) if "chart3" in unit.id]


def test_chart_without_embedded_workbook_is_listed_for_review(tmp_path: Path) -> None:
    source = tmp_path / "linked.pptx"
    _rewrite(
        GRAPHICS,
        source,
        "ppt/charts/_rels/chart3.xml.rels",
        lambda data: re.sub(rb'Target="[^"]+"', b'Target="file:///C:/data.xlsx" TargetMode="External"', data),
    )
    result, out = _write(tmp_path, source=source)
    assert [(item.context, item.reason) for item in result.review] == [("第 3 页 · Chart 2", REASON_DATA_NOT_SYNCED)]
    assert "Main-d’œuvre" in _cached(out, "chart3.xml")


def test_formula_cells_are_left_and_listed(tmp_path: Path) -> None:
    with zipfile.ZipFile(GRAPHICS) as package:
        rels = etree.fromstring(package.read("ppt/charts/_rels/chart3.xml.rels"))
        name = "ppt/" + rels[0].get("Target").removeprefix("../")
    source = tmp_path / "formula.pptx"

    def with_formula(data: bytes) -> bytes:
        book = openpyxl.load_workbook(io.BytesIO(data))
        book.active["A2"] = '="人工"&"费"'
        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    _rewrite(GRAPHICS, source, name, with_formula)
    result, out = _write(tmp_path, source=source)
    assert [item.reason for item in result.review] == [REASON_DATA_NOT_SYNCED]
    sheet = _workbook(out, "chart3.xml").active
    assert sheet["A2"].value == '="人工"&"费"'
    assert sheet["A3"].value == "Matériaux"


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("Sheet1!$A$2:$A$5", ("Sheet1", 2, 1, 5, 1)),
        ("'销售 数据'!$B$1", ("销售 数据", 1, 2, 1, 2)),
        ("'It''s'!A1:C1", ("It's", 1, 1, 1, 3)),
        ("(Sheet1!$A$2,Sheet1!$A$4)", None),
        ("[1]Sheet1!$A$1", None),
        ("Sales", None),
    ],
)
def test_parse_range(formula: str, expected) -> None:
    rng = _parse_range(formula)
    assert (None if rng is None else (rng.sheet, rng.r1, rng.c1, rng.r2, rng.c2)) == expected


def test_patch_workbook_refuses_duplicate_table_names() -> None:
    with zipfile.ZipFile(GRAPHICS) as package:
        rels = etree.fromstring(package.read("ppt/charts/_rels/chart1.xml.rels"))
        data = package.read("ppt/" + rels[0].get("Target").removeprefix("../"))
    patched, missed = patch_workbook(
        data, {("Sheet1", 1, 2): ("计划产值", "Production"), ("Sheet1", 1, 3): ("实际产值", "production")}
    )
    assert patched is None
    assert missed == {("Sheet1", 1, 2), ("Sheet1", 1, 3)}
    patched, missed = patch_workbook(data, {("Sheet1", 2, 1): ("不是这个字", "X")})
    assert patched is None and missed == {("Sheet1", 2, 1)}


def test_untouched_parts_stay_byte_for_byte(tmp_path: Path) -> None:
    copy = tmp_path / "copy.pptx"
    shutil.copyfile(GRAPHICS, copy)
    write_file(copy, tmp_path / "same.pptx", {}, target_lang="fr")
    with zipfile.ZipFile(GRAPHICS) as before, zipfile.ZipFile(tmp_path / "same.pptx") as after:
        assert [(info.filename, info.CRC) for info in before.infolist()] == [
            (info.filename, info.CRC) for info in after.infolist()
        ]
