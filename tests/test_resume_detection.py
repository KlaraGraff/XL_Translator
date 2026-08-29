"""core/resume_detection.py 的契约测试。

被测对象是纯只读的 detect_previous_output + match_previous_output（见该模块顶部
docstring）。夹具尽量复用生产代码的产物写出函数（write_bilingual_file /
write_bilingual_docx / write_pdf_manifest_and_report），而不是手搓字节——这样
命名规则一旦漂移，夹具会跟着漂移，测试也就跟着失真警报，不会盲测。

候选目录名的时间戳直接编进目录名（core.file_scanner.GENERATED_OUTPUT_DIR_MARKER
分隔），不依赖真实系统时间，排序测试才能稳定。
"""
from __future__ import annotations

import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from PIL import Image

from core.bilingual_writer import bilingual_output_name, write_bilingual_file
from core.language_registry import get_target_lang_display
from core.file_scanner import GENERATED_OUTPUT_DIR_MARKER
from core.pdf_image_translation import (
    PDF_MANIFEST_FILENAME,
    PDF_OUTPUT_STATE_COMPLETED,
    PdfFileRecord,
    PdfPageRecord,
    PdfTaskSummary,
    page_image_name,
    resolve_pdf_page_archive_dirs,
    reusable_pdf_pages,
    translated_pdf_base_name,
    write_pdf_manifest_and_report,
)
from core.resume_detection import (
    baseline_missing_source_texts,
    detect_previous_output,
    match_previous_output,
)
from core.translation_coverage import COVERAGE_COVERED, CoverageUnit
from core.word_document import write_bilingual_docx
from settings import AppSettings


# ---------------------------------------------------------------------------
# 夹具
# ---------------------------------------------------------------------------


def _history_dir(root: Path, timestamp: str, *, name_root: Path | None = None) -> Path:
    """在 root 下建一个历史输出目录，目录名取自 name_root（默认就是 root 自己）。

    custom_output_root 场景下，历史目录建在别处，但目录名仍然是"源文件夹名"。
    """
    base = name_root or root
    path = root / f"{base.name}{GENERATED_OUTPUT_DIR_MARKER}{timestamp}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _make_excel_source(path: Path, texts: list[str]) -> Path:
    wb = Workbook()
    ws = wb.active
    for i, text in enumerate(texts, start=1):
        ws.cell(row=i, column=1, value=text)
    wb.save(path)
    return path


def _write_excel_bilingual(
    source: Path,
    output_dir: Path,
    translations: dict[str, str],
    target_lang: str = "en",
) -> Path:
    return write_bilingual_file(
        source_path=source,
        output_dir=output_dir,
        translations=translations,
        target_lang=target_lang,
        keep_original_sheets=False,
        formula_display_value_backfill=True,
    )


def _make_word_source(path: Path, paragraphs: list[str]) -> Path:
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(path)
    return path


def _write_word_bilingual(
    source: Path,
    output_dir: Path,
    translations: dict[str, str],
    target_lang: str = "en",
) -> Path:
    return write_bilingual_docx(
        source_path=source,
        output_dir=output_dir,
        translations=translations,
        target_lang=target_lang,
    )


def _pdf_output_filename(
    relative_path: str, target_lang: str, variant_label: str | None = None
) -> str:
    """独立于被测模块，用生产函数算出"真实"产物名，供夹具与断言双方共用同一把尺。"""
    settings = AppSettings(target_lang=target_lang)
    return translated_pdf_base_name(
        Path(relative_path).name, target_lang, settings, variant_label=variant_label
    )


def _tiny_png_bytes() -> bytes:
    """一张真能被 PIL 打开的最小 PNG——resume_detection 复用页判定要求页图能通过
    ``Image.open`` 的文件头解析（reusable_pdf_pages(verify_images=False) 跳过整图
    解码，但 open() 本身的头部解析跳不过），随手写的字节（如 b"fake-png"）过不了。
    """
    buf = io.BytesIO()
    Image.new("RGB", (1, 1), color=(255, 255, 255)).save(buf, format="PNG")
    return buf.getvalue()


def _write_pdf_history(
    root: Path,
    *,
    timestamp: str,
    relative_path: str,
    page_count: int,
    page_statuses: list[str],
    target_lang: str = "en",
    stopped: bool | None = False,
    missing_png_pages: frozenset[int] = frozenset(),
    name_root: Path | None = None,
    source_pdf_size_bytes: int = 0,
) -> Path:
    """用生产的 manifest 写出函数造一份历史 PDF 产物包。

    stopped=None 表示手工把 "stopped" 键从写出的 manifest 里删掉，模拟没有这个
    字段的旧 manifest；missing_png_pages 让某些页"manifest 记成功但 png 缺失"。
    """
    history = _history_dir(root, timestamp, name_root=name_root)
    pages = [
        PdfPageRecord(
            page_number=i,
            source_image_path="",
            file_name=Path(relative_path).name,
            status=status,
            # runner 复用页面要求页面几何存在（page_width_pt/page_height_pt > 0）——
            # 没有尺寸的页装订出来是一张 0×0 的废页。数值本身不影响判定，随手给
            # 一个 A4 尺寸占位。
            page_width_pt=595.0,
            page_height_pt=842.0,
        )
        for i, status in enumerate(page_statuses, start=1)
    ]
    record = PdfFileRecord(
        name=Path(relative_path).name,
        source_path=str(root / relative_path),
        relative_path=relative_path,
        page_count=page_count,
        pages=pages,
        source_pdf_size_bytes=source_pdf_size_bytes,
    )
    summary = PdfTaskSummary(
        status=PDF_OUTPUT_STATE_COMPLETED,
        output_dir=str(history),
        target_lang=target_lang,
        target_lang_label=target_lang,
        started_at="2026-01-01T00:00:00",
        completed_at="2026-01-01T00:01:00",
        elapsed_sec=60.0,
        file_count=1,
        total_page_count=page_count,
        generated_pdf_count=1,
        placeholder_page_count=0,
        emergency_ratio_normalized_count=0,
        retry_count=0,
        stopped=bool(stopped) if stopped is not None else False,
        files=[record],
    )
    write_pdf_manifest_and_report(summary)

    if stopped is None:
        manifest_path = history / PDF_MANIFEST_FILENAME
        data = json.loads(manifest_path.read_text(encoding="utf-8"))
        del data["stopped"]
        manifest_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    _, translated_dir = resolve_pdf_page_archive_dirs(history, Path(relative_path))
    translated_dir.mkdir(parents=True, exist_ok=True)
    for i, status in enumerate(page_statuses, start=1):
        if status in ("success", "emergency_normalized") and i not in missing_png_pages:
            (translated_dir / page_image_name(i, page_count)).write_bytes(_tiny_png_bytes())

    return history


# ---------------------------------------------------------------------------
# detect_previous_output —— excel / word 状态判定
# ---------------------------------------------------------------------------


def test_single_folder_hit_full_translation(tmp_path):
    """单文件夹场景：命中历史目录，全部翻完判 full。"""
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称", "施工内容"])
    history = _write_excel_bilingual(
        source, _history_dir(root, "20260101_010101"),
        {"项目名称": "Project name", "施工内容": "Construction scope"},
    )

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result is not None
    assert result["selected_dir"] == str(history.parent)
    assert len(result["candidates"]) == 1
    assert result["candidates"][0]["dir"] == str(history.parent)
    entry = result["files"][0]
    assert entry["path"] == str(source)
    assert entry["status"] == "full"
    assert entry["untranslated_count"] == 0
    assert entry["matched_output"] is not None
    assert result["summary"] == {
        "full": 1, "partial": 0, "none": 0, "found": 0, "diverged": 0,
        "page_done": None, "page_total": None,
        "last_task_stopped": None,
        "last_time_label": "01-01 01:01",
    }


def test_excel_partial_translation_counts_untranslated(tmp_path):
    """部分翻完：先用 write_bilingual_file 翻一部分，剩下的不给译文，让覆盖率计划去数。"""
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称", "施工内容", "设备安装"])
    history_dir = _history_dir(root, "20260101_010101")
    # 只翻一格，另外两格留白——build_excel_coverage_plan 应把它们计入 source_only。
    _write_excel_bilingual(source, history_dir, {"项目名称": "Project name"})

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "partial"
    assert entry["untranslated_count"] == 2
    assert result["summary"]["partial"] == 1
    assert result["summary"]["full"] == 0


def test_excel_ambiguous_cells_are_not_counted_as_untranslated(tmp_path):
    """底稿里剩下的都是 AMBIGUOUS（多行歧义格，拆不出原文/译文边界）时不算未译。

    补译写回只处理 source_only 的 unit，AMBIGUOUS 碰都不碰——算进未译数会出现
    「剩 1 处 → 续译 → 还是剩 1 处」的死循环观感。这里底稿有一格正常译完
    （covered），另一格是多行歧义内容（"A\\nB"，不含中文也凑不出目标语言证据，
    分类器只能判 ambiguous），源文件里对应位置内容和底稿一模一样（这份歧义
    内容从没被翻译流程动过）。untranslated_count 必须是 0，状态 full。
    """
    root = tmp_path / "docs"
    root.mkdir()
    source = root / "a.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="项目名称")
    ws.cell(row=2, column=1, value="A\nB")
    wb.save(source)

    history_dir = _history_dir(root, "20260101_010101")
    baseline = history_dir / bilingual_output_name("a.xlsx", get_target_lang_display("en"))
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.cell(row=1, column=1, value="项目名称\nProject name")
    ws2.cell(row=2, column=1, value="A\nB")  # 歧义格未被翻译流程处理，原样留在底稿里
    wb2.save(baseline)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "full"
    assert entry["untranslated_count"] == 0


def test_excel_source_grown_since_baseline_is_reported_as_diverged(tmp_path):
    """底稿本身已经全翻完了，但源文件在上次翻译之后新长出了一格没见过的内容。

    这种分歧 runner 会整个拒用底稿、按源文件整份重翻（已有译文靠翻译记忆复用）。
    检测必须原样上报 diverged，不能把「整份重翻」折算成一个小小的补译数字——
    旧契约（判 partial、untranslated_count=1）承诺的是一次并不存在的省钱，已废除。
    """
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称", "施工内容"])
    history_dir = _history_dir(root, "20260101_010101")
    _write_excel_bilingual(
        source, history_dir, {"项目名称": "Project name", "施工内容": "Construction scope"}
    )

    # 源文件翻完之后又长了一格新内容，底稿完全没见过它。
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="项目名称")
    ws.cell(row=2, column=1, value="施工内容")
    ws.cell(row=3, column=1, value="设备安装")
    wb.save(source)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "diverged"
    assert entry["untranslated_count"] is None
    assert result["summary"]["diverged"] == 1
    assert result["summary"]["partial"] == 0


def test_excel_new_row_reusing_existing_text_is_diverged(tmp_path):
    """新增的行哪怕逐字复用了底稿里出现过的文本，也是底稿没有的内容。

    比对必须按出现次数（多重集）而不是文本集合：集合比对会认为「这句话见过」
    而放行，续译换底稿后这一行既不翻译也不进产物——静默丢内容。
    """
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    _write_excel_bilingual(
        source, _history_dir(root, "20260101_010101"), {"项目名称": "Project name"}
    )

    # 又长出一行，文字和第一行一模一样。
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="项目名称")
    ws.cell(row=2, column=1, value="项目名称")
    wb.save(source)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "diverged"
    assert entry["untranslated_count"] is None


def test_excel_text_redistributed_to_new_sheet_is_diverged(tmp_path):
    """总次数不变、只是挪进新表的文本也判 diverged——比对键必须带分表名。

    旧表两行「项目名称」都已译好；这次源文件把其中一行挪进整张新增的「新增分册」
    表。不带分表名的多重集比对会看到总次数相等（2 对 2）而放行，整张新表被静默
    判定为「已覆盖」丢掉；带分表名后新表那一格找不到归属，必须判 diverged。
    """
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称", "项目名称"])
    _write_excel_bilingual(
        source,
        _history_dir(root, "20260101_010101"),
        {"项目名称": "Project name"},
    )

    # 一行留在原表，另一行挪进整张新增的表——文本总次数与旧表持平。
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="项目名称")
    extra = wb.create_sheet("新增分册")
    extra.cell(row=1, column=1, value="项目名称")
    wb.save(source)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "diverged"
    assert entry["untranslated_count"] is None
    assert result["summary"]["diverged"] == 1


def test_word_full_and_partial_translation(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()

    full_source = _make_word_source(root / "full.docx", ["项目名称", "施工内容"])
    _write_word_bilingual(
        full_source,
        _history_dir(root, "20260101_010101"),
        {"项目名称": "Project name", "施工内容": "Construction scope"},
    )

    result = detect_previous_output(
        surface="word",
        items=[{"path": str(full_source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )
    entry = result["files"][0]
    assert entry["status"] == "full"
    assert entry["untranslated_count"] == 0

    # 部分翻完场景在独立的根目录下做，避免和上面的历史目录混在一起判定候选时间戳。
    root2 = tmp_path / "docs2"
    root2.mkdir()
    partial_source = _make_word_source(root2 / "partial.docx", ["项目名称", "施工内容"])
    _write_word_bilingual(
        partial_source, _history_dir(root2, "20260101_010101"), {"项目名称": "Project name"}
    )
    result2 = detect_previous_output(
        surface="word",
        items=[{"path": str(partial_source)}],
        scan_roots=[str(root2)],
        custom_output_root=None,
        target_lang="en",
    )
    entry2 = result2["files"][0]
    assert entry2["status"] == "partial"
    assert entry2["untranslated_count"] == 1


def test_target_lang_mismatch_yields_status_none(tmp_path):
    """历史产物是按 en 翻的，这次扫描目标语言换成 fr——产物名对不上，判 none。"""
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    _write_excel_bilingual(
        source, _history_dir(root, "20260101_010101"), {"项目名称": "Project name"}, target_lang="en"
    )

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="fr",
    )

    entry = result["files"][0]
    assert entry["status"] == "none"
    assert entry["matched_output"] is None
    assert result["summary"]["none"] == 1


def test_renamed_source_file_does_not_match(tmp_path):
    """历史产物是按"a.xlsx"这个名字产出的；这次源文件改名成"b.xlsx"，镜像文件名对不上。"""
    root = tmp_path / "docs"
    root.mkdir()
    source_a = _make_excel_source(root / "a.xlsx", ["项目名称"])
    _write_excel_bilingual(
        source_a, _history_dir(root, "20260101_010101"), {"项目名称": "Project name"}
    )
    source_b = root / "b.xlsx"
    source_a.rename(source_b)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source_b)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "none"
    assert entry["matched_output"] is None


# ---------------------------------------------------------------------------
# 候选目录：排序 / 上限 / preferred_dir
# ---------------------------------------------------------------------------


def test_candidates_sorted_newest_first_and_capped_at_five(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.xlsx").write_bytes(b"placeholder")  # 候选目录列举不看内容，随便一个源文件占位即可
    timestamps = [f"2026010{i}_010101" for i in range(1, 8)]  # 7 个，超过上限 5
    for ts in timestamps:
        _history_dir(root, ts)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(root / "a.xlsx")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert len(result["candidates"]) == 5
    expected_order = sorted(timestamps, reverse=True)[:5]
    got_order = [c["dir"].split(GENERATED_OUTPUT_DIR_MARKER)[-1] for c in result["candidates"]]
    assert got_order == expected_order
    # 未指定 preferred_dir 时，统计目录取最新的那个。
    assert result["selected_dir"].endswith(expected_order[0])


def test_output_dir_lookup_treats_brackets_as_literal_characters(tmp_path):
    """根文件夹名带 glob 元字符（方括号）时，目录发现按字面量比对，不当 glob 模式解释。

    旧实现用 ``parent.glob(f"{root_name}{MARKER}*")``：根名里的 "[2024]" 会被
    fnmatch 解释成「匹配 2/0/4 中任意一个字符」的字符类，于是"报价[2024]_翻译输出_*"
    这个模式反而会去匹配一个毫不相关、字面量叫"报价2_翻译输出_..."的目录（"2" 落在
    字符类 [2024] 里）——自己的产物目录找不到，不相关的目录却认成了自己的。改成
    iterdir + startswith 字面量比对后两者都不会再发生。
    """
    root = tmp_path / "报价[2024]"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    true_output = _write_excel_bilingual(
        source, _history_dir(root, "20250101_120000"), {"项目名称": "Project name"}
    ).parent

    # 诱饵一：字面量前缀之外多一个字符——不能被当成"报价[2024]"自己的产物目录。
    decoy_prefix_suffix = root / f"报价[2024]x{GENERATED_OUTPUT_DIR_MARKER}20250101_120000"
    decoy_prefix_suffix.mkdir()
    # 诱饵二：旧 glob 实现下 "报价2" 会落进字符类 [2024] 而被误认成"报价[2024]"的产物。
    decoy_character_class = root / f"报价2{GENERATED_OUTPUT_DIR_MARKER}20250101_120000"
    decoy_character_class.mkdir()

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result is not None
    assert result["selected_dir"] == str(true_output)
    assert len(result["candidates"]) == 1
    candidate_dirs = {c["dir"] for c in result["candidates"]}
    assert candidate_dirs == {str(true_output)}
    assert str(decoy_prefix_suffix) not in candidate_dirs
    assert str(decoy_character_class) not in candidate_dirs
    assert result["files"][0]["status"] == "full"


def test_preferred_dir_takes_priority_over_latest(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    older = _write_excel_bilingual(
        source, _history_dir(root, "20260101_010101"), {"项目名称": "Project name"}
    ).parent
    _history_dir(root, "20260201_010101")  # 更新，但不含任何匹配产物

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
        preferred_dir=str(older),
    )

    assert result["selected_dir"] == str(older)
    # preferred 指向的旧目录里有匹配产物，判定 full；如果错用了最新目录（没有产物）会判 none。
    assert result["files"][0]["status"] == "full"


def test_invalid_preferred_dir_is_ignored(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    newest = _write_excel_bilingual(
        source, _history_dir(root, "20260101_010101"), {"项目名称": "Project name"}
    ).parent
    not_a_dir = root / "not_a_real_dir"  # 不存在
    wrong_name_dir = root / "随便起的名字"  # 存在但不含 _翻译输出_ 标记
    wrong_name_dir.mkdir()

    for bad_preferred in (str(not_a_dir), str(wrong_name_dir)):
        result = detect_previous_output(
            surface="excel",
            items=[{"path": str(source)}],
            scan_roots=[str(root)],
            custom_output_root=None,
            target_lang="en",
            preferred_dir=bad_preferred,
        )
        assert result["selected_dir"] == str(newest)
        assert result["files"][0]["status"] == "full"


# ---------------------------------------------------------------------------
# 多根扫描
# ---------------------------------------------------------------------------


def test_multi_root_scan_has_empty_candidates_and_independent_latest_dirs(tmp_path):
    """多选零散文件（多根）：candidates=[]，且每个文件独立取"自己所在根"的最新目录。"""
    root_a = tmp_path / "folder_a"
    root_a.mkdir()
    root_b = tmp_path / "folder_b"
    root_b.mkdir()

    source_a = _make_excel_source(root_a / "a.xlsx", ["项目名称"])
    # root_a 有两个历史目录：旧的没有匹配产物，新的才有——用来确认"取各自最新"。
    _history_dir(root_a, "20260101_010101")
    _write_excel_bilingual(
        source_a, _history_dir(root_a, "20260102_010101"), {"项目名称": "Project name"}
    )

    source_b = _make_excel_source(root_b / "b.xlsx", ["设备安装"])
    _write_excel_bilingual(
        source_b, _history_dir(root_b, "20260101_010101"), {"设备安装": "Equipment installation"}
    )

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source_a)}, {"path": str(source_b)}],
        scan_roots=[str(root_a), str(root_b)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result["candidates"] == []
    by_path = {entry["path"]: entry for entry in result["files"]}
    assert by_path[str(source_a)]["status"] == "full"
    assert by_path[str(source_b)]["status"] == "full"


# ---------------------------------------------------------------------------
# custom_output_root
# ---------------------------------------------------------------------------


def test_custom_output_root_directory_is_discovered(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    custom_root = tmp_path / "custom_out"
    custom_root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    # 历史目录建在 custom_root 下，但目录名仍然是"源文件夹名_翻译输出_时间戳"。
    history = _write_excel_bilingual(
        source, _history_dir(custom_root, "20260101_010101", name_root=root),
        {"项目名称": "Project name"},
    ).parent

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=str(custom_root),
        target_lang="en",
    )

    assert result["selected_dir"] == str(history)
    assert result["files"][0]["status"] == "full"


# ---------------------------------------------------------------------------
# PDF：manifest 降级 / full / partial / stopped
# ---------------------------------------------------------------------------


def test_pdf_corrupted_manifest_degrades_to_found(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4 placeholder")
    history = _history_dir(root, "20260101_010101")
    (history / PDF_MANIFEST_FILENAME).write_text("{not valid json", encoding="utf-8")
    # manifest 坏了之后，检测退化成"直接找产物文件"——文件确实存在时应当是 found。
    (history / _pdf_output_filename("a.pdf", "en")).write_bytes(b"%PDF-1.4 translated")

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "found"
    assert entry["matched_output"] is not None
    assert entry["page_done"] is None
    assert entry["page_total"] is None


def test_pdf_manifest_missing_and_no_matched_file_is_none(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4 placeholder")
    _history_dir(root, "20260101_010101")  # 空目录，什么产物都没有

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result["files"][0]["status"] == "none"


def test_pdf_partial_translation_reports_page_progress(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=3,
        page_statuses=["success", "success", "pending"],
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "partial"
    assert entry["page_done"] == 2
    assert entry["page_total"] == 3
    assert result["summary"]["page_done"] == 2
    assert result["summary"]["page_total"] == 3


def test_pdf_manifest_language_mismatch_ignores_manifest(tmp_path):
    """manifest 是按 en 翻的，这次扫描换成 fr——整份清单对本次毫无用处。

    PDF runner 的语言闸会整批拒绝复用旧页面；检测若照旧按清单报「上次已翻完
    N/N 页」，弹窗就在承诺一次 runner 根本不会兑现的复用。语言不一致时清单当
    不存在处理，退到按本次语言的产物命名规则找文件——找不到就是 none；清单里
    的 stopped 状态说的也是旧语言那次任务，同样不得外泄。
    """
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=2,
        page_statuses=["success", "success"],
        target_lang="en",
        stopped=True,
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="fr",
    )

    entry = result["files"][0]
    assert entry["status"] == "none"
    assert entry["page_done"] is None
    assert entry["page_total"] is None
    assert result["summary"]["last_task_stopped"] is None


def test_pdf_source_size_change_is_diverged(tmp_path):
    """源 PDF 在上次翻译之后被改过（体积对不上 manifest 记录）——按旧记录报
    页进度就是在替一份过期产物背书。runner 的体积闸会整份重新生成，检测同口径
    判 diverged。"""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=2,
        page_statuses=["success", "success"],
        source_pdf_size_bytes=999_999,  # 与磁盘上 8 字节的占位源文件对不上
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "diverged"
    assert entry["page_done"] is None
    assert entry["page_total"] is None
    assert result["summary"]["diverged"] == 1
    assert result["summary"]["page_done"] is None


def test_pdf_source_size_match_keeps_page_progress(tmp_path):
    """体积闸不误伤：manifest 记录的体积与磁盘一致时，页进度照常上报。"""
    root = tmp_path / "docs"
    root.mkdir()
    payload = b"%PDF-1.4"
    (root / "a.pdf").write_bytes(payload)
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=2,
        page_statuses=["success", "success"],
        source_pdf_size_bytes=len(payload),
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "full"
    assert entry["page_done"] == 2
    assert entry["page_total"] == 2


def test_emergency_normalized_pages_do_not_count_as_done(tmp_path):
    """应急比例归一化页跑完了但走了应急降级，续译判定不能把它算成"翻完了"——
    真跑续译时 runner 的 reusable_pdf_pages 只认 status=="success"，emergency_normalized
    会被重新生成。这里 1 success + 1 emergency_normalized：page_done 只数那 1 页
    success，总数没到 2，判 partial（不是旧口径下的 full）。"""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=2,
        page_statuses=["success", "emergency_normalized"],
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["status"] == "partial"
    assert entry["page_done"] == 1
    assert entry["page_total"] == 2


def test_pdf_success_page_with_missing_png_does_not_count_as_done(tmp_path):
    """manifest 说这一页成功了，但 translated_pages 里的 png 实际不存在——绝不能算完成。"""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=2,
        page_statuses=["success", "success"],
        missing_png_pages=frozenset({2}),
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["page_done"] == 1
    assert entry["page_total"] == 2
    assert entry["status"] == "partial"


def test_pdf_page_done_matches_reusable_pdf_pages_exactly(tmp_path):
    """反漂移锚点：detect_previous_output 报的 page_done 必须和 runner 真正复用时会
    用到的 reusable_pdf_pages 数出来的页数一字不差——两边一旦有一个改了判定标准
    没改另一个，这个测试就会红。

    一份 manifest 里塞 5 页，覆盖 reusable_pdf_pages 排除的每一种理由：
      1. success + 几何 + 真图 —— 唯一该算完成的页
      2. success 但带 quality_flags —— 内容质检有疑点，不算
      3. emergency_normalized —— 应急降级，不算
      4. success 但译后图缺失 —— 不算
      5. success 但 review_status=="failed" —— 审核判失败，不算
    """
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    history = _history_dir(root, "20260101_010101")
    relative_path = "a.pdf"
    page_count = 5

    pages = [
        PdfPageRecord(
            page_number=1, source_image_path="", file_name="a.pdf", status="success",
            page_width_pt=595.0, page_height_pt=842.0,
        ),
        PdfPageRecord(
            page_number=2, source_image_path="", file_name="a.pdf", status="success",
            page_width_pt=595.0, page_height_pt=842.0, quality_flags=["near_blank"],
        ),
        PdfPageRecord(
            page_number=3, source_image_path="", file_name="a.pdf", status="emergency_normalized",
            page_width_pt=595.0, page_height_pt=842.0,
        ),
        PdfPageRecord(
            page_number=4, source_image_path="", file_name="a.pdf", status="success",
            page_width_pt=595.0, page_height_pt=842.0,
        ),
        PdfPageRecord(
            page_number=5, source_image_path="", file_name="a.pdf", status="success",
            page_width_pt=595.0, page_height_pt=842.0, review_status="failed",
        ),
    ]
    record = PdfFileRecord(
        name="a.pdf", source_path=str(root / relative_path), relative_path=relative_path,
        page_count=page_count, pages=pages,
    )
    summary = PdfTaskSummary(
        status=PDF_OUTPUT_STATE_COMPLETED, output_dir=str(history), target_lang="en",
        target_lang_label="en", started_at="2026-01-01T00:00:00", completed_at="2026-01-01T00:01:00",
        elapsed_sec=60.0, file_count=1, total_page_count=page_count, generated_pdf_count=1,
        placeholder_page_count=0, emergency_ratio_normalized_count=1, retry_count=0,
        stopped=False, files=[record],
    )
    write_pdf_manifest_and_report(summary)

    _, translated_dir = resolve_pdf_page_archive_dirs(history, Path(relative_path))
    translated_dir.mkdir(parents=True, exist_ok=True)
    # 页 4（missing image）故意不写；其余"manifest 记成功/应急"的页都给真图，
    # 确保它们被排除是因为判定标准本身，而不是碰巧图也没有。
    for page_number in (1, 2, 3, 5):
        (translated_dir / page_image_name(page_number, page_count)).write_bytes(_tiny_png_bytes())

    manifest = json.loads((history / PDF_MANIFEST_FILENAME).read_text(encoding="utf-8"))
    expected_pages = reusable_pdf_pages(manifest, history, verify_images=False).get(relative_path, [])

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    entry = result["files"][0]
    assert entry["page_done"] == len(expected_pages)
    # 明确锚定具体数字，不能只满足"两边相等"就放过两边一起漂移到别的值。
    assert entry["page_done"] == 1


def test_pdf_page_count_dirty_value_degrades_without_crashing(tmp_path):
    """manifest 里一个文件的 page_count 是脏值（"unknown"，比如手工改坏或历史遗留）：
    _safe_int 把它兜成 0，这个文件的分类走"found"分支而不是让 int() 直接抛出去；
    同一批次里另一个健康文件不受影响，各自独立分类（逐文件隔离）。"""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    (root / "b.pdf").write_bytes(b"%PDF-1.4")
    history = _history_dir(root, "20260101_010101")

    def _one_page_record(name: str) -> PdfFileRecord:
        pages = [
            PdfPageRecord(
                page_number=1, source_image_path="", file_name=name, status="success",
                page_width_pt=595.0, page_height_pt=842.0,
            )
        ]
        return PdfFileRecord(
            name=name, source_path=str(root / name), relative_path=name,
            page_count=1, pages=pages,
        )

    bad_record = _one_page_record("a.pdf")
    good_record = _one_page_record("b.pdf")
    summary = PdfTaskSummary(
        status=PDF_OUTPUT_STATE_COMPLETED, output_dir=str(history), target_lang="en",
        target_lang_label="en", started_at="2026-01-01T00:00:00", completed_at="2026-01-01T00:01:00",
        elapsed_sec=60.0, file_count=2, total_page_count=2, generated_pdf_count=2,
        placeholder_page_count=0, emergency_ratio_normalized_count=0, retry_count=0,
        stopped=False, files=[bad_record, good_record],
    )
    write_pdf_manifest_and_report(summary)

    for name in ("a.pdf", "b.pdf"):
        _, translated_dir = resolve_pdf_page_archive_dirs(history, Path(name))
        translated_dir.mkdir(parents=True, exist_ok=True)
        (translated_dir / page_image_name(1, 1)).write_bytes(_tiny_png_bytes())

    # 手工把 a.pdf 的 page_count 改成一个脏字符串，模拟坏字段。
    manifest_path = history / PDF_MANIFEST_FILENAME
    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    for file_entry in data["files"]:
        if file_entry["relative_path"] == "a.pdf":
            file_entry["page_count"] = "unknown"
    manifest_path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}, {"path": str(root / "b.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result is not None
    by_path = {entry["path"]: entry for entry in result["files"]}
    bad_entry = by_path[str(root / "a.pdf")]
    # 脏 page_count 被 _safe_int 兜成 0：不崩、也不误报 full/partial。
    assert bad_entry["status"] == "found"
    assert bad_entry["page_total"] == 0
    good_entry = by_path[str(root / "b.pdf")]
    assert good_entry["status"] == "full"
    assert good_entry["page_done"] == 1
    assert good_entry["page_total"] == 1


def test_pdf_last_task_stopped_flag_and_last_time_label(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260828_154200",
        relative_path="a.pdf",
        page_count=2,
        page_statuses=["success", "pending"],
        stopped=True,
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result["summary"]["last_task_stopped"] is True
    assert result["summary"]["last_time_label"] == "08-28 15:42"


def test_pdf_manifest_without_stopped_field_reports_none(tmp_path):
    """老 manifest 没有 stopped 字段时，不能瞎猜成 False——必须是 None。"""
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.pdf").write_bytes(b"%PDF-1.4")
    _write_pdf_history(
        root,
        timestamp="20260101_010101",
        relative_path="a.pdf",
        page_count=1,
        page_statuses=["success"],
        stopped=None,
    )

    result = detect_previous_output(
        surface="pdf",
        items=[{"path": str(root / "a.pdf")}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result["summary"]["last_task_stopped"] is None


# ---------------------------------------------------------------------------
# 异常兜底
# ---------------------------------------------------------------------------


def test_detect_previous_output_swallows_internal_exception(tmp_path, monkeypatch):
    """检测内部任何一步炸了，detect_previous_output 必须吞掉异常返回 None，不能拖垮扫描。"""
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])
    _write_excel_bilingual(
        source, _history_dir(root, "20260101_010101"), {"项目名称": "Project name"}
    )

    import core.resume_detection as resume_detection_module

    def _boom(*_args, **_kwargs):
        raise RuntimeError("simulated internal failure")

    monkeypatch.setattr(resume_detection_module, "_root_output_dirs", _boom)

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result is None


def test_detect_previous_output_returns_none_when_no_history_at_all(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    source = _make_excel_source(root / "a.xlsx", ["项目名称"])

    result = detect_previous_output(
        surface="excel",
        items=[{"path": str(source)}],
        scan_roots=[str(root)],
        custom_output_root=None,
        target_lang="en",
    )

    assert result is None


# ---------------------------------------------------------------------------
# match_previous_output —— 直接单测（PDF 修订号 / 形态标签顺序）
# ---------------------------------------------------------------------------


def test_match_previous_output_excel_xls_source_maps_to_xlsx_product(tmp_path):
    resume_dir = tmp_path / "resume"
    resume_dir.mkdir()
    # 产物永远是 .xlsx，即便源文件是老式 .xls。命名走生产同款函数，不手搓。
    expected = resume_dir / bilingual_output_name("报表.xlsx", get_target_lang_display("en"))
    expected.write_bytes(b"placeholder")

    found = match_previous_output(resume_dir, None, "报表.xls", "en", "excel")
    assert found == expected


def test_match_previous_output_word_normalizes_doc_extension(tmp_path):
    resume_dir = tmp_path / "resume"
    resume_dir.mkdir()
    expected = resume_dir / bilingual_output_name("合同.docx", get_target_lang_display("en"))
    expected.write_bytes(b"placeholder")

    found = match_previous_output(resume_dir, None, "合同.doc", "en", "word")
    assert found == expected


def test_match_previous_output_pdf_picks_highest_revision(tmp_path):
    resume_dir = tmp_path / "resume"
    resume_dir.mkdir()
    base_name = _pdf_output_filename("a.pdf", "en", variant_label="高清")
    stem, suffix = Path(base_name).stem, Path(base_name).suffix
    (resume_dir / f"{stem}{suffix}").write_bytes(b"r0")
    (resume_dir / f"{stem}_R1{suffix}").write_bytes(b"r1")
    highest = resume_dir / f"{stem}_R3{suffix}"
    highest.write_bytes(b"r3")

    found = match_previous_output(resume_dir, None, "a.pdf", "en", "pdf")
    assert found == highest


def test_match_previous_output_pdf_falls_back_across_variant_labels(tmp_path):
    """高清变体没有时，退到压缩变体；都没有再退到不带形态标签的纯语言文件。"""
    resume_dir = tmp_path / "resume"
    resume_dir.mkdir()
    compressed_name = _pdf_output_filename("a.pdf", "en", variant_label="压缩")
    expected = resume_dir / compressed_name
    expected.write_bytes(b"placeholder")

    found = match_previous_output(resume_dir, None, "a.pdf", "en", "pdf")
    assert found == expected


def test_match_previous_output_uses_relative_path_to_mirror_subdirectories(tmp_path):
    resume_dir = tmp_path / "resume"
    (resume_dir / "sub").mkdir(parents=True)
    expected = resume_dir / "sub" / bilingual_output_name("a.xlsx", get_target_lang_display("en"))
    expected.write_bytes(b"placeholder")

    found = match_previous_output(resume_dir, Path("sub"), "a.xlsx", "en", "excel")
    assert found == expected
    # 相对路径对不上（当作直接落在 resume_dir 根下）就找不到。
    assert match_previous_output(resume_dir, None, "a.xlsx", "en", "excel") is None


def test_match_previous_output_returns_none_for_unknown_surface(tmp_path):
    resume_dir = tmp_path / "resume"
    resume_dir.mkdir()
    assert match_previous_output(resume_dir, None, "a.txt", "en", "unknown") is None


# ---------------------------------------------------------------------------
# baseline_missing_source_texts —— Word 编号前缀容差（单测，直接传 baseline_plan）
# ---------------------------------------------------------------------------


def test_baseline_missing_source_texts_word_tolerates_numbering_prefix(tmp_path):
    """Word 预处理会给段落物化编号前缀："1.1 概述" 对源文件里的 "概述"。

    底稿是上次翻译时生成的，那时源文本已经带着编号前缀（"1.1 概述"）；这次重新
    从源文件提取，python-docx 读到的纯段落文字不含编号（"概述"）。两者其实是
    同一段内容，只是编号前缀有没有被物化的差异——baseline_missing_source_texts
    对 Word 表面有前缀容差（底稿文本以源文本结尾就算已覆盖），不能因为这个差异
    就误判成"源文件比底稿多了内容"。

    直接传 baseline_plan（伪造一个只含一条已覆盖 unit 的假底稿计划），跳过底稿
    docx 本身要不要真实存在——baseline_plan 给了就不会去开 baseline_path。
    """
    source = tmp_path / "source.docx"
    _make_word_source(source, ["概述"])

    baseline_plan = SimpleNamespace(
        units=[
            CoverageUnit(
                source_text="1.1 概述",
                target_text="1.1 Overview",
                status=COVERAGE_COVERED,
                location="body.paragraph[0]",
                kind="paragraph",
                reason="测试底稿：编号前缀已物化进源文本。",
            )
        ]
    )

    missing = baseline_missing_source_texts(
        source,
        tmp_path / "baseline_placeholder.docx",  # 不会被打开：baseline_plan 已给出
        surface="word",
        target_lang="en",
        source_lang="zh",
        baseline_plan=baseline_plan,
    )
    assert missing == []


def test_baseline_missing_source_texts_word_true_divergence_survives_prefix_tolerance(tmp_path):
    """前缀容差不能变成万能挡箭牌：源文件真的比底稿多出来的段落，照样要报出来。"""
    source = tmp_path / "source.docx"
    _make_word_source(source, ["概述", "新增章节"])

    baseline_plan = SimpleNamespace(
        units=[
            CoverageUnit(
                source_text="1.1 概述",
                target_text="1.1 Overview",
                status=COVERAGE_COVERED,
                location="body.paragraph[0]",
                kind="paragraph",
                reason="测试底稿：只覆盖了概述一段，没见过新增章节。",
            )
        ]
    )

    missing = baseline_missing_source_texts(
        source,
        tmp_path / "baseline_placeholder.docx",
        surface="word",
        target_lang="en",
        source_lang="zh",
        baseline_plan=baseline_plan,
    )
    assert missing == ["新增章节"]


def test_baseline_missing_source_texts_word_prefix_rescue_consumed_once(tmp_path):
    """前缀容差是消耗式的：一条带编号的底稿文本只能救一个源段落。

    源文件里有两段「概述」——第一段就是当年被物化成「1.1 概述」的那段（前缀
    容差理应放行），第二段是这次真正新增的短段落。旧实现的 endswith 扫描不消耗
    底稿条目，同一条「1.1 概述」把两段全救了，新增段落静默蒙混过关；消耗式比对
    下第二段必须报缺。
    """
    source = tmp_path / "source.docx"
    _make_word_source(source, ["概述", "概述"])

    baseline_plan = SimpleNamespace(
        units=[
            CoverageUnit(
                source_text="1.1 概述",
                target_text="1.1 Overview",
                status=COVERAGE_COVERED,
                location="body.paragraph[0]",
                kind="paragraph",
                reason="测试底稿：只有一段带编号的概述。",
            )
        ]
    )

    missing = baseline_missing_source_texts(
        source,
        tmp_path / "baseline_placeholder.docx",
        surface="word",
        target_lang="en",
        source_lang="zh",
        baseline_plan=baseline_plan,
    )
    assert missing == ["概述"]


def test_baseline_missing_source_texts_word_numbered_paragraphs_have_no_rescue_cap(tmp_path):
    """编号段落逐条走前缀容差是常态，数量不设上限。

    write_bilingual_docx 会把自动编号物化进底稿文本（「12. 条款说明12」对
    「条款说明12」），于是重扫一份完全译好的编号文档时，每个编号段都要靠前缀
    容差匹配。曾有实现给这条路径设 200 次上限，超过 200 个编号段的正常文档会被
    误判「源文件已变化」而丢掉底稿；现在按尾部字符分桶建索引，不设限。
    """
    n = 250
    source = tmp_path / "source.docx"
    _make_word_source(source, [f"条款说明{i}" for i in range(n)])

    baseline_plan = SimpleNamespace(
        units=[
            CoverageUnit(
                source_text=f"{i + 1}. 条款说明{i}",
                target_text=f"{i + 1}. Clause {i}",
                status=COVERAGE_COVERED,
                location=f"body.paragraph[{i}]",
                kind="paragraph",
                reason="测试底稿：物化了编号前缀的段落。",
            )
            for i in range(n)
        ]
    )

    missing = baseline_missing_source_texts(
        source,
        tmp_path / "baseline_placeholder.docx",
        surface="word",
        target_lang="en",
        source_lang="zh",
        baseline_plan=baseline_plan,
    )
    assert missing == []

    # 真正新增的段落不能被 250 个编号段「淹没」——照样逐条报出来。
    source_plus = tmp_path / "source_plus.docx"
    _make_word_source(
        source_plus, [f"条款说明{i}" for i in range(n)] + ["这次新增的段落"]
    )
    missing_plus = baseline_missing_source_texts(
        source_plus,
        tmp_path / "baseline_placeholder.docx",
        surface="word",
        target_lang="en",
        source_lang="zh",
        baseline_plan=baseline_plan,
    )
    assert missing_plus == ["这次新增的段落"]
