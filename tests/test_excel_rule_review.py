"""Excel 端到端：规则没过的译文交给模型复查（真跑 TaskRunner._run() + 真实写盘）。

只把翻译引擎和记忆库换成桩：
  - 复查判意思一致的（「共3层」→ Three floors in total）写进文件，不涂复核底色，报告里一行
    「复查后采用」（severity=resolved），不存进记忆库；
  - 判不一致的带着问题重译，重译稿通过规则就用；
  - 拿不准的照旧恢复原文、记进「质量校验回退」并带上原因。
"""

from __future__ import annotations

import json
import os
import tempfile
import threading
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from core.api_config_check import ApiConfigCheckResult
from core.file_scanner import FileItem
from core.task_runner import DoneMsg, TaskRunner
from settings import AppSettings, EngineSettings
from tests.rate_pins import pinned_surface_rate

ACCEPT = "共3层"
FIX = "25厚1:2.5水泥砂浆找平"
UNSURE = "设备3台"
PLAIN = "门"


class _ReviewingEngine:
    engine_name = "excel-review-test/mock"

    def __init__(self) -> None:
        self.lock = threading.Lock()
        self.prompts: list[str] = []
        self.reviewed: list[str] = []
        self.first_fix = True

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        out = {}
        with self.lock:
            self.prompts.append(system_prompt)
            for text in texts:
                if text == ACCEPT:
                    out[text] = "Three floors in total"
                elif text == FIX:
                    out[text] = "Mortar screed 1:2.5" if "【重译要求】" not in system_prompt else "Mortar screed 1:2.5, 25 mm thick"
                elif text == UNSURE:
                    out[text] = "Equipment"
                else:
                    out[text] = "Door"
        return out

    def chat(self, system: str, user: str) -> str:
        payload = json.loads(user)
        if "items" not in payload:  # 其它对话请求（语种探测等）
            raise NotImplementedError
        verdicts = {ACCEPT: "equivalent", FIX: "not_equivalent", UNSURE: "uncertain"}
        results = []
        with self.lock:
            for item in payload["items"]:
                self.reviewed.append(item["source"])
                results.append(
                    {"id": item["id"], "verdict": verdicts.get(item["source"], "uncertain"), "reason": "理由"}
                )
        return json.dumps({"results": results}, ensure_ascii=False)


class ExcelRuleReviewTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name) / "docs"
        self.root.mkdir()
        self._env = patch.dict(os.environ, {"TRANSLATOR_APP_DATA_DIR": str(Path(self._tmp.name) / "appdata")})
        self._env.start()

    def tearDown(self) -> None:
        self._env.stop()
        self._tmp.cleanup()

    def test_failed_rules_are_reviewed_fixed_or_kept(self) -> None:
        source = self.root / "清单.xlsx"
        wb = Workbook()
        ws = wb.active
        for row, text in enumerate([ACCEPT, FIX, UNSURE, PLAIN], start=1):
            ws.cell(row=row, column=1, value=text)
        wb.save(source)

        settings = AppSettings(
            engine=EngineSettings(
                mode="cloud", cloud_provider="custom_openai", cloud_model="excel-review-test",
                cloud_base_url="https://example.invalid/v1", concurrency=1, batch_size=10,
            ),
            source_lang="zh",
            target_lang="en",
            excel_domain_preset="无",
            excel_smart_context=False,
        )
        settings.excel_output.enable_excel_autofit = False
        runner = TaskRunner([FileItem(path=source, name=source.stem, size_kb=1.0)], settings, source_root=self.root)
        engine = _ReviewingEngine()
        inserted: list = []
        with ExitStack() as stack:
            stack.enter_context(patch("core.task_runner.TaskLogger", return_value=MagicMock(task_id="excel-review")))
            stack.enter_context(
                patch("core.task_runner.check_translation_api_config", return_value=ApiConfigCheckResult(ok=True))
            )
            stack.enter_context(patch("core.task_runner.build_engine", return_value=engine))
            stack.enter_context(patch("core.task_runner.get_system_prompt", return_value="system"))
            stack.enter_context(patch("core.task_runner.resolve_effective_model_config", return_value=object()))
            stack.enter_context(
                patch("core.task_runner.resolve_surface_rate", side_effect=pinned_surface_rate(batch_items=10, concurrency=1))
            )
            stack.enter_context(patch("core.task_runner.tm_manager.lookup_batch", return_value={}))
            stack.enter_context(
                patch(
                    "core.task_runner.tm_manager.insert_batch",
                    side_effect=lambda pairs, *args, **kwargs: inserted.extend(pairs) or len(pairs),
                )
            )
            runner._run()

        done = [m for m in list(runner._queue.queue) if isinstance(m, DoneMsg)]
        self.assertEqual(len(done), 1)
        result = done[0]
        output = Path(result.file_results[0]["output_path"])
        sheet = load_workbook(output).active
        cells = {sheet.cell(row=row, column=1).value.split("\n")[0]: sheet.cell(row=row, column=1) for row in range(1, 5)}

        self.assertIn("Three floors in total", cells[ACCEPT].value)
        self.assertIn("Mortar screed 1:2.5, 25 mm thick", cells[FIX].value)
        self.assertEqual(cells[UNSURE].value, UNSURE)
        # 复查采用的不涂底色；拿不准的照旧标需复核
        self.assertFalse(cells[ACCEPT].fill and cells[ACCEPT].fill.fgColor and cells[ACCEPT].fill.fill_type)
        self.assertTrue(cells[UNSURE].fill.fill_type)

        issues = {issue.get("type"): issue for issue in result.issues}
        accepted = issues["model_review_accepted"]
        self.assertEqual(accepted["severity"], "resolved")
        self.assertEqual(accepted["count"], 1)
        reset = issues["quality_filter_reset"]
        self.assertEqual(reset["count"], 1)
        self.assertIn("数字", reset["failed_sources"][0]["error"])
        self.assertIn("复查：理由", reset["failed_sources"][0]["error"])

        # 只有没过规则的三条送去复查（重译后那一稿通过了规则，不再复查）
        self.assertEqual(sorted(set(engine.reviewed)), sorted([ACCEPT, FIX, UNSURE]))
        # 记忆库：复查采用的不存
        stored = {pair[0] for pair in inserted}
        self.assertNotIn(ACCEPT, stored)
        self.assertIn(FIX, stored)


if __name__ == "__main__":
    unittest.main(verbosity=2)
