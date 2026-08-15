"""Excel 补译复核：改判之后新译文要真的落进单元格，可疑的旧译文一个字不删。

Word 侧的同一套回归在 tests/test_coverage_arbitration.py。两边共用
core.coverage_review 的判定，差别只在写入器和报告条目——这个文件盯的正是那半。
"""

from __future__ import annotations

import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl import Workbook, load_workbook

from core import coverage_review, task_runner
from core.coverage_arbitration import (
    RETRANSLATE_MODEL,
    RETRANSLATE_UNCERTAIN,
    ArbitrationOutcome,
    PairReview,
    apply_arbitration,
    review_coverage_pairs,
)
from core.excel_coverage import build_excel_coverage_plan, write_untranslated_excel_file
from core.mixed_language import MIXED_MARK_FOREIGN_NOISE
from core.translation_coverage import COVERAGE_COVERED, COVERAGE_SOURCE_ONLY

SOURCE = (
    "本工程受甲供短柱供货严重滞后影响，土建结构施工无法按原计划穿插进行，"
    "经与监理及业主协商，工期顺延一百六十九天。"
)
# 长度正常、句子通顺，但讲的完全是另一件事——这一格配错了对。
MISMATCHED = (
    "The present chapter sets out the safety measures for work at height and the "
    "personal protective equipment required on site."
)
CORRECT = (
    "Owing to the severe delay in the owner-supplied short columns, the structural "
    "works could not be interleaved as originally planned; after consultation with "
    "the supervisor and the owner, the schedule is extended by one hundred and "
    "sixty-nine days."
)


def _workbook(root: Path, cells: dict[str, str]) -> Path:
    path = root / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"
    for coordinate, value in cells.items():
        ws[coordinate] = value
    wb.save(path)
    wb.close()
    return path


def _batch_arbitrate(verdict: str):
    def arbitrate(pairs):
        return {pair.id: verdict for pair in pairs}

    return arbitrate


class WriterTests(unittest.TestCase):
    def test_flipped_cell_keeps_the_suspect_translation_and_gains_a_new_one(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = _workbook(Path(tmp), {"A1": f"{SOURCE}\n{MISMATCHED}"})

            plan = build_excel_coverage_plan(source, target_lang="en", source_lang="zh")
            by_location = {unit.location: unit for unit in plan.units}
            self.assertEqual(by_location["Sheet!A1"].status, COVERAGE_COVERED)
            self.assertEqual(plan.source_texts, [])

            outcome = review_coverage_pairs(
                plan.units, arbitrate=_batch_arbitrate("not_equivalent")
            )
            apply_arbitration(outcome)
            # 改判要体现在写入器读的那个属性上，而不只是留在 outcome 里。
            self.assertEqual(by_location["Sheet!A1"].status, COVERAGE_SOURCE_ONLY)
            self.assertEqual(plan.source_texts, [SOURCE])

            out_path = write_untranslated_excel_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                plan=plan,
                translations={SOURCE: CORRECT},
                target_lang="en",
                source_lang="zh",
                keep_original_sheets=False,
            )

            wb = load_workbook(out_path)
            try:
                value = wb["Sheet"]["A1"].value
            finally:
                wb.close()
            # 原文、判可疑的旧译文、补上的新译文，三段都在，顺序不变。
            self.assertEqual(value, f"{SOURCE}\n{MISMATCHED}\n{CORRECT}")

    def test_a_flipped_cell_can_be_highlighted(self) -> None:
        """底色要落在这一格上。改判过的格子键是整格文字，底色不跟着换键就涂空。"""
        with tempfile.TemporaryDirectory() as tmp:
            source = _workbook(Path(tmp), {"A1": f"{SOURCE}\n{MISMATCHED}"})
            plan = build_excel_coverage_plan(source, target_lang="en", source_lang="zh")
            apply_arbitration(
                review_coverage_pairs(
                    plan.units, arbitrate=_batch_arbitrate("not_equivalent")
                )
            )

            positions: list[dict[str, str]] = []
            write_untranslated_excel_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                plan=plan,
                translations={SOURCE: CORRECT},
                target_lang="en",
                source_lang="zh",
                keep_original_sheets=False,
                review_marks={SOURCE: MIXED_MARK_FOREIGN_NOISE},
                review_mark_colors={MIXED_MARK_FOREIGN_NOISE: "#FFC7CE"},
                mark_review_items=True,
                review_positions=positions,
            )

            self.assertEqual(
                [
                    (p["worksheet"], p["cell"], p["category"], p["action"])
                    for p in positions
                ],
                [("Sheet", "A1", MIXED_MARK_FOREIGN_NOISE, "marked_fill")],
            )

    def test_untouched_covered_cells_are_still_left_alone(self) -> None:
        """没被改判的格子一个字都不能动——复核只加，不改。"""
        with tempfile.TemporaryDirectory() as tmp:
            source = _workbook(
                Path(tmp),
                {"A1": "项目名称\nProject name", "A2": "施工内容"},
            )
            plan = build_excel_coverage_plan(source, target_lang="en", source_lang="zh")
            apply_arbitration(
                review_coverage_pairs(
                    plan.units, arbitrate=_batch_arbitrate("equivalent")
                )
            )

            out_path = write_untranslated_excel_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                plan=plan,
                translations={"施工内容": "Construction scope"},
                target_lang="en",
                source_lang="zh",
                keep_original_sheets=False,
            )

            wb = load_workbook(out_path)
            try:
                self.assertEqual(wb["Sheet"]["A1"].value, "项目名称\nProject name")
                self.assertEqual(
                    wb["Sheet"]["A2"].value, "施工内容\nConstruction scope"
                )
            finally:
                wb.close()


class RunnerGlueTests(unittest.TestCase):
    """Excel 侧只负责把结论翻成报告条目和底色，判定本身是共用的。"""

    def _runner(self, logs: list[tuple[str, str]]):
        runner = task_runner.TaskRunner.__new__(task_runner.TaskRunner)
        runner._log = lambda level, message: logs.append((level, message))
        runner._stop_event = threading.Event()
        return runner

    def _unit(self, *, index: int = 1):
        from core.translation_coverage import CoverageUnit

        return CoverageUnit(
            source_text=SOURCE,
            target_text=MISMATCHED,
            status=COVERAGE_COVERED,
            location=f"Sheet!A{index}",
            kind="cell",
            reason="同一单元格已包含源文和目标语言译文。",
            data={
                "sheet": "Sheet",
                "coordinate": f"A{index}",
                "cell_text": f"{SOURCE}\n{MISMATCHED}",
            },
        )

    def _arbitrate(self, outcome, *, issues, marks, logs):
        plan = SimpleNamespace(units=[review.unit for review in outcome.reviews])
        with mock.patch.object(
            coverage_review, "review_coverage_pairs", lambda *a, **kw: outcome
        ):
            self._runner(logs)._arbitrate_excel_coverage_pairs(
                plan,
                engine=object(),
                api_scheduler=None,
                target_lang="en",
                source_lang="zh",
                lang_pair=None,
                concurrency=4,
                file_name="进度表.xlsx",
                quality_issues=issues,
                review_marks=marks,
            )

    def test_a_flipped_cell_is_reported_with_its_sheet_and_coordinate(self) -> None:
        logs: list[tuple[str, str]] = []
        issues: list[dict] = []
        marks: dict[str, str] = {}
        unit = self._unit()
        self._arbitrate(
            ArbitrationOutcome(
                reviews=[PairReview(unit=unit, trusted=False, reason=RETRANSLATE_MODEL)],
                model_check_count=1,
                model_batch_count=1,
            ),
            issues=issues,
            marks=marks,
            logs=logs,
        )

        self.assertEqual(unit.status, COVERAGE_SOURCE_ONLY)
        self.assertEqual(marks, {SOURCE: MIXED_MARK_FOREIGN_NOISE})
        self.assertEqual(len(issues), 1)
        # 报告里得给得出"去哪一格看"——分表名加坐标，不是内部 location 串。
        self.assertEqual(issues[0]["sheet"], "Sheet")
        self.assertEqual(issues[0]["cell"], "A1")
        self.assertEqual(issues[0]["severity"], "needs_review")
        self.assertIn("人工核对", issues[0]["status"])

    def test_cells_flipped_only_because_no_verdict_came_back_are_not_highlighted(
        self,
    ) -> None:
        """接口抖一下就是整整一批落到 uncertain——全涂上底色，底色就废了。"""
        logs: list[tuple[str, str]] = []
        issues: list[dict] = []
        marks: dict[str, str] = {}
        units = [self._unit(index=i + 1) for i in range(3)]
        self._arbitrate(
            ArbitrationOutcome(
                reviews=[
                    PairReview(unit=unit, trusted=False, reason=RETRANSLATE_UNCERTAIN)
                    for unit in units
                ],
                model_check_count=3,
                model_batch_count=1,
            ),
            issues=issues,
            marks=marks,
            logs=logs,
        )

        self.assertTrue(all(unit.status == COVERAGE_SOURCE_ONLY for unit in units))
        self.assertEqual(marks, {})
        self.assertEqual([issue["severity"] for issue in issues], ["resolved"] * 3)
        self.assertIn("未取得判定结果", issues[0]["status"])

    def test_arbitration_failure_leaves_the_file_on_the_original_verdict(self) -> None:
        logs: list[tuple[str, str]] = []
        issues: list[dict] = []
        unit = self._unit()
        plan = SimpleNamespace(units=[unit])

        def boom(*args, **kwargs):
            raise RuntimeError("接口限流")

        with mock.patch.object(coverage_review, "review_coverage_pairs", boom):
            self._runner(logs)._arbitrate_excel_coverage_pairs(
                plan,
                engine=object(),
                api_scheduler=None,
                target_lang="en",
                source_lang="zh",
                lang_pair=None,
                concurrency=4,
                file_name="进度表.xlsx",
                quality_issues=issues,
                review_marks={},
            )

        self.assertEqual(unit.status, COVERAGE_COVERED)
        self.assertEqual(issues, [])
        self.assertTrue(
            any(level == "WARNING" and "接口限流" in message for level, message in logs),
            logs,
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
