"""Output names stay unique after legacy conversion and within a batch."""

from __future__ import annotations

import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from core.file_scanner import FileItem
from core.output_name_translation import avoid_bilingual_name_collision
from core.task_runner import TaskRunner
from core.word_document import WordFileItem
from core.word_task_runner import WordTaskRunner
from tests.test_excel_resume import (
    _fake_translate,
    _make_xlsx,
    _pipeline_patches,
    _run_and_get_done,
    _settings,
)
from tests import test_phase5_word_contracts as word_contracts


class OutputNameCollisionTests(unittest.TestCase):
    def test_legacy_conversion_and_existing_files_are_reserved(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            reserved: set[Path] = set()
            (output / "报告_英文_双语.xlsx").write_bytes(b"existing")

            excel_legacy = avoid_bilingual_name_collision(output, "报告.xls", "en", reserved)
            excel_current = avoid_bilingual_name_collision(output, "报告.xlsx", "en", reserved)
            self.assertEqual(excel_legacy, "报告_2.xlsx")
            self.assertEqual(excel_current, "报告_3.xlsx")
            self.assertEqual((output / "报告_英文_双语.xlsx").read_bytes(), b"existing")

            word_legacy = avoid_bilingual_name_collision(output, "合同.doc", "en", reserved)
            word_current = avoid_bilingual_name_collision(output, "合同.docx", "en", reserved)
            self.assertEqual(word_legacy, "合同.docx")
            self.assertEqual(word_current, "合同_2.docx")

    def test_same_translated_stem_is_unique_before_any_write(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            reserved: set[Path] = set()
            first = avoid_bilingual_name_collision(output, "Report.docx", "en", reserved)
            second = avoid_bilingual_name_collision(output, "Report.docx", "en", reserved)
            self.assertEqual(first, "Report.docx")
            self.assertEqual(second, "Report_2.docx")

    def test_excel_batch_keeps_both_same_translated_names(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            sources = [
                _make_xlsx(root / f"{stem}.xlsx", {"A1": value})
                for stem, value in (("报告", "甲"), ("报表", "乙"))
            ]
            settings = _settings()
            settings.excel_output.translate_output_filename = True
            settings.excel_output.translate_sheet_names = False
            with _pipeline_patches(translate_side_effect=_fake_translate([])):
                with patch(
                    "core.task_runner.translate_names",
                    side_effect=lambda _engine, names, *_args, **_kwargs: {
                        name: "Report" for name in names
                    },
                ):
                    runner = TaskRunner(
                        [FileItem(path=path, name=path.stem, size_kb=1.0) for path in sources],
                        settings,
                        source_root=root,
                    )
                    done = _run_and_get_done(runner)

            self.assertTrue(all(result["success"] for result in done.file_results))
            paths = [Path(result["output_path"]) for result in done.file_results]
            self.assertEqual([path.name for path in paths], [
                "Report_英文_双语.xlsx", "Report_2_英文_双语.xlsx"
            ])
            self.assertTrue(all(path.exists() for path in paths))
            for path, expected in zip(paths, ("甲\nT:甲", "乙\nT:乙")):
                book = load_workbook(path)
                try:
                    self.assertEqual(book.active["A1"].value, expected)
                finally:
                    book.close()

    def test_word_batch_default_names_distinguish_doc_and_docx(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            legacy = root / "合同.doc"
            legacy.write_bytes(b"legacy fixture; conversion is mocked")
            current = root / "合同.docx"
            document = Document()
            document.add_paragraph("施工范围")
            document.save(current)
            converted = root / "converted.docx"
            document.save(converted)
            prepared = {
                path: SimpleNamespace(
                    path=source,
                    method="test",
                    temp_paths=(),
                    fallback_messages=(),
                    labels_seen=0,
                    labels_prepended=0,
                    conversion_method="test" if path == legacy else "not_required",
                    conversion_fidelity="test" if path == legacy else "not_required",
                    numbering_method="not_needed",
                    numbering_fallback_messages=(),
                )
                for path, source in ((legacy, converted), (current, current))
            }
            runner = WordTaskRunner(
                [
                    WordFileItem(path=legacy, name=legacy.name, size_kb=1.0, format="doc", needs_conversion=True),
                    WordFileItem(path=current, name=current.name, size_kb=1.0),
                ],
                word_contracts.WordTaskResultContractTests._settings(),
                source_root=root,
            )
            with ExitStack() as stack:
                writer = word_contracts.WordTaskResultContractTests()._runner_patches(
                    stack, root=root, prepared_by_path=prepared
                )
                runner._run()
            names = [call.kwargs["output_name"] for call in writer.call_args_list]
            self.assertEqual(names, ["合同.docx", "合同_2.docx"])


if __name__ == "__main__":
    unittest.main()
