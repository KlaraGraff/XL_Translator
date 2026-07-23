"""Reusable fixtures for the Phase 0 contract and compatibility checks.

The fixture intentionally creates all files under a caller-provided temporary
directory.  It must not read or write the real application data directory.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from PIL import Image

from core.language_registry import build_custom_target_lang_code


_MINIMAL_PDF = b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 180] >>
endobj
trailer
<< /Root 1 0 R >>
%%EOF
"""


@dataclass(frozen=True)
class Phase0Fixtures:
    """Paths and data used by foundation tests and later phase tests."""

    root: Path
    app_data: Path
    legacy_data: Path
    excel: Path
    word: Path
    pdf: Path
    image: Path
    tm_export: Path
    custom_target_code: str


def create_phase0_fixtures(root: Path) -> Phase0Fixtures:
    """Create deterministic sample inputs beneath ``root``.

    The generated files are deliberately small.  They exercise file discovery
    and contract plumbing without requiring Office, a PDF viewer, or network
    credentials.
    """

    root = Path(root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    app_data = root / "app-data"
    legacy_data = root / "legacy-data"
    app_data.mkdir()
    legacy_data.mkdir()

    excel = root / "sample.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Source"
    worksheet["A1"] = "Hello"
    worksheet["B1"] = "Bonjour"
    workbook.save(excel)

    word = root / "sample.docx"
    document = Document()
    document.add_paragraph("Hello from the Phase 0 fixture.")
    document.save(word)

    pdf = root / "sample.pdf"
    pdf.write_bytes(_MINIMAL_PDF)

    image = root / "sample.png"
    Image.new("RGB", (8, 8), (255, 255, 255)).save(image)

    custom_target_code = build_custom_target_lang_code("Test language")
    tm_payload: dict[str, Any] = {
        "schema_version": 3,
        "custom_target_languages": [
            {
                "code": custom_target_code,
                "name": "Test language",
                "description": "Phase 0 fixture language",
            }
        ],
        "lang_pair": f"en-{custom_target_code}",
        "entries": [
            {
                "source_text": "Hello",
                "target_text": "Test greeting",
                "source_lang": "en",
                "target_lang": custom_target_code,
                "origin": "manual",
            }
        ],
    }
    tm_export = root / "tm-export.json"
    tm_export.write_text(
        json.dumps(tm_payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    # A sentinel makes accidental legacy probing observable in tests.
    (legacy_data / "settings.json").write_text(
        '{"sentinel":"legacy-data-must-stay-untouched"}\n',
        encoding="utf-8",
    )
    return Phase0Fixtures(
        root=root,
        app_data=app_data,
        legacy_data=legacy_data,
        excel=excel,
        word=word,
        pdf=pdf,
        image=image,
        tm_export=tm_export,
        custom_target_code=custom_target_code,
    )


class MockTranslationProvider:
    """Deterministic provider contract used without network/API credentials."""

    def __init__(self) -> None:
        self.calls: list[dict[str, str]] = []

    def translate(self, text: str, *, source_lang: str, target_lang: str) -> dict[str, str]:
        call = {
            "text": str(text),
            "source_lang": str(source_lang),
            "target_lang": str(target_lang),
        }
        self.calls.append(call)
        return {
            "translation": f"[{target_lang}] {text}",
            "source_lang": source_lang,
            "target_lang": target_lang,
        }
