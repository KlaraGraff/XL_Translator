"""CAD 译文对照表：查看、修改、Excel 导出导入，按修改只重新写回图纸、不再调用模型。"""
from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest

import core.process_artifacts as process_artifacts
from api.cad_review_service import memory_pair
from core.cad_review import (
    CadReviewError,
    clear_applied_edits,
    export_workbook,
    load_review,
    parse_workbook,
    rewrite_files,
    save_edits,
)
from core.cad_task_runner import CadTaskRunner
from core.cad_translation import CadPipelineOptions
from tests.test_cad_new_features import _drawing, _entities, _field, _text

TRANSLATIONS = {"墙体拆除图": "Plan de démolition", "天花图": "Plan de plafond"}


def _run_task(tmp_path, monkeypatch, *, drawing=None, translator=None, options=None) -> Path:
    monkeypatch.setattr(process_artifacts, "APP_DATA_DIR", tmp_path / "app-data")
    source = tmp_path / "src" / "plan.dxf"
    source.parent.mkdir(parents=True, exist_ok=True)
    source.write_text(
        drawing or _drawing(*_text("20", "墙体拆除图"), *_text("21", "天花图", y=10), *_text("22", "墙体拆除图", y=20)),
        encoding="utf-8",
    )
    calls: list[list[str]] = []

    def translate(texts, glossary, **_):
        calls.append(list(texts))
        return {text: TRANSLATIONS[text] for text in texts if text in TRANSLATIONS}

    runner = CadTaskRunner(
        files=[SimpleNamespace(path=source, name=source.name, format="dxf")],
        source_root=source.parent,
        output_dir=tmp_path / "results",
        converter=None,
        translator=translator or translate,
        options=options or CadPipelineOptions(target_lang="fr", translation_font="original"),
    )
    runner.start()
    while runner.needs_poll():
        runner.get_message(timeout=1)
    manifests = list((tmp_path / "app-data").rglob("manifest.json"))
    assert len(manifests) == 1, manifests
    return manifests[0]


def _texts(path: Path) -> list[str]:
    return [_field(body, 1) for body in _entities(path, "TEXT")]


def test_review_lists_each_source_once_with_its_translation(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    review = load_review(batch)
    assert review["target_lang"] == "fr"
    [file] = review["files"]
    assert file["name"] == "plan.dxf" and file["available"] and file["rewritable"], file
    rows = {row["source"]: row for row in file["rows"]}
    assert rows["墙体拆除图"]["count"] == 2
    assert rows["墙体拆除图"]["translation"] == "Plan de démolition"
    assert rows["天花图"]["status"] == "translated"
    assert not any(row["edited"] for row in file["rows"])


def test_edits_are_saved_and_rewritten_without_calling_the_model(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    [file] = load_review(batch)["files"]
    saved = save_edits(batch, file["key"], {"天花图": "Plan du faux plafond", "墙体拆除图": "Plan de démolition"})
    # 和原译文一样的不算改动。
    assert saved == {"天花图": "Plan du faux plafond"}
    [file] = load_review(batch)["files"]
    assert file["edited"] == 1
    assert {row["source"]: row["edited"] for row in file["rows"]} == {"墙体拆除图": False, "天花图": True}

    output = tmp_path / "results" / file["output_name"]
    assert "Plan de plafond" in _texts(output)
    [result] = rewrite_files(batch, [file["key"]], converter=None)
    assert result["ok"], result
    assert _texts(output) == ["Plan de démolition", "Plan du faux plafond", "Plan de démolition"]

    clear_applied_edits(batch, [file["key"]])
    [file] = load_review(batch)["files"]
    # 写回之后，新的译文就是对照表里的「原译文」，修改记录清空。
    assert file["edited"] == 0
    assert {row["source"]: row["translation"] for row in file["rows"]}["天花图"] == "Plan du faux plafond"


def test_clearing_an_edit_restores_the_first_translation(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    [file] = load_review(batch)["files"]
    save_edits(batch, file["key"], {"天花图": "X"})
    assert save_edits(batch, file["key"], {"天花图": ""}) == {}
    assert load_review(batch)["files"][0]["edited"] == 0


def test_unresolved_text_can_be_filled_in_by_hand(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch, translator=lambda texts, glossary, **_: {})
    [file] = load_review(batch)["files"]
    assert {row["status"] for row in file["rows"]} == {"unresolved"}
    save_edits(batch, file["key"], {"墙体拆除图": "Démolition"})
    [result] = rewrite_files(batch, [file["key"]], converter=None)
    assert result["ok"] and result["untranslated"] == 1
    output = tmp_path / "results" / file["output_name"]
    assert _texts(output) == ["Démolition", "天花图", "Démolition"]


def test_bilingual_rewrite_does_not_stack_translations(tmp_path, monkeypatch):
    batch = _run_task(
        tmp_path, monkeypatch,
        options=CadPipelineOptions(target_lang="fr", display_mode="source_first", translation_font="original"),
    )
    [file] = load_review(batch)["files"]
    output = tmp_path / "results" / file["output_name"]
    before = len(_entities(output, "TEXT"))
    save_edits(batch, file["key"], {"天花图": "Faux plafond"})
    [result] = rewrite_files(batch, [file["key"]], converter=None)
    assert result["ok"], result
    texts = _texts(output)
    assert len(texts) == before
    assert "Faux plafond" in texts and "Plan de plafond" not in texts


def test_changed_source_drawing_is_not_rewritten(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    [file] = load_review(batch)["files"]
    output = tmp_path / "results" / file["output_name"]
    before = output.read_bytes()
    source = tmp_path / "src" / "plan.dxf"
    source.write_text(_drawing(*_text("20", "新的图名")), encoding="utf-8")
    [result] = rewrite_files(batch, [file["key"]], converter=None)
    assert not result["ok"] and "改过" in result["message"]
    assert output.read_bytes() == before


def test_missing_source_and_old_batches_are_view_only(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    (tmp_path / "src" / "plan.dxf").unlink()
    [file] = load_review(batch)["files"]
    assert file["available"] and not file["rewritable"] and "移动或删除" in file["reason"]

    payload = json.loads(batch.read_text(encoding="utf-8"))
    payload.pop("pipeline_options")
    batch.write_text(json.dumps(payload), encoding="utf-8")
    [file] = load_review(batch)["files"]
    assert not file["rewritable"] and "旧版本" in file["reason"]
    with pytest.raises(CadReviewError):
        rewrite_files(batch, [file["key"]], converter=None)


def test_purged_records_explain_themselves(tmp_path):
    with pytest.raises(CadReviewError, match="已被清理"):
        load_review(tmp_path / "missing" / "manifest.json")


def test_dwg_needs_the_converter(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    payload = json.loads(batch.read_text(encoding="utf-8"))
    entry = payload["files"][0]
    source = Path(entry["source_path"]).with_suffix(".dwg")
    Path(entry["source_path"]).rename(source)
    entry["source_path"] = str(source)
    manifest = json.loads(Path(entry["manifest_path"]).read_text(encoding="utf-8"))
    batch.write_text(json.dumps(payload), encoding="utf-8")
    assert manifest["source_sha256"]
    [result] = rewrite_files(batch, [entry["manifest_path"]], converter=None)
    assert not result["ok"] and "转换工具" in result["message"]


def test_excel_round_trip_imports_only_changed_rows(tmp_path, monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    batch = _run_task(tmp_path, monkeypatch)
    review = load_review(batch)
    workbook = load_workbook(BytesIO(export_workbook(review)))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert header == ["图纸", "原文", "译文", "状态", "出现次数"]
    rows = {row[1].value: row for row in sheet.iter_rows(min_row=2)}
    assert rows["墙体拆除图"][4].value == 2
    rows["天花图"][2].value = "Plafond suspendu"
    # 别的图纸名、空译文、不认识的原文都不算。
    sheet.append(["other.dxf", "墙体拆除图", "Ignored", "", 1])
    sheet.append(["plan.dxf", "不存在的原文", "Ignored", "", 1])
    sheet.append(["plan.dxf", "墙体拆除图", "", "", 1])
    buffer = BytesIO()
    workbook.save(buffer)
    changes = parse_workbook(review, buffer.getvalue())
    assert changes == {review["files"][0]["key"]: {"天花图": "Plafond suspendu"}}


def test_a_file_that_is_not_the_export_is_rejected(tmp_path, monkeypatch):
    from io import BytesIO

    from openpyxl import Workbook

    batch = _run_task(tmp_path, monkeypatch)
    review = load_review(batch)
    with pytest.raises(CadReviewError, match="读不出"):
        parse_workbook(review, b"not a workbook")
    workbook = Workbook()
    workbook.active.append(["A", "B"])
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(CadReviewError, match="表头"):
        parse_workbook(review, buffer.getvalue())


def test_memory_pair_detects_the_source_language():
    assert memory_pair("墙体拆除图", "auto", "fr") == "zh-fr"
    assert memory_pair("Plan de démolition des murs", "auto", "zh") == "fr-zh"
    assert memory_pair("墙体", "zh", "fr") == "zh-fr"
    assert memory_pair("墙体", "zh", "zh") is None
    assert memory_pair("3100X2100", "auto", "fr") is None
    assert memory_pair("墙体", "auto", "") is None


# ── 接口 ──────────────────────────────────────────────────────


class _Tasks:
    def __init__(self, manifest: Path, *, terminal: bool = True, surface: str = "cad"):
        self.payload = {"surface": surface, "terminal": terminal, "result": {"manifest_path": str(manifest)}}

    def task_status(self, task_id):
        from api.task_manager import TaskNotFoundError

        if task_id != "t1":
            raise TaskNotFoundError(task_id)
        return self.payload

    def shutdown(self):
        pass


def _client(tasks):
    from fastapi.testclient import TestClient

    from api.app import create_app

    return TestClient(create_app(task_manager=tasks))


def test_api_edit_remember_export_import_and_rewrite(tmp_path, monkeypatch):
    import base64
    import time
    from io import BytesIO

    from openpyxl import load_workbook

    from core import tm_manager

    batch = _run_task(tmp_path, monkeypatch)
    monkeypatch.setattr(tm_manager, "DB_PATH", tmp_path / "tm.db")
    monkeypatch.setattr("api.cad_review_service.probe_status", lambda: {"converter": None})
    with _client(_Tasks(batch)) as client:
        review = client.get("/api/tasks/t1/cad-review").json()
        key = review["files"][0]["key"]
        assert review["terminal"] and review["rewrite"] == {"active": False}

        response = client.put("/api/tasks/t1/cad-review/edits", json={"file": key, "edits": {"天花图": "Faux plafond"}})
        assert response.json() == {"saved": 1, "remembered": 1}, response.text
        assert tm_manager.lookup_batch(["天花图"], "zh-fr")["天花图"] == "Faux plafond"

        exported = client.get("/api/tasks/t1/cad-review/export")
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        rows = {row[1].value: row for row in workbook.active.iter_rows(min_row=2)}
        assert rows["天花图"][2].value == "Faux plafond" and rows["天花图"][3].value == "已修改"
        rows["墙体拆除图"][2].value = "Démolition"
        buffer = BytesIO()
        workbook.save(buffer)
        imported = client.post(
            "/api/tasks/t1/cad-review/import", json={"data": base64.b64encode(buffer.getvalue()).decode()},
        )
        assert imported.json() == {"files": 1, "rows": 1, "remembered": 1}, imported.text
        assert client.post("/api/tasks/t1/cad-review/import", json={"data": "!!"}).status_code == 422

        started = client.post("/api/tasks/t1/cad-review/rewrite", json={"files": [key]})
        assert started.json()["accepted"] == 1, started.text
        for _ in range(100):
            rewrite = client.get("/api/tasks/t1/cad-review").json()["rewrite"]
            if not rewrite["active"]:
                break
            time.sleep(0.05)
        assert rewrite["results"][0]["ok"], rewrite
        output = tmp_path / "results" / review["files"][0]["output_name"]
        assert _texts(output) == ["Démolition", "Faux plafond", "Démolition"]
        assert client.get("/api/tasks/t1/cad-review").json()["files"][0]["edited"] == 0


def test_api_refuses_unfinished_tasks_and_other_surfaces(tmp_path, monkeypatch):
    batch = _run_task(tmp_path, monkeypatch)
    with _client(_Tasks(batch, terminal=False)) as client:
        key = client.get("/api/tasks/t1/cad-review").json()["files"][0]["key"]
        response = client.post("/api/tasks/t1/cad-review/rewrite", json={"files": [key]})
        assert response.status_code == 409 and "还没有结束" in response.json()["detail"]
        assert client.get("/api/tasks/nope/cad-review").status_code == 404
    with _client(_Tasks(batch, surface="excel")) as client:
        assert client.get("/api/tasks/t1/cad-review").status_code == 422


def test_importing_an_older_export_keeps_edits_made_since(tmp_path, monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    batch = _run_task(tmp_path, monkeypatch)
    review = load_review(batch)
    workbook = load_workbook(BytesIO(export_workbook(review)))
    key = review["files"][0]["key"]
    # 导出之后，在应用里又改了一行；拿到表的人改的是另一行。
    save_edits(batch, key, {"天花图": "Faux plafond"})
    for row in workbook.active.iter_rows(min_row=2):
        if row[1].value == "墙体拆除图":
            row[2].value = "Démolition"
    buffer = BytesIO()
    workbook.save(buffer)
    assert parse_workbook(load_review(batch), buffer.getvalue()) == {key: {"墙体拆除图": "Démolition"}}
