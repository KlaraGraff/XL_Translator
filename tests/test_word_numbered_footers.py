"""页眉页脚里只差数字的几行（各页页码）按同一种写法输出。

本机实测（法译中，PDF 转 Word）：三页页脚「Page 1」「Page 2」「Page 3」各自翻译，
成了「第 1 页」「第2页」「第3页」。现在以数字最小那行的译文为样板，其余页换上各自的数字。
"""
from __future__ import annotations

import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_SECTION

from core.task_runner import LogMsg
from core.word_document import WordFileItem
from core.word_task_runner import WordTaskRunner, _fill_numbered_line, _numbered_header_footer_fills
from engines.base_engine import TranslationEngine
from settings import AppSettings
from tests import test_word_resume as word_helpers
from tests.app_data_isolation import IsolatedAppDataTestCase

FOOTER = [{"kind": "footer"}]


class FillNumberedLineTests(unittest.TestCase):
    def test_other_page_takes_the_template_with_its_own_number(self) -> None:
        self.assertEqual(_fill_numbered_line("Page 1", "第 1 页", "Page 3"), "第 3 页")

    def test_several_numbers_map_by_value(self) -> None:
        self.assertEqual(_fill_numbered_line("Page 1 sur 3", "第 1 页，共 3 页", "Page 3 sur 3"), "第 3 页，共 3 页")
        self.assertEqual(_fill_numbered_line("Page 1 of 3", "共 3 页，第 1 页", "Page 2 of 3"), "共 3 页，第 2 页")

    def test_template_with_numbers_written_out_is_not_used(self) -> None:
        self.assertIsNone(_fill_numbered_line("Page 1", "第一页", "Page 2"))

    def test_template_with_an_extra_number_is_not_used(self) -> None:
        self.assertIsNone(_fill_numbered_line("Page 1", "第 1 页 2026", "Page 2"))


class NumberedHeaderFooterFillTests(unittest.TestCase):
    def test_page_numbers_follow_the_first_page(self) -> None:
        translations = {"Page 1": "第 1 页", "Page 2": "第2页", "Page 3": "第3页", "Titre": "标题"}
        locations = {"Page 1": FOOTER, "Page 2": FOOTER, "Page 3": FOOTER, "Titre": FOOTER}
        self.assertEqual(
            _numbered_header_footer_fills(translations, locations), {"Page 2": "第 2 页", "Page 3": "第 3 页"}
        )

    def test_whole_footer_line_with_a_page_number(self) -> None:
        translations = {
            "Méthodologie d’intervention\tPage 2": "干预方法\t第2页",
            "Méthodologie d’intervention\tPage 1": "干预方法\t第 1 页",
        }
        locations = dict.fromkeys(translations, FOOTER)
        self.assertEqual(
            _numbered_header_footer_fills(translations, locations),
            {"Méthodologie d’intervention\tPage 2": "干预方法\t第 2 页"},
        )

    def test_body_paragraphs_are_left_alone(self) -> None:
        translations = {"Article 1": "第一条", "Article 2": "第 2 条"}
        locations = dict.fromkeys(translations, [{"kind": "paragraph"}])
        self.assertEqual(_numbered_header_footer_fills(translations, locations), {})

    def test_line_whose_template_does_not_fit_keeps_its_translation(self) -> None:
        translations = {"Page 1": "第一页", "Page 2": "第2页"}
        locations = dict.fromkeys(translations, FOOTER)
        self.assertEqual(_numbered_header_footer_fills(translations, locations), {})


class InconsistentPageEngine(TranslationEngine):
    @property
    def engine_name(self) -> str:
        return "fake/page-numbers"

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        results = {}
        for text in texts:
            if text.startswith("Page "):
                number = text.split()[-1]
                results[text] = f"第 {number} 页" if number == "1" else f"第{number}页"
            else:
                results[text] = "施工内容"
        return results


class NumberedFooterRunnerTests(IsolatedAppDataTestCase):
    def test_each_section_footer_uses_the_first_page_format(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "pages.docx"
            doc = Document()
            for index in range(3):
                if index:
                    doc.add_section(WD_SECTION.NEW_PAGE)
                section = doc.sections[-1]
                section.footer.is_linked_to_previous = False
                section.footer.paragraphs[0].text = f"Page {index + 1}"
                doc.add_paragraph("Contenu des travaux")
            doc.save(path)
            settings = AppSettings(source_lang="fr", target_lang="zh")
            settings.word_output.output_translation_only = True
            runner = WordTaskRunner(
                [WordFileItem(path=path, name=path.name, size_kb=1.0)],
                settings,
                source_root=root,
                source_lang="fr",
                translate_headers_footers=True,
            )
            with ExitStack() as stack:
                word_helpers.WordResumeTests._runner_patches(self, stack, root=root)
                stack.enter_context(patch("core.word_task_runner.build_engine", return_value=InconsistentPageEngine()))
                stack.enter_context(patch("core.word_task_runner.tm_manager.lookup_batch", return_value={}))
                runner._run()
            output = next((root / "out").rglob("*.docx"))
            footers = [section.footer.paragraphs[0].text for section in Document(output).sections]
        self.assertEqual(footers, ["第 1 页", "第 2 页", "第 3 页"])
        logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
        self.assertTrue([line for line in logs if "2 行页码已按第一页的写法统一" in line], logs)


if __name__ == "__main__":
    unittest.main(verbosity=2)
