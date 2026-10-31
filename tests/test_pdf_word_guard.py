"""转 Word 的正文段：认出 pdf2docx 会转坏的页、转坏了的页，以及页数暴增的段。"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml

from core.pdf_word_guard import (
    MAX_ANCHORS_PER_PAGE,
    MAX_SAME_IMAGE_PER_PAGE,
    broken_docx_pages,
    docx_page_stats,
    segment_grew_too_much,
)
from tests.test_pdf_image_translation import _png_bytes

A4 = (595.276, 841.89)
_ANCHOR = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
           '<w:drawing><wp:anchor/></w:drawing></w:r>')


class DocxPageStatsTests(unittest.TestCase):
    def _docx(self, pages) -> Path:
        """pages: 每页 (同一张图放几次, 浮动块几个, 页内再分几栏节)。"""
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        root = Path(tmp.name)
        image = root / "chart.png"
        image.write_bytes(_png_bytes(20, 20))
        doc = Document()
        for index, (images, anchors, columns) in enumerate(pages):
            if index:
                doc.add_section(WD_SECTION.NEW_PAGE)
            for column in range(columns + 1):
                if column:
                    doc.add_section(WD_SECTION.CONTINUOUS)
                paragraph = doc.add_paragraph(f"page {index + 1}")
                for _ in range(images if column == 0 else 0):
                    paragraph.add_run().add_picture(str(image))
                for _ in range(anchors if column == 0 else 0):
                    paragraph._p.append(parse_xml(_ANCHOR))
        path = root / "converted.docx"
        doc.save(path)
        return path

    def test_counts_per_page_across_continuous_sections(self):
        path = self._docx([(1, 0, 0), (3, 2, 2), (0, 0, 0)])
        self.assertEqual(docx_page_stats(path), [(1, 0), (3, 2), (0, 0)])

    def test_finds_the_pages_pdf2docx_broke(self):
        path = self._docx([(2, 0, 0), (MAX_SAME_IMAGE_PER_PAGE + 1, 0, 1), (0, MAX_ANCHORS_PER_PAGE + 1, 0),
                           (4, 21, 0)])
        self.assertEqual(broken_docx_pages(path, 4), [2, 3])

    def test_a_logo_on_every_page_is_not_a_broken_page(self):
        path = self._docx([(1, 0, 0)] * (MAX_SAME_IMAGE_PER_PAGE + 5))
        self.assertEqual(broken_docx_pages(path, MAX_SAME_IMAGE_PER_PAGE + 5), [])

    def test_sections_that_do_not_match_the_page_count_are_not_judged(self):
        path = self._docx([(MAX_SAME_IMAGE_PER_PAGE + 1, 0, 0), (1, 0, 0)])
        self.assertEqual(broken_docx_pages(path, 3), [])


class SegmentGrowthTests(unittest.TestCase):
    def test_reflow_is_fine_but_a_blow_up_is_not(self):
        self.assertFalse(segment_grew_too_much(260, 292))    # 中文换行，表格挤到下一页
        self.assertFalse(segment_grew_too_much(1, 3))
        self.assertTrue(segment_grew_too_much(322, 2581))
        self.assertTrue(segment_grew_too_much(1, 48))


class ImageGridPageTests(unittest.TestCase):
    fitz = None

    @classmethod
    def setUpClass(cls):
        cls.fitz = pytest.importorskip("fitz")

    def _kind(self, *, lines: int, image_share: float = 0.6):
        from core.pdf_routes import classify_page

        fitz = self.fitz
        doc = fitz.open()
        page = doc.new_page(width=A4[0], height=A4[1])
        page.insert_text((60, 60), "Sondage SC56  Profondeur finale 15 m", fontsize=11)
        height = A4[1] * image_share
        box = fitz.Rect(30, 100, A4[0] - 30, 100 + height * A4[0] / (A4[0] - 60))
        page.insert_image(box, stream=_png_bytes(200, 200))
        for index in range(lines):
            y = box.y0 + (index + 1) * box.height / (lines + 1)
            page.draw_line((box.x0, y), (box.x1, y), width=0.5)
        try:
            return classify_page(page, 1, oversized=False)
        finally:
            doc.close()

    def test_full_page_image_with_table_lines_over_it(self):
        kind = self._kind(lines=40)
        self.assertEqual(kind.kind, "text")
        self.assertTrue(kind.image_grid)

    def test_a_big_photo_without_lines_still_goes_to_word(self):
        self.assertFalse(self._kind(lines=0).image_grid)

    def test_a_small_image_with_lines_still_goes_to_word(self):
        self.assertFalse(self._kind(lines=40, image_share=0.2).image_grid)


if __name__ == "__main__":
    unittest.main()
