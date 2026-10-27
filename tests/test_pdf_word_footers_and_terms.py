"""PDF 转 Word 实测问题（法译中，3 页施工方法说明）里的另外三项。

1. 每页底部的页脚是正文里一行表格，前面垫一个固定行高的空段顶到页底；译文稍高一点，
   页脚就挤到下一页，3 页变 6 页。现在挪进这一页的页脚区。
2. 译文里冒出原文和目标语言都不用的文字（俄文词「предусмотр」）：首次译文要重译。
3. 大标题「MÉTHODOLOGIE D’INTERVENTION」译成「实施方法」，页脚同一词组却是「干预方法」。
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.pdf_word_route import move_running_footers
from core.word_task_runner import (
    _evaluate_word_translation,
    _short_text_term_conflicts,
    _unexpected_script_fragments,
    _uses_short_text_terms,
)

SPACER = '<w:p {ns}><w:pPr><w:spacing w:line="2556" w:lineRule="exact"/></w:pPr></w:p>'
TINY = '<w:p {ns}><w:pPr><w:spacing w:line="14" w:lineRule="exact"/>{sect}</w:pPr></w:p>'
FOOTER = (
    '<w:tbl {ns}><w:tblPr/><w:tblGrid><w:gridCol w:w="6000"/><w:gridCol w:w="2000"/></w:tblGrid>'
    '<w:tr><w:trPr><w:trHeight w:hRule="atLeast" w:val="272"/></w:trPr>'
    '<w:tc><w:p><w:r><w:t>Méthodologie d’intervention – Expertise structurelle</w:t></w:r></w:p></w:tc>'
    '<w:tc><w:p><w:r><w:t>Page {page}</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
)
SECT = (
    '<w:sectPr {ns}><w:pgSz w:w="11906" w:h="16838"/>'
    '<w:pgMar w:top="524" w:right="934" w:bottom="236" w:left="954" w:header="720" w:footer="720" w:gutter="0"/>'
    '</w:sectPr>'
)


def _page_docx(path: Path, pages: int, *, footer_text=None) -> None:
    doc = Document()
    body = doc.element.body
    for child in list(body):
        body.remove(child)
    ns = nsdecls("w")
    for page in range(1, pages + 1):
        body.append(parse_xml(f'<w:p {ns}><w:r><w:t>Jour {page} : contrôle visuel.</w:t></w:r></w:p>'))
        body.append(parse_xml(SPACER.format(ns=ns)))
        footer = FOOTER if footer_text is None else FOOTER.replace("Méthodologie d’intervention – Expertise structurelle", footer_text(page))
        body.append(parse_xml(footer.format(ns=ns, page=page)))
        last = page == pages
        body.append(parse_xml(TINY.format(ns=ns, sect="" if last else SECT.format(ns=""))))
    body.append(parse_xml(SECT.format(ns=ns)))
    doc.save(str(path))


class RunningFooterTests(unittest.TestCase):
    def setUp(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.path = Path(tmp.name) / "converted.docx"

    def test_bottom_table_on_every_page_moves_into_the_page_footer(self) -> None:
        _page_docx(self.path, 3)
        self.assertEqual(move_running_footers(self.path), 3)
        doc = Document(str(self.path))
        self.assertEqual(len(doc.tables), 0, "页脚表格不再留在正文里")
        body_texts = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(body_texts, [f"Jour {n} : contrôle visuel." for n in (1, 2, 3)])
        spacers = [
            p for p in doc.paragraphs
            if p._p.pPr is not None and p._p.pPr.find(qn("w:spacing")) is not None
            and p._p.pPr.find(qn("w:spacing")).get(qn("w:line")) == "2556"
        ]
        self.assertEqual(spacers, [], "垫高用的空段去掉")
        self.assertEqual(len(doc.sections), 3)
        for number, section in enumerate(doc.sections, start=1):
            footer = section.footer
            self.assertFalse(footer.is_linked_to_previous)
            cells = [cell.text for cell in footer.tables[0].rows[0].cells]
            self.assertEqual(cells, ["Méthodologie d’intervention – Expertise structurelle", f"Page {number}"])
            margins = section._sectPr.find(qn("w:pgMar"))
            self.assertEqual(margins.get(qn("w:footer")), "236")

    def test_single_page_is_left_alone(self) -> None:
        _page_docx(self.path, 1)
        self.assertEqual(move_running_footers(self.path), 0)
        self.assertEqual(len(Document(str(self.path)).tables), 1)

    def test_bottom_tables_that_differ_from_page_to_page_stay(self) -> None:
        _page_docx(self.path, 2, footer_text=lambda page: ["Synthèse des résultats", "Annexe photographique"][page - 1])
        self.assertEqual(move_running_footers(self.path), 0)
        self.assertEqual(len(Document(str(self.path)).tables), 2)


class UnexpectedScriptTests(unittest.TestCase):
    TRANSLATION = "• 试验计划安排，尤其是在计划中 предусмотр时确定混凝土抗压强度。"
    SOURCE = "• Planification des essais, notamment la détermination de la résistance du béton si elle est prévue."

    def test_cyrillic_word_in_a_french_to_chinese_translation(self) -> None:
        self.assertEqual(_unexpected_script_fragments(self.SOURCE, self.TRANSLATION, "zh"), ["предусмотр"])

    def test_first_translation_is_retried_and_a_retry_is_kept_for_review(self) -> None:
        first = _evaluate_word_translation(
            self.SOURCE, self.TRANSLATION, source_lang="fr", target_lang="zh", allow_recovery=False
        )
        self.assertFalse(first.accepted)
        retry = _evaluate_word_translation(
            self.SOURCE, self.TRANSLATION, source_lang="fr", target_lang="zh", allow_recovery=True
        )
        self.assertTrue(retry.accepted)
        self.assertTrue(retry.validation.needs_review)
        self.assertIn("предусмотр", retry.validation.review_fragments)

    def test_symbols_scripts_in_the_source_and_the_target_script_are_fine(self) -> None:
        self.assertEqual(_unexpected_script_fragments("Résistance 30 MPa, épaisseur 20 μm", "厚度 20 μm", "zh"), [])
        self.assertEqual(_unexpected_script_fragments("Le béton ΣΑΒ", "混凝土 ΣΑΒ", "zh"), [])
        self.assertEqual(_unexpected_script_fragments("混凝土强度", "Прочность бетона", "ru"), [])
        self.assertEqual(_unexpected_script_fragments("混凝土", "Béton ΣΑΒ", "custom-lang"), [])


class ShortTextTermTests(unittest.TestCase):
    TITLE = "MÉTHODOLOGIE D’INTERVENTION"
    FOOTER = "Méthodologie d’intervention – Expertise structurelle"

    def test_footer_that_does_not_use_the_title_translation_is_picked(self) -> None:
        translations = {
            self.TITLE: "实施方法",
            self.FOOTER: "干预方法——结构检测与鉴定",
            "Planning prévisionnel des interventions": "干预预定计划",
        }
        self.assertEqual(_short_text_term_conflicts(translations), {self.FOOTER: [(self.TITLE, "实施方法")]})

    def test_consistent_translations_are_left_alone(self) -> None:
        translations = {self.TITLE: "实施方法", self.FOOTER: "实施方法——结构检测与鉴定"}
        self.assertEqual(_short_text_term_conflicts(translations), {})

    def test_retry_must_use_the_title_translation(self) -> None:
        terms = [(self.TITLE, "实施方法")]
        self.assertTrue(_uses_short_text_terms("实施方法——结构检测与鉴定", terms, previous="干预方法——结构检测与鉴定"))
        self.assertFalse(_uses_short_text_terms("干预方法——结构鉴定", terms, previous="干预方法——结构检测与鉴定"))

    def test_phrase_used_in_many_paragraphs_is_left_to_document_terms(self) -> None:
        translations = {"béton armé": "钢筋混凝土"}
        for index in range(6):
            translations[f"Le béton armé du niveau {chr(65 + index)} est fissuré"] = "钢砼开裂"
        self.assertEqual(_short_text_term_conflicts(translations), {})


if __name__ == "__main__":
    unittest.main(verbosity=2)
