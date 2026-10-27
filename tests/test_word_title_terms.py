"""大标题、小标题、页眉页脚里含同一个词的几句，翻完后请模型复查一次译法。

本机实测（法译中，PDF 转 Word）：大标题和页脚里的「Méthodologie d’intervention」译成「施工方法」，
第 2 节小标题「Planning prévisionnel des interventions」那一批却译成「干预预安排」。四处 intervention
分在 9 批里翻，各批互相看不到。正文里的「avant intervention」译「进场作业前」是对的，不在复查之列。
"""
from __future__ import annotations

import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from docx import Document

from core.pdf_word_route import translate_pdf_via_word
from core.task_runner import DoneMsg, LogMsg
from core.word_document import WordFileItem
from core.word_task_runner import WordTaskRunner, _is_title_like, _title_review_change_ok, _title_term_groups
from engines.base_engine import TranslationEngine
from settings import AppSettings
from tests import test_word_resume as word_helpers
from tests.app_data_isolation import IsolatedAppDataTestCase

TITLE = "MÉTHODOLOGIE D’INTERVENTION"
HEADING = "2. Planning prévisionnel des interventions"
BODY = "• Validation du programme : les points de contrôle doivent être validés avant intervention."
OTHER = "3. Essais en laboratoire"
TRANSLATIONS = {
    TITLE: "施工方法",
    HEADING: "2. 干预预安排",
    BODY: "• 计划确认：检查点必须在进场作业前确认。",
    OTHER: "3. 实验室试验",
}
PARAGRAPH = [{"kind": "paragraph"}]


class TitleLikeTests(unittest.TestCase):
    def test_headings_and_footers_count_as_titles(self) -> None:
        self.assertTrue(_is_title_like(TITLE, {"paragraph"}))
        self.assertTrue(_is_title_like(HEADING, {"paragraph"}))
        self.assertTrue(_is_title_like("Méthodologie d’intervention – Expertise structurelle", {"footer"}))

    def test_sentences_bullets_and_table_cells_do_not(self) -> None:
        self.assertFalse(_is_title_like(BODY, {"paragraph"}))
        self.assertFalse(_is_title_like("Les essais sont réalisés sur site.", {"paragraph"}))
        self.assertFalse(_is_title_like("Planning des interventions", {"table_cell"}))


class TitleTermGroupTests(unittest.TestCase):
    def test_title_and_heading_share_intervention_but_the_body_bullet_is_left_out(self) -> None:
        groups = _title_term_groups(TRANSLATIONS, dict.fromkeys(TRANSLATIONS, PARAGRAPH))
        self.assertEqual(groups, {"intervention": [TITLE, HEADING]})

    def test_chinese_titles_share_the_longest_common_piece(self) -> None:
        translations = {"基坑支护施工方案": "Plan de soutènement", "二、基坑支护监测": "2. Suivi du soutènement"}
        groups = _title_term_groups(translations, dict.fromkeys(translations, PARAGRAPH))
        self.assertEqual(groups, {"基坑支护": list(translations)})

    def test_a_single_title_has_nothing_to_compare(self) -> None:
        translations = {TITLE: "施工方法", OTHER: "3. 实验室试验"}
        self.assertEqual(_title_term_groups(translations, dict.fromkeys(translations, PARAGRAPH)), {})


class ChangeCheckTests(unittest.TestCase):
    def test_numbering_must_stay(self) -> None:
        kwargs = {"source_lang": "fr", "target_lang": "zh"}
        self.assertTrue(_title_review_change_ok(HEADING, "2. 干预预安排", "2. 施工计划安排", **kwargs))
        self.assertFalse(_title_review_change_ok(HEADING, "2. 干预预安排", "施工计划安排", **kwargs))
        self.assertFalse(_title_review_change_ok(HEADING, "2. 干预预安排", "2. 干预预安排", **kwargs))


class ReviewingEngine(TranslationEngine):
    def __init__(self, reply: str) -> None:
        self.reply = reply
        self.reviewed: list[dict] = []

    @property
    def engine_name(self) -> str:
        return "fake/title-review"

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        return {text: TRANSLATIONS[text] for text in texts}

    def chat(self, system: str, user: str) -> str:
        self.reviewed.append(json.loads(user))
        return self.reply


class TitleReviewRunnerTests(IsolatedAppDataTestCase):
    def _run(self, engine):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "methodologie.docx"
            doc = Document()
            for text in TRANSLATIONS:
                doc.add_paragraph(text)
            doc.save(path)
            settings = AppSettings(source_lang="fr", target_lang="zh")
            settings.word_output.output_translation_only = True
            runner = WordTaskRunner(
                [WordFileItem(path=path, name=path.name, size_kb=1.0)], settings, source_root=root, source_lang="fr"
            )
            with ExitStack() as stack:
                word_helpers.WordResumeTests._runner_patches(self, stack, root=root)
                stack.enter_context(patch("core.word_task_runner.build_engine", return_value=engine))
                stack.enter_context(patch("core.word_task_runner.tm_manager.lookup_batch", return_value={}))
                runner._run()
            output = next((root / "out").rglob("*.docx"))
            texts = [paragraph.text for paragraph in Document(output).paragraphs]
        logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
        return texts, logs

    def test_heading_takes_the_reviewed_translation(self) -> None:
        engine = ReviewingEngine(json.dumps({"changes": [{"id": 2, "translation": "2. 施工计划安排"}]}))
        texts, logs = self._run(engine)
        self.assertEqual(len(engine.reviewed), 1)
        reviewed = engine.reviewed[0]
        self.assertEqual([item["source"] for item in reviewed["items"]], [TITLE, HEADING])
        self.assertEqual(reviewed["shared_words"], [{"word": "intervention", "ids": [1, 2]}])
        self.assertIn("2. 施工计划安排", texts)
        self.assertIn("施工方法", texts)
        self.assertIn(TRANSLATIONS[BODY], texts, "正文不在复查之列")
        self.assertTrue([line for line in logs if "1 句已改为与其他标题一致的译法" in line], logs)

    def test_unusable_reply_keeps_the_translation(self) -> None:
        texts, logs = self._run(ReviewingEngine("not json"))
        self.assertIn("2. 干预预安排", texts)
        self.assertTrue([line for line in logs if "保持原译文" in line], logs)

    def test_change_that_drops_the_numbering_is_ignored(self) -> None:
        texts, _logs = self._run(ReviewingEngine(json.dumps({"changes": [{"id": 2, "translation": "施工计划安排"}]})))
        self.assertIn("2. 干预预安排", texts)


class FakeWordRunner:
    """Stands in for the Word run inside a PDF task: three log lines, then done."""

    def __init__(self, items, settings, **_kwargs) -> None:
        self.output = Path(settings.word_output.custom_output_dir) / "out.docx"
        self.messages = [
            LogMsg(level="INFO", message="读取完成：全部文档合计 4 处待翻译文本"),
            LogMsg(level="INFO", message="标题、页眉页脚：2 句含有相同的词，正在请模型复查这些句子的译法是否一致。", relay=True),
            LogMsg(level="OK", message="1 句已改为与其他标题一致的译法。", relay=True),
        ]

    def start(self) -> None:
        Document().save(self.output)
        self.messages.append(
            DoneMsg(
                output_dir=str(self.output.parent),
                file_results=[{"success": True, "output": str(self.output)}],
                elapsed_sec=0.0,
                tm_hit_count=0,
                api_call_count=0,
            )
        )

    def needs_poll(self) -> bool:
        return bool(self.messages)

    def get_message(self, timeout: float = 0.0):
        return self.messages.pop(0) if self.messages else None

    def stop(self) -> None:
        pass


class PdfRouteRelayTests(unittest.TestCase):
    def test_pdf_task_log_shows_the_title_review_lines(self) -> None:
        logged: list[str] = []
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)

            def fake_convert(pdf_path, docx, **_kwargs):
                doc = Document()
                doc.add_paragraph(TITLE)
                doc.save(docx)

            with patch("core.pdf_word_route.convert_pdf_to_docx", fake_convert), \
                    patch("core.word_task_runner.WordTaskRunner", FakeWordRunner):
                result = translate_pdf_via_word(
                    root / "plan.pdf",
                    root / "work",
                    root / "final.docx",
                    AppSettings(source_lang="fr", target_lang="zh"),
                    source_lang="fr",
                    should_stop=lambda: False,
                    log=lambda level, message: logged.append(message),
                )
        self.assertTrue(result.ok, result.error)
        self.assertIn("1 句已改为与其他标题一致的译法。", logged)
        self.assertTrue(any("正在请模型复查" in line for line in logged), logged)
        self.assertFalse(any("读取完成" in line for line in logged), "内层 Word 的读取说明照旧不转")


if __name__ == "__main__":
    unittest.main(verbosity=2)
