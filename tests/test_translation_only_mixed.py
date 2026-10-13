"""仅输出译文时，「中文＋外文」的内容要只剩目标语言（反馈 FB-20261013-UVW95F 复测）。

V10.1.0 签名包在 Mac 上用真实模型复测，中译法、仅译文：
「钢筋 Rebar HRB400 Ø12」→「Armatures Rebar HRB400 Ø12」、「混凝土 C30 Concrete」→
「Béton C30 Concrete」——中文翻了，夹着的英文原样留下；「合计 Total」→「Total Total」、
「验收 Réception」→「Réception Réception」；而且一格都没标需复核。

这里的假模型照着真实模型的毛病写：第一遍总把英文原样留着、把对照写两遍、把中法对照判成
「已有双语」；只有收到「上一轮结果未通过质量校验」的重试提示才改对（或者怎么都不改）。
"""
from __future__ import annotations

import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

from core import mixed_language
from core.file_scanner import FileItem
from core.mixed_language import (
    MIXED_ACTION_EXISTING_BILINGUAL,
    MIXED_ACTION_TRANSLATE,
    MIXED_MARK_UNRESOLVED,
    _build_mixed_translation_prompt,
    collapse_echoed_repeat,
    leftover_foreign_fragments,
    translate_mixed_language_texts,
)
from core.task_runner import LogMsg, TaskRunner
from core.word_document import WordFileItem
from core.word_task_runner import WordTaskRunner
from engines.base_engine import TranslationEngine
from settings import AppSettings
from tests import test_excel_resume as excel_helpers
from tests import test_word_resume as word_helpers

REBAR = "钢筋 Rebar HRB400 Ø12"
CONCRETE = "混凝土 C30 Concrete"
FIRE_DOOR = "Supply and install 防火门 fire door"
TOTAL = "合计 Total"
RECEPTION = "验收 Réception"

# 真实模型第一遍的样子 → 收到重试提示后改对的样子。
FIRST = {
    REBAR: (MIXED_ACTION_TRANSLATE, "Armatures Rebar HRB400 Ø12"),
    CONCRETE: (MIXED_ACTION_TRANSLATE, "Béton C30 Concrete"),
    FIRE_DOOR: (MIXED_ACTION_TRANSLATE, "Fourniture et pose de la porte coupe-feu fire door"),
    TOTAL: (MIXED_ACTION_TRANSLATE, "Total Total"),
    RECEPTION: (MIXED_ACTION_EXISTING_BILINGUAL, ""),
}
FIXED = {
    REBAR: (MIXED_ACTION_TRANSLATE, "Armatures HRB400 Ø12"),
    CONCRETE: (MIXED_ACTION_TRANSLATE, "Béton C30"),
    FIRE_DOOR: (MIXED_ACTION_TRANSLATE, "Fourniture et pose de la porte coupe-feu"),
    TOTAL: (MIXED_ACTION_TRANSLATE, "Total Total"),
    RECEPTION: (MIXED_ACTION_TRANSLATE, "Réception Réception"),
}


class KeepsForeignEngine(TranslationEngine):
    """Mixed-path model that keeps foreign words until told to retry."""

    def __init__(self, *, fixes_on_retry: bool) -> None:
        self.fixes_on_retry = fixes_on_retry
        self.systems: list[str] = []

    @property
    def engine_name(self) -> str:
        return "fake/keeps-foreign"

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        return {text: text for text in texts}

    def chat(self, system: str, user: str) -> str:
        self.systems.append(system)
        payload = json.loads(user)
        if isinstance(payload, dict):
            # 语义复查：真实模型常常放行，这里一律拿不准，逼程序自己把关。
            return json.dumps({"verdict": "uncertain"})
        retry = "上一轮结果未通过质量校验" in system
        table = FIXED if retry and self.fixes_on_retry else FIRST
        results = []
        for item in payload:
            action, translation = table.get(item["text"], (MIXED_ACTION_TRANSLATE, "Traduction"))
            results.append({"id": item["id"], "action": action, "translation": translation, "note": ""})
        return json.dumps(results, ensure_ascii=False)


def _translate(texts, *, translation_only, fixes_on_retry):
    engine = KeepsForeignEngine(fixes_on_retry=fixes_on_retry)
    results = translate_mixed_language_texts(
        texts,
        engine=engine,
        target_lang="fr",
        system_prompt="base",
        source_lang="zh",
        retry_attempts=2,
        translation_only=translation_only,
    )
    return results, engine


class TranslationOnlyMixedPathTests(unittest.TestCase):
    def test_prompt_forbids_kept_bilingual_and_leftover_foreign_words(self) -> None:
        normal = _build_mixed_translation_prompt(target_lang="fr", source_lang="zh", base_prompt="", retry_hint=False)
        only = _build_mixed_translation_prompt(
            target_lang="fr", source_lang="zh", base_prompt="", retry_hint=False, translation_only=True
        )
        self.assertIn("不得原样留在译文里", normal)
        self.assertIn("不得重复", normal)
        self.assertNotIn("不得使用 existing_bilingual", normal)
        self.assertIn("不得使用 existing_bilingual", only)

    def test_leftovers_are_retried_and_the_fixed_answer_is_used(self) -> None:
        results, engine = _translate([REBAR, CONCRETE, FIRE_DOOR, TOTAL, RECEPTION], translation_only=True, fixes_on_retry=True)
        self.assertEqual(results[REBAR].translation, "Armatures HRB400 Ø12")
        self.assertEqual(results[CONCRETE].translation, "Béton C30")
        self.assertEqual(results[FIRE_DOOR].translation, "Fourniture et pose de la porte coupe-feu")
        # 写两遍的对照直接合并，不必为此再问一次模型。
        self.assertEqual(results[TOTAL].translation, "Total")
        # 仅译文模式下「已有双语」不算数，重试拿到译文。
        self.assertEqual(results[RECEPTION].action, MIXED_ACTION_TRANSLATE)
        self.assertEqual(results[RECEPTION].translation, "Réception")
        for result in results.values():
            self.assertIsNone(result.mark_kind, result)
        self.assertIn("不得使用 existing_bilingual", engine.systems[0])

    def test_a_model_that_never_drops_the_english_gets_marked_for_review(self) -> None:
        results, _engine = _translate([REBAR, CONCRETE, TOTAL], translation_only=True, fixes_on_retry=False)
        # 译文照写（总比整格中文强），但要标需复核，不能悄悄交出去。
        self.assertEqual(results[CONCRETE].translation, "Béton C30 Concrete")
        self.assertEqual(results[CONCRETE].mark_kind, MIXED_MARK_UNRESOLVED)
        self.assertEqual(results[REBAR].mark_kind, MIXED_MARK_UNRESOLVED)
        self.assertEqual(results[TOTAL].translation, "Total")
        self.assertIsNone(results[TOTAL].mark_kind)

    def test_bilingual_output_still_keeps_existing_bilingual(self) -> None:
        results, _engine = _translate([RECEPTION, CONCRETE], translation_only=False, fixes_on_retry=True)
        self.assertEqual(results[RECEPTION].action, MIXED_ACTION_EXISTING_BILINGUAL)
        # 原文＋译文模式不做残留检查（原文就在旁边），只合并重复。
        self.assertEqual(results[CONCRETE].translation, "Béton C30 Concrete")
        self.assertIsNone(results[CONCRETE].mark_kind)

    def test_collapse_only_touches_phrases_echoed_from_the_source(self) -> None:
        self.assertEqual(collapse_echoed_repeat("Réception Réception", RECEPTION), "Réception")
        self.assertEqual(collapse_echoed_repeat("Total Total", TOTAL), "Total")
        self.assertEqual(collapse_echoed_repeat("très très bien", "非常好"), "très très bien")
        self.assertEqual(collapse_echoed_repeat("Niveau 1 1", "1 1 层"), "Niveau 1 1")

    def test_leftover_check_ignores_codes_target_words_and_english_targets(self) -> None:
        self.assertEqual(leftover_foreign_fragments(REBAR, "Armatures HRB400 Ø12", target_lang="fr"), [])
        self.assertEqual(leftover_foreign_fragments(RECEPTION, "Réception", target_lang="fr"), [])
        self.assertEqual(leftover_foreign_fragments(CONCRETE, "C30 Concrete", target_lang="en"), [])
        self.assertEqual(leftover_foreign_fragments(CONCRETE, "Béton C30 Concrete", target_lang="fr"), ["Concrete"])


def _with_keeps_foreign_engine(fixes_on_retry: bool, seen: dict):
    real = mixed_language.translate_mixed_language_texts

    def wrapper(texts, **kwargs):
        seen["translation_only"] = kwargs.get("translation_only")
        kwargs["engine"] = KeepsForeignEngine(fixes_on_retry=fixes_on_retry)
        return real(texts, **kwargs)

    return wrapper


class ExcelTranslationOnlyMixedTests(unittest.TestCase):
    def _run(self, *, translation_only: bool, fixes_on_retry: bool):
        seen: dict = {}
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = excel_helpers._make_xlsx(
                root / "boq.xlsx", {"A1": REBAR, "A2": CONCRETE, "A3": TOTAL, "A4": RECEPTION, "A5": "混凝土垫层"}
            )
            settings = excel_helpers._settings(source_lang="zh", target_lang="fr")
            settings.excel_output.output_translation_only = translation_only
            settings.excel_output.keep_original_sheets = False
            with excel_helpers._pipeline_patches(
                translate_side_effect=lambda texts, *a, **k: {text: "Béton de propreté" for text in texts}
            ), patch(
                "core.task_runner.translate_mixed_language_texts",
                side_effect=_with_keeps_foreign_engine(fixes_on_retry, seen),
            ):
                runner = TaskRunner([FileItem(path=path, name="boq", size_kb=1.0)], settings, source_root=root)
                runner._run()
            outputs = [p for p in root.rglob("*.xlsx") if p != path]
            self.assertEqual(len(outputs), 1, outputs)
            sheet = load_workbook(outputs[0]).worksheets[0]
            values = {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value}
            logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
        return seen, values, logs

    def test_translation_only_cells_end_up_in_french_only(self) -> None:
        seen, values, logs = self._run(translation_only=True, fixes_on_retry=True)
        self.assertTrue(seen["translation_only"])
        texts = set(values.values())
        self.assertTrue({"Armatures HRB400 Ø12", "Béton C30", "Total", "Réception"} <= texts, values)
        self.assertFalse([v for v in texts if any(w in str(v) for w in ("Concrete", "Rebar", "Total Total", "合计", "验收"))], values)
        self.assertFalse([line for line in logs if "仍夹带未翻译的外文" in line], logs)

    def test_stubborn_model_leaves_a_review_mark_and_a_log_line(self) -> None:
        _seen, values, logs = self._run(translation_only=True, fixes_on_retry=False)
        self.assertIn("Béton C30 Concrete", set(values.values()))
        self.assertTrue([line for line in logs if "2 处译文中仍夹带未翻译的外文" in line], logs)

    def test_bilingual_output_passes_the_flag_off(self) -> None:
        seen, _values, _logs = self._run(translation_only=False, fixes_on_retry=True)
        self.assertFalse(seen["translation_only"])


class WordTranslationOnlyMixedTests(unittest.TestCase):
    def _run(self, *, fixes_on_retry: bool):
        seen: dict = {}
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "boq.docx"
            word_helpers.WordResumeTests._make_docx(path, [REBAR, CONCRETE, TOTAL, RECEPTION, "混凝土垫层"])
            settings = AppSettings(source_lang="zh", target_lang="fr")
            settings.word_output.output_translation_only = True
            runner = WordTaskRunner([WordFileItem(path=path, name=path.name, size_kb=1.0)], settings, source_root=root)
            with ExitStack() as stack:
                word_helpers.WordResumeTests._runner_patches(self, stack, root=root)
                stack.enter_context(patch(
                    "core.word_task_runner.translate_mixed_language_texts",
                    side_effect=_with_keeps_foreign_engine(fixes_on_retry, seen),
                ))
                runner._run()
            outputs = list((root / "out").rglob("*.docx"))
            self.assertEqual(len(outputs), 1, outputs)
            paragraphs = [p.text for p in Document(outputs[0]).paragraphs if p.text.strip()]
            logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
        return seen, paragraphs, logs

    def test_translation_only_paragraphs_end_up_in_french_only(self) -> None:
        seen, paragraphs, logs = self._run(fixes_on_retry=True)
        self.assertTrue(seen["translation_only"])
        for expected in ("Armatures HRB400 Ø12", "Béton C30", "Total", "Réception"):
            self.assertIn(expected, paragraphs)
        self.assertFalse([p for p in paragraphs if any(w in p for w in ("Concrete", "Rebar", "Total Total", "合计", "验收"))], paragraphs)
        self.assertFalse([line for line in logs if "仍夹带未翻译的外文" in line], logs)

    def test_stubborn_model_leaves_a_review_mark_and_a_log_line(self) -> None:
        _seen, paragraphs, logs = self._run(fixes_on_retry=False)
        self.assertIn("Béton C30 Concrete", paragraphs)
        self.assertTrue([line for line in logs if "2 处译文中仍夹带未翻译的外文" in line], logs)


if __name__ == "__main__":
    unittest.main(verbosity=2)
