"""Excel 输出的三处修复：原文副本不成组、仅译文文件名不带「双语」、空填充不再让扫描跳过。"""

from __future__ import annotations

import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from core.bilingual_writer import bilingual_output_name, write_bilingual_file
from core.file_scanner import scan_excel_sources
from core.openpyxl_compat import load_workbook
from core.output_name_translation import avoid_bilingual_name_collision
from core.resume_detection import match_previous_output


def _tab_selected_sheets(path: Path) -> list[str]:
    """按工作簿里的表名返回带 tabSelected="1" 的分表。"""
    with zipfile.ZipFile(path) as package:
        workbook = package.read("xl/workbook.xml").decode("utf-8")
        rels = package.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        targets.update(
            {rid: target for target, rid in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)}
        )
        selected = []
        for name, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', workbook):
            part = "xl/" + targets[rid].lstrip("/").removeprefix("xl/")
            if 'tabSelected="1"' in package.read(part).decode("utf-8"):
                selected.append(name)
        return selected


def _with_empty_fill(source: Path, target: Path) -> None:
    """把 styles.xml 的填充表末尾加一个空的 <fill/>（Excel 能打开，openpyxl 原本读不了）。"""
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/styles.xml":
                text = data.decode("utf-8")
                count = int(re.search(r'<fills count="(\d+)"', text).group(1))
                text = text.replace(f'<fills count="{count}">', f'<fills count="{count + 1}">')
                text = text.replace("</fills>", "<fill/></fills>")
                data = text.encode("utf-8")
            zout.writestr(info, data)


class OriginalSheetGroupingTests(unittest.TestCase):
    def test_cloned_original_sheet_is_not_selected_with_its_source(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "产值.xlsx"
            wb = Workbook()
            wb.active.title = "产值"
            wb.active["A1"] = "项目"
            wb.active.sheet_view.tabSelected = True  # Excel 保存时当前表都会带
            wb.save(source)
            self.assertEqual(_tab_selected_sheets(source), ["产值"])

            out_path = write_bilingual_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                translations={"项目": "Projet"},
                target_lang="fr",
                keep_original_sheets=True,
                formula_display_value_backfill=True,
            )

            self.assertEqual(_tab_selected_sheets(out_path), ["产值"])
            self.assertEqual(load_workbook(out_path).sheetnames, ["产值", "产值_原文"])

    def test_resumed_output_from_old_version_is_ungrouped(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "旧产物.xlsx"
            wb = Workbook()
            wb.active.title = "产值"
            wb.active["A1"] = "项目"
            wb.active.sheet_view.tabSelected = True
            clone = wb.create_sheet("产值_原文")
            clone["A1"] = "项目"
            clone.sheet_view.tabSelected = True
            wb.save(source)
            self.assertEqual(_tab_selected_sheets(source), ["产值", "产值_原文"])

            out_path = write_bilingual_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                translations={"项目": "Projet"},
                target_lang="fr",
                keep_original_sheets=True,
                formula_display_value_backfill=True,
            )

            self.assertEqual(_tab_selected_sheets(out_path), ["产值"])
            self.assertEqual(load_workbook(out_path).sheetnames, ["产值", "产值_原文"])


class TranslationOnlyNameTests(unittest.TestCase):
    def test_name_drops_bilingual_label_only_for_translation_only(self) -> None:
        self.assertEqual(bilingual_output_name("报价.xlsx", "法文"), "报价_法文_双语.xlsx")
        self.assertEqual(
            bilingual_output_name("报价.xlsx", "法文", translation_only=True), "报价_法文.xlsx"
        )

    def test_excel_writer_uses_translation_only_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "报价.xlsx"
            wb = Workbook()
            wb.active["A1"] = "项目"
            wb.save(source)

            out_path = write_bilingual_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                translations={"项目": "Projet"},
                target_lang="fr",
                keep_original_sheets=False,
                formula_display_value_backfill=True,
                output_translation_only=True,
            )

            self.assertEqual(out_path.name, "报价_法文.xlsx")
            self.assertEqual(load_workbook(out_path).active["A1"].value, "Projet")

    def test_collision_check_looks_at_translation_only_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            (output / "报价_法文.xlsx").write_bytes(b"existing")
            self.assertEqual(
                avoid_bilingual_name_collision(output, "报价.xlsx", "fr", translation_only=True),
                "报价_2.xlsx",
            )
            self.assertEqual(avoid_bilingual_name_collision(output, "报价.xlsx", "fr"), "报价.xlsx")

    def test_resume_finds_translation_only_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            previous = Path(tmp) / "报价_法文.xlsx"
            previous.write_bytes(b"x")
            self.assertEqual(match_previous_output(tmp, ".", "报价.xlsx", "fr", "excel"), previous)


class EmptyFillTests(unittest.TestCase):
    def test_workbook_with_empty_fill_is_scanned_and_written(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp) / "base.xlsx"
            wb = Workbook()
            wb.active["A1"] = "项目"
            wb.save(base)
            folder = Path(tmp) / "in"
            folder.mkdir()
            source = folder / "产值.xlsx"
            _with_empty_fill(base, source)
            base.unlink()

            result = scan_excel_sources(folder)
            self.assertEqual([item.path.name for item in result.items], ["产值.xlsx"])
            self.assertEqual(result.skipped, [])

            out_path = write_bilingual_file(
                source_path=source,
                output_dir=Path(tmp) / "out",
                translations={"项目": "Projet"},
                target_lang="fr",
                keep_original_sheets=True,
                formula_display_value_backfill=True,
                review_marks={"项目": "unresolved"},
            )
            with zipfile.ZipFile(out_path) as package:
                self.assertIn("<fill/>", package.read("xl/styles.xml").decode("utf-8"))
            self.assertEqual(load_workbook(out_path).active["A1"].value, "项目\nProjet")


if __name__ == "__main__":
    unittest.main()
