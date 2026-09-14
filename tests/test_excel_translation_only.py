"""Excel translation-only output contracts for full and supplementary writes."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.bilingual_writer import write_bilingual_file
from core.excel_coverage import ExcelCoveragePlan, write_untranslated_excel_file
from core.translation_coverage import COVERAGE_SOURCE_ONLY, CoverageUnit


def _plan(*items: tuple[str, str, str]) -> ExcelCoveragePlan:
    return ExcelCoveragePlan(
        path=Path("unused.xlsx"),
        units=[
            CoverageUnit(
                source_text=text,
                status=COVERAGE_SOURCE_ONLY,
                location=f"{sheet}!{coordinate}",
                reason="test",
                data={"sheet": sheet, "coordinate": coordinate},
            )
            for sheet, coordinate, text in items
        ],
        sheet_count=1,
    )


class ExcelTranslationOnlyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / "source.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "报价"
        ws["A1"] = "施工内容"
        ws["A2"] = "待补译"
        ws["A3"] = "原文\nOld translation"
        ws["B1"] = "=A1"
        wb.save(self.source)
        wb.close()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_main_writeback_defaults_to_bilingual_and_can_emit_translation_only(self) -> None:
        bilingual = write_bilingual_file(
            source_path=self.source,
            output_dir=self.root / "main-bilingual",
            translations={"施工内容": "Construction scope"},
            target_lang="en",
            source_lang="zh",
            keep_original_sheets=True,
            formula_display_value_backfill=True,
        )
        only = write_bilingual_file(
            source_path=self.source,
            output_dir=self.root / "main-only",
            translations={"施工内容": "Construction scope"},
            target_lang="en",
            source_lang="zh",
            keep_original_sheets=True,
            formula_display_value_backfill=True,
            output_translation_only=True,
        )
        for path, expected in ((bilingual, "施工内容\nConstruction scope"), (only, "Construction scope")):
            wb = load_workbook(path, data_only=False)
            try:
                self.assertEqual(wb["报价"]["A1"].value, expected)
                self.assertIn("报价_原文", wb.sheetnames)
                self.assertEqual(wb["报价_原文"]["A1"].value, "施工内容")
                self.assertEqual(wb["报价"]["B1"].value, "=A1")
            finally:
                wb.close()

    def test_supplementary_writeback_defaults_to_bilingual_and_supports_translation_only(self) -> None:
        plan = _plan(("报价", "A1", "施工内容"), ("报价", "A3", "原文\nOld translation"))
        bilingual = write_untranslated_excel_file(
            source_path=self.source,
            output_dir=self.root / "supp-bilingual",
            plan=plan,
            translations={"施工内容": "Construction scope", "原文\nOld translation": "New translation"},
            target_lang="en",
            keep_original_sheets=True,
        )
        only = write_untranslated_excel_file(
            source_path=self.source,
            output_dir=self.root / "supp-only",
            plan=plan,
            translations={"施工内容": "Construction scope", "原文\nOld translation": "New translation"},
            target_lang="en",
            keep_original_sheets=True,
            output_translation_only=True,
        )
        for path, expected in ((bilingual, "施工内容\nConstruction scope"), (only, "Construction scope")):
            wb = load_workbook(path, data_only=False)
            try:
                self.assertEqual(wb["报价"]["A1"].value, expected)
                expected_a3 = (
                    "New translation"
                    if expected == "Construction scope"
                    else "原文\nOld translation\nNew translation"
                )
                self.assertEqual(wb["报价"]["A3"].value, expected_a3)
                self.assertIn("报价_原文", wb.sheetnames)
                self.assertEqual(wb["报价_原文"]["A1"].value, "施工内容")
                self.assertEqual(wb["报价"]["B1"].value, "=A1")
            finally:
                wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
