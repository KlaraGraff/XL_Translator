"""PDF 翻译方式（Beta 分支）：设置、文件分类、每页走哪条路线。"""
from __future__ import annotations

import tempfile
import unittest

import pytest
from pathlib import Path


from core.pdf_routes import (
    KIND_BLANK,
    KIND_DRAWING,
    KIND_MIXED,
    KIND_SCAN,
    KIND_TEXT,
    classify_document,
    exceeds_word_page,
    page_route,
)
from core.pdf_word_route import normalize_fonts_xml
from settings import AppSettings, PdfSettings

# PyMuPDF 属于「PDF 处理组件」，不随安装包分发；没装它的环境跳过这些测试。
fitz = pytest.importorskip("fitz")

A4 = (595.276, 841.89)
A3 = (841.89, 1190.551)
A1 = (1683.78, 2383.94)
SENTENCE = "The applicant must present a valid passport and proof of travel insurance before departure."


def _route(pdf: PdfSettings, *, a3=False, big=False, kind=KIND_TEXT) -> str:
    return page_route(pdf, oversized_a3=a3, oversized=big, page_kind=kind)


class SettingsTests(unittest.TestCase):
    def test_defaults_keep_the_original_behaviour(self) -> None:
        pdf = PdfSettings()
        self.assertEqual(pdf.translation_route, "image")
        self.assertTrue(pdf.skip_oversized_pages)
        self.assertEqual(_route(pdf, a3=True, big=True, kind=KIND_DRAWING), "skip")
        self.assertEqual(_route(pdf, kind=KIND_TEXT), "image")

    def test_old_settings_file_still_loads(self) -> None:
        # 主线版本存的设置没有这几个字段；分支早期版本存过的旧字段要被忽略
        pdf = PdfSettings.model_validate({"skip_oversized_pages": False, "oversized_page_mode": "inplace"})
        self.assertEqual(pdf.translation_route, "image")
        self.assertFalse(pdf.skip_oversized_pages)

    def test_saved_through_the_settings_api(self) -> None:
        from api.routes.settings import _deep_merge

        base = AppSettings().model_dump(mode="json")
        for route in ("image", "auto"):
            merged = _deep_merge(base, {"pdf": {"translation_route": route, "auto_text_route": "inplace",
                                                "auto_scan_route": "skip", "auto_fallback": "skip"}})
            pdf = AppSettings.model_validate(merged).pdf
            self.assertEqual(pdf.translation_route, route)
            self.assertEqual((pdf.auto_text_route, pdf.auto_scan_route, pdf.auto_fallback), ("inplace", "skip", "skip"))


class PageRouteTests(unittest.TestCase):
    def test_image_route_honours_the_skip_switch(self) -> None:
        pdf = PdfSettings(skip_oversized_pages=False)
        self.assertEqual(_route(pdf, a3=True, big=True, kind=KIND_DRAWING), "image")

    def test_auto_route_prefers_one_method_per_page_type(self) -> None:
        pdf = PdfSettings(translation_route="auto")
        self.assertEqual(_route(pdf, kind=KIND_SCAN), "image")
        self.assertEqual(_route(pdf, kind=KIND_TEXT), "word")
        self.assertEqual(_route(pdf, kind=KIND_DRAWING), "inplace")
        self.assertEqual(_route(pdf, big=True, kind=KIND_DRAWING), "inplace")
        self.assertEqual(_route(pdf, kind=KIND_BLANK), "skip")
        # 先看内容再看尺寸：扫描页不论多大都按扫描页，文字页放不进 Word 才按图纸
        self.assertEqual(_route(pdf, a3=True, big=True, kind=KIND_SCAN), "image")
        self.assertEqual(_route(pdf, a3=True, kind=KIND_TEXT), "word")
        self.assertEqual(_route(pdf, a3=True, big=True, kind=KIND_TEXT), "inplace")
        custom = PdfSettings(translation_route="auto", auto_scan_route="skip", auto_text_route="inplace",
                             auto_drawing_route="skip")
        self.assertEqual(_route(custom, kind=KIND_SCAN), "skip")
        self.assertEqual(_route(custom, kind=KIND_TEXT), "inplace")
        self.assertEqual(_route(custom, big=True, kind=KIND_DRAWING), "skip")

    def test_safety_net(self) -> None:
        from core.pdf_routes import fallback_route

        pdf = PdfSettings(translation_route="auto")
        self.assertEqual(fallback_route(pdf, oversized=False, image_available=True), "image")
        self.assertEqual(fallback_route(pdf, oversized=True, image_available=True), "skip")   # no blurry drawings
        self.assertEqual(fallback_route(pdf, oversized=False, image_available=False), "skip")
        keep = PdfSettings(translation_route="auto", auto_fallback="skip")
        self.assertEqual(fallback_route(keep, oversized=False, image_available=True), "skip")

    def test_earlier_single_routes_become_the_same_combination(self) -> None:
        inplace = PdfSettings(translation_route="inplace")
        self.assertEqual((inplace.translation_route, inplace.auto_text_route, inplace.auto_drawing_route),
                         ("auto", "inplace", "inplace"))
        word = PdfSettings(translation_route="word")
        self.assertEqual((word.translation_route, word.auto_text_route, word.auto_drawing_route),
                         ("auto", "word", "skip"))
        word_no_skip = PdfSettings(translation_route="word", skip_oversized_pages=False)
        self.assertEqual(word_no_skip.auto_drawing_route, "image")

    def test_unknown_route_values_fall_back_to_defaults(self) -> None:
        # a newer build may write options this one does not know: never refuse the file
        pdf = PdfSettings.model_validate({"translation_route": "auto", "auto_text_route": "ocr",
                                          "oversized_page_threshold": "A5", "auto_fallback": "skip"})
        self.assertEqual((pdf.translation_route, pdf.auto_text_route, pdf.oversized_page_threshold,
                          pdf.auto_fallback), ("auto", "word", "A3", "skip"))



def _text_page(doc, size=A4) -> None:
    page = doc.new_page(width=size[0], height=size[1])
    y = 80
    for _ in range(12):
        page.insert_text((50, y), SENTENCE[:90], fontsize=9)
        y += 16


def _drawing_page(doc, size=A4) -> None:
    page = doc.new_page(width=size[0], height=size[1])
    shape = page.new_shape()
    for i in range(700):
        shape.draw_line((20 + i % 50 * 10, 100 + i // 50 * 20), (30 + i % 50 * 10, 110 + i // 50 * 20))
        shape.draw_circle((40 + i % 40 * 12, 500 + i // 40 * 12), 4)
    shape.finish(color=(0, 0, 0), width=0.3)
    shape.commit()
    for i, label in enumerate(("Mur coupe feu", "Plan de toiture", "Détail A")):
        page.insert_text((60, 60 + i * 14), label, fontsize=8)


def _scan_page(doc, size=A4) -> None:
    page = doc.new_page(width=size[0], height=size[1])
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 80), 0)
    pix.clear_with(200)
    page.insert_image(page.rect, pixmap=pix)


class ClassifyTests(unittest.TestCase):
    def _classify(self, builders) -> str:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "doc.pdf"
        doc = fitz.open()
        for build in builders:
            build(doc)
        doc.save(path)
        return classify_document(path, is_oversized=exceeds_word_page)

    def test_kinds(self) -> None:
        self.assertEqual(self._classify([_text_page, _text_page]).kind, KIND_TEXT)
        self.assertEqual(self._classify([_drawing_page]).kind, KIND_DRAWING)
        self.assertEqual(self._classify([lambda d: _drawing_page(d, A1)]).kind, KIND_DRAWING)
        self.assertEqual(self._classify([_scan_page]).kind, KIND_SCAN)
        self.assertEqual(self._classify([_text_page, _scan_page]).kind, KIND_MIXED)
        blank = self._classify([_text_page, lambda d: d.new_page()])
        self.assertEqual(blank.kind, KIND_TEXT)
        self.assertEqual(blank.pages[1].kind, KIND_BLANK)

    def test_content_first_then_size(self) -> None:
        self.assertEqual(self._classify([lambda d: _text_page(d, A3)]).kind, KIND_TEXT)     # fits in Word
        self.assertEqual(self._classify([lambda d: _text_page(d, A1)]).kind, KIND_DRAWING)  # too large for Word
        self.assertEqual(self._classify([lambda d: _scan_page(d, A1)]).kind, KIND_SCAN)     # a scan stays a scan
        self.assertEqual(self._classify([lambda d: d.new_page(width=A1[0], height=A1[1])]).kind, KIND_SCAN)

    def test_ocr_only_where_it_can_find_something(self) -> None:
        from core.pdf_routes import needs_ocr

        big_text = self._classify([lambda d: _text_page(d, A1)]).pages[0]
        self.assertEqual(big_text.kind, KIND_DRAWING)
        self.assertFalse(needs_ocr(big_text))                     # too big for Word, text all readable
        self.assertFalse(needs_ocr(self._classify([_text_page]).pages[0]))
        self.assertTrue(needs_ocr(self._classify([_drawing_page]).pages[0]))   # dense lines
        self.assertTrue(needs_ocr(self._classify([lambda d: _scan_page(d, A1)]).pages[0]))

    def test_lines_without_text_are_a_drawing(self) -> None:
        def outlined(doc) -> None:   # CAD export with the text drawn as lines
            page = doc.new_page(width=A3[0], height=A3[1])
            shape = page.new_shape()
            for i in range(1100):
                shape.draw_line((20 + i % 50 * 10, 100 + i // 50 * 20), (28 + i % 50 * 10, 112 + i // 50 * 20))
            shape.finish(color=(0, 0, 0), width=0.3)
            shape.commit()

        self.assertEqual(self._classify([outlined]).kind, KIND_DRAWING)


class FontNormalizationTests(unittest.TestCase):
    def test_subset_and_cid_fonts_get_real_names_and_an_east_asian_font(self) -> None:
        xml = ('<w:rPr><w:rFonts w:ascii="ABCDEF+Arial-BoldMT" w:hAnsi="CIDFont+F3" w:eastAsia="Arial"/>'
               '<w:lang w:val="en-US" w:eastAsia="zh-CN"/></w:rPr>')
        out = normalize_fonts_xml(xml)
        self.assertIn('w:ascii="Arial-BoldMT"', out)
        self.assertIn('w:hAnsi="Arial"', out)
        self.assertIn('<w:rFonts w:eastAsia="Microsoft YaHei"', out)
        self.assertIn('<w:lang w:val="en-US" w:eastAsia="zh-CN"/>', out)  # untouched


if __name__ == "__main__":
    unittest.main()


class WritebackGeometryTests(unittest.TestCase):
    def test_white_text_on_a_dark_band_counts_as_ink(self) -> None:
        from PIL import Image, ImageDraw

        from core.pdf_vector.writeback import ink_count, is_light

        im = Image.new("L", (100, 20), 30)            # dark title bar
        ImageDraw.Draw(im).rectangle((10, 5, 40, 14), fill=255)  # white "text"
        self.assertGreater(ink_count(im, (0, 0, 100, 20), light=True)[0], 200)
        self.assertEqual(ink_count(Image.new("L", (100, 20), 30), (0, 0, 100, 20), light=True)[0], 0)
        self.assertTrue(is_light((1.0, 1.0, 1.0)))
        self.assertFalse(is_light((0.0, 0.2, 0.5)))

    def test_content_box_is_the_outermost_ink(self) -> None:
        from PIL import Image, ImageDraw

        from core.pdf_vector.writeback import Layout, content_box

        im = Image.new("L", (600, 800), 255)
        ImageDraw.Draw(im).rectangle((80, 100, 520, 700), outline=0)
        box = content_box(im)
        self.assertEqual(box, (80, 100, 521, 701))
        layout = Layout(Image.new("L", (600, 800), 255), 72, bounds=box)
        self.assertTrue(layout.blank((100, 200, 300, 220)))
        self.assertFalse(layout.blank((500, 200, 560, 220)))  # would run into the margin


class LightTextGroundTests(unittest.TestCase):
    def test_light_ink_only_on_a_dark_ground(self) -> None:
        from PIL import Image

        from core.pdf_vector.writeback import dark_ground

        self.assertTrue(dark_ground(Image.new("L", (50, 20), 30), (0, 0, 50, 20)))
        self.assertFalse(dark_ground(Image.new("L", (50, 20), 250), (0, 0, 50, 20)))  # yellow text on white paper


class WordConversionCleanupTests(unittest.TestCase):
    def _doc(self, lines):
        from docx import Document
        from docx.shared import Pt

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "conv.docx"
        doc = Document()
        for text, size, bold in lines:
            run = doc.add_paragraph().add_run(text)
            run.font.size = Pt(size)
            run.bold = bold
        doc.save(path)
        return path

    def test_wrapped_title_lines_become_one_paragraph(self) -> None:
        from docx import Document

        from core.pdf_word_route import merge_wrapped_lines

        path = self._doc([
            ("Early-Age Strength of Mortars Prepared with ", 16, True),
            ("Washed River Sand", 16, True),
            ("1. Introduction", 11, True),               # different size: a new block
            ("Mortars use cement.", 10, False),
            ("- first item", 10, False),                 # list item: kept apart
        ])
        self.assertEqual(merge_wrapped_lines(path), 1)
        texts = [p.text for p in Document(str(path)).paragraphs]
        self.assertEqual(texts[0], "Early-Age Strength of Mortars Prepared with Washed River Sand")
        self.assertEqual(texts[1:], ["1. Introduction", "Mortars use cement.", "- first item"])

    def test_dropped_word_spaces_are_restored(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        from core.pdf_word_route import restore_word_spaces

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "spaces.docx"
        doc = Document()
        p = doc.add_paragraph()
        for text in ("Aspects of the shell’s", "syntax and command language.", "A", "number of builtins such as"):
            p.add_run(text)
        p.add_run("cd").bold = True
        p.add_run("and ")
        p.add_run("exec").bold = True
        p.add_run("(see")
        p.add_run("1").font.superscript = True
        p.add_run(")")
        line = doc.add_paragraph()
        line.add_run("first line")
        line.runs[0].add_break(WD_BREAK.LINE)
        line.add_run("second")
        doc.add_paragraph().add_run("中文")
        doc.paragraphs[-1].add_run("内容")
        split = doc.add_paragraph()
        split.add_run("the ev")
        split.add_run("al builtin")     # one word split across two runs: no space
        doc.save(path)

        self.assertEqual(restore_word_spaces(path, {"the", "eval", "builtin"}), 6)
        texts = [p.text for p in Document(str(path)).paragraphs]
        self.assertEqual(texts[0], "Aspects of the shell’s syntax and command language. A number of builtins such as cd and exec (see1)")
        self.assertEqual(texts[1], "first line\nsecond")
        self.assertEqual(texts[2], "中文内容")
        self.assertEqual(texts[3], "the eval builtin")

    def test_restored_space_stays_outside_a_formatted_run(self) -> None:
        from docx import Document

        from core.pdf_word_route import restore_word_spaces

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "marks.docx"
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Press the")
        p.add_run("STOP").bold = True
        p.add_run("button.")
        doc.save(path)
        restore_word_spaces(path, set())
        runs = [run.text for run in Document(str(path)).paragraphs[0].runs]
        self.assertEqual(runs, ["Press the ", "STOP", " button."])

    def test_soft_break_after_a_trailing_space_adds_no_second_space(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK

        from core.pdf_word_route import merge_wrapped_lines

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "br.docx"
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("may open the cabinet and replace the ")
        p.add_run().add_break(WD_BREAK.LINE)
        p.add_run("fuses after isolation.")
        doc.save(path)
        merge_wrapped_lines(path)
        self.assertEqual(Document(str(path)).paragraphs[0].text,
                         "may open the cabinet and replace the fuses after isolation.")

    def test_row_heights_are_loosened(self) -> None:
        from core.pdf_word_route import loosen_row_heights_xml

        self.assertIn('w:hRule="atLeast"', loosen_row_heights_xml('<w:trHeight w:hRule="exact" w:val="306"/>'))


class DominantRunShapeTests(unittest.TestCase):
    def test_a_short_bold_label_does_not_make_the_translation_bold(self) -> None:
        from docx import Document

        from core.word_document import _replace_paragraph_text

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Abstract — ").bold = True
        p.add_run("This study investigates the compressive strength of masonry mortars.")
        _replace_paragraph_text(p, "摘要——本研究考察了砌筑砂浆的抗压强度。", target_lang="zh")
        self.assertEqual(p.text, "摘要——本研究考察了砌筑砂浆的抗压强度。")
        self.assertFalse(any(run.bold for run in p.runs if run.text))

    def test_a_mostly_bold_paragraph_stays_bold(self) -> None:
        from docx import Document

        from core.word_document import _replace_paragraph_text

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("General ICAO passport photo requirements").bold = True
        p.add_run(" (2024)")
        _replace_paragraph_text(p, "ICAO 护照照片一般要求（2024）", target_lang="zh")
        self.assertTrue(next(run for run in p.runs if run.text).bold)


    def test_bold_first_line_of_a_converted_paragraph_does_not_win(self) -> None:
        # pdf2docx: one run per line; the bold first line is the longest single run,
        # but most of the paragraph is regular text.
        from docx import Document

        from core.word_document import _replace_paragraph_text

        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Warning: the cabinet must only be opened by qualified staff members").bold = True
        p.add_run("after the supply has been isolated and locked out at the ")
        p.add_run("main switch, and the residual voltage has been verified as safe.")
        _replace_paragraph_text(p, "警告：……", target_lang="zh")
        self.assertFalse(next(run for run in p.runs if run.text).bold)

    def test_a_leading_link_does_not_style_the_whole_translation(self) -> None:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        from core.word_document import _replace_paragraph_text

        doc = Document()
        p = doc.add_paragraph()
        link = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        style = OxmlElement("w:rStyle")
        style.set(qn("w:val"), "Hyperlink")
        rpr.append(style)
        run.append(rpr)
        text = OxmlElement("w:t")
        text.text = "Installation and maintenance manual for the pump"
        run.append(text)
        link.append(run)
        p._p.append(link)
        p.add_run(" p. 12")
        _replace_paragraph_text(p, "泵的安装和维护手册 第 12 页", target_lang="zh")
        self.assertNotIn("Hyperlink", p._p.xml.split("泵的安装")[0].rsplit("<w:r>", 1)[-1])


class EastAsianLangTests(unittest.TestCase):
    def test_default_east_asian_language_follows_the_target(self) -> None:
        from core.pdf_word_route import east_asian_locale, set_default_east_asian_lang_xml

        xml = ('<w:docDefaults><w:rPrDefault><w:rPr><w:sz w:val="22"/>'
               '<w:lang w:val="en-US" w:eastAsia="en-US" w:bidi="ar-SA"/></w:rPr></w:rPrDefault></w:docDefaults>')
        out = set_default_east_asian_lang_xml(xml, east_asian_locale("zh"))
        self.assertIn('w:eastAsia="zh-CN"', out)
        self.assertIn('w:val="en-US"', out)
        bare = '<w:rPrDefault><w:rPr><w:sz w:val="22"/></w:rPr></w:rPrDefault>'
        self.assertIn('<w:lang w:eastAsia="ja-JP"/>', set_default_east_asian_lang_xml(bare, east_asian_locale("ja")))
        self.assertEqual(set_default_east_asian_lang_xml(xml, east_asian_locale("fr")), xml)


class NumberedLineMergeTests(unittest.TestCase):
    def test_numbered_question_and_soft_break(self) -> None:
        from docx import Document
        from docx.enum.text import WD_BREAK
        from docx.shared import Pt

        from core.pdf_word_route import merge_wrapped_lines

        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "q.docx"
        doc = Document()
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Pt(30))
        p.add_run("4 ").bold = True
        p.add_run().add_tab()
        r = p.add_run("I am a family member. Are the standard requirements ")
        r.bold = True
        r.font.size = Pt(12)
        p.runs[0].font.size = Pt(12)
        q = doc.add_paragraph().add_run("applicable to me?")
        q.bold = True
        q.font.size = Pt(12)
        caption = doc.add_paragraph()
        c1 = caption.add_run("Other people are visible in the")
        c1.add_break(WD_BREAK.LINE)
        caption.add_run("image")
        dashes = doc.add_paragraph()
        d1 = dashes.add_run("-")
        d1.add_break(WD_BREAK.LINE)
        dashes.add_run("-")
        doc.save(path)
        merge_wrapped_lines(path)
        texts = [p.text for p in Document(str(path)).paragraphs]
        self.assertEqual(texts[0], "4 \tI am a family member. Are the standard requirements applicable to me?")
        self.assertEqual(texts[1], "Other people are visible in the image")
        self.assertEqual(texts[2], "-\n-")


class OfficePdfTests(unittest.TestCase):
    def _pdf(self, render_mode: int) -> Path:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / f"cjk_{render_mode}.pdf"
        doc = fitz.open()
        page = doc.new_page(width=A4[0], height=A4[1])
        for i in range(6):
            page.insert_text((60, 100 + i * 30), "砌筑砂浆在中小型施工现场的水泥消耗中占很大比例", fontsize=14,
                             fontname="china-s", render_mode=render_mode)
        doc.save(path)
        return path

    def test_chinese_that_renders_blank_is_caught(self) -> None:
        from core.pdf_word_route import cjk_text_visible

        self.assertTrue(cjk_text_visible(self._pdf(0)))
        self.assertFalse(cjk_text_visible(self._pdf(3)))   # text layer present, nothing drawn

    def test_route_capabilities_shape(self) -> None:
        from unittest.mock import patch

        from api.routes.pdf_routes import route_capabilities

        with patch("core.pdf_word_route.office_pdf_engine", return_value=("", "none")):
            caps = route_capabilities()
        self.assertEqual(set(caps), {"image_model_ready", "image_model_message", "word_component",
                                     "inplace_component", "inplace_component_message",
                                     "office_engine", "office_message", "inplace_target_ok",
                                     "inplace_unsupported_targets"})
        self.assertEqual(caps["office_engine"], "")
        self.assertIn("ar", caps["inplace_unsupported_targets"])
        self.assertNotIn("fr", caps["inplace_unsupported_targets"])
