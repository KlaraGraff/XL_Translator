"""Excel 单元格局部格式：原文片段照抄、译文按标记分片段写回，以及保存后核对。

一格「注意：严禁带电作业」只有「严禁」加粗标红。以前双语和仅译文都整格写成一个纯文本，
双语里原文那半截的格式也丢了。
"""
from __future__ import annotations

import os
import shutil
import tempfile
import unittest
import zipfile
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import MagicMock, patch

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from core.xlsx_patcher import PackageVerificationError, _Package, write_bilingual_workbook
from core.xlsx_rich_text import collect_marked_sources, rich_cell, string_item_runs

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
SOURCE = "注意：严禁带电作业"
TRANSLATION = "Attention : interdiction de travailler sous tension"
MARKED_SOURCE = "注意：<g1>严禁</g1>带电作业"
MARKED_TRANSLATION = "Attention : <g1>interdiction</g1> de travailler sous tension"
LONG_TAIL = " ; couper l'alimentation et vérifier l'absence de tension avant toute intervention" * 4


def _plain() -> InlineFont:
    return InlineFont(sz=11, rFont="宋体")


def _red_bold() -> InlineFont:
    return InlineFont(b=True, color="FFFF0000", sz=11, rFont="宋体")


def _blocks(path: Path, coordinate: str = "A1") -> list[tuple[str, bool, str | None]]:
    """(文字, 是否加粗, 颜色)；纯文本单元格返回一项 (文字, False, None)。"""
    value = load_workbook(path, rich_text=True).active[coordinate].value
    if isinstance(value, str):
        return [(value, False, None)]
    out = []
    for block in value:
        if isinstance(block, str):
            out.append((block, False, None))
        else:
            color = block.font.color.rgb if block.font.color is not None else None
            out.append((block.text, bool(block.font.b), color))
    return out


class RichTextFixture(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def _workbook(self, value=None) -> Path:
        wb = Workbook()
        ws = wb.active
        ws["A1"] = value if value is not None else CellRichText(
            [TextBlock(_plain(), "注意："), TextBlock(_red_bold(), "严禁"), TextBlock(_plain(), "带电作业")]
        )
        ws["A2"] = "普通文本"
        path = self.root / "src.xlsx"
        wb.save(path)
        return path

    def _write(self, source: Path, *, translation_only: bool, marked=None, translations=None, **kwargs) -> Path:
        out = self.root / ("only.xlsx" if translation_only else "bi.xlsx")
        shutil.copy(source, out)
        self.stats: dict = {}
        write_bilingual_workbook(
            out,
            translations=translations or {SOURCE: TRANSLATION, "普通文本": "Texte"},
            target_lang="fr",
            output_translation_only=translation_only,
            marked_translations=marked,
            stats=self.stats,
            **kwargs,
        )
        return out


class RichCellMarksTest(RichTextFixture):
    def test_collects_marks_for_mixed_format_cells_only(self) -> None:
        source = self._workbook()
        self.assertEqual(collect_marked_sources(source), {SOURCE: MARKED_SOURCE})

    def test_group_spanning_a_line_break_gets_one_pair_per_line(self) -> None:
        item = etree.fromstring(
            f'<si xmlns="{NS[1:-1]}"><r><t>甲</t></r><r><rPr><b/></rPr><t>乙\n丙</t></r><r><t>丁</t></r></si>'
        )
        cell = rich_cell(string_item_runs(item))
        self.assertEqual(cell.marked_text, "甲<g1>乙</g1>\n<g2>丙</g2>丁")

    def test_runs_that_only_differ_by_font_name_need_no_marks(self) -> None:
        item = etree.fromstring(
            f'<si xmlns="{NS[1:-1]}"><r><rPr><rFont val="宋体"/></rPr><t>型号</t></r>'
            '<r><rPr><rFont val="Arial"/></rPr><t>ABC</t></r></si>'
        )
        cell = rich_cell(string_item_runs(item))
        self.assertEqual(cell.groups, {})
        self.assertEqual(cell.marked_text, "型号ABC")


class RichCellWriteTest(RichTextFixture):
    def test_bilingual_keeps_source_runs_and_marks_translation(self) -> None:
        out = self._write(self._workbook(), translation_only=False, marked={SOURCE: (MARKED_SOURCE, MARKED_TRANSLATION)})
        self.assertEqual(
            _blocks(out),
            [
                ("注意：", False, None),
                ("严禁", True, "FFFF0000"),
                ("带电作业", False, None),
                ("\n", False, None),
                ("Attention : ", False, None),
                ("interdiction", True, "FFFF0000"),
                (" de travailler sous tension", False, None),
            ],
        )
        self.assertEqual(self.stats["rich_cells"], 1)
        self.assertEqual(self.stats["marked_cells"], 1)
        # 普通格照旧是一段纯文本
        self.assertEqual(_blocks(out, "A2"), [("普通文本\nTexte", False, None)])

    def test_translation_only_marks_translation(self) -> None:
        out = self._write(self._workbook(), translation_only=True, marked={SOURCE: (MARKED_SOURCE, MARKED_TRANSLATION)})
        self.assertEqual(
            _blocks(out),
            [
                ("Attention : ", False, None),
                ("interdiction", True, "FFFF0000"),
                (" de travailler sous tension", False, None),
            ],
        )

    def test_without_marks_bilingual_still_keeps_source_format(self) -> None:
        out = self._write(self._workbook(), translation_only=False, marked=None)
        blocks = _blocks(out)
        self.assertEqual(blocks[:3], [("注意：", False, None), ("严禁", True, "FFFF0000"), ("带电作业", False, None)])
        # 译文写成一段，不带加粗也不带红色
        self.assertEqual("".join(text for text, _b, _c in blocks[3:]), "\n" + TRANSLATION)
        self.assertTrue(all(not bold and color is None for _t, bold, color in blocks[3:]))
        self.assertEqual(self.stats["marked_cells"], 0)

    def test_stale_marked_translation_is_not_used(self) -> None:
        # 翻译后被修复改过的译文和带标记版本对不上：译文写成一段，不硬套标记
        marked = {SOURCE: (MARKED_SOURCE, "Attention : <g1>interdit</g1> de travailler")}
        out = self._write(self._workbook(), translation_only=True, marked=marked)
        self.assertEqual(_blocks(out), [(TRANSLATION, False, None)])

    def test_long_formatted_part_does_not_paint_whole_translation(self) -> None:
        # 红字那截比普通那截长：标记对不上时也不能整段变红
        source = self._workbook(
            CellRichText([TextBlock(_plain(), "注："), TextBlock(_red_bold(), "严禁在未断电情况下带电作业")])
        )
        translations = {"注：严禁在未断电情况下带电作业": "Note : travail sous tension interdit"}
        out = self._write(source, translation_only=True, translations=translations)
        self.assertEqual(_blocks(out), [("Note : travail sous tension interdit", False, None)])

    def test_uniform_rich_cell_keeps_its_run_format_in_translation_only(self) -> None:
        source = self._workbook(CellRichText([TextBlock(_red_bold(), "严禁带电作业")]))
        out = self._write(source, translation_only=True, translations={"严禁带电作业": "Travail sous tension interdit"})
        self.assertEqual(_blocks(out), [("Travail sous tension interdit", True, "FFFF0000")])

    def test_shared_strings_part_stays_byte_identical(self) -> None:
        source = self._workbook()
        with zipfile.ZipFile(source) as archive:
            names = archive.namelist()
            before = archive.read("xl/sharedStrings.xml") if "xl/sharedStrings.xml" in names else None
        out = self._write(source, translation_only=False, marked={SOURCE: (MARKED_SOURCE, MARKED_TRANSLATION)})
        if before is not None:
            with zipfile.ZipFile(out) as archive:
                self.assertEqual(archive.read("xl/sharedStrings.xml"), before)

    def test_shared_string_rich_cell_is_written_inline(self) -> None:
        """Excel 自己存的富文本在共享字符串表里：照样按片段写成 inlineStr。"""
        source = self._workbook()
        sheet = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{NS[1:-1]}"><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row>'
            "</sheetData></worksheet>"
        )
        shared = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<sst xmlns="{NS[1:-1]}" count="1" uniqueCount="1"><si>'
            '<r><rPr><sz val="11"/></rPr><t>注意：</t></r>'
            '<r><rPr><b/><color rgb="FFFF0000"/><sz val="11"/></rPr><t>严禁</t></r>'
            '<r><rPr><sz val="11"/></rPr><t>带电作业</t></r>'
            '<phoneticPr fontId="1" type="noConversion"/></si></sst>'
        )
        with zipfile.ZipFile(source) as src:
            parts = {name: src.read(name) for name in src.namelist()}
        parts["xl/worksheets/sheet1.xml"] = sheet.encode()
        parts["xl/sharedStrings.xml"] = shared.encode()
        # 只有 Content_Types 和 rels 声明了，Excel 才认这个共享字符串表
        ct = etree.fromstring(parts["[Content_Types].xml"])
        override = etree.SubElement(ct, "{http://schemas.openxmlformats.org/package/2006/content-types}Override")
        override.set("PartName", "/xl/sharedStrings.xml")
        override.set("ContentType", "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
        parts["[Content_Types].xml"] = etree.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
        rels = etree.fromstring(parts["xl/_rels/workbook.xml.rels"])
        rel = etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
        rel.set("Id", "rIdShared")
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings")
        rel.set("Target", "sharedStrings.xml")
        parts["xl/_rels/workbook.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
        rebuilt = self.root / "shared.xlsx"
        with zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as dst:
            for name, data in parts.items():
                dst.writestr(name, data)
        self.assertEqual(collect_marked_sources(rebuilt), {SOURCE: MARKED_SOURCE})
        out = self._write(rebuilt, translation_only=False, marked={SOURCE: (MARKED_SOURCE, MARKED_TRANSLATION)})
        self.assertEqual([item[:2] for item in _blocks(out)][:3], [("注意：", False), ("严禁", True), ("带电作业", False)])
        self.assertIn(("interdiction", True, "FFFF0000"), _blocks(out))

    def test_red_font_review_mark_reaches_runs_with_own_color(self) -> None:
        from core.mixed_language import MIXED_MARK_UNRESOLVED
        from openpyxl.styles import PatternFill

        wb = Workbook()
        ws = wb.active
        ws["A1"] = CellRichText(
            [TextBlock(InlineFont(color="FF0000FF", sz=11), "注意："), TextBlock(_red_bold(), "严禁"),
             TextBlock(InlineFont(color="FF0000FF", sz=11), "带电作业")]
        )
        ws["A1"].fill = PatternFill("solid", fgColor="FFFFFF00")
        source = self.root / "filled.xlsx"
        wb.save(source)
        out = self._write(
            source,
            translation_only=True,
            marked={SOURCE: (MARKED_SOURCE, MARKED_TRANSLATION)},
            review_marks={SOURCE: MIXED_MARK_UNRESOLVED},
            existing_fill_policy="red_font",
        )
        colors = {color for _t, _b, color in _blocks(out)}
        self.assertEqual(colors, {"FFFF0000"})

    def test_locked_row_height_shrinks_run_sizes(self) -> None:
        out = self._write(
            self._workbook(),
            translation_only=True,
            marked={SOURCE: (MARKED_SOURCE, MARKED_TRANSLATION + LONG_TAIL)},
            translations={SOURCE: TRANSLATION + LONG_TAIL, "普通文本": "Texte"},
            lock_row_height=True,
        )
        # 片段自带的字号会盖过单元格样式，缩字时要一并缩小；加粗那截仍加粗
        value = load_workbook(out, rich_text=True).active["A1"].value
        self.assertNotIsInstance(value, str)
        sizes = {block.font.sz for block in value if not isinstance(block, str)}
        self.assertTrue(sizes and all(size < 11 for size in sizes), sizes)
        self.assertTrue(any(block.font.b for block in value if not isinstance(block, str)))


class PackageVerificationTest(RichTextFixture):
    def test_corrupted_rewritten_part_is_refused_and_source_kept(self) -> None:
        source = self._workbook()
        before = source.read_bytes()
        package = _Package(source)
        package.write("xl/worksheets/sheet1.xml", b"<worksheet><sheetData>")
        with self.assertRaises(PackageVerificationError) as raised:
            package.save()
        package.close()
        self.assertEqual(raised.exception.part, "xl/worksheets/sheet1.xml")
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse(source.with_name(source.name + ".patching").exists())

    def test_untouched_part_that_changed_on_the_way_is_refused(self) -> None:
        source = self._workbook()
        package = _Package(source)
        real_read = package._zip.read

        def tampered(name):
            data = real_read(name)
            return data + b" " if name == "xl/styles.xml" else data

        package._zip.read = tampered
        with self.assertRaises(PackageVerificationError) as raised:
            package.save()
        package.close()
        self.assertEqual(raised.exception.part, "xl/styles.xml")

    def test_normal_save_passes(self) -> None:
        out = self._write(self._workbook(), translation_only=False)
        with zipfile.ZipFile(out) as archive:
            self.assertIsNone(archive.testzip())


class ExcelRunnerRichTextTest(unittest.TestCase):
    """真跑 TaskRunner：带局部格式的格以带标记的原文发给模型，写回时按片段保留格式。"""

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name) / "docs"
        self.root.mkdir()
        self._env = patch.dict(os.environ, {"TRANSLATOR_APP_DATA_DIR": str(Path(self._tmp.name) / "appdata")})
        self._env.start()

    def tearDown(self) -> None:
        self._env.stop()
        self._tmp.cleanup()

    def _run(self, engine, *, tm_hits=None, translation_only=False) -> Path:
        from core.api_config_check import ApiConfigCheckResult
        from core.file_scanner import FileItem
        from core.task_runner import DoneMsg, TaskRunner
        from settings import AppSettings, EngineSettings
        from tests.rate_pins import pinned_surface_rate

        source = self.root / "安全须知.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = CellRichText(
            [TextBlock(_plain(), "注意："), TextBlock(_red_bold(), "严禁"), TextBlock(_plain(), "带电作业")]
        )
        ws["A2"] = "配电箱"
        wb.save(source)

        settings = AppSettings(
            engine=EngineSettings(
                mode="cloud",
                cloud_provider="custom_openai",
                cloud_model="rich-text-test",
                cloud_base_url="https://example.invalid/v1",
                concurrency=1,
                batch_size=10,
            ),
            source_lang="zh",
            target_lang="fr",
            excel_domain_preset="无",
            excel_smart_context=False,
        )
        settings.excel_output.enable_excel_autofit = False
        settings.excel_output.output_translation_only = translation_only
        runner = TaskRunner([FileItem(path=source, name=source.stem, size_kb=1.0)], settings, source_root=self.root)
        with ExitStack() as stack:
            stack.enter_context(patch("core.task_runner.TaskLogger", return_value=MagicMock(task_id="rich-text")))
            stack.enter_context(
                patch("core.task_runner.check_translation_api_config", return_value=ApiConfigCheckResult(ok=True))
            )
            stack.enter_context(patch("core.task_runner.build_engine", return_value=engine))
            stack.enter_context(patch("core.task_runner.get_system_prompt", return_value="system"))
            stack.enter_context(patch("core.task_runner.resolve_effective_model_config", return_value=object()))
            stack.enter_context(
                patch(
                    "core.task_runner.resolve_surface_rate",
                    side_effect=pinned_surface_rate(batch_items=10, concurrency=1),
                )
            )
            stack.enter_context(patch("core.task_runner.tm_manager.lookup_batch", return_value=tm_hits or {}))
            stack.enter_context(patch("core.task_runner.tm_manager.insert_batch", return_value=0))
            runner._run()
        done = [m for m in list(runner._queue.queue) if isinstance(m, DoneMsg)]
        self.assertEqual(len(done), 1, list(runner._queue.queue))
        return Path(done[0].files[0]["output_path"])

    def test_marked_source_goes_to_model_and_comes_back_as_runs(self) -> None:
        class Engine:
            engine_name = "rich-text/mock"

            def __init__(self) -> None:
                self.sent: list[str] = []

            def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
                self.sent.extend(texts)
                answers = {MARKED_SOURCE: MARKED_TRANSLATION, "配电箱": "Coffret électrique"}
                return {text: answers[text] for text in texts}

        engine = Engine()
        out = self._run(engine)
        self.assertIn(MARKED_SOURCE, engine.sent)
        blocks = _blocks(out)
        self.assertEqual(blocks[1], ("严禁", True, "FFFF0000"))
        self.assertIn(("interdiction", True, "FFFF0000"), blocks)
        self.assertEqual("".join(text for text, _b, _c in blocks), SOURCE + "\n" + TRANSLATION)
        self.assertEqual(_blocks(out, "A2"), [("配电箱\nCoffret électrique", False, None)])

    def test_memory_hit_gets_marks_placed_by_chat(self) -> None:
        class Engine:
            engine_name = "rich-text/mock"

            def __init__(self) -> None:
                self.chat_payloads: list[str] = []

            def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
                return {text: "Coffret électrique" for text in texts}

            def chat(self, system: str, user: str) -> str:
                import json

                self.chat_payloads.append(user)
                return json.dumps([MARKED_TRANSLATION], ensure_ascii=False)

        engine = Engine()
        with patch("engines.base_engine.engine_supports_chat", return_value=True):
            out = self._run(engine, tm_hits={SOURCE: TRANSLATION}, translation_only=True)
        self.assertEqual(len(engine.chat_payloads), 1)
        self.assertIn(("interdiction", True, "FFFF0000"), _blocks(out))


if __name__ == "__main__":
    unittest.main()
