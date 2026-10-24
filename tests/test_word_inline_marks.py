"""Word 段内局部格式：标记生成、译文标记校验、按标记写回、出错退回整段写法。

docx 都用 python-docx 真实构造（含 w:hyperlink），断言落在输出文件的 run 上。
"""

from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from core.language_preflight import TranslationLanguageResult
from core.word_document import extract_word_segments, write_bilingual_docx
from core.word_inline_marks import (
    MARK_PROMPT_RULE,
    PLACE_MARKS_PROMPT,
    clean_marks,
    InlineMarkEngine,
    add_boundary_spaces,
    marks_match,
    paragraph_marks,
    parse_marked,
    place_marks,
    placeholder_texts,
    split_marked_parts,
    strip_marks,
    unmark,
)
from engines.base_engine import engine_supports_chat


def _bold_label_paragraph(doc, label="注意：", body="施工前必须完成技术交底。"):
    paragraph = doc.add_paragraph()
    paragraph.add_run(label).bold = True
    paragraph.add_run(body)
    return paragraph


def _add_hyperlink(paragraph, text: str, url: str):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _numbered_heading_paragraph(doc):
    """「1.2⇥施工范围见附图。」，「施工范围」加粗。"""
    paragraph = doc.add_paragraph()
    paragraph.add_run("1.2")
    paragraph.add_run().add_tab()
    paragraph.add_run("施工范围").bold = True
    paragraph.add_run("见附图。")
    return paragraph


def _soft_break_paragraph(doc):
    """「第一行↵第二行重点」，第二行红字。"""
    paragraph = doc.add_paragraph()
    paragraph.add_run("第一行")
    paragraph.add_run().add_break()
    paragraph.add_run("第二行重点").font.color.rgb = RGBColor(0xC0, 0, 0)
    return paragraph


def _add_footnote_reference(paragraph, note_id: str = "1"):
    run = paragraph.add_run()
    reference = OxmlElement("w:footnoteReference")
    reference.set(qn("w:id"), note_id)
    run._r.append(reference)
    return run


def _sequence(paragraph):
    """段落内容按顺序：文字写成文字，制表符 ⇥、软换行 ↵、脚注编号 ¹；红字、加粗分别包在
    【红|…】【粗|…】里，便于整段比较。"""
    out = []
    for node in paragraph._p.iter(qn("w:r")):
        rpr = node.find(qn("w:rPr"))
        flag = ""
        if rpr is not None and rpr.find(qn("w:b")) is not None:
            flag = "粗"
        elif rpr is not None and rpr.find(qn("w:color")) is not None:
            flag = "红"
        for child in node:
            if child.tag == qn("w:t") and child.text:
                out.append(f"【{flag}|{child.text}】" if flag else child.text)
            elif child.tag in (qn("w:tab"), qn("w:br"), qn("w:footnoteReference")):
                out.append({qn("w:tab"): "⇥", qn("w:br"): "↵", qn("w:footnoteReference"): "¹"}[child.tag])
    return "".join(out)


def _runs(paragraph):
    """[(文字, 加粗, 斜体, 是否在链接里)]"""
    out = []
    for node in paragraph._p.iter(qn("w:r")):
        text = "".join(t.text or "" for t in node.iter(qn("w:t")))
        if not text:
            continue
        rpr = node.find(qn("w:rPr"))
        bold = rpr is not None and rpr.find(qn("w:b")) is not None
        italic = rpr is not None and rpr.find(qn("w:i")) is not None
        out.append((text, bold, italic, node.getparent().tag == qn("w:hyperlink")))
    return out


_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _add_math(paragraph, formula: str, *, display: bool = False):
    """段落里接一个 Word 公式（OMML），和 w:r 平级。"""
    from docx.oxml import parse_xml

    math = f'<m:oMath xmlns:m="{_MATH_NS}"><m:r><m:t>{formula}</m:t></m:r></m:oMath>'
    if display:
        math = f'<m:oMathPara xmlns:m="{_MATH_NS}">{math}</m:oMathPara>'
    paragraph._p.append(parse_xml(math))


def _with_math(paragraph):
    """段落内容按顺序，公式写成 {公式}。"""
    out = []
    for node in paragraph._p.iter(qn("w:t"), f"{{{_MATH_NS}}}t"):
        out.append(node.text if node.tag == qn("w:t") else "{" + node.text + "}")
    return "".join(out)


def _formula_paragraph(doc):
    """「混凝土强度按 [f=N/A] 计算。」"""
    paragraph = doc.add_paragraph()
    paragraph.add_run("混凝土强度按")
    _add_math(paragraph, "f=N/A")
    paragraph.add_run("计算。")
    return paragraph


class ParagraphMarksTests(unittest.TestCase):
    def test_bold_label_is_marked_and_base_is_the_longer_run(self) -> None:
        marks = paragraph_marks(_bold_label_paragraph(Document()))
        self.assertEqual(marks.marked_text, "<g1>注意：</g1>施工前必须完成技术交底。")
        self.assertEqual(list(marks.groups), [1])
        self.assertIsNone(marks.base_rpr.find(qn("w:b")) if marks.base_rpr is not None else None)

    def test_uniform_paragraph_needs_no_marks(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("本工程位于")
        paragraph.add_run("卡萨布兰卡。")  # 格式相同、只是拆成两个 run
        marks = paragraph_marks(paragraph)
        self.assertEqual(marks.groups, {})
        self.assertEqual(marks.marked_text, "本工程位于卡萨布兰卡。")

    def test_font_name_and_language_differences_do_not_count(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("采用")
        latin = paragraph.add_run("C30")
        latin.font.name = "Times New Roman"
        paragraph.add_run("混凝土。")
        self.assertEqual(paragraph_marks(paragraph).groups, {})

    def test_middle_word_italic_and_hyperlink(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("详见")
        paragraph.add_run("技术规范").italic = True
        paragraph.add_run("和")
        _add_hyperlink(paragraph, "官网", "https://example.com")
        paragraph.add_run("。")
        marks = paragraph_marks(paragraph)
        self.assertEqual(marks.marked_text, "详见<g1>技术规范</g1>和<g2>官网</g2>。")
        self.assertIsNotNone(marks.groups[2].hyperlink)

    def test_whitespace_between_two_bold_words_joins_the_group(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("Very").bold = True
        paragraph.add_run(" ")
        paragraph.add_run("important").bold = True
        paragraph.add_run(" notice for all workers on site.")
        self.assertEqual(paragraph_marks(paragraph).marked_text, "<g1>Very important</g1> notice for all workers on site.")

    def test_literal_tags_and_fields_are_not_marked(self) -> None:
        doc = Document()
        literal = _bold_label_paragraph(doc, body="文中出现 <g1> 字样。")
        slot_literal = _bold_label_paragraph(doc, body="文中出现 <x1/> 字样。")
        field = _bold_label_paragraph(doc)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        field.add_run()._r.append(begin)
        for paragraph in (literal, slot_literal, field):
            self.assertIsNone(paragraph_marks(paragraph))

    def test_tab_and_soft_break_become_placeholders(self) -> None:
        doc = Document()
        numbered = _numbered_heading_paragraph(doc)
        marks = paragraph_marks(numbered)
        self.assertEqual(marks.marked_text, "1.2<x1/><g1>施工范围</g1>见附图。")
        self.assertEqual(marks.placeholder_texts, {1: "\t"})
        broken = _soft_break_paragraph(doc)
        marks = paragraph_marks(broken)
        # 红字那截字多，是主体；没变色的「第一行」加标记
        self.assertEqual(marks.marked_text, "<g1>第一行</g1><x1/>第二行重点")
        self.assertEqual(marks.placeholder_texts, {1: "\n"})

    def test_placeholder_inside_a_bold_span_stays_inside_the_group(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("编号").bold = True
        paragraph.add_run().add_tab()
        paragraph.add_run("名称").bold = True
        paragraph.add_run("及规格型号说明。")
        self.assertEqual(paragraph_marks(paragraph).marked_text, "<g1>编号<x1/>名称</g1>及规格型号说明。")

    def test_adjacent_layout_nodes_share_one_placeholder(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("名称")
        run = paragraph.add_run()
        run.add_tab()
        run.add_tab()
        paragraph.add_run("数量")
        marks = paragraph_marks(paragraph)
        self.assertEqual(marks.marked_text, "名称<x1/>数量")
        self.assertEqual(marks.placeholder_texts, {1: "\t\t"})

    def test_leading_tab_and_trailing_note_reference_stay_out_of_the_text(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run().add_tab()
        paragraph.add_run("注：").bold = True
        paragraph.add_run("按规范执行。")
        _add_footnote_reference(paragraph)
        marks = paragraph_marks(paragraph)
        self.assertEqual(marks.marked_text, "<g1>注：</g1>按规范执行。")
        self.assertEqual(marks.placeholders, {})
        self.assertEqual(len(marks.lead), 1)
        self.assertEqual(len(marks.trail), 1)

    def test_uniform_paragraph_with_inline_note_reference_is_marked(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("混凝土强度")
        _add_footnote_reference(paragraph)
        paragraph.add_run("应满足设计要求。")
        marks = paragraph_marks(paragraph)
        self.assertEqual(marks.groups, {})
        self.assertTrue(marks.is_marked)
        self.assertEqual(marks.marked_text, "混凝土强度<x1/>应满足设计要求。")
        self.assertEqual(marks.placeholder_texts, {1: ""})


class ParseMarkedTests(unittest.TestCase):
    def test_reordered_tags_are_accepted(self) -> None:
        pieces = parse_marked("Voir le <g2>site</g2> et les <g1>CCTP</g1>.", [1, 2])
        self.assertEqual(pieces, [("Voir le ", None), ("site", 2), (" et les ", None), ("CCTP", 1), (".", None)])

    def test_broken_tags_are_rejected(self) -> None:
        for text in (
            "Attention : travaux.",  # 标记丢了
            "<g1>Attention</g1> : <g1>travaux</g1>.",  # 重复
            "<g1>Attention : <g2>x</g2></g1>",  # 嵌套
            "</g1>Attention<g1>",  # 先闭后开
            "<g1>Attention",  # 没闭合
            "<g3>Attention</g3>",  # 多出来的组号
        ):
            with self.subTest(text=text):
                self.assertIsNone(parse_marked(text, [1] if "g2" not in text else [1]))

    def test_tolerates_spaces_inside_tags(self) -> None:
        self.assertEqual(parse_marked("< g1 >Attention :</ g1 > texte", [1]), [("Attention :", 1), (" texte", None)])
        self.assertEqual(strip_marks("< g1 >A</ g1 >"), "A")

    def test_placeholders_must_come_back_exactly_once(self) -> None:
        pieces = parse_marked("1.2<x1/><g1>Périmètre</g1> : voir plan.", [1], [1])
        self.assertEqual(pieces[1][0].id, 1)
        self.assertEqual(parse_marked("1.2< x1 />A <g1>B</g1>", [1], [1])[1][0].id, 1)
        for text in (
            "1.2 <g1>Périmètre</g1> : voir plan.",  # 丢了
            "1.2<x1/><g1>Périmètre</g1><x1/> : voir plan.",  # 复制了
            "1.2<x2/><g1>Périmètre</g1> : voir plan.",  # 编号不对
        ):
            with self.subTest(text=text):
                self.assertIsNone(parse_marked(text, [1], [1]))
        # 原文没有占位符时，译文里冒出来的也不认
        self.assertIsNone(parse_marked("A<x1/><g1>B</g1>", [1]))

    def test_placeholders_unmark_to_the_original_characters(self) -> None:
        marked, plain = "1.2<x1/><g1>施工范围</g1>见附图<x2/>。", "1.2\t施工范围见附图。"
        self.assertEqual(placeholder_texts(marked, plain), {1: "\t", 2: ""})
        self.assertEqual(unmark("1.2<x1/><g1>Périmètre</g1> : voir plan<x2/>.", marked, plain), "1.2\tPérimètre : voir plan.")
        self.assertIsNone(placeholder_texts(marked, "别的原文"))
        self.assertEqual(strip_marks("A<x1/>B"), "AB")

    def test_placeholder_only_sources_match(self) -> None:
        self.assertTrue(marks_match("混凝土强度<x1/>应满足。", "La résistance<x1/> doit."))
        self.assertFalse(marks_match("混凝土强度<x1/>应满足。", "La résistance doit."))
        self.assertTrue(marks_match("<g1>编号</g1><x1/>名称\n单位<x2/>米", "<g1>N°</g1><x1/>Nom\nUnité<x2/>m"))
        self.assertFalse(marks_match("<g1>编号</g1><x1/>名称\n单位<x2/>米", "<g1>N°</g1>Nom\nUnité<x1/><x2/>m"))

    def test_cell_lines_must_pair_up(self) -> None:
        self.assertTrue(marks_match("<g1>注：</g1>甲\n乙", "<g1>Note :</g1> A\nB"))
        self.assertFalse(marks_match("<g1>注：</g1>甲\n乙", "<g1>Note :</g1> A B"))
        self.assertFalse(marks_match("<g1>注：</g1>甲\n乙", "Note : A\n<g1>B</g1>"))


def _bare_runner(marked_sources, *, retries: int = 3):
    """只够跑 _place_missing_inline_marks 的 WordTaskRunner。"""
    import threading
    from types import SimpleNamespace

    from core.word_task_runner import WordTaskRunner

    runner = WordTaskRunner.__new__(WordTaskRunner)
    runner._marked_sources = dict(marked_sources)
    runner._slot_hints = {}
    runner._formula_flags = {}
    runner._stop_event = threading.Event()
    runner._inline_mark_engine = None
    runner._word_batch = SimpleNamespace(strict_retry_attempts=retries)
    runner._rate = SimpleNamespace(concurrency=1)
    return runner


class _FakeEngine:
    engine_name = "fake"

    def __init__(self, reply):
        self.reply = reply
        self.calls = []

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        self.calls.append((list(texts), system_prompt))
        return {text: self.reply(text) for text in texts}

    def translate_batch_with_sources(self, texts, target_lang, system_prompt, source_lang="auto"):
        self.calls.append((list(texts), system_prompt))
        return [TranslationLanguageResult(source_text=t, translation=self.reply(t), source_lang="zh") for t in texts]

    def chat(self, system, user):
        return "ok"


class InlineMarkEngineTests(unittest.TestCase):
    MARKED = {"注意：施工前必须完成技术交底。": "<g1>注意：</g1>施工前必须完成技术交底。"}

    def test_marked_text_goes_out_and_plain_comes_back(self) -> None:
        inner = _FakeEngine(lambda t: "<g1>Attention :</g1> briefing obligatoire." if "<g1>" in t else "Autre.")
        engine = InlineMarkEngine(inner, self.MARKED)
        result = engine.translate_batch(["注意：施工前必须完成技术交底。", "其他。"], "fr", "SYS", source_lang="zh")
        self.assertEqual(result, {"注意：施工前必须完成技术交底。": "Attention : briefing obligatoire.", "其他。": "Autre."})
        sent, prompt = inner.calls[0]
        self.assertEqual(sent[0], "<g1>注意：</g1>施工前必须完成技术交底。")
        self.assertIn(MARK_PROMPT_RULE, prompt)
        self.assertEqual(
            engine.marked_translations["注意：施工前必须完成技术交底。"],
            ("<g1>注意：</g1>施工前必须完成技术交底。", "<g1>Attention :</g1> briefing obligatoire."),
        )

    def test_lost_tags_fall_back_and_forget_earlier_result(self) -> None:
        replies = iter(["<g1>Attention :</g1> x.", "Attention : y."])
        inner = _FakeEngine(lambda t: next(replies))
        engine = InlineMarkEngine(inner, self.MARKED)
        engine.translate_batch(["注意：施工前必须完成技术交底。"], "fr", "SYS")
        result = engine.translate_batch(["注意：施工前必须完成技术交底。"], "fr", "SYS")
        self.assertEqual(result["注意：施工前必须完成技术交底。"], "Attention : y.")
        self.assertNotIn("注意：施工前必须完成技术交底。", engine.marked_translations)

    def test_batches_without_marks_pass_through_untouched(self) -> None:
        inner = _FakeEngine(lambda t: "x")
        InlineMarkEngine(inner, self.MARKED).translate_batch(["其他。"], "fr", "SYS")
        self.assertEqual(inner.calls[0], (["其他。"], "SYS"))

    def test_source_language_protocol_and_capabilities(self) -> None:
        inner = _FakeEngine(lambda t: "<g1>Attention :</g1> x." if "<g1>" in t else "y")
        engine = InlineMarkEngine(inner, self.MARKED)
        results = engine.translate_batch_with_sources(["注意：施工前必须完成技术交底。"], "fr", "SYS")
        self.assertEqual(results[0].source_text, "注意：施工前必须完成技术交底。")
        self.assertEqual(results[0].translation, "Attention : x.")
        self.assertEqual(engine.engine_name, "fake")
        self.assertTrue(engine_supports_chat(engine) is engine_supports_chat(inner))


    def test_missing_space_after_tagged_colon_is_restored(self) -> None:
        inner = _FakeEngine(lambda t: "<g1>Attention :</g1>briefing obligatoire.")
        engine = InlineMarkEngine(inner, self.MARKED)
        result = engine.translate_batch(["注意：施工前必须完成技术交底。"], "fr", "SYS")
        self.assertEqual(result["注意：施工前必须完成技术交底。"], "Attention : briefing obligatoire.")
        self.assertEqual(
            engine.marked_translations["注意：施工前必须完成技术交底。"][1], "<g1>Attention :</g1> briefing obligatoire."
        )


class BoundarySpaceTests(unittest.TestCase):
    def test_space_goes_outside_the_formatting(self) -> None:
        self.assertEqual(add_boundary_spaces("<g1>Note :</g1>les cotes.", "fr"), "<g1>Note :</g1> les cotes.")
        self.assertEqual(add_boundary_spaces("Note :<g1>les cotes</g1>.", "fr"), "Note : <g1>les cotes</g1>.")
        self.assertEqual(add_boundary_spaces("fini.<g1>Attention</g1> !", "fr"), "fini. <g1>Attention</g1> !")

    def test_leaves_correct_or_intended_joins_alone(self) -> None:
        for text in (
            "<g1>Note :</g1> les cotes.",
            "120 m<g1>2</g1>.",
            "<g1>pré</g1>fabriqué",
            "www.example.<g1>com</g1>",
            "3,<g1>5</g1> m",
            "Note :les cotes.",
        ):
            self.assertEqual(add_boundary_spaces(text, "en"), text)

    def test_space_after_a_note_reference_between_two_words(self) -> None:
        texts = {1: "", 2: "\t"}
        self.assertEqual(add_boundary_spaces("béton<x1/>doit", "fr", texts), "béton<x1/> doit")
        self.assertEqual(add_boundary_spaces("béton<x1/><g1>doit</g1>", "fr", texts), "béton<x1/> <g1>doit</g1>")
        self.assertEqual(add_boundary_spaces("fini.<x1/>Le", "fr", texts), "fini.<x1/> Le")
        for text in ("béton<x1/> doit", "béton <x1/>, doit", "1.2<x2/>Portée"):
            self.assertEqual(add_boundary_spaces(text, "fr", texts), text)
        self.assertEqual(add_boundary_spaces("混凝土<x1/>强度", "zh", texts), "混凝土<x1/>强度")

    def test_languages_written_without_spaces_are_untouched(self) -> None:
        for lang in ("zh", "zh-CN", "ja", "ko"):
            self.assertEqual(add_boundary_spaces("<g1>注:</g1>尺寸", lang), "<g1>注:</g1>尺寸")


class _ChatEngine(_FakeEngine):
    def __init__(self, answers):
        super().__init__(lambda t: t)
        self.answers = answers
        self.prompts = []

    def chat(self, system, user):
        self.prompts.append((system, user))
        return self.answers


class PlaceMarksTests(unittest.TestCase):
    ITEMS = [
        ("注意：施工前完成交底。", "<g1>注意：</g1>施工前完成交底。", "Attention : briefing avant travaux."),
        ("严禁堆放。", "<g1>严禁</g1>堆放。", "Stockage interdit."),
        ("注：单位毫米。", "<g1>注：</g1>单位毫米。", "Note : unités en mm."),
    ]

    def test_only_answers_that_keep_the_final_text_are_used(self) -> None:
        engine = _ChatEngine(
            '["<g1>Attention :</g1> briefing avant travaux.", '
            '"Stockage <g1>strictement</g1> interdit.", '  # 改了译文
            '"Note : unités en mm."]'  # 没插标记
        )
        placed = place_marks(engine, self.ITEMS)
        self.assertEqual(placed, {"注意：施工前完成交底。": ("<g1>注意：</g1>施工前完成交底。", "<g1>Attention :</g1> briefing avant travaux.")})
        self.assertIn("原文", engine.prompts[0][1])

    def test_tabs_are_seeded_without_asking_the_model(self) -> None:
        engine = _ChatEngine("[]")
        items = [("编号\t名称", "编号<x1/>名称", "N°\tDésignation")]
        self.assertEqual(place_marks(engine, items), {"编号\t名称": ("编号<x1/>名称", "N°<x1/>Désignation")})
        self.assertEqual(engine.prompts, [])

    def test_model_places_note_reference_and_keeps_seeded_tab(self) -> None:
        engine = _ChatEngine('["1.2<x1/><g1>Périmètre</g1><x2/> : voir plan."]')
        items = [("1.2\t施工范围见附图。", "1.2<x1/><g1>施工范围</g1><x2/>见附图。", "1.2\tPérimètre : voir plan.")]
        placed = place_marks(engine, items)
        self.assertEqual(placed["1.2\t施工范围见附图。"][1], "1.2<x1/><g1>Périmètre</g1><x2/> : voir plan.")
        self.assertIn("1.2<x1/>Périmètre : voir plan.", engine.prompts[0][1])

    def test_a_tab_the_model_moves_stays_where_the_final_text_has_it(self) -> None:
        """只取这次要插的标记的位置，已经放好的（按定稿找到的制表符）不跟着模型挪。"""
        engine = _ChatEngine('["1.2 <g1>Périmètre</g1><x1/><x2/>: voir plan."]')
        items = [("1.2\t施工范围见附图。", "1.2<x1/><g1>施工范围</g1><x2/>见附图。", "1.2\tPérimètre : voir plan.")]
        self.assertEqual(place_marks(engine, items)["1.2\t施工范围见附图。"][1], "1.2<x1/><g1>Périmètre</g1><x2/> : voir plan.")

    def test_placeholder_may_replace_the_space_a_stored_translation_has_for_a_tab(self) -> None:
        """记忆库里的定稿把制表符存成了空格：占位符顶替这个空格也算没改译文。"""
        engine = _ChatEngine('["1.2<x1/><g1>Portée des travaux</g1> : voir le dessin annexe."]')
        items = [("1.2\t施工范围见附图。", "1.2<x1/><g1>施工范围</g1>见附图。",
                  "1.2 Portée des travaux : voir le dessin annexe.")]
        placed = place_marks(engine, items)
        self.assertEqual(placed["1.2\t施工范围见附图。"][1], "1.2<x1/><g1>Portée des travaux</g1> : voir le dessin annexe.")
        # 模型顺手改了别处的字也不要紧：只取标记的位置，文字仍是定稿
        engine = _ChatEngine('["1.2<x1/><g1>Portée des travaux</g1> : voir le dessin."]')
        self.assertEqual(place_marks(engine, items)["1.2\t施工范围见附图。"][1],
                         "1.2<x1/><g1>Portée des travaux</g1> : voir le dessin annexe.")

    def test_note_reference_that_took_the_space_between_two_words_still_lands_in_the_sentence(self) -> None:
        """模型把脚注编号插在空格上（「béton<x1/>doit」「béton <x1/>doit」）：编号留在句中，后面补回空格。"""
        source, marked = "混凝土强度应满足设计要求。", "混凝土<x1/>强度应满足设计要求。"
        translation = "La résistance du béton doit satisfaire aux exigences de conception."
        for answer in ("La résistance du béton<x1/>doit satisfaire aux exigences de conception.",
                       "La résistance du béton <x1/>doit satisfaire aux exigences de conception.",
                       "La résistance du béton<x1/> doit satisfaire aux exigences de conception."):
            placed = place_marks(_ChatEngine(json.dumps([answer], ensure_ascii=False)), [(source, marked, translation)])
            self.assertEqual(placed[source][1], "La résistance du béton<x1/> doit satisfaire aux exigences de conception.")
        # 定稿那里本来没有空格的不补，补了就和定稿对不上
        placed = place_marks(_ChatEngine('["béton<x1/>armé"]'), [("钢筋混凝土", "钢筋<x1/>混凝土", "bétonarmé")])
        self.assertEqual(placed["钢筋混凝土"][1], "béton<x1/>armé")
        # 制表符的占位符有字符，不按这条补空格
        placed = place_marks(_ChatEngine('["No<x1/>Désignation <g1>x</g1>"]'),
                             [("编号\t名称甲", "编号<x1/>名称<g1>甲</g1>", "No Désignation x")])
        self.assertEqual(placed["编号\t名称甲"][1], "No<x1/>Désignation <g1>x</g1>")

    def test_picture_with_a_stray_space_before_punctuation_keeps_its_place_and_the_bold(self) -> None:
        """模型在图片占位符前后多写了空格（「figure <x1/>, 」）：图片和加粗照样按回答写回。"""
        source, marked = "如图所示，严禁堆放。", "如图<x1/>所示，<g1>严禁</g1>堆放。"
        translation = "Comme indiqué sur la figure, strictement interdit de stocker."
        for answer in ("Comme indiqué sur la figure <x1/>, <g1>strictement interdit</g1> de stocker.",
                       "Comme indiqué sur la figure<x1/> , <g1>strictement interdit</g1> de stocker.",
                       "Comme indiqué sur la figure <x1/> , <g1>strictement interdit</g1> de stocker."):
            placed = place_marks(_ChatEngine(json.dumps([answer], ensure_ascii=False)), [(source, marked, translation)])
            self.assertEqual(placed[source][1], "Comme indiqué sur la figure<x1/>, <g1>strictement interdit</g1> de stocker.")
        # 定稿冒号前有空格的照留：「Note<x1/> : 」
        placed = place_marks(_ChatEngine('["<g1>Note</g1> <x1/>: unités en mm."]'),
                             [("注图：单位毫米。", "<g1>注</g1>图<x1/>：单位毫米。", "Note : unités en mm.")])
        self.assertEqual(placed["注图：单位毫米。"][1], "<g1>Note</g1><x1/> : unités en mm.")
        # 模型多写了几个字：标记位置照样换算到定稿上
        placed = place_marks(_ChatEngine('["Comme indiqué sur la figure <x1/>, il est <g1>strictement interdit</g1> de stocker."]'),
                             [(source, marked, translation)])
        self.assertEqual(placed[source][1], "Comme indiqué sur la figure<x1/>, <g1>strictement interdit</g1> de stocker.")
        # 改得太多（和定稿差得远）的不收
        placed = place_marks(_ChatEngine('["Il est <g1>interdit</g1> de stocker quoi que ce soit ici <x1/>."]'),
                             [(source, marked, translation)])
        self.assertEqual(placed, {})

    def test_bad_reply_and_engines_without_chat_place_nothing(self) -> None:
        self.assertEqual(place_marks(_ChatEngine("not json"), self.ITEMS), {})

        class NoChat:
            engine_name = "x"

        self.assertEqual(place_marks(NoChat(), self.ITEMS), {})


class _GroupEngine(_FakeEngine):
    """补标记时按「插入」里列出的标记，从 answers（标记 → 插好这一个标记的定稿译文）里取位置
    插进这次的译文；refuse 里的标记每次都不放（回原样的译文）。"""

    def __init__(self, answers, refuse=()):
        super().__init__(lambda t: t)
        self.answers = answers
        self.refuse = set(refuse)
        self.requests = []

    def chat(self, system, user):
        from core.word_inline_marks import _insert_mark, _mark_tokens

        assert system == PLACE_MARKS_PROMPT
        items = json.loads(user)
        self.requests.append(items)
        out = []
        for item in items:
            text = item["译文"]
            for mark in item["插入"].split("、"):
                if mark in self.refuse:
                    continue
                for name, position, raw in _mark_tokens(self.answers[mark]):
                    text = _insert_mark(text, position, raw, name, {})
            out.append(text)
        return json.dumps(out, ensure_ascii=False)


class FormulaPlacementTests(unittest.TestCase):
    """反馈 #19 新方案（maple 10-24）：整句带标记翻一次，标记没回齐时译文不动，缺的标记按分句
    分组、每组几个，让模型只说位置；一组放不进去按「单段重试次数」重试（第 1 次原样、第 2 次
    拆两半、之后再问），试满仍不行的公式放到所在那一句的句末，写回时标黄加批注。"""

    MARKED = ("where <x1/> is the stress, <x2/> is the strain, <x3/> is the modulus. "
              "Then <g1>we</g1> have <x4/> and <x5/>, so <x6/>.")
    PLAIN = strip_marks(MARKED)
    FINAL = "其中为应力，为应变，为模量。于是我们有和，所以。"
    ANSWERS = {
        "<x1/>": "其中<x1/>为应力，为应变，为模量。于是我们有和，所以。",
        "<x2/>": "其中为应力，<x2/>为应变，为模量。于是我们有和，所以。",
        "<x3/>": "其中为应力，为应变，<x3/>为模量。于是我们有和，所以。",
        "<g1>…</g1>": "其中为应力，为应变，为模量。于是<g1>我们</g1>有和，所以。",
        "<x4/>": "其中为应力，为应变，为模量。于是我们有<x4/>和，所以。",
        "<x5/>": "其中为应力，为应变，为模量。于是我们有和<x5/>，所以。",
        "<x6/>": "其中为应力，为应变，为模量。于是我们有和，所以<x6/>。",
    }
    PLACED = "其中<x1/>为应力，<x2/>为应变，<x3/>为模量。于是<g1>我们</g1>有<x4/>和<x5/>，所以<x6/>。"

    def test_garbled_marks_are_cleaned_and_lost_ones_are_not_guessed(self) -> None:
        self.assertEqual(clean_marks(self.MARKED, "其中<x1>为应力，<x2/><x2/>为应变，<x9/>为模量。于是<g1>我们有。"),
                         "其中<x1/>为应力，<x2/>为应变，为模量。于是我们有。")
        self.assertEqual(clean_marks("a<x1/>b\nc", "a<x1/>b"), "ab")  # 行数对不上：全部去掉

    def test_lost_marks_are_not_translated_again(self) -> None:
        inner = _FakeEngine(lambda text: "其中<x1/>为应力，为应变，<x3/>为模量。于是我们有和，所以。")
        engine = InlineMarkEngine(inner, {self.PLAIN: self.MARKED})
        result = engine.translate_batch([self.PLAIN], "zh", "SYS")
        self.assertEqual(result[self.PLAIN], self.FINAL)
        self.assertEqual(len(inner.calls), 1)
        self.assertEqual(engine.partial_translations[self.PLAIN][1],
                         "其中<x1/>为应力，为应变，<x3/>为模量。于是我们有和，所以。")
        self.assertNotIn(self.PLAIN, engine.marked_translations)

    def test_missing_marks_are_grouped_and_packed(self) -> None:
        engine = _GroupEngine(self.ANSWERS)
        stats: dict = {}
        placed = place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], group_size=3, request_marks=3,
                             stats=stats)
        self.assertEqual(placed[self.PLAIN][1], self.PLACED)
        # 每组最多 3 个，在分句处断开；一次请求只放一组
        self.assertEqual([[item["插入"] for item in request] for request in engine.requests],
                         [["<x1/>、<x2/>、<x3/>"], ["<g1>…</g1>、<x4/>、<x5/>"], ["<x6/>"]])
        engine = _GroupEngine(self.ANSWERS)
        place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], group_size=3, request_marks=20)
        self.assertEqual(len(engine.requests), 1)
        self.assertEqual(len(engine.requests[0]), 3)

    def test_marks_that_came_back_are_kept_and_only_the_rest_asked(self) -> None:
        engine = _GroupEngine(self.ANSWERS)
        seed = "其中<x1/>为应力，<x2/>为应变，<x3/>为模量。于是<g1>我们</g1>有<x4/>和<x5/>，所以。"
        placed = place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], seeds={self.PLAIN: seed})
        self.assertEqual(placed[self.PLAIN][1], self.PLACED)
        self.assertEqual([item["插入"] for item in engine.requests[0]], ["<x6/>"])

    def test_formula_text_is_shown_to_the_model(self) -> None:
        engine = _GroupEngine(self.ANSWERS)
        place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], hints={self.PLAIN: {1: "公式 σ=M/W"}})
        self.assertEqual(engine.requests[0][0]["说明"], {"<x1/>": "公式 σ=M/W"})

    def test_retry_then_split_then_retry_then_settle_at_the_sentence_end(self) -> None:
        engine = _GroupEngine(self.ANSWERS, refuse={"<x2/>"})
        flags: dict = {}
        stats: dict = {}
        placed = place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], group_size=3, request_marks=20,
                             retries=3, flags=flags, stats=stats)
        asked = [[item["插入"] for item in request] for request in engine.requests]
        self.assertEqual(asked, [
            ["<x1/>、<x2/>、<x3/>", "<g1>…</g1>、<x4/>、<x5/>", "<x6/>"],
            ["<x2/>"],  # 放好的不再问；第 1 次重试：原样
            ["<x2/>"],  # 第 2 次：只剩一个，拆不开
            ["<x2/>"],  # 第 3 次
        ])
        # 其余都放好了；x2 放在它所在那一句的句末（句号前），这一句标出来
        self.assertEqual(placed[self.PLAIN][1],
                         "其中<x1/>为应力，为应变，<x3/>为模量<x2/>。于是<g1>我们</g1>有<x4/>和<x5/>，所以<x6/>。")
        self.assertEqual(flags, {self.PLAIN: [(0, 0, 14)]})
        self.assertEqual(stats["settled"], 1)
        # 不给 flags（Excel、扩展）时放不全的不算补上
        engine = _GroupEngine(self.ANSWERS, refuse={"<x2/>"})
        self.assertEqual(place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], retries=1), {})

    def test_a_group_the_model_cannot_do_at_once_is_split(self) -> None:
        class Picky(_GroupEngine):
            def chat(self, system, user):
                items = json.loads(user)
                if any("、" in item["插入"] for item in items):
                    self.requests.append(items)
                    return json.dumps([item["译文"] for item in items], ensure_ascii=False)
                return super().chat(system, user)

        engine = Picky(self.ANSWERS)
        marked, plain = "where <x1/> is the stress, <x2/> is the strain.", "where  is the stress,  is the strain."
        placed = place_marks(engine, [(plain, marked, "其中为应力，为应变。")], retries=3, flags={})
        self.assertEqual([[item["插入"] for item in request] for request in engine.requests],
                         [["<x1/>、<x2/>"], ["<x1/>、<x2/>"], ["<x1/>", "<x2/>"]])  # 第 2 次重试拆成两组
        self.assertEqual(placed[plain][1], "其中<x1/>为应力，<x2/>为应变。")

    def test_format_marks_that_never_fit_are_dropped(self) -> None:
        engine = _GroupEngine(self.ANSWERS, refuse={"<g1>…</g1>"})
        flags: dict = {}
        placed = place_marks(engine, [(self.PLAIN, self.MARKED, self.FINAL)], retries=0, flags=flags)
        self.assertEqual(placed[self.PLAIN][1], self.PLACED.replace("<g1>", "").replace("</g1>", ""))
        self.assertEqual(flags, {})

    def test_math_is_written_as_linear_text(self) -> None:
        from docx.oxml import parse_xml

        from core.word_inline_marks import _math_linear

        m = f'xmlns:m="{_MATH_NS}"'
        node = parse_xml(
            f'<m:oMath {m}><m:r><m:t>σ=</m:t></m:r><m:f><m:num><m:r><m:t>M</m:t></m:r></m:num>'
            f'<m:den><m:r><m:t>W</m:t></m:r></m:den></m:f><m:r><m:t>+</m:t></m:r>'
            f'<m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath>'
        )
        self.assertEqual(_math_linear(node), "σ=M/W+x^2")

    def test_runner_flags_the_sentence_and_the_writer_marks_it(self) -> None:
        source = "混凝土强度按计算，取值见表。"
        marked = "混凝土强度按<x1/>计算，取值见表。"
        final = "La résistance du béton est calculée, voir le tableau."
        runner = _bare_runner({source: marked}, retries=1)
        engine = _GroupEngine({}, refuse={"<x1/>"})
        runner._place_missing_inline_marks(engine, {source: final}, target_lang="fr")
        entry = runner._inline_mark_engine.marked_translations[source]
        self.assertEqual(entry[1], "La résistance du béton est calculée, voir le tableau<x1/>.")
        self.assertEqual(runner._formula_flags, {source: [(0, 0, len(final))]})
        doc = Document()
        _formula_paragraph(doc).runs[1].text = "计算，取值见表。"
        issues = []
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "src.docx"
            doc.save(path)
            out = write_bilingual_docx(
                source_path=path, output_dir=Path(temporary) / "out", translations={source: final},
                target_lang="fr", source_lang="zh", output_translation_only=True,
                marked_translations={source: entry}, formula_flags=runner._formula_flags,
                issue_callback=issues.append,
            )
            written = Document(str(out))
            paragraph = written.paragraphs[0]
            self.assertEqual(_with_math(paragraph), "La résistance du béton est calculée, voir le tableau{f=N/A}.")
            highlights = [run.font.highlight_color for run in paragraph.runs if run.text]
            self.assertTrue(highlights and all(color == 7 for color in highlights), highlights)  # 7 = YELLOW
            comments = list(written.comments)
            self.assertEqual([comment.text for comment in comments], ["公式位置待核对"])
        self.assertEqual(issues, [{"kind": "formula_position", "location": "body.paragraph[0]", "source": source}])

    def test_a_soft_line_break_does_not_split_the_paragraph_into_two_lines(self) -> None:
        """本机实测：段里有软换行（占位内容就是换行）时，定稿译文按换行切成两行，和原文一行
        对不上，整段的公式补不上。"""
        marked = "第一行<x1/>第二行按<x2/>计算。"
        source = "第一行\n第二行按计算。"
        final = "Line one\nline two uses the formula."
        engine = _GroupEngine({"<x2/>": "Line oneline two uses the formula<x2/>."})
        placed = place_marks(engine, [(source, marked, final)], flags={})
        self.assertEqual(placed[source], (marked, "Line one<x1/>line two uses the formula<x2/>."))

    def test_runner_groups_by_the_formula_batch_setting(self) -> None:
        runner = _bare_runner({self.PLAIN: self.MARKED})
        runner._rate.formula_batch = 2
        engine = _GroupEngine(self.ANSWERS)
        runner._place_missing_inline_marks(engine, {self.PLAIN: self.FINAL}, target_lang="zh")
        sizes = [len(item["插入"].split("、")) for request in engine.requests for item in request]
        self.assertTrue(sizes and max(sizes) <= 2, sizes)
        self.assertTrue(all(len(request) == 1 for request in engine.requests))  # 一组一次请求，不打包
        self.assertEqual(sum(sizes), 7)
        self.assertEqual(runner._inline_mark_engine.marked_translations[self.PLAIN][1],
                         "其中<x1/>为应力，<x2/>为应变，<x3/>为模量。于是<g1>我们</g1>有<x4/>和<x5/>，所以<x6/>。")


class LongParagraphPartsTests(unittest.TestCase):
    """反馈 #19 第三轮本机实测：超过 800 字的段落 word_batching 按句子切块发，块不在带标记
    原文表里，原样发出、没有占位符，四段公式最多的（24、13、9、6 个）全挤到段末。"""

    MARKED = "where <x1/> is the stress. Then <g1>we have <x2/> and. More</g1> <x3/> is x."

    def test_marked_source_is_cut_where_the_plain_text_was_cut(self) -> None:
        plain = strip_marks(self.MARKED)
        parts = ["where  is", "the stress.", "Then we have  and.", "More  is x."]
        self.assertEqual(
            split_marked_parts(self.MARKED, plain, parts),
            ["where <x1/> is", "the stress.", "Then <g1>we have <x2/> and.</g1>", "<g1>More</g1> <x3/> is x."],
        )
        self.assertIsNone(split_marked_parts(self.MARKED, plain, ["where  is the stress.", "nope"]))

    def test_placeholder_at_a_cut_goes_with_its_own_side(self) -> None:
        marked = "公式如下<x1/>。<x2/>为应力。"
        plain = strip_marks(marked)
        self.assertEqual(split_marked_parts(marked, plain, ["公式如下。", "为应力。"]), ["公式如下<x1/>。", "<x2/>为应力。"])
        marked = "Value is <x1/>. <x2/> is next."
        self.assertEqual(split_marked_parts(marked, strip_marks(marked), ["Value is .", "is next."]),
                         ["Value is <x1/>.", "<x2/> is next."])

    def test_parts_go_out_marked_and_are_joined_back(self) -> None:
        from core.word_batching import _join_translated_parts

        plain = strip_marks(self.MARKED)
        answers = {
            "where <x1/> is the stress.": "其中<x1/>为应力。",
            "Then <g1>we have <x2/> and. More</g1> <x3/> is x.": "<g1>于是有<x2/>。还有</g1><x3/>为x。",
        }
        inner = _FakeEngine(lambda text: answers.get(text, text))
        engine = InlineMarkEngine(inner, {plain: self.MARKED})
        engine.split_long_sources(lambda text: ["where  is the stress.", "Then we have  and. More  is x."])
        pieces = [piece for piece, _marked in engine.parts[plain]]
        self.assertEqual(pieces, ["where  is the stress.", "Then we have  and. More  is x."])
        engine.translate_batch(pieces, "zh", "SYS")
        self.assertEqual(inner.calls[0][0], list(answers))
        marked, joined = engine.assemble_parts(plain, lambda parts: _join_translated_parts(parts, target_lang="zh"))
        self.assertEqual((marked, joined), (self.MARKED, "其中<x1/>为应力。<g1>于是有<x2/>。还有</g1><x3/>为x。"))

    def test_format_tag_split_across_parts_is_merged_at_the_seam(self) -> None:
        from core.word_batching import _join_translated_parts

        plain = strip_marks(self.MARKED)
        engine = InlineMarkEngine(_FakeEngine(lambda text: {
            "where <x1/> is the stress. Then <g1>we have <x2/> and.</g1>": "Où <x1/> est la contrainte. Puis <g1>on a <x2/> et.</g1>",
            "<g1>More</g1> <x3/> is x.": "<g1>Encore</g1> <x3/> vaut x.",
        }.get(text, text)), {plain: self.MARKED})
        engine.split_long_sources(lambda text: ["where  is the stress. Then we have  and.", "More  is x."])
        engine.translate_batch([piece for piece, _marked in engine.parts[plain]], "fr", "SYS")
        _marked, joined = engine.assemble_parts(plain, lambda parts: _join_translated_parts(parts, target_lang="fr"))
        self.assertEqual(joined, "Où <x1/> est la contrainte. Puis <g1>on a <x2/> et. Encore</g1> <x3/> vaut x.")
        self.assertTrue(marks_match(self.MARKED, joined))

    def test_a_part_that_lost_its_marks_gets_them_back_on_its_own(self) -> None:
        """本机第四轮：24 个公式那段有一块的标记丢了，整段拼不回，整段二十几个公式一起补又
        补不上。现在只把缺的那一块交给模型补，补完再拼回。"""

        from core.word_batching import _join_translated_parts

        plain = strip_marks(self.MARKED)
        part2 = "Then <g1>we have <x2/> and. More</g1> <x3/> is x."
        asked = []

        class Engine(_FakeEngine):
            def chat(self, system, user):
                asked.append(json.loads(user))
                return json.dumps(["<g1>于是有<x2/>。还有</g1><x3/>为x。"], ensure_ascii=False)

        inner = Engine(lambda text: {"where <x1/> is the stress.": "其中<x1/>为应力。", part2: "于是有。还有为x。"}[text])
        runner = _bare_runner({plain: self.MARKED})
        runner._inline_mark_engine = engine = InlineMarkEngine(inner, runner._marked_sources)
        engine.split_long_sources(lambda text: ["where  is the stress.", "Then we have  and. More  is x."])
        out = engine.translate_batch([piece for piece, _marked in engine.parts[plain]], "zh", "SYS")
        final = _join_translated_parts(list(out.values()), target_lang="zh")
        self.assertIsNone(engine.assemble_parts(plain, lambda parts: _join_translated_parts(parts, target_lang="zh")))
        runner._place_missing_inline_marks(inner, {plain: final}, target_lang="zh")
        self.assertEqual([item["原文"] for item in asked[0]], [part2])
        self.assertEqual(engine.marked_translations[plain],
                         (self.MARKED, "其中<x1/>为应力。<g1>于是有<x2/>。还有</g1><x3/>为x。"))


class WriteBackTests(unittest.TestCase):
    SOURCE = "注意：施工前必须完成技术交底。"
    MARKED_SOURCE = "<g1>注意：</g1>施工前必须完成技术交底。"
    MARKED_FR = "<g1>Attention :</g1> le briefing technique est obligatoire avant les travaux."
    PLAIN_FR = "Attention : le briefing technique est obligatoire avant les travaux."

    def _write(self, doc, translations, marked, *, translation_only: bool):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        source = root / "src.docx"
        doc.save(source)
        out = write_bilingual_docx(
            source_path=source,
            output_dir=root / "out",
            translations=translations,
            target_lang="fr",
            source_lang="zh",
            output_translation_only=translation_only,
            marked_translations=marked,
        )
        return Document(str(out)), source

    def test_extraction_carries_marked_source(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "src.docx"
            doc = Document()
            _bold_label_paragraph(doc)
            doc.add_paragraph("普通段落。")
            doc.save(source)
            segments = {s.source: s.marked_source for s in extract_word_segments(source, target_lang="fr")}
        self.assertEqual(segments, {self.SOURCE: self.MARKED_SOURCE, "普通段落。": ""})

    def test_translation_only_keeps_bold_label(self) -> None:
        doc = Document()
        _bold_label_paragraph(doc)
        out, _ = self._write(
            doc, {self.SOURCE: self.PLAIN_FR}, {self.SOURCE: (self.MARKED_SOURCE, self.MARKED_FR)}, translation_only=True
        )
        self.assertEqual(
            _runs(out.paragraphs[0]),
            [("Attention :", True, False, False), (" le briefing technique est obligatoire avant les travaux.", False, False, False)],
        )
        self.assertEqual(out.paragraphs[0].text, self.PLAIN_FR)

    def test_bilingual_keeps_bold_label_in_the_translation_line(self) -> None:
        doc = Document()
        _bold_label_paragraph(doc)
        out, _ = self._write(
            doc, {self.SOURCE: self.PLAIN_FR}, {self.SOURCE: (self.MARKED_SOURCE, self.MARKED_FR)}, translation_only=False
        )
        self.assertEqual(out.paragraphs[0].text, self.SOURCE)
        self.assertEqual([r[:2] for r in _runs(out.paragraphs[1])],
                         [("Attention :", True), (" le briefing technique est obligatoire avant les travaux.", False)])

    def test_hyperlink_survives_and_points_to_the_same_target(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("详见")
        _add_hyperlink(paragraph, "官网", "https://example.com")
        paragraph.add_run("。")
        source, marked = "详见官网。", "详见<g1>官网</g1>。"
        out, _ = self._write(
            doc, {source: "Voir le site officiel."}, {source: (marked, "Voir le <g1>site officiel</g1>.")},
            translation_only=True,
        )
        self.assertEqual(_runs(out.paragraphs[0])[1], ("site officiel", False, False, True))
        link = out.paragraphs[0]._p.find(qn("w:hyperlink"))
        self.assertEqual(out.part.rels[link.get(qn("r:id"))].target_ref, "https://example.com")

    def test_falls_back_when_final_text_was_changed_later(self) -> None:
        """复查或修数字改过译文：标记版本对不上最终译文，按整段写法写，格式取主体那一截。"""
        doc = Document()
        _bold_label_paragraph(doc)
        changed = "Attention : briefing obligatoire (version corrigée)."
        out, _ = self._write(
            doc, {self.SOURCE: changed}, {self.SOURCE: (self.MARKED_SOURCE, self.MARKED_FR)}, translation_only=True
        )
        self.assertEqual(_runs(out.paragraphs[0]), [(changed, False, False, False)])

    def test_falls_back_when_marked_source_differs(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("注意：施工前").bold = True  # 同样的文字，加粗范围不同
        paragraph.add_run("必须完成技术交底。")
        out, _ = self._write(
            doc, {self.SOURCE: self.PLAIN_FR}, {self.SOURCE: (self.MARKED_SOURCE, self.MARKED_FR)}, translation_only=True
        )
        self.assertEqual(len(_runs(out.paragraphs[0])), 1)

    def _write_case(self, doc, source, marked_source, marked_fr, *, translation_only):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "p.docx"
            doc.save(path)
            segments = {s.source: s.marked_source for s in extract_word_segments(path, target_lang="fr")}
        self.assertEqual(segments[source], marked_source)
        plain = unmark(marked_fr, marked_source, source)
        out, _ = self._write(doc, {source: plain}, {source: (marked_source, marked_fr)}, translation_only=translation_only)
        return out, plain

    def test_numbered_heading_keeps_tab_and_bold(self) -> None:
        for translation_only, index in ((True, 0), (False, 1)):
            with self.subTest(translation_only=translation_only):
                doc = Document()
                _numbered_heading_paragraph(doc)
                out, plain = self._write_case(
                    doc, "1.2\t施工范围见附图。", "1.2<x1/><g1>施工范围</g1>见附图。",
                    "1.2<x1/><g1>Périmètre des travaux</g1> : voir le plan joint.", translation_only=translation_only,
                )
                self.assertEqual(plain, "1.2\tPérimètre des travaux : voir le plan joint.")
                self.assertEqual(_sequence(out.paragraphs[index]), "1.2⇥【粗|Périmètre des travaux】 : voir le plan joint.")

    def test_soft_break_keeps_red_on_the_second_line_only(self) -> None:
        for translation_only, index in ((True, 0), (False, 1)):
            with self.subTest(translation_only=translation_only):
                doc = Document()
                _soft_break_paragraph(doc)
                out, _plain = self._write_case(
                    doc, "第一行\n第二行重点", "<g1>第一行</g1><x1/>第二行重点",
                    "<g1>Première ligne</g1><x1/>Deuxième ligne importante", translation_only=translation_only,
                )
                self.assertEqual(_sequence(out.paragraphs[index]), "Première ligne↵【红|Deuxième ligne importante】")

    def test_note_reference_moves_with_the_translation(self) -> None:
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("混凝土强度")
        _add_footnote_reference(paragraph)
        paragraph.add_run("应满足设计要求。")
        marked_fr = "La résistance du béton<x1/> doit satisfaire aux exigences."
        out, _ = self._write_case(
            doc, "混凝土强度应满足设计要求。", "混凝土强度<x1/>应满足设计要求。", marked_fr, translation_only=True
        )
        self.assertEqual(_sequence(out.paragraphs[0]), "La résistance du béton¹ doit satisfaire aux exigences.")

        # 双语：脚注编号只在原文段里留一份，译文段不复制
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("混凝土强度")
        _add_footnote_reference(paragraph)
        paragraph.add_run("应满足设计要求。")
        out, _ = self._write_case(
            doc, "混凝土强度应满足设计要求。", "混凝土强度<x1/>应满足设计要求。", marked_fr, translation_only=False
        )
        self.assertEqual(_sequence(out.paragraphs[0]), "混凝土强度¹应满足设计要求。")
        self.assertEqual(_sequence(out.paragraphs[1]), "La résistance du béton doit satisfaire aux exigences.")

    def test_formula_stays_in_the_sentence(self) -> None:
        """反馈 #18：段落里的 Word 公式读不进纯文本，整段写法的译文里公式没了（双语）或挪到
        段尾（仅译文）。现在公式作为占位符随原文发出，译文里放回原处，双语的译文段也有一份。"""
        source, marked = "混凝土强度按计算。", "混凝土强度按<x1/>计算。"
        marked_fr = "La résistance du béton est calculée selon <x1/>."
        for translation_only, expected in ((True, ["La résistance du béton est calculée selon {f=N/A}."]),
                                           (False, ["混凝土强度按{f=N/A}计算。",
                                                    "La résistance du béton est calculée selon {f=N/A}."])):
            with self.subTest(translation_only=translation_only):
                doc = Document()
                _formula_paragraph(doc)
                out, plain = self._write_case(doc, source, marked, marked_fr, translation_only=translation_only)
                self.assertEqual(plain, "La résistance du béton est calculée selon .")
                self.assertEqual([_with_math(item) for item in out.paragraphs], expected)

    def test_formula_is_kept_when_the_placeholder_is_lost(self) -> None:
        """模型弄丢了占位符：按整段写法写回，公式仍在——双语的译文段把它接在句末，仅译文留在原段。"""
        source = "混凝土强度按计算。"
        for translation_only, expected in ((True, ["Résistance calculée.{f=N/A}"]),
                                           (False, ["混凝土强度按{f=N/A}计算。", "Résistance calculée.{f=N/A}"])):
            with self.subTest(translation_only=translation_only):
                doc = Document()
                _formula_paragraph(doc)
                out, _ = self._write(doc, {source: "Résistance calculée."},
                                     {source: ("混凝土强度按<x1/>计算。", "Résistance calculée.")},
                                     translation_only=translation_only)
                self.assertEqual([_with_math(item) for item in out.paragraphs], expected)

    def test_repaired_translation_without_format_tags_keeps_formulas_in_place(self) -> None:
        """反馈 #19：补救时格式标记整套去掉了，公式照样放回句中，文字用主体格式。"""
        doc = Document()
        paragraph = doc.add_paragraph()
        paragraph.add_run("注").bold = True
        paragraph.add_run("：强度按")
        _add_math(paragraph, "f=N/A")
        paragraph.add_run("计算，面积为")
        _add_math(paragraph, "A")
        paragraph.add_run("。")
        source, marked = "注：强度按计算，面积为。", "<g1>注</g1>：强度按<x1/>计算，面积为<x2/>。"
        repaired = "Note : la résistance est calculée selon <x1/>, la surface vaut <x2/>."
        for translation_only, expected in (
            (True, ["Note : la résistance est calculée selon {f=N/A}, la surface vaut {A}."]),
            (False, ["注：强度按{f=N/A}计算，面积为{A}。",
                     "Note : la résistance est calculée selon {f=N/A}, la surface vaut {A}."]),
        ):
            with self.subTest(translation_only=translation_only):
                out, _ = self._write(doc, {source: strip_marks(repaired)}, {source: (marked, repaired)},
                                     translation_only=translation_only)
                self.assertEqual([_with_math(item) for item in out.paragraphs], expected)

    def test_paragraph_with_many_formulas_is_marked(self) -> None:
        """反馈 #19：一段 24 个行内公式超过原来 16 个占位符的上限，整段退回写法、公式全到段末。"""
        doc = Document()
        paragraph = doc.add_paragraph()
        for index in range(24):
            paragraph.add_run(f"项{index}")
            _add_math(paragraph, f"a{index}")
        paragraph.add_run("。")
        marks = paragraph_marks(paragraph)
        self.assertIsNotNone(marks)
        self.assertEqual(len(marks.placeholders), 24)

    def test_formula_in_a_table_cell_and_display_formula_at_the_end(self) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run("面积")
        _add_math(cell.paragraphs[0], "A=πr^2")
        cell.paragraphs[0].add_run("取整")
        doc.add_paragraph("式中：")
        _add_math(doc.add_paragraph(), "σ=M/W", display=True)
        source, marked = "面积取整", "面积<x1/>取整"
        out, _ = self._write_case(doc, source, marked, "Surface <x1/> arrondie", translation_only=False)
        self.assertEqual([_with_math(item) for item in out.tables[0].cell(0, 0).paragraphs],
                         ["面积{A=πr^2}取整", "Surface {A=πr^2} arrondie"])
        # 文末单独一行的公式不是空段，不能被收尾时删掉
        self.assertEqual(_with_math(out.paragraphs[-1]), "{σ=M/W}")

    def test_stored_translation_with_a_space_for_the_tab_still_keeps_tab_and_bold(self) -> None:
        """本机实测：双语跑过一次后记忆库存的是「1.2 Portée…」（制表符成了空格），再跑仅译文
        命中这条、补上标记后，写回仍是 1.2⇥加粗标题，段尾没有多出来的制表符。"""
        doc = Document()
        _numbered_heading_paragraph(doc)
        source = "1.2\t施工范围见附图。"
        stored = "1.2 Portée des travaux : voir le dessin annexe."
        marked = "1.2<x1/><g1>施工范围</g1>见附图。"
        out, _ = self._write(
            doc, {source: stored},
            {source: (marked, "1.2<x1/><g1>Portée des travaux</g1> : voir le dessin annexe.")},
            translation_only=True,
        )
        self.assertEqual(_sequence(out.paragraphs[0]), "1.2⇥【粗|Portée des travaux】 : voir le dessin annexe.")

    def test_lost_placeholder_falls_back_to_whole_paragraph(self) -> None:
        doc = Document()
        _numbered_heading_paragraph(doc)
        source = "1.2\t施工范围见附图。"
        marked = "1.2<x1/><g1>施工范围</g1>见附图。"
        out, _ = self._write(
            doc, {source: "1.2\tPérimètre : voir plan."},
            {source: (marked, "1.2 <g1>Périmètre</g1> : voir plan.")},  # 占位符丢了，对不上最终译文
            translation_only=True,
        )
        self.assertEqual(out.paragraphs[0].text, "1.2\tPérimètre : voir plan.")

    def test_table_cell_paragraphs_keep_their_own_marks(self) -> None:
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        first = cell.paragraphs[0]
        first.add_run("注：").bold = True
        first.add_run("尺寸以毫米计。")
        cell.add_paragraph("标高以米计。")
        source = "注：尺寸以毫米计。\n标高以米计。"
        marked_source = "<g1>注：</g1>尺寸以毫米计。\n标高以米计。"
        marked_fr = "<g1>Note :</g1> dimensions en millimètres.\nCotes de niveau en mètres."
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "c.docx"
            doc.save(path)
            segments = {s.source: s.marked_source for s in extract_word_segments(path, target_lang="fr")}
        self.assertEqual(segments[source], marked_source)
        out, _ = self._write(doc, {source: strip_marks(marked_fr)}, {source: (marked_source, marked_fr)}, translation_only=True)
        out_cell = out.tables[0].cell(0, 0)
        self.assertEqual([r[:2] for r in _runs(out_cell.paragraphs[0])], [("Note :", True), (" dimensions en millimètres.", False)])
        self.assertEqual(out_cell.paragraphs[1].text, "Cotes de niveau en mètres.")


if __name__ == "__main__":
    unittest.main()
