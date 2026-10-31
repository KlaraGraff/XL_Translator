"""记忆库导入：用户自己整理的词表（列名、编码、分隔符、空行、多余列各不相同）也能认出来。"""

from __future__ import annotations

import base64
import io
import json
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook

from api.app import create_app
from core.tm_import_parse import build_template_xlsx, parse_import_file


def _pairs(result: dict) -> list[tuple[str, str, str]]:
    mapping = result["mapping"]
    out = []
    for row in result["rows"]:
        src = row[mapping["source"]].strip()
        tgt = row[mapping["target"]].strip()
        note = row[mapping["note"]].strip() if "note" in mapping else ""
        if src and tgt:
            out.append((src, tgt, note))
    return out


class TmImportParseTests(unittest.TestCase):
    def test_gbk_csv_with_chinese_headers_serial_column_and_blank_rows(self) -> None:
        data = "序号,中文,英文,说明\r\n1,桥墩,pier,\r\n\r\n2,主梁,main girder,桥梁\r\n3,盖梁,,\r\n".encode("gb18030")
        result = parse_import_file("词表.csv", data, "zh-en")
        self.assertEqual(result["encoding"], "GBK")
        self.assertEqual(result["mapping"], {"source": 1, "target": 2, "note": 3})
        self.assertEqual(result["skipped_blank"], 1)
        self.assertEqual(_pairs(result), [("桥墩", "pier", ""), ("主梁", "main girder", "桥梁")])

    def test_semicolon_csv_with_native_language_header_suggests_pair(self) -> None:
        data = "STT;中文;Tiếng Việt;Ghi chú\n1;桥墩;trụ cầu;\n2;主梁;dầm chính;\n".encode("utf-8")
        result = parse_import_file("glossary.csv", data, "zh-en")
        self.assertEqual(result["delimiter"], ";")
        self.assertEqual(result["mapping"]["source"], 1)
        self.assertEqual(result["mapping"]["target"], 2)
        self.assertEqual(result["suggested_lang_pair"], "zh-vi")

    def test_utf16_tab_text_without_header(self) -> None:
        data = "桥墩\tpier\n主梁\tmain girder\n".encode("utf-16")
        result = parse_import_file("terms.txt", data, "zh-en")
        self.assertFalse(result["has_header"])
        self.assertEqual(_pairs(result), [("桥墩", "pier", ""), ("主梁", "main girder", "")])

    def test_xlsx_title_row_extra_columns_and_cover_sheet(self) -> None:
        workbook = Workbook()
        cover = workbook.active
        cover.title = "说明"
        cover.append(["本表仅供参考"])
        sheet = workbook.create_sheet("词表")
        sheet.append(["XX 项目术语表"])
        sheet.append(["No.", "中文名称", "Français", "Remarque", "审核人"])
        sheet.append([1, "桥墩", "pile", None, "张三"])
        sheet.append([None, None, None, None, None])
        sheet.append([2, "主梁", "poutre principale", "pont", None])
        buffer = io.BytesIO()
        workbook.save(buffer)
        result = parse_import_file("术语.xlsx", buffer.getvalue(), "zh-fr")
        self.assertEqual(result["sheet"], "词表")
        self.assertEqual(result["sheets_skipped"], 1)
        self.assertEqual(result["skipped_above_header"], 1)
        self.assertEqual(result["mapping"], {"source": 1, "target": 2, "note": 3})
        self.assertEqual(_pairs(result), [("桥墩", "pile", ""), ("主梁", "poutre principale", "pont")])

    def test_other_json_is_refused_with_plain_message(self) -> None:
        data = json.dumps({"桥墩": "pier"}, ensure_ascii=False).encode()
        with self.assertRaisesRegex(ValueError, "只支持本软件导出的 JSON"):
            parse_import_file("a.json", data, "zh-en")

    def test_template_round_trips_through_detection(self) -> None:
        from openpyxl import load_workbook

        for locale in ("zh", "en", "fr"):
            workbook = load_workbook(io.BytesIO(build_template_xlsx(locale)))
            workbook.active.append(["桥墩", "pile", "桥梁"])
            buffer = io.BytesIO()
            workbook.save(buffer)
            result = parse_import_file("模板.xlsx", buffer.getvalue(), "zh-fr")
            self.assertEqual(result["mapping"], {"source": 0, "target": 1, "note": 2}, locale)
            self.assertEqual(_pairs(result), [("桥墩", "pile", "桥梁")])

    def test_own_csv_export_round_trip(self) -> None:
        data = "﻿source_text,target_text,word_type,pinned,updated_at\r\n桥墩,pier,manual,1,2026-10-31\r\n".encode("utf-8")
        result = parse_import_file("translator-tm.csv", data, "zh-en")
        self.assertEqual(result["mapping"], {"source": 0, "target": 1, "word_type": 2, "pinned": 3})

    def test_broken_excel_gives_plain_message(self) -> None:
        with self.assertRaisesRegex(ValueError, "无法打开这个 Excel 文件"):
            parse_import_file("坏文件.xlsx", b"PK\x03\x04garbage", "zh-en")


class TmImportParseRouteTests(unittest.TestCase):
    def test_route_parses_base64_file(self) -> None:
        client = TestClient(create_app())
        data = base64.b64encode("中文,English\n桥墩,pier\n".encode("utf-8")).decode()
        response = client.post("/api/tm/import/parse", json={"file_name": "x.csv", "data": data, "lang_pair": "zh-en"})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["rows"], [["桥墩", "pier"]])

    def test_template_route_returns_xlsx(self) -> None:
        response = TestClient(create_app()).get("/api/tm/import/template?locale=fr")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))

    def test_route_rejects_bad_base64(self) -> None:
        client = TestClient(create_app())
        response = client.post("/api/tm/import/parse", json={"file_name": "x.csv", "data": "@@@", "lang_pair": "zh-en"})
        self.assertEqual(response.status_code, 422)


if __name__ == "__main__":
    unittest.main()
