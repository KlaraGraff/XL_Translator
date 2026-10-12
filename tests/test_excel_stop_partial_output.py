"""Excel 中途停止：已翻好的部分照常写进文件，停止口径与续译检测一致。

真跑 ``TaskRunner._run()`` + 真实写盘 + 真实续译检测，只把翻译引擎和记忆库换成桩：
  - 停止后输出文件里一半是双语、一半保留原文，逐文件结果是「已生成」并带
    unfinished_text_count，横幅用的 kpi.unstarted_text_count 与重新扫描后续译弹窗
    的「补 N 格」是同一个数；
  - 停止时没翻的格子不标黄、不计「需复核」；
  - 停止打断的在途请求不算「接口异常」（不出 api_unavailable 那条「请检查 API 密钥」）。
"""

from __future__ import annotations

import os
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from core.api_config_check import ApiConfigCheckResult
from core.file_scanner import FileItem
from core.model_throughput import EffectiveModelThroughput
from core.resume_detection import detect_previous_output
from core.task_runner import StoppedMsg, TaskRunner
from settings import AppSettings, EngineSettings

_SOURCES = [f"第{index}号施工段混凝土浇筑" for index in range(1, 31)]


class _StoppingEngine:
    """第 ``stop_on_call`` 次请求时按下停止；``fail_after_stop`` 时这一次请求被打断。"""

    engine_name = "excel-stop-test/mock"

    def __init__(self, *, stop_on_call: int, fail_after_stop: bool) -> None:
        self.runner: TaskRunner | None = None
        self.calls = 0
        self.chat_calls = 0
        self._stop_on_call = stop_on_call
        self._fail_after_stop = fail_after_stop

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        self.calls += 1
        if self.calls == self._stop_on_call:
            assert self.runner is not None
            self.runner.stop()
            if self._fail_after_stop:
                raise ConnectionError("connection aborted while stopping")
        return {text: f"Concrete pouring section {_SOURCES.index(text) + 1}" for text in texts}

    def chat(self, system: str, user: str) -> str:
        # 停止之后不该再发任何请求（包括「失败后语言方向」那一次探测）。
        self.chat_calls += 1
        raise NotImplementedError


def _settings() -> AppSettings:
    settings = AppSettings(
        engine=EngineSettings(
            mode="cloud",
            cloud_provider="custom_openai",
            cloud_model="excel-stop-test",
            cloud_base_url="https://example.invalid/v1",
            concurrency=1,
            batch_size=10,
        ),
        source_lang="zh",
        target_lang="en",
        # 这里数的是语种探测请求；「自动识别场景」的文档说明请求另有测试。
        excel_domain_preset="无",
        excel_smart_context=False,
    )
    settings.excel_output.enable_excel_autofit = False
    return settings


class ExcelStopWritesTranslatedPart(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name) / "docs"
        self.root.mkdir()
        self._env = patch.dict(os.environ, {"TRANSLATOR_APP_DATA_DIR": str(Path(self._tmp.name) / "appdata")})
        self._env.start()

    def tearDown(self) -> None:
        self._env.stop()
        self._tmp.cleanup()

    def _run(
        self,
        *,
        fail_after_stop: bool,
        stop_on_call: int | None = None,
        sources: list[str] = _SOURCES,
        sheet_count: int = 1,
        translate_names=None,
    ) -> tuple[StoppedMsg, Path]:
        source = self.root / "工程量清单.xlsx"
        wb = Workbook()
        ws = wb.active
        for row, text in enumerate(sources, start=1):
            ws.cell(row=row, column=1, value=text)
        for index in range(2, sheet_count + 1):
            wb.create_sheet(f"分表{index}").cell(row=1, column=1, value=sources[0])
        wb.save(source)

        settings = _settings()
        runner = TaskRunner(
            [FileItem(path=source, name=source.stem, size_kb=1.0)],
            settings,
            source_root=self.root,
        )
        if stop_on_call is None:
            stop_on_call = 1 if not fail_after_stop else 2
        engine = _StoppingEngine(stop_on_call=stop_on_call, fail_after_stop=fail_after_stop)
        engine.runner = runner
        self.engine = engine
        with ExitStack() as stack:
            stack.enter_context(
                patch("core.task_runner.TaskLogger", return_value=MagicMock(task_id="excel-stop-test"))
            )
            stack.enter_context(
                patch("core.task_runner.check_translation_api_config", return_value=ApiConfigCheckResult(ok=True))
            )
            stack.enter_context(patch("core.task_runner.build_engine", return_value=engine))
            stack.enter_context(patch("core.task_runner.get_system_prompt", return_value="system"))
            stack.enter_context(patch("core.task_runner.resolve_effective_model_config", return_value=object()))
            stack.enter_context(
                patch(
                    "core.task_runner.get_model_throughput",
                    return_value=EffectiveModelThroughput(profile_key="t", batch_size=10, concurrency=1),
                )
            )
            stack.enter_context(patch("core.task_runner.tm_manager.lookup_batch", return_value={}))
            stack.enter_context(patch("core.task_runner.tm_manager.insert_batch", return_value=0))
            if translate_names is not None:
                stack.enter_context(
                    patch(
                        "core.task_runner.translate_names",
                        side_effect=lambda *args, **kwargs: translate_names(runner, *args, **kwargs),
                    )
                )
            runner._run()

        messages = [m for m in list(runner._queue.queue) if isinstance(m, StoppedMsg)]
        self.assertEqual(len(messages), 1)
        return messages[0], source

    def _assert_partial_output(self, stopped: StoppedMsg, source: Path, *, translated: int) -> None:
        untranslated = len(_SOURCES) - translated
        entry = stopped.files[0]
        self.assertTrue(entry["success"], entry)
        self.assertEqual(entry["status"], "succeeded")
        self.assertEqual(entry["unfinished_text_count"], untranslated)
        self.assertEqual(stopped.kpi["unstarted_text_count"], untranslated)
        # 停止时没翻的格子不是「需复核」。
        self.assertEqual(stopped.review["total_count"], 0)
        self.assertIn(f"尚有 {untranslated} 个单元格未翻译", stopped.message)

        output = Path(entry["output_path"])
        cells = [
            str(row[0].value or "")
            for row in load_workbook(output).active.iter_rows(max_col=1)
        ]
        bilingual = [value for value in cells if "\n" in value]
        self.assertEqual(len(bilingual), translated)
        self.assertEqual(len([value for value in cells if value in _SOURCES]), untranslated)

        # 重新扫描后的续译弹窗说的是同一个数。
        detected = detect_previous_output(
            surface="excel",
            items=[{"path": str(source)}],
            scan_roots=[str(self.root)],
            custom_output_root=None,
            target_lang="en",
            source_lang="zh",
        )
        self.assertIsNotNone(detected)
        self.assertEqual(detected["files"][0]["status"], "partial")
        self.assertEqual(detected["files"][0]["untranslated_count"], untranslated)

    def test_stop_between_batches_writes_translated_part(self) -> None:
        stopped, source = self._run(fail_after_stop=False)
        self._assert_partial_output(stopped, source, translated=10)

    def test_request_cut_off_by_stop_is_not_reported_as_api_failure(self) -> None:
        stopped, source = self._run(fail_after_stop=True)
        self._assert_partial_output(stopped, source, translated=10)
        types = {str(issue.get("type")) for issue in stopped.issues}
        self.assertNotIn("api_unavailable", types)
        self.assertFalse(
            any("API 密钥" in str(issue.get("message")) for issue in stopped.issues),
            stopped.issues,
        )
        # 停止交回的原文占位不是「质量校验回退」，也没有格子被标记待复核。
        self.assertNotIn("quality_filter_reset", types)
        self.assertNotIn("suspected_source_language_mismatch", types)
        self.assertEqual(self.engine.chat_calls, 0)

    def test_stop_on_first_request_makes_no_language_probe(self) -> None:
        # 小任务（一批就装下）第一批就被停止打断：全部条目都是停止占位。以前它们
        # 被当成「质量校验回退」，占比过线后还会在停止之后多发一次 chat 探测语言方向。
        stopped, _source = self._run(
            fail_after_stop=True, stop_on_call=1, sources=_SOURCES[:10]
        )
        types = {str(issue.get("type")) for issue in stopped.issues}
        self.assertNotIn("quality_filter_reset", types)
        self.assertNotIn("suspected_source_language_mismatch", types)
        self.assertNotIn("api_unavailable", types)
        self.assertEqual(self.engine.chat_calls, 0)

    def test_stop_during_sheet_name_translation_still_writes_file(self) -> None:
        # 超过 24 个工作表名要分多批翻。第一批翻完时按下停止：以前下一批前的
        # _raise_if_stopped() 抛 TaskStopped，格子明明都翻好了，一个文件都不写。
        name_calls: list[list[str]] = []

        def stop_on_first_name_batch(runner, engine, names, target_lang, source_lang, *, kind):
            name_calls.append(list(names))
            runner.stop()
            return {name: f"Sheet {name}" for name in names}

        stopped, _source = self._run(
            fail_after_stop=False,
            stop_on_call=10_000,
            sheet_count=30,
            translate_names=stop_on_first_name_batch,
        )
        self.assertEqual(len(name_calls), 1)
        entry = stopped.files[0]
        self.assertTrue(entry["success"], entry)
        output = Path(entry["output_path"])
        self.assertTrue(output.exists())
        workbook = load_workbook(output)
        # 第一批的名称已经换上，没轮到的保持原样。
        self.assertIn("Sheet 分表2", workbook.sheetnames)
        self.assertIn("分表30", workbook.sheetnames)
        cells = [str(row[0].value or "") for row in workbook.worksheets[0].iter_rows(max_col=1)]
        self.assertEqual(len([value for value in cells if "\n" in value]), len(_SOURCES))

    def test_uncountable_output_leaves_unfinished_kpi_unset(self) -> None:
        # 产物数不出来（返回 None）时以前 kpi 记 0，界面就说「停止前内容已经全部翻完」
        # 并藏起「接着翻译」。现在不写这个键，由界面按「没数清」处理。
        with patch(
            "core.resume_detection.count_untranslated_in_output", return_value=None
        ):
            stopped, _source = self._run(fail_after_stop=False)
        entry = stopped.files[0]
        self.assertTrue(entry["success"], entry)
        self.assertNotIn("unstarted_text_count", stopped.kpi)
        self.assertNotIn("unfinished_text_count", entry)
        self.assertNotIn("尚有", stopped.message)


if __name__ == "__main__":  # pragma: no cover
    unittest.main()
