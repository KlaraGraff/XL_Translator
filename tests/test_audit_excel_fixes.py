"""2026-08-29 审查报告 · Excel 板块（A9）的回归测试。

盯住的条目：

* **中-28** 共享公式让渡是 O(n²)：每让渡一个主控格就把整张分表重扫一遍。实测每翻
  一倍行数耗时涨四倍（200 行 0.19s → 800 行 2.91s → 1600 行 11.66s），上万行的
  长公式列在用户眼里就是任务假死。
* **中-29** 目标语是中文时，复核改判过的格子按整格「原文＋可疑译文」建键，这串必然
  含中文 → 写入端重跑一次 ``should_translate`` 判 False → 译文和复核底色一起丢。
* **中-30** ``.xls`` 兼容转换只取裸值不看 ``ctype``：日期变序列号、布尔变 1/0。
* **低** 锁定行高时小于 6pt 的字号被「缩」到 6pt（反向放大）。
* **低** 共享公式让渡失败只写 loguru，任务日志无痕，且留下「标了色没译文」的格子。
* **低** 让渡后 ``ref`` 左上角指着已经被改写掉的旧主控格。
* **低** ``build_excel_coverage_plan`` 第二次 load 抛错时泄掉第一个工作簿句柄。
* **低** ``.xls`` 转换失败留下半个 .xlsx，没有任何人回收。

第二轮（对抗审查开出的整改）：

* **XLS-ERR-TEXT** 错误值单元格（``#N/A`` / ``#REF!`` …）被当成普通文本送进翻译
  管线，最后整格被改写成 ``#N/A\\n不适用``，错误格降级成文本格。
* **XLS-CLEANUP-ORDER** ``convert_with_excel`` 在 ``except`` 里删半成品、``wb.close()``
  却在其后的 ``finally``：删文件时工作簿还开在 Excel 进程里，Windows 上 unlink 直接
  抛 ``PermissionError``，半截文件照样留下。
"""

from __future__ import annotations

import datetime
import importlib.util
import re
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

import openpyxl
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core import excel_coverage, xls_converter
from core.coverage_arbitration import apply_arbitration, review_coverage_pairs
from core.excel_coverage import build_excel_coverage_plan, write_untranslated_excel_file
from core.mixed_language import MIXED_MARK_FOREIGN_NOISE
from core.translation_coverage import COVERAGE_COVERED, COVERAGE_SOURCE_ONLY
from core.xls_converter import _cell_value_by_ctype, convert_with_excel, convert_with_fallback
from core.xlsx_patcher import (
    NS_MAIN,
    _shrink_font_for_locked_row,
    write_bilingual_workbook,
)

M = f"{{{NS_MAIN}}}"

HAS_XLWT = importlib.util.find_spec("xlwt") is not None


# ── 夹具：手写分表 XML ────────────────────────────────────────────────────────
# openpyxl 写不出共享公式（``<f t="shared" si= ref=>``），真实 Excel 存的每一张
# 长公式列都是那个样子。所以这里用 openpyxl 建一个合法的 xlsx 包，再把分表 XML
# 整个换成手写的——包里其余部件（styles、rels、contentTypes）保持真实。
def _package_with_sheet(root: Path, name: str, rows_xml: str) -> Path:
    base = root / f"base_{name}.xlsx"
    workbook = Workbook()
    workbook.active.title = "S"
    workbook["S"]["A1"] = "placeholder"
    workbook.save(base)
    workbook.close()

    with zipfile.ZipFile(base) as archive:
        parts = {item: archive.read(item) for item in archive.namelist()}
    parts["xl/worksheets/sheet1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{NS_MAIN}"><sheetData>{rows_xml}</sheetData></worksheet>'
    ).encode("utf-8")

    target = root / f"{name}.xlsx"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for item, data in parts.items():
            archive.writestr(item, data)
    return target


def _one_group_per_row(rows: int) -> str:
    """每行一个共享公式组：B 是主控格（带 ref），C 是从属格。

    B 的缓存显示值是待译文本，所以每一行都会触发一次主控权让渡。
    """
    return "".join(
        f'<row r="{r}">'
        f'<c r="A{r}" t="str"><v>数据{r}</v></c>'
        f'<c r="B{r}" t="str"><f t="shared" si="{r}" ref="B{r}:C{r}">A{r}&amp;"合计"</f>'
        f"<v>施工内容</v></c>"
        f'<c r="C{r}" t="str"><f t="shared" si="{r}"/><v>其它</v></c>'
        f"</row>"
        for r in range(1, rows + 1)
    )


def _sheet_cells(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    return {cell.get("r"): cell for cell in root.iter(f"{M}c")}


def _formula(cell) -> object | None:
    return None if cell is None else cell.find(f"{M}f")


def _translate_sheet(path: Path, **kwargs) -> None:
    write_bilingual_workbook(
        path,
        translations={"施工内容": "[EN]content"},
        target_lang="en",
        source_lang="zh",
        keep_original_sheets=False,
        formula_display_value_backfill=True,
        mark_review_items=False,
        **kwargs,
    )


class SharedFormulaDonationTests(unittest.TestCase):
    """中-28 + 两条低危：让渡的复杂度、让渡结果的合法性、让渡失败的交代。"""

    def setUp(self) -> None:
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)

    def tearDown(self) -> None:
        self._temp.cleanup()

    def _time_translation(self, rows: int) -> float:
        """跑三遍取最快的一遍，尽量摘掉机器抖动。"""
        source = _package_with_sheet(self.root, f"n{rows}", _one_group_per_row(rows))
        best = float("inf")
        for attempt in range(3):
            target = self.root / f"timed_{rows}_{attempt}.xlsx"
            target.write_bytes(source.read_bytes())
            start = time.perf_counter()
            _translate_sheet(target)
            best = min(best, time.perf_counter() - start)
        return best

    def test_donation_cost_no_longer_grows_quadratically(self) -> None:
        """行数翻两番，耗时也只该翻两番左右——而不是十六番。

        断的是**增长关系**不是绝对秒数：绝对值跟机器走，比例不跟。修复前实测
        400 行 0.73s、1600 行 11.66s（16.0 倍，标准的 O(n²)）；修复后 0.015s /
        0.057s（3.8 倍）。阈值取 8 倍，两边都留足了余量。
        """
        small = self._time_translation(400)
        large = self._time_translation(1600)
        self.assertGreater(small, 0.0)
        self.assertLess(
            large,
            small * 8,
            f"1600 行耗时 {large:.3f}s 是 400 行 {small:.3f}s 的 "
            f"{large / small:.1f} 倍，让渡又退回了整表重扫",
        )

    def test_every_group_still_hands_the_master_over(self) -> None:
        """快不能以少干活为代价：每一组的主控权都要真的交到从属格手上。"""
        source = _package_with_sheet(self.root, "groups", _one_group_per_row(3))
        target = self.root / "groups_out.xlsx"
        target.write_bytes(source.read_bytes())
        _translate_sheet(target)

        cells = _sheet_cells(target)
        for row in (1, 2, 3):
            self.assertIsNone(
                _formula(cells[f"B{row}"]),
                f"B{row} 被改写成双语文本，<f> 该跟着消失",
            )
            handed_over = _formula(cells[f"C{row}"])
            self.assertIsNotNone(handed_over, f"C{row} 没接过主控权，整组公式失效")
            self.assertEqual(handed_over.get("si"), str(row))
            self.assertEqual(handed_over.get("ref"), f"C{row}:C{row}")
            # 主控格换了位置，公式要跟着平移：B{r} 上的 A{r}&"合计" 到 C{r} 是 B{r}&"合计"。
            self.assertEqual(handed_over.text, f'B{row}&"合计"')

    def test_shared_ref_never_points_at_a_cell_that_left_the_group(self) -> None:
        """二维共享组：主控 B2，从属 C2/B3/C3。

        剩下的三格摆不成一个「左上角是新主控格」的矩形。旧写法照样写 ref="B2:C3"
        —— 左上角 B2 这时已经是纯文本、不在组里了，共享公式的从属格全靠「相对主控
        格的偏移」推自己的公式，左上角指着组外的格子，文件在 Excel 眼里是坏的。
        """
        rows_xml = (
            '<row r="2">'
            '<c r="B2" t="str"><f t="shared" si="7" ref="B2:C3">A2&amp;"合计"</f>'
            "<v>施工内容</v></c>"
            '<c r="C2" t="str"><f t="shared" si="7"/><v>其它</v></c>'
            "</row>"
            '<row r="3">'
            '<c r="B3" t="str"><f t="shared" si="7"/><v>其它</v></c>'
            '<c r="C3" t="str"><f t="shared" si="7"/><v>其它</v></c>'
            "</row>"
        )
        source = _package_with_sheet(self.root, "twodim", rows_xml)
        target = self.root / "twodim_out.xlsx"
        target.write_bytes(source.read_bytes())
        _translate_sheet(target)

        cells = _sheet_cells(target)
        self.assertIsNone(_formula(cells["B2"]))

        remaining = {name: _formula(cells[name]) for name in ("C2", "B3", "C3")}
        for name, formula_el in remaining.items():
            self.assertIsNotNone(formula_el, f"{name} 的公式整个不见了")
            ref = formula_el.get("ref")
            if ref:
                # 还想当共享组的话，ref 左上角必须就是这一格自己。
                self.assertEqual(ref.split(":")[0], name)
            else:
                # 让不出合法的组就整组拆开，每格写回自己的完整公式。
                self.assertIsNone(formula_el.get("si"))
                self.assertTrue((formula_el.text or "").strip())

        # 不管走哪条路，每一格算出来的东西都必须和原来一样（相对偏移不变）。
        expected = {"C2": 'B2&"合计"', "B3": 'A3&"合计"', "C3": 'B3&"合计"'}
        for name, text in expected.items():
            if not remaining[name].get("si"):
                self.assertEqual(remaining[name].text, text)

    def test_failed_donation_reports_itself_and_leaves_no_orphan_fill(self) -> None:
        """让渡失败 → 整格不动，底色也不涂，理由要出现在任务日志里。

        旧写法先按「要改写」把底色涂了，之后才发现公式让不出去、把译文丢掉：
        用户拿到的是一格无缘无故变了色、却一个译文都没有的单元格，而放弃的原因
        只写进了 loguru，任务日志里一个字都没有。
        """
        # 主控格的 <f> 没有公式文本 → 无法平移给从属格 → 让渡必然失败。
        rows_xml = (
            '<row r="1">'
            '<c r="B1" t="str"><f t="shared" si="3" ref="B1:C1"></f><v>施工内容</v></c>'
            '<c r="C1" t="str"><f t="shared" si="3"/><v>其它</v></c>'
            "</row>"
        )
        source = _package_with_sheet(self.root, "failed", rows_xml)
        target = self.root / "failed_out.xlsx"
        target.write_bytes(source.read_bytes())

        logs: list[str] = []
        positions: list[dict[str, str]] = []
        write_bilingual_workbook(
            target,
            translations={"施工内容": "[EN]content"},
            target_lang="en",
            source_lang="zh",
            keep_original_sheets=False,
            formula_display_value_backfill=True,
            review_marks={"施工内容": MIXED_MARK_FOREIGN_NOISE},
            review_mark_colors={MIXED_MARK_FOREIGN_NOISE: "#FFC7CE"},
            mark_review_items=True,
            review_positions=positions,
            log_callback=logs.append,
        )

        self.assertEqual(positions, [], "译文没写成，底色不该单独涂上去")
        warning = [line for line in logs if "B1" in line and "共享公式" in line]
        self.assertTrue(warning, f"任务日志里没有交代放弃的原因：{logs}")

        cells = _sheet_cells(target)
        self.assertIsNotNone(_formula(cells["B1"]), "放弃改写就该原样保留公式")
        self.assertIsNone(cells["B1"].get("s"), "没写译文的格子不该被换上带底色的样式")


class LockedRowHeightFontTests(unittest.TestCase):
    """低危：锁定行高模式下，小字号被「缩」成大字号。"""

    def test_tiny_font_is_never_enlarged_by_the_floor(self) -> None:
        """原字号 4pt < 触底阈值 6pt：装不下时旧写法直接把它设成 6pt。

        锁行高的整个用意是「宁可字小也不让行变高」，把 4pt 放大到 6pt 正好反过来
        —— 行高锁死、字更大，本来显示得下的内容被挤没了。
        """
        size, reached_floor = _shrink_font_for_locked_row(
            "很长很长的一段双语文本" * 20,
            col_width=8.0,
            row_height=12.0,
            font_size=4.0,
        )
        self.assertTrue(reached_floor)
        # 断言必须无条件跑：``size is None`` 是「这一格没被动过」，等价于仍是 4pt。
        # 旧写法把 assertLessEqual 包在 ``if size is not None`` 里，修复版走的正是
        # None 那一支，于是真正生效的只剩 reached_floor 一条——将来 size 因别的原因
        # 变 None，这条断言就静默空过。
        effective = 4.0 if size is None else size
        self.assertLessEqual(effective, 4.0, "本该缩小的一格被放大了")

    def test_normal_font_still_shrinks(self) -> None:
        """别把上一条修成「永远不动字号」——正常字号该缩还得缩。"""
        size, _reached_floor = _shrink_font_for_locked_row(
            "很长很长的一段双语文本" * 20,
            col_width=8.0,
            row_height=12.0,
            font_size=11.0,
        )
        self.assertIsNotNone(size)
        self.assertLess(size, 11.0)

    def test_end_to_end_locked_row_keeps_tiny_font_tiny(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "tiny.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "S"
            sheet["A1"] = "施工内容"
            sheet["A1"].font = Font(size=4)
            sheet.column_dimensions["A"].width = 8
            sheet.row_dimensions[1].height = 12
            workbook.save(path)
            workbook.close()

            write_bilingual_workbook(
                path,
                translations={"施工内容": "Construction scope description " * 6},
                target_lang="en",
                source_lang="zh",
                keep_original_sheets=False,
                lock_row_height=True,
                mark_review_items=False,
            )

            workbook = load_workbook(path)
            try:
                self.assertLessEqual(workbook["S"]["A1"].font.size, 4.0)
            finally:
                workbook.close()


class ReviewFlippedCellIntoChineseTests(unittest.TestCase):
    """中-29：目标语是中文时，改判过的格子整格被静默丢弃。"""

    SOURCE = (
        "The contractor shall complete the structural works before the rainy season "
        "and submit a revised programme to the engineer for approval."
    )
    # 长度像样、读着通顺，但讲的完全是另一件事——这一格配错了对。
    MISMATCHED = "本章规定了高处作业的安全措施以及现场必须配备的个人防护用品。"
    CORRECT = (
        "承包商应在雨季前完成结构工程，并向工程师提交修订后的进度计划以供批准。"
    )

    def _flipped_plan(self, root: Path):
        source = root / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet"
        sheet["A1"] = f"{self.SOURCE}\n{self.MISMATCHED}"
        workbook.save(source)
        workbook.close()

        plan = build_excel_coverage_plan(source, target_lang="zh", source_lang="en")
        unit = next(item for item in plan.units if item.location == "Sheet!A1")
        self.assertEqual(unit.status, COVERAGE_COVERED)

        apply_arbitration(
            review_coverage_pairs(
                plan.units,
                arbitrate=lambda pairs: {pair.id: "not_equivalent" for pair in pairs},
            )
        )
        self.assertEqual(unit.status, COVERAGE_SOURCE_ONLY)
        # 改判后建键用的是整格文字，不是格里的原文那一半。
        self.assertEqual(unit.data["cell_text"], f"{self.SOURCE}\n{self.MISMATCHED}")
        return source, plan

    def test_new_chinese_translation_lands_in_the_flipped_cell(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, plan = self._flipped_plan(root)

            out_path = write_untranslated_excel_file(
                source_path=source,
                output_dir=root / "out",
                plan=plan,
                translations={self.SOURCE: self.CORRECT},
                target_lang="zh",
                source_lang="en",
                keep_original_sheets=False,
            )

            workbook = load_workbook(out_path)
            try:
                value = workbook["Sheet"]["A1"].value
            finally:
                workbook.close()
            # 原文、判可疑的旧译文、补上的新译文，三段都在，顺序不变。
            self.assertEqual(
                value, f"{self.SOURCE}\n{self.MISMATCHED}\n{self.CORRECT}"
            )

    def test_flipped_cell_into_chinese_still_gets_its_fill(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, plan = self._flipped_plan(root)

            positions: list[dict[str, str]] = []
            write_untranslated_excel_file(
                source_path=source,
                output_dir=root / "out",
                plan=plan,
                translations={self.SOURCE: self.CORRECT},
                target_lang="zh",
                source_lang="en",
                keep_original_sheets=False,
                review_marks={self.SOURCE: MIXED_MARK_FOREIGN_NOISE},
                review_mark_colors={MIXED_MARK_FOREIGN_NOISE: "#FFC7CE"},
                mark_review_items=True,
                review_positions=positions,
            )

            self.assertEqual(
                [(item["cell"], item["action"]) for item in positions],
                [("A1", "marked_fill")],
            )


class XlsValueTypeTests(unittest.TestCase):
    """中-30：.xls 兼容转换必须按 ctype 还原日期与布尔。"""

    @staticmethod
    def _cell(ctype: int, value):
        return SimpleNamespace(ctype=ctype, value=value)

    def test_serial_number_becomes_a_real_date(self) -> None:
        import xlrd

        value = _cell_value_by_ctype(self._cell(xlrd.XL_CELL_DATE, 45122.0), 0)
        self.assertEqual(value, datetime.date(2023, 7, 15))

    def test_date_with_a_clock_keeps_the_clock(self) -> None:
        import xlrd

        value = _cell_value_by_ctype(self._cell(xlrd.XL_CELL_DATE, 45122.5), 0)
        self.assertEqual(value, datetime.datetime(2023, 7, 15, 12, 0, 0))

    def test_time_only_cell_stays_a_time(self) -> None:
        import xlrd

        value = _cell_value_by_ctype(self._cell(xlrd.XL_CELL_DATE, 0.25), 0)
        self.assertEqual(value, datetime.time(6, 0, 0))

    def test_boolean_does_not_degrade_into_one_and_zero(self) -> None:
        import xlrd

        self.assertIs(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_BOOLEAN, 1.0), 0), True)
        self.assertIs(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_BOOLEAN, 0.0), 0), False)

    def test_error_code_becomes_readable_text(self) -> None:
        import xlrd

        self.assertEqual(
            _cell_value_by_ctype(self._cell(xlrd.XL_CELL_ERROR, 0x07), 0), "#DIV/0!"
        )

    def test_plain_values_are_untouched(self) -> None:
        import xlrd

        self.assertEqual(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_TEXT, "施工"), 0), "施工")
        self.assertEqual(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_NUMBER, 12.5), 0), 12.5)
        self.assertEqual(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_NUMBER, 5.0), 0), 5)
        self.assertIsNone(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_EMPTY, ""), 0))
        self.assertIsNone(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_TEXT, ""), 0))

    def test_unconvertible_serial_falls_back_to_the_raw_number(self) -> None:
        import xlrd

        # 负序列号不是任何日期，xlrd 会抛 XLDateError——退回数字总好过整格丢失。
        self.assertEqual(_cell_value_by_ctype(self._cell(xlrd.XL_CELL_DATE, -1.0), 0), -1.0)

    @unittest.skipUnless(HAS_XLWT, "本机没有 xlwt，无法现场生成 .xls 夹具")
    def test_end_to_end_conversion_preserves_dates_and_booleans(self) -> None:
        import xlwt

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "legacy.xls"
            book = xlwt.Workbook()
            sheet = book.add_sheet("报价")
            date_style = xlwt.XFStyle()
            date_style.num_format_str = "YYYY-MM-DD"
            sheet.write(0, 0, "开工日期")
            sheet.write(0, 1, datetime.datetime(2023, 7, 15), date_style)
            sheet.write(1, 0, "是否验收")
            sheet.write(1, 1, True)
            sheet.write(2, 0, "金额")
            sheet.write(2, 1, 1250.5)
            book.save(str(path))

            out_path = convert_with_fallback(path)
            self.addCleanup(lambda: out_path.unlink(missing_ok=True))

            workbook = load_workbook(out_path)
            try:
                sheet_out = workbook["报价"]
                self.assertEqual(
                    sheet_out["B1"].value, datetime.datetime(2023, 7, 15, 0, 0)
                )
                self.assertIs(sheet_out["B2"].value, True)
                self.assertEqual(sheet_out["B3"].value, 1250.5)
                self.assertEqual(sheet_out["A1"].value, "开工日期")
            finally:
                workbook.close()


class XlsConversionCleanupTests(unittest.TestCase):
    """低危：转换半途失败留下的半个 .xlsx 没有任何人回收。"""

    def setUp(self) -> None:
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)

    def tearDown(self) -> None:
        self._temp.cleanup()

    @unittest.skipUnless(HAS_XLWT, "本机没有 xlwt，无法现场生成 .xls 夹具")
    def test_fallback_conversion_removes_the_half_written_output(self) -> None:
        import xlwt

        path = self.root / "legacy.xls"
        book = xlwt.Workbook()
        book.add_sheet("S").write(0, 0, "施工内容")
        book.save(str(path))

        seen: list[Path] = []

        def _explode(self_, target):  # noqa: ANN001 - 替身签名跟着 openpyxl
            out = Path(target)
            out.write_bytes(b"half written")  # 落盘落到一半
            seen.append(out)
            raise OSError("磁盘写满了")

        with mock.patch.object(openpyxl.Workbook, "save", _explode):
            with self.assertRaises(OSError):
                convert_with_fallback(path)

        self.assertEqual(len(seen), 1)
        self.assertFalse(seen[0].exists(), "转换失败后半个 .xlsx 还留在临时目录里")

    def test_excel_conversion_removes_the_half_written_output(self) -> None:
        seen: list[Path] = []

        class _Book:
            def __init__(self) -> None:
                self.closed = False

            def save(self, target: str) -> None:
                out = Path(target)
                out.write_bytes(b"half written")
                seen.append(out)
                raise RuntimeError("Excel 中途退出了")

            def close(self) -> None:
                self.closed = True

        book = _Book()
        app = SimpleNamespace(books=SimpleNamespace(open=lambda _path: book))

        with self.assertRaises(xls_converter.XlwingsUnavailableError):
            convert_with_excel(app, self.root / "legacy.xls")

        self.assertTrue(book.closed, "失败路径也要关掉工作簿")
        self.assertEqual(len(seen), 1)
        self.assertFalse(seen[0].exists(), "转换失败后半个 .xlsx 还留在临时目录里")

    def test_excel_conversion_deletes_the_output_only_after_the_book_is_closed(self) -> None:
        """XLS-CLEANUP-ORDER：删半成品必须排在 ``wb.close()`` 之后。

        真实场景在 Windows + 本机 Excel（这正是 xlwings 主路径）：``SaveAs`` 中途失败
        时 Excel 已经建好 ``out_path`` 并仍持有它，这时 unlink 抛 ``PermissionError``
        （``OSError`` 子类）→ 被 ``_discard_partial_output`` 吞成一条 debug 日志 →
        半截文件照样留在临时目录里没人回收。

        替身在这里替 Windows 当那把锁：``close()`` 之前删这个文件一律 PermissionError。
        顺序错了，最后 ``exists()`` 就是 True。
        """
        seen: list[Path] = []

        class _LockingBook:
            def __init__(self) -> None:
                self.closed = False

            def save(self, target: str) -> None:
                out = Path(target)
                out.write_bytes(b"half written")
                seen.append(out)
                raise RuntimeError("Excel 中途退出了")

            def close(self) -> None:
                self.closed = True

        book = _LockingBook()
        app = SimpleNamespace(books=SimpleNamespace(open=lambda _path: book))
        real_unlink = Path.unlink
        order: list[str] = []

        def _guarded_unlink(self_path, *args, **kwargs):
            if seen and self_path == seen[0]:
                order.append(f"unlink(closed={book.closed})")
                if not book.closed:
                    raise PermissionError(
                        32, "The process cannot access the file because it is being "
                        "used by another process"
                    )
            return real_unlink(self_path, *args, **kwargs)

        with mock.patch.object(Path, "unlink", _guarded_unlink):
            with self.assertRaises(xls_converter.XlwingsUnavailableError):
                convert_with_excel(app, self.root / "legacy.xls")

        self.assertTrue(book.closed)
        self.assertEqual(len(seen), 1)
        self.assertEqual(
            order,
            ["unlink(closed=True)"],
            "半成品是在工作簿还开着的时候删的，Windows 上这一刀砍空",
        )
        self.assertFalse(seen[0].exists(), "转换失败后半个 .xlsx 还留在临时目录里")


class ErrorCellsStayOutOfTranslationTests(unittest.TestCase):
    """XLS-ERR-TEXT：错误值单元格不许进翻译管线，更不许被改写成双语文本。

    ``.xls`` 兼容转换按 ctype 还原错误值之后，``#N/A`` 这类文本第一次成了「openpyxl
    读得到的字符串」。它在 .xlsx 里是合法的 ``<c t="e">`` 错误格，可整条管线只认
    「值是不是 str」：抽取端收下它、``should_translate('#N/A')`` 判真、写入端把整格
    换成 ``#N/A\\n不适用`` 的 inlineStr——错误格降级成文本格，公式的计算结果被改写。
    原生 .xlsx 里的错误格一直有同样的毛病，这次一起挡住。
    """

    # xlrd 的错误码：0x2A = #N/A，0x17 = #REF!。
    NA_CODE = 0x2A

    def setUp(self) -> None:
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)

    def tearDown(self) -> None:
        self._temp.cleanup()

    def _workbook_with_error_cell(self, name: str) -> Path:
        """A1 是错误值（走一遍转换器的还原逻辑），A2 是普通待译文本。"""
        import xlrd

        path = self.root / name
        error_text = _cell_value_by_ctype(
            SimpleNamespace(ctype=xlrd.XL_CELL_ERROR, value=self.NA_CODE), 0
        )
        self.assertEqual(error_text, "#N/A")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "S"
        sheet["A1"] = error_text
        sheet["A2"] = "施工内容"
        workbook.save(path)
        workbook.close()

        check = load_workbook(path)
        try:
            # 前提确认：openpyxl 把 '#N/A' 存成了真正的错误格，不是普通文本。
            self.assertEqual(check["S"]["A1"].data_type, "e")
        finally:
            check.close()
        return path

    def test_error_cell_is_not_rewritten_into_bilingual_text(self) -> None:
        path = self._workbook_with_error_cell("errors.xlsx")

        write_bilingual_workbook(
            path,
            translations={"#N/A": "不适用", "施工内容": "[EN]content"},
            target_lang="zh",
            source_lang="en",
            keep_original_sheets=False,
            mark_review_items=False,
        )

        workbook = load_workbook(path)
        try:
            sheet = workbook["S"]
            self.assertEqual(sheet["A1"].value, "#N/A", "错误格被改写成了双语文本")
            self.assertEqual(sheet["A1"].data_type, "e", "错误格降级成了普通文本格")
            # 别把这一条修成「整张表都不写了」。
            self.assertEqual(sheet["A2"].value, "施工内容\n[EN]content")
        finally:
            workbook.close()

    def test_error_cell_is_reported_in_the_task_log(self) -> None:
        """默默不写也不行：用户得知道那一格为什么没跟着译。"""
        path = self._workbook_with_error_cell("errors_log.xlsx")

        logs: list[str] = []
        write_bilingual_workbook(
            path,
            translations={"#N/A": "不适用", "施工内容": "[EN]content"},
            target_lang="zh",
            source_lang="en",
            keep_original_sheets=False,
            mark_review_items=False,
            log_callback=logs.append,
        )

        hits = [line for line in logs if "错误值" in line and "S!A1" in line]
        self.assertTrue(hits, f"任务日志里没有交代错误值单元格：{logs}")

    def test_formula_cell_with_a_cached_error_keeps_its_formula(self) -> None:
        """公式算出来是 #N/A 时，回填显示值那条路同样不能改写这一格。"""
        rows_xml = (
            '<row r="1">'
            '<c r="A1" t="e"><f>NA()</f><v>#N/A</v></c>'
            '<c r="B1" t="str"><v>施工内容</v></c>'
            "</row>"
        )
        source = _package_with_sheet(self.root, "formula_error", rows_xml)
        target = self.root / "formula_error_out.xlsx"
        target.write_bytes(source.read_bytes())

        write_bilingual_workbook(
            target,
            translations={"#N/A": "不适用", "施工内容": "[EN]content"},
            target_lang="zh",
            source_lang="en",
            keep_original_sheets=False,
            formula_display_value_backfill=True,
            mark_review_items=False,
        )

        cells = _sheet_cells(target)
        self.assertIsNotNone(_formula(cells["A1"]), "错误格的公式被删掉了")
        self.assertEqual(cells["A1"].get("t"), "e")
        self.assertIsNone(cells["B1"].find(f"{M}f"))

    def test_error_cell_never_enters_the_coverage_plan(self) -> None:
        """补译计划也不该收下它——收下就意味着一次真金白银的 API 调用。"""
        path = self._workbook_with_error_cell("errors_plan.xlsx")

        plan = build_excel_coverage_plan(path, target_lang="zh", source_lang="en")

        self.assertNotIn("#N/A", plan.source_texts, "错误值被排进了待译词条")
        self.assertNotIn(
            "S!A1",
            [unit.location for unit in plan.source_units],
            "错误格被判成了未翻译内容",
        )

    def test_cached_formula_error_never_enters_the_coverage_plan(self) -> None:
        """回填显示值那条路同理：公式算出来是 #N/A，缓存值也是一串字符串。

        公式闸门已按「公式显示值回填」开关放宽（开着时公式格按显示值进补译
        候选），这一格现在靠的是显示值映射（``_DisplayValues``）对错误格
        （``data_type == "e"``）的过滤——错误码那串符号不是正文，进了候选就是
        白花一次 API 调用。
        """
        rows_xml = (
            '<row r="1">'
            '<c r="A1" t="e"><f>NA()</f><v>#N/A</v></c>'
            '<c r="B1" t="str"><f>C1</f><v>施工内容</v></c>'
            "</row>"
        )
        path = _package_with_sheet(self.root, "plan_formula_error", rows_xml)

        plan = build_excel_coverage_plan(
            path,
            target_lang="zh",
            source_lang="en",
            formula_display_value_backfill=True,
        )

        self.assertNotIn("#N/A", plan.source_texts, "错误公式格被排进了待译词条")
        self.assertNotIn(
            "S!A1",
            [unit.location for unit in plan.source_units],
            "错误公式格被判成了未翻译内容",
        )

    def test_unknown_error_code_is_not_invented_into_text(self) -> None:
        """认不出的错误码不许编一个 '#ERR' 出来——那不是 Excel 的错误码，
        落进 .xlsx 就是一个普通文本格，照样会被送译。"""
        import xlrd

        self.assertIsNone(
            _cell_value_by_ctype(SimpleNamespace(ctype=xlrd.XL_CELL_ERROR, value=0x99), 0)
        )


class CoveragePlanHandleTests(unittest.TestCase):
    """低危：第二次 load 抛错时，第一个工作簿句柄没人关。"""

    def test_first_workbook_is_closed_when_the_second_load_fails(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "施工内容"
            workbook.save(source)
            workbook.close()

            closed: list[bool] = []
            real_load = openpyxl.load_workbook

            class _Tracked:
                """真工作簿的透明代理，只额外记一笔「关过没有」。"""

                def __init__(self, inner) -> None:
                    self._inner = inner

                def __getattr__(self, name):
                    return getattr(self._inner, name)

                def close(self) -> None:
                    closed.append(True)
                    self._inner.close()

            calls: list[int] = []

            def _load(*args, **kwargs):
                calls.append(1)
                if len(calls) == 1:
                    return _Tracked(real_load(*args, **kwargs))
                raise MemoryError("第二次 load 炸了")

            with mock.patch.object(openpyxl, "load_workbook", _load):
                with self.assertRaises(MemoryError):
                    excel_coverage.build_excel_coverage_plan(
                        source,
                        target_lang="en",
                        source_lang="zh",
                        formula_display_value_backfill=True,
                    )

            self.assertEqual(closed, [True], "第一个工作簿句柄泄了")


_UI_SRC = Path(__file__).resolve().parents[1] / "ui" / "src"


def _read_ui(relative: str) -> str:
    return (_UI_SRC / relative).read_text(encoding="utf-8")


def _function_body(source: str, name: str) -> str:
    """按「function NAME(...) ... 列首 }」截取函数体（见 B2 集群先例）。"""
    match = re.search(rf"(?:async )?function {re.escape(name)}\(.*?\n\}}", source, re.S)
    assert match, f"在源码里找不到函数 {name}"
    return match.group(0)


class XlsCompatibilityWarningWordingTests(unittest.TestCase):
    """后果导向文案：告警必须说清兼容转换的真实后果，且不改动原始文件。

    LibreOffice 接入后「公式变数值、样式丢失」不再对所有机器成立——按本机有没有
    LibreOffice，这里分两组钉：没有 LO 的机器维持原话，装了 LO 的机器换成新话，
    两组都必须留住「原始文件不会被改动」这句安抚（core/xls_converter.py 的
    describe_xls_compatibility_consequence 是唯一出处，见该函数文档）。
    """

    def test_aggregate_scan_message_states_the_consequence_without_libreoffice(self) -> None:
        from core.file_scanner import ExcelScanResult, FileItem

        # ExcelScanResult.risk 是个每次访问都重算的 property（不缓存），必须在同一次
        # 打桩窗口内把要断言的值一次取全，出了 with 块再访问就是在读真机的探测结果。
        item = FileItem(path=Path("legacy.xls"), name="legacy", size_kb=1.0, format="xls")
        with mock.patch("core.word_converter._find_soffice", return_value=None):
            result = ExcelScanResult(root=Path("."), items=[item])
            risk = result.risk
        self.assertFalse(risk["libreoffice_available"])
        self.assertIn("公式会变成算好的数值", risk["message"])
        self.assertIn("原始文件不会被改动", risk["message"])

    def test_aggregate_scan_message_switches_when_libreoffice_is_available(self) -> None:
        from core.file_scanner import ExcelScanResult, FileItem

        item = FileItem(path=Path("legacy.xls"), name="legacy", size_kb=1.0, format="xls")
        with mock.patch(
            "core.word_converter._find_soffice",
            return_value="/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ):
            result = ExcelScanResult(root=Path("."), items=[item])
            risk = result.risk
        self.assertTrue(risk["libreoffice_available"])
        self.assertIn("LibreOffice", risk["message"])
        self.assertIn("原始文件不会被改动", risk["message"])
        # 装了 LO 之后不该再吓唬用户说公式一定会被拍死成数值。
        self.assertNotIn("公式会变成算好的数值", risk["message"])

    @unittest.skipUnless(HAS_XLWT, "本机没有 xlwt，无法现场生成 .xls 夹具")
    def test_single_file_risk_message_states_the_consequence_without_libreoffice(self) -> None:
        import xlwt

        from core.file_scanner import _build_file_item

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy.xls"
            book = xlwt.Workbook()
            book.add_sheet("S").write(0, 0, "施工内容")
            book.save(str(path))

            with mock.patch("core.word_converter._find_soffice", return_value=None):
                item = _build_file_item(path)
            message = item.risk["message"]
            self.assertIn("公式会变成算好的数值", message)
            self.assertIn("原始文件不会被改动", message)

    def test_ui_modal_pins_both_excel_wording_variants_and_keeps_word_wording(self) -> None:
        source = _read_ui("views/workspace.ts")
        body = _function_body(source, "showCompatibilityModal")
        # 互审抓过两处假钉子：整个函数体里搜 ".doc" 连变量声明都能满足，Word 文案
        # 全删都不红；不切分支的话 Excel/Word 文案整个调包也不红。所以这里必须
        # ①钉死三元写的是 ===（写成 !== 等于把两套文案调包），②按分支切开各钉各的。
        self.assertIn('surface === "excel"', body)
        _, _, arms = body.partition('surface === "excel"')
        excel_arm, sep, word_arm = arms.partition(": [")
        self.assertTrue(sep, "在函数体里找不到三元的 Word 分支")
        self.assertIn("原始文件不会被改动", excel_arm)
        # 「允许兼容转换」这一段本身又按 libreofficeAvailable 分了两个变体，两句都
        # 必须原样留在 Excel 分支里——只留一句就是把某一类机器的用户晾在旧文案上。
        self.assertIn("本机检测到 LibreOffice", excel_arm)
        self.assertIn("公式会变成算好的数值", excel_arm)
        self.assertIn(
            "复杂样式、合并单元格、图片、图表和宏可能无法完整保留",
            word_arm,
            "Word 分支的既有文案不该被误改",
        )

    def test_backend_error_messages_call_the_shared_consequence_helper(self) -> None:
        """弹窗后紧接着可能出现的两条后端报错必须和弹窗说同一套话。

        互审抓到的矛盾：用户刚在弹窗里读完一套说法，几秒后预检失败的报错又说
        另一套——同一个选择两种说法。LibreOffice 接入后已经没有唯一一句「写死
        的话」可以整段 grep 了（后果按本机有没有 LO 分叉，见
        describe_xls_compatibility_consequence），所以这里改成两层校验：
        ①共享函数本身的两个分支都留着「原始文件不会被改动」；
        ②两处报错源文件真的在调用这个共享函数，而不是各自写死一份新句子。
        """
        for has_libreoffice in (True, False):
            consequence = xls_converter.describe_xls_compatibility_consequence(
                has_libreoffice=has_libreoffice
            )
            with self.subTest(has_libreoffice=has_libreoffice):
                self.assertIn("原始文件不会被改动", consequence)

        repo = Path(__file__).resolve().parents[1]
        for rel in ("api/task_manager.py", "core/xls_converter.py"):
            source = (repo / rel).read_text(encoding="utf-8")
            with self.subTest(rel):
                self.assertIn("describe_xls_compatibility_consequence(", source)
                self.assertNotIn("兼容转换可能损失", source)


if __name__ == "__main__":
    unittest.main()
