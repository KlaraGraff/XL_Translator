"""Word 译文排版三处小修：句中软换行并句、固定行高放开、东亚断行语言。

docx 都用 python-docx 真实构造，断言落在输出文件的 XML 上。
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from core.word_document import extract_word_segments, write_bilingual_docx
from core.word_layout_fixes import (
    _is_wrapped_break,
    apply_east_asian_language,
    east_asian_locale,
    join_wrapped_line_breaks,
)
from core.word_task_runner import _prepare_word_source_for_translation


def _paragraph_with_breaks(doc, *lines: str, break_type=None):
    paragraph = doc.add_paragraph()
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        if index < len(lines) - 1:
            if break_type is None:
                run.add_break()
            else:
                run.add_break(break_type)
    return paragraph


LONG_FR_HEAD = "Les travaux de terrassement seront réalisés conformément aux"
LONG_ZH_HEAD = "本工程土方开挖应严格按照经审批的专项施工方案组织实施并做好基坑"


class WrappedBreakRuleTests(unittest.TestCase):
    def test_mid_sentence_wrap_is_joined(self) -> None:
        self.assertTrue(_is_wrapped_break(LONG_FR_HEAD, "prescriptions du CCTP."))
        self.assertTrue(_is_wrapped_break(LONG_ZH_HEAD, "监测工作。"))
        # 后一行小写开头，前一行不必排满
        self.assertTrue(_is_wrapped_break("La présente note décrit", "les mesures prises."))

    def test_intentional_breaks_are_kept(self) -> None:
        cases = [
            ("12 rue de la Paix", "75002 Paris"),  # 地址
            ("Le Directeur", "Jean Dupont"),  # 落款
            ("名称", "Name"),  # 双语表头
            ("Chapitre 3", "Installation électrique"),  # 标题
            (LONG_FR_HEAD + ".", "Les prescriptions suivantes"),  # 前一行已是整句
            (LONG_FR_HEAD, "- premier point"),  # 列表项
            (LONG_ZH_HEAD, "1. 第一条"),
            (LONG_ZH_HEAD, "（一）总则"),
            ("Le présent document est établi conformément à l'instal-", "lation"),  # 断词
            ("项目名称", "盛泰纺织"),  # 短的中文行
        ]
        for head, tail in cases:
            with self.subTest(head=head, tail=tail):
                self.assertFalse(_is_wrapped_break(head, tail))


class JoinWrappedLineBreaksTests(unittest.TestCase):
    def test_joins_only_mid_sentence_breaks(self) -> None:
        doc = Document()
        wrapped = _paragraph_with_breaks(doc, LONG_FR_HEAD, "prescriptions du CCTP.")
        zh = _paragraph_with_breaks(doc, LONG_ZH_HEAD, "监测工作。")
        address = _paragraph_with_breaks(doc, "12 rue de la Paix", "75002 Paris")
        page = _paragraph_with_breaks(doc, LONG_FR_HEAD, "prescriptions.", break_type=WD_BREAK.PAGE)
        table = doc.add_table(rows=1, cols=1)
        cell_paragraph = table.cell(0, 0).paragraphs[0]
        cell_paragraph.add_run(LONG_FR_HEAD).add_break()
        cell_paragraph.add_run("prescriptions du CCTP.")

        self.assertEqual(join_wrapped_line_breaks(doc), 3)
        self.assertEqual(wrapped.text, LONG_FR_HEAD + " prescriptions du CCTP.")
        self.assertEqual(zh.text, LONG_ZH_HEAD + "监测工作。")
        self.assertEqual(address.text, "12 rue de la Paix\n75002 Paris")
        self.assertEqual(len(page._p.findall(".//" + qn("w:br"))), 1)
        self.assertEqual(cell_paragraph.text, LONG_FR_HEAD + " prescriptions du CCTP.")

    def test_tab_aligned_lines_and_fields_are_left_alone(self) -> None:
        doc = Document()
        tabbed = _paragraph_with_breaks(doc, "Maître d'ouvrage\tSociété Générale de Travaux du Maroc", "suite")
        field = _paragraph_with_breaks(doc, LONG_FR_HEAD, "prescriptions du CCTP.")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        field._p.append(fld)

        self.assertEqual(join_wrapped_line_breaks(doc), 0)
        self.assertIn("\n", tabbed.text)
        self.assertIn("\n", field.text)

    def test_prepared_working_copy_carries_whole_sentences(self) -> None:
        """抽取和写回读的是同一份工作副本：模型拿到整句，原件不改。"""
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "wrapped.docx"
            doc = Document()
            _paragraph_with_breaks(doc, LONG_FR_HEAD, "prescriptions du CCTP.")
            doc.save(source)

            prepared = _prepare_word_source_for_translation(source, use_native_preprocessing=False)
            try:
                segments = extract_word_segments(prepared.path, target_lang="zh", source_lang="fr")
                self.assertEqual([s.source for s in segments], [LONG_FR_HEAD + " prescriptions du CCTP."])
                self.assertIn("\n", Document(str(source)).paragraphs[0].text)
            finally:
                for path in prepared.temp_paths:
                    path.unlink(missing_ok=True)


def _exact_row_table(doc, text: str):
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = text
    table.cell(0, 1).text = "0001"
    table.rows[0].height = Pt(14)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.find(qn("w:trHeight")).set(qn("w:hRule"), "exact")
    return table


def _row_rule(doc, index: int = 0) -> str | None:
    return doc.tables[index].rows[0]._tr.find(qn("w:trPr")).find(qn("w:trHeight")).get(qn("w:hRule"))


class RowHeightTests(unittest.TestCase):
    def _write(self, root: Path, doc, translations: dict[str, str], *, translation_only: bool) -> Document:
        source = root / "rows.docx"
        doc.save(source)
        out = write_bilingual_docx(
            source_path=source,
            output_dir=root / "out",
            translations=translations,
            target_lang="fr",
            source_lang="zh",
            output_translation_only=translation_only,
        )
        return Document(str(out))

    def test_translated_rows_grow_with_the_translation(self) -> None:
        for translation_only in (False, True):
            with self.subTest(translation_only=translation_only), tempfile.TemporaryDirectory() as temporary:
                doc = Document()
                _exact_row_table(doc, "基坑支护")
                _exact_row_table(doc, "0002")  # 没有要翻译的内容
                out = self._write(
                    Path(temporary), doc, {"基坑支护": "Soutènement de la fouille"}, translation_only=translation_only
                )
                self.assertEqual(_row_rule(out, 0), "atLeast")
                self.assertEqual(_row_rule(out, 1), "exact")
                height = out.tables[0].rows[0]._tr.find(qn("w:trPr")).find(qn("w:trHeight"))
                self.assertEqual(height.get(qn("w:val")), str(14 * 20))


class EastAsianLanguageTests(unittest.TestCase):
    def test_locale_mapping(self) -> None:
        self.assertEqual(east_asian_locale("zh"), "zh-CN")
        self.assertEqual(east_asian_locale("zh_TW"), "zh-TW")
        self.assertEqual(east_asian_locale("ja"), "ja-JP")
        self.assertEqual(east_asian_locale("ko"), "ko-KR")
        self.assertEqual(east_asian_locale("fr"), "")
        self.assertEqual(east_asian_locale(""), "")

    def _default_east_asian(self, doc) -> str | None:
        lang = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:lang"))
        return None if lang is None else lang.get(qn("w:eastAsia"))

    def test_chinese_output_uses_chinese_line_breaking(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            doc = Document()
            doc.add_paragraph("The contractor shall submit the method statement.")
            # 正文 run 上显式写了西文东亚语言的，也一并改
            run = doc.add_paragraph().add_run("Safety first.")
            run_lang = OxmlElement("w:lang")
            run_lang.set(qn("w:val"), "en-GB")
            run_lang.set(qn("w:eastAsia"), "en-GB")
            run._r.get_or_add_rPr().append(run_lang)
            source = root / "en.docx"
            doc.save(source)
            self.assertNotEqual(self._default_east_asian(Document(str(source))), "zh-CN")

            out = Document(
                str(
                    write_bilingual_docx(
                        source_path=source,
                        output_dir=root / "out",
                        translations={
                            "The contractor shall submit the method statement.": "承包商应提交施工方案。",
                            "Safety first.": "安全第一。",
                        },
                        target_lang="zh",
                        source_lang="en",
                    )
                )
            )
            self.assertEqual(self._default_east_asian(out), "zh-CN")
            body_langs = [lang.get(qn("w:eastAsia")) for lang in out.element.body.iter(qn("w:lang"))]
            self.assertIn("zh-CN", body_langs)
            self.assertNotIn("en-GB", body_langs)

    def test_latin_targets_are_untouched(self) -> None:
        doc = Document()
        before = self._default_east_asian(doc)
        self.assertEqual(apply_east_asian_language(doc, "fr"), 0)
        self.assertEqual(self._default_east_asian(doc), before)

    def test_defaults_follow_target_but_east_asian_runs_keep_their_language(self) -> None:
        doc = Document()
        lang = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:lang"))
        lang.set(qn("w:eastAsia"), "ja-JP")
        run = doc.add_paragraph().add_run("本工程")
        run_lang = OxmlElement("w:lang")
        run_lang.set(qn("w:eastAsia"), "zh-CN")
        run._r.get_or_add_rPr().append(run_lang)

        apply_east_asian_language(doc, "ko")
        self.assertEqual(self._default_east_asian(doc), "ko-KR")
        self.assertEqual(run_lang.get(qn("w:eastAsia")), "zh-CN")

    def test_missing_doc_defaults_are_created(self) -> None:
        doc = Document()
        defaults = doc.styles.element.find(qn("w:docDefaults"))
        doc.styles.element.remove(defaults)
        self.assertGreater(apply_east_asian_language(doc, "zh"), 0)
        self.assertEqual(self._default_east_asian(doc), "zh-CN")


if __name__ == "__main__":
    unittest.main()
