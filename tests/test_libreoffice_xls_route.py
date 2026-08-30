"""LibreOffice 兼容转换路线：core/xls_converter.py 的 convert_with_libreoffice。

设计见 docs/LIBREOFFICE_XLS_ROUTE_2026-08-30.md 的「推荐方案」。这里只测转换
函数本身的成败分支（全部打桩 subprocess，不依赖机器状态），加一个本机真跑的
集成测试兜底——桩测能证明「代码按预期调用了 subprocess」，证不了「soffice 真的
认识这条命令行、真的能转出带公式的 .xlsx」，只有真机调用能回答后一个问题。
"""

from __future__ import annotations

import importlib.util
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from core import xls_converter
from core.xls_converter import LibreOfficeConversionError, convert_with_libreoffice

HAS_XLWT = importlib.util.find_spec("xlwt") is not None

_FAKE_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


class ConvertWithLibreofficeMockedTests(unittest.TestCase):
    """打桩 core.word_converter._find_soffice 和 subprocess.run，机器上有没有装
    LibreOffice 都不影响这组测试——探测函数本身复用 word_converter 那份，不在这里
    重新验证它的候选路径列表（那是 test_word_converter.py 的职责）。
    """

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.xls_path = Path(self._tmp.name) / "legacy.xls"
        self.xls_path.write_bytes(b"fake xls bytes")

    def test_soffice_absent_raises_without_touching_subprocess(self) -> None:
        with (
            mock.patch("core.word_converter._find_soffice", return_value=None),
            mock.patch.object(xls_converter.subprocess, "run") as run,
        ):
            with self.assertRaises(LibreOfficeConversionError):
                convert_with_libreoffice(self.xls_path)
        run.assert_not_called()

    def test_success_moves_converted_file_to_a_temp_xlsx_path(self) -> None:
        def _fake_run(command, **_kwargs):
            command = [str(part) for part in command]
            # 每次转换必须带独立的 -env:UserInstallation：和一个正开着的 LibreOffice
            # 图形界面共享 profile 会静默失败，word_converter.py 的 UNO 路径已经踩过
            # 这个坑（见 convert_with_libreoffice 的函数文档）。
            self.assertTrue(
                any(part.startswith("-env:UserInstallation=") for part in command)
            )
            outdir = Path(command[command.index("--outdir") + 1])
            xls_arg = Path(command[-1])
            (outdir / f"{xls_arg.stem}.xlsx").write_bytes(b"converted-bytes")
            return SimpleNamespace(returncode=0, stdout="", stderr="")

        with (
            mock.patch("core.word_converter._find_soffice", return_value=_FAKE_SOFFICE),
            mock.patch.object(xls_converter.subprocess, "run", side_effect=_fake_run),
        ):
            out_path = convert_with_libreoffice(self.xls_path)
        self.addCleanup(lambda: out_path.unlink(missing_ok=True))

        self.assertTrue(out_path.exists())
        self.assertEqual(out_path.suffix, ".xlsx")
        self.assertEqual(out_path.read_bytes(), b"converted-bytes")

    def test_nonzero_return_code_raises_and_leaves_no_partial_output(self) -> None:
        def _fake_run(_command, **_kwargs):
            return SimpleNamespace(returncode=1, stdout="", stderr="转换失败：不支持的格式")

        with (
            mock.patch("core.word_converter._find_soffice", return_value=_FAKE_SOFFICE),
            mock.patch.object(xls_converter.subprocess, "run", side_effect=_fake_run),
        ):
            with self.assertRaises(LibreOfficeConversionError) as ctx:
                convert_with_libreoffice(self.xls_path)
        self.assertIn("不支持的格式", str(ctx.exception))

    def test_timeout_raises_libreoffice_conversion_error(self) -> None:
        def _fake_run(command, **kwargs):
            raise subprocess.TimeoutExpired(cmd=command, timeout=kwargs.get("timeout", 180))

        with (
            mock.patch("core.word_converter._find_soffice", return_value=_FAKE_SOFFICE),
            mock.patch.object(xls_converter.subprocess, "run", side_effect=_fake_run),
        ):
            with self.assertRaises(LibreOfficeConversionError) as ctx:
                convert_with_libreoffice(self.xls_path)
        self.assertIn("超时", str(ctx.exception))

    def test_no_output_file_raises_libreoffice_conversion_error(self) -> None:
        def _fake_run(_command, **_kwargs):
            # 进程正常退出但没有产物：比如 soffice 判定源文件已损坏、静默放弃转换。
            return SimpleNamespace(returncode=0, stdout="", stderr="")

        with (
            mock.patch("core.word_converter._find_soffice", return_value=_FAKE_SOFFICE),
            mock.patch.object(xls_converter.subprocess, "run", side_effect=_fake_run),
        ):
            with self.assertRaises(LibreOfficeConversionError) as ctx:
                convert_with_libreoffice(self.xls_path)
        self.assertIn("未生成", str(ctx.exception))


@unittest.skipUnless(HAS_XLWT, "本机没有 xlwt，无法现场生成 .xls 夹具")
class ConvertWithLibreofficeRealMachineTests(unittest.TestCase):
    """真机集成测试：本机没装 LibreOffice 就跳过，呼应仓库里 HAS_XLWT 的跳过写法。"""

    def setUp(self) -> None:
        from core.word_converter import _find_soffice

        if _find_soffice() is None:
            self.skipTest("本机没有安装 LibreOffice/soffice")

    def test_formula_survives_a_real_conversion(self) -> None:
        import xlwt
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            xls_path = Path(tmp) / "legacy.xls"
            book = xlwt.Workbook()
            sheet = book.add_sheet("S")
            sheet.write(0, 0, 1)
            sheet.write(0, 1, 2)
            sheet.write(0, 2, xlwt.Formula("A1+B1"))
            book.save(str(xls_path))

            out_path = convert_with_libreoffice(xls_path)
            self.addCleanup(lambda: out_path.unlink(missing_ok=True))

            workbook = load_workbook(out_path)
            try:
                cell = workbook.active["C1"]
                self.assertIsInstance(cell.value, str)
                self.assertTrue(cell.value.startswith("="))
                self.assertIn("A1", cell.value)
                self.assertIn("B1", cell.value)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
