"""审计批次 2 · 第③条：replace 模式含域的正文段落改成原地替换。

``_replace_paragraph_text_around_fields`` 已经在页眉页脚路径上用了一轮（见
tests/test_audit_word_doc_fixes.py 里的 HeaderFieldReplacementTests /
FieldSplitAmbiguityTests），带着超链接保护、制表位保护、域配平校验这些从
WD1 复审里揪出来的修复。这次是把正文 replace 分支从"含域一律插行"改成
"配平就原地替换、配不平才插行"，复用的是同一个函数——这里要验证的不是函数
本身（已经测过），而是正文分流接进来之后：域还在、周边文字被译文顶替、
插行兜底没丢、review mark 和编号前缀这两个正文特有的伴随动作在新路径上
仍然生效。docx 全部用真实 XML 结构构造（w:fldChar 三件套、真实 w:hyperlink），
断言落在 XML 和 python-docx 视图两层。
"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.mixed_language import MIXED_MARK_UNRESOLVED
from core.translation_protocol import REPLACE_TRANSLATION_PREFIX
from core.word_document import write_bilingual_docx


def _replace(text: str) -> str:
    return f"{REPLACE_TRANSLATION_PREFIX}{text}"


def _append_field(paragraph, instruction: str, *, result: str) -> None:
    """在段落末尾接一个真正的域：begin / instrText / separate / 缓存结果 / end。"""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    code_run = paragraph.add_run()
    code_run._r.append(begin)
    code_run._r.append(instr)
    code_run._r.append(separate)
    paragraph.add_run(result)
    end_run = paragraph.add_run()
    end_run._r.append(end)


def _add_hyperlink(paragraph, text: str, url: str):
    """在段落末尾追加一个真正的 w:hyperlink（带外部关系），而不是伪造的 run。"""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    text_element.set(qn("xml:space"), "preserve")
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _paragraph_highlight_values(paragraph) -> set[str]:
    values: set[str] = set()
    for run in paragraph._p.iter(qn("w:r")):
        r_pr = run.find(qn("w:rPr"))
        if r_pr is None:
            continue
        highlight = r_pr.find(qn("w:highlight"))
        if highlight is not None:
            values.add(str(highlight.get(qn("w:val")) or ""))
    return values


class BodyFieldReplacementTests(unittest.TestCase):
    """含域正文段落：域配平就原地替换，域和作者文字各归各位。"""

    def _build_body_with_seq_caption(self, path: Path) -> str:
        doc = Document()
        doc.add_paragraph("正文说明")
        caption = doc.add_paragraph()
        caption.add_run("表 ")
        _append_field(caption, " SEQ Table \\* ARABIC ", result="1")
        caption.add_run(" 施工进度总表")
        doc.save(str(path))
        return Document(str(path)).paragraphs[1].text

    def test_body_paragraph_with_seq_field_is_replaced_in_place(self) -> None:
        """域仍旧原地生效，作者写的那部分文字被译文顶掉，不再在下一行插一段译文。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "seq_caption.docx"
            line = self._build_body_with_seq_caption(source_path)
            self.assertEqual(line, "表 1 施工进度总表")

            paragraph_count_before = len(Document(str(source_path)).paragraphs)

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Table 1 Calendrier general des travaux")},
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))
            caption = out_doc.paragraphs[1]

            self.assertEqual(caption.text, "Table 1 Calendrier general des travaux")
            self.assertNotIn("施工进度总表", caption.text)
            # 原地替换：不许在题注后面多出一段插入的译文段落。
            self.assertEqual(len(out_doc.paragraphs), paragraph_count_before)
            # 域本体（SEQ 指令 + begin/end）必须原样还在，否则编号变成一次性死文本。
            xml = caption._p.xml
            self.assertIn("SEQ", xml)
            self.assertIn('w:fldCharType="begin"', xml)
            self.assertIn('w:fldCharType="end"', xml)

    def test_unbalanced_field_in_body_falls_back_to_inserting_a_line(self) -> None:
        """域的 begin/end 不配平时读不准边界，正文一样退回插行，不强行原地改写。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "broken_seq.docx"
            doc = Document()
            doc.add_paragraph("正文说明")
            caption = doc.add_paragraph()
            # 域一旦 begin 没有匹配的 end，_paragraph_literal_text 的 depth 从此锁死在 1，
            # 域后面的文字全部被算作"域内"而不计入可译字数——域前的文字得自己够
            # _FIELD_PARAGRAPH_MIN_CJK_CHARS（4 个）的门槛，这段落才不会在
            # _is_toc_or_field_paragraph 那道闸门被当成纯域段落直接跳过。
            caption.add_run("工程进度统计表 第 ")
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instr = OxmlElement("w:instrText")
            instr.text = " SEQ Table \\* ARABIC "
            broken = caption.add_run()
            broken._r.append(begin)
            broken._r.append(instr)
            caption.add_run("1 号")
            doc.save(str(source_path))
            line = Document(str(source_path)).paragraphs[1].text

            paragraph_count_before = len(Document(str(source_path)).paragraphs)

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Table 1 Calendrier general des travaux")},
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))

            # 插行兜底：原段落原文照留，译文另起一段紧跟在后面。
            self.assertEqual(len(out_doc.paragraphs), paragraph_count_before + 1)
            self.assertIn("工程进度统计表", out_doc.paragraphs[1].text)
            self.assertIn("1 号", out_doc.paragraphs[1].text)
            self.assertIn("SEQ", out_doc.paragraphs[1]._p.xml)
            self.assertEqual(
                out_doc.paragraphs[2].text,
                "Table 1 Calendrier general des travaux",
            )

    def test_hyperlink_wrapping_a_field_in_body_falls_back_instead_of_corrupting_link(
        self,
    ) -> None:
        """域被包在超链接内部时不硬替换：超链接连同域必须一字不改地留在原位。

        这正是 WD1 复审追加修复的那类风险（"域段落替换绕过超链接保护"）——
        ``_iter_paragraph_content_nodes`` 会把超链接内的 run 摊平成顶层节点，可
        ``_paragraph_text_anchor_run`` 的插入点在整个 w:hyperlink 之前，域之后的
        文字会被错误地搬到域前面。函数对这种结构直接返回 False，交回插行兜底，
        正文分流要保证真的落到这条兜底上，而不是把超链接结构写坏。
        """
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "hyperlink_around_field.docx"
            doc = Document()
            doc.add_paragraph("正文说明")
            caption = doc.add_paragraph()
            hyperlink = _add_hyperlink(caption, "前缀 ", "https://example.com/spec")
            _append_field(caption, " SEQ Table \\* ARABIC ", result="1")
            tail = caption.add_run(" 尾巴")
            # 域三件套和尾巴文字先落在段落级，再一起搬进超链接内部，构造出
            # "域被超链接整个包住"的真实结构。
            for run in list(caption._p.iterchildren(qn("w:r"))):
                hyperlink.append(run)
            self.assertIs(tail._r.getparent(), hyperlink)
            doc.save(str(source_path))
            line = Document(str(source_path)).paragraphs[1].text
            self.assertEqual(line, "前缀 1 尾巴")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Prefixe 1 queue")},
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))
            caption = out_doc.paragraphs[1]

            # 走的是追加回退：原文（含超链接）原样留着，译文另起一段。
            self.assertTrue(caption.text.startswith("前缀 1 尾巴"))
            self.assertEqual(len(caption._p.findall(qn("w:hyperlink"))), 1)
            self.assertEqual(out_doc.paragraphs[2].text, "Prefixe 1 queue")

    def test_hyperlink_beside_a_field_does_not_become_a_translated_link(self) -> None:
        """超链接不包住域、只是紧挨着域时：原地替换生效，译文不会被写进链接里。

        译文对应超链接那一段的部分必须落进段落级的普通 run（而不是直接写进
        w:hyperlink 内部的 run），否则整句译文就变成了一个指向原 URL、语义不相干
        的可点击区域——``_paragraph_text_anchor_run`` 专门防的就是这个。
        """
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "hyperlink_beside_field.docx"
            doc = Document()
            doc.add_paragraph("正文说明")
            caption = doc.add_paragraph()
            _add_hyperlink(caption, "标准 ", "https://example.com/std")
            _append_field(caption, " SEQ Table \\* ARABIC ", result="2")
            caption.add_run(" 条款")
            doc.save(str(source_path))
            line = Document(str(source_path)).paragraphs[1].text
            self.assertEqual(line, "标准 2 条款")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Norme 2 clause")},
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))
            caption = out_doc.paragraphs[1]

            self.assertEqual(caption.text, "Norme 2 clause")
            self.assertIn("SEQ", caption._p.xml)
            # 超链接那部分内容已经被顶替译文覆盖、清空后按既有规则删除——留下一个
            # 零长度却仍可点击的区域比没有链接更糟。数量为 0 同时保证了「译文不会
            # 残留在超链接内部的 run 里」：连 w:hyperlink 元素都没有了。
            self.assertEqual(len(caption._p.findall(qn("w:hyperlink"))), 0)

    def test_review_mark_applies_to_field_paragraph_replaced_in_place(self) -> None:
        """复核底色在新路径上照样生效：原地替换之后再补一遍高亮，而不是被域绕开。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "seq_caption_review.docx"
            line = self._build_body_with_seq_caption(source_path)

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Table 1 Calendrier general des travaux")},
                target_lang="fr",
                source_lang="zh",
                review_marks={line: MIXED_MARK_UNRESOLVED},
            )
            out_doc = Document(str(out_path))
            caption = out_doc.paragraphs[1]

            self.assertEqual(caption.text, "Table 1 Calendrier general des travaux")
            self.assertIn("SEQ", caption._p.xml)
            highlights = _paragraph_highlight_values(caption)
            self.assertTrue(
                highlights and "none" not in {value.lower() for value in highlights},
                f"含域正文段原地替换后应带复核高亮，实际 highlight 值：{highlights!r}",
            )

    def test_leading_prefix_survives_field_paragraph_replaced_in_place(self) -> None:
        """段落原本带的行首空白前缀，在新路径上照样补回译文前面，不会被域分流丢掉。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "seq_caption_prefix.docx"
            doc = Document()
            doc.add_paragraph("正文说明")
            caption = doc.add_paragraph()
            # 行首留一个空格模拟"人工敲的缩进前缀"：_paragraph_source_text 用 strip()
            # 求翻译查表的 key，这个空格因此不在 key 里，只能靠 leading_prefix 补回来。
            caption.add_run(" 表 ")
            _append_field(caption, " SEQ Table \\* ARABIC ", result="1")
            caption.add_run(" 施工进度总表")
            doc.save(str(source_path))
            raw_line = Document(str(source_path)).paragraphs[1].text
            self.assertTrue(raw_line.startswith(" "))
            lookup_key = raw_line.strip()

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={
                    lookup_key: _replace("Table 1 Calendrier general des travaux")
                },
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))
            caption = out_doc.paragraphs[1]

            self.assertEqual(
                caption.text,
                " Table 1 Calendrier general des travaux",
            )
            self.assertIn("SEQ", caption._p.xml)


if __name__ == "__main__":
    unittest.main()
