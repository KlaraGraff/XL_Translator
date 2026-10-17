"""Excel 文本框和批注的抽取与写回（core/xlsx_embedded_text.py）。"""

from __future__ import annotations

import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import Workbook
from openpyxl.comments import Comment

from core.xlsx_embedded_text import (
    KIND_COMMENT,
    KIND_SHAPE,
    KIND_THREADED_COMMENT,
    REVIEW_CATEGORY_SHAPE_OVERFLOW,
    EmbeddedTextOptions,
    collect_embedded_units,
)
from core.xlsx_patcher import write_bilingual_workbook
from tests.rate_pins import pinned_surface_rate

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_THREADED = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
REL_THREADED = "http://schemas.microsoft.com/office/2017/10/relationships/threadedComment"

THREADED_FALLBACK = (
    "[Threaded comment]\n\nYour version of Excel allows you to read this threaded "
    "comment; however, any edits to it will get removed if the file is opened in a "
    "newer version of Excel. Learn more: https://go.microsoft.com/fwlink/?linkid=870924"
    "\n\nComment:\n    对话批注原文"
)


def _drawing_xml(
    paragraphs: list[str],
    *,
    name: str = "TextBox 1",
    box: tuple[int, int] = (1828800, 762000),
    list_style: str = "",
    anchor_box: tuple[int, int] | None = None,
) -> bytes:
    anchor_box = anchor_box or box
    body = "".join(
        f'<a:p><a:r><a:rPr lang="zh-CN" sz="1100"/><a:t>{text}</a:t></a:r></a:p>'
        for text in paragraphs
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<xdr:wsDr xmlns:xdr="{NS_XDR}" xmlns:a="{NS_A}" xmlns:r="{NS_REL}">'
        '<xdr:twoCellAnchor><xdr:from><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff>'
        "<xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
        # 右下角锚点按框的大小放：Excel 按锚点定两格锚定的形状有多大。
        f"<xdr:to><xdr:col>3</xdr:col><xdr:colOff>{anchor_box[0]}</xdr:colOff>"
        f"<xdr:row>1</xdr:row><xdr:rowOff>{anchor_box[1]}</xdr:rowOff></xdr:to>"
        '<xdr:sp macro="" textlink="">'
        f'<xdr:nvSpPr><xdr:cNvPr id="2" name="{name}"/><xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
        '<xdr:spPr><a:xfrm><a:off x="0" y="0"/>'
        f'<a:ext cx="{box[0]}" cy="{box[1]}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
        f'<xdr:txBody><a:bodyPr wrap="square"/><a:lstStyle>{list_style}</a:lstStyle>{body}</xdr:txBody>'
        "</xdr:sp><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>"
    ).encode()


def _threaded_xml() -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<ThreadedComments xmlns="{NS_THREADED}">'
        '<threadedComment ref="C3" dT="2026-01-01T00:00:00Z" personId="{P}" id="{C1}">'
        "<text>对话批注原文</text></threadedComment></ThreadedComments>"
    ).encode()


def build_workbook(
    path: Path,
    *,
    threaded: bool = True,
    box: tuple[int, int] = (1828800, 762000),
    list_style: str = "",
    cells: dict[str, str] | None = None,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "表1"
    ws["A1"] = "单元格原文"
    for ref, value in (cells or {}).items():
        ws[ref] = value
    ws["B2"] = "有批注"
    ws["B2"].comment = Comment("张三:\n这里要核对", "张三")
    if threaded:
        ws["C3"] = "有对话批注"
        ws["C3"].comment = Comment(THREADED_FALLBACK, "tc={C1}")
    wb.save(path)

    with zipfile.ZipFile(path) as source:
        items = {name: source.read(name) for name in source.namelist()}

    sheet_part = "xl/worksheets/sheet1.xml"
    rels_part = "xl/worksheets/_rels/sheet1.xml.rels"
    sheet = etree.fromstring(items[sheet_part])
    drawing = etree.Element(f"{{{NS_MAIN}}}drawing")
    drawing.set(f"{{{NS_REL}}}id", "rIdDr1")
    legacy = sheet.find(f"{{{NS_MAIN}}}legacyDrawing")
    legacy.addprevious(drawing)
    items[sheet_part] = etree.tostring(sheet, xml_declaration=True, encoding="UTF-8", standalone=True)

    rels = etree.fromstring(items[rels_part])
    rel = etree.SubElement(rels, f"{{{NS_PKG_REL}}}Relationship")
    rel.set("Id", "rIdDr1")
    rel.set("Type", f"{NS_REL}/drawing")
    rel.set("Target", "../drawings/drawing1.xml")
    if threaded:
        rel = etree.SubElement(rels, f"{{{NS_PKG_REL}}}Relationship")
        rel.set("Id", "rIdTc1")
        rel.set("Type", REL_THREADED)
        rel.set("Target", "../threadedComments/threadedComment1.xml")
    items[rels_part] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    items["xl/drawings/drawing1.xml"] = _drawing_xml(["图例说明", "红色表示待施工区域"], box=box, list_style=list_style)
    content_types = items["[Content_Types].xml"].decode()
    overrides = (
        '<Override PartName="/xl/drawings/drawing1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
    )
    if threaded:
        items["xl/threadedComments/threadedComment1.xml"] = _threaded_xml()
        overrides += (
            '<Override PartName="/xl/threadedComments/threadedComment1.xml" '
            'ContentType="application/vnd.ms-excel.threadedcomments+xml"/>'
        )
    items["[Content_Types].xml"] = content_types.replace("</Types>", overrides + "</Types>").encode()

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in items.items():
            target.writestr(name, data)


TRANSLATIONS = {
    "单元格原文": "Texte de cellule",
    "图例说明": "Légende",
    "红色表示待施工区域": "Le rouge indique les zones où les travaux restent à réaliser",
    "这里要核对": "À vérifier ici",
    "对话批注原文": "Commentaire de discussion",
}


def _part(path: Path, name: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(name).decode("utf-8")


def _row_heights(path: Path, part: str = "xl/worksheets/sheet1.xml") -> dict[int, float]:
    root = etree.fromstring(_part(path, part).encode())
    return {
        int(row.get("r")): float(row.get("ht"))
        for row in root.iter(f"{{{NS_MAIN}}}row")
        if row.get("ht")
    }


def _shape_paragraphs(path: Path, part: str = "xl/drawings/drawing1.xml") -> list[str]:
    root = etree.fromstring(_part(path, part).encode())
    result = []
    for paragraph in root.iter(f"{{{NS_A}}}p"):
        result.append("".join(t.text or "" for t in paragraph.iter(f"{{{NS_A}}}t")))
    return result


class EmbeddedTextTests(unittest.TestCase):
    def setUp(self) -> None:
        self.root = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.root, True)
        self.source = self.root / "source.xlsx"
        build_workbook(self.source)

    def _write(
        self, *, translation_only=False, keep=False, review_positions=None, stats=None, allowed=None, lock=False
    ):
        output = self.root / f"out-{translation_only}-{keep}.xlsx"
        shutil.copy2(self.source, output)
        write_bilingual_workbook(
            output,
            translations={"单元格原文": TRANSLATIONS["单元格原文"]},
            target_lang="fr",
            keep_original_sheets=keep,
            lock_row_height=lock,
            output_translation_only=translation_only,
            review_positions=review_positions,
            stats=stats,
            embedded_text=EmbeddedTextOptions(
                shapes=True, comments=True, translations=dict(TRANSLATIONS), allowed_keys=allowed
            ),
        )
        return output

    def test_collect_lists_shape_paragraphs_and_comments_once(self) -> None:
        units = collect_embedded_units(self.source, shapes=True, comments=True)
        by_kind = {}
        for unit in units:
            by_kind.setdefault(unit.kind, []).append(unit.text)
        self.assertEqual(by_kind[KIND_SHAPE], ["图例说明", "红色表示待施工区域"])
        # 传统批注去掉了作者署名；对话式批注的兼容副本不单独列出。
        self.assertEqual(by_kind[KIND_COMMENT], ["这里要核对"])
        self.assertEqual(by_kind[KIND_THREADED_COMMENT], ["对话批注原文"])
        self.assertEqual({unit.sheet for unit in units}, {"表1"})
        self.assertEqual(collect_embedded_units(self.source, shapes=False, comments=False), [])
        only_shapes = collect_embedded_units(self.source, shapes=True, comments=False)
        self.assertEqual({unit.kind for unit in only_shapes}, {KIND_SHAPE})

    def test_bilingual_write_adds_translations_and_shrinks_text(self) -> None:
        stats: dict[str, object] = {}
        review: list[dict[str, str]] = []
        output = self._write(stats=stats, review_positions=review)
        self.assertEqual(
            _shape_paragraphs(output),
            [
                "图例说明",
                "Légende",
                "红色表示待施工区域",
                "Le rouge indique les zones où les travaux restent à réaliser",
            ],
        )
        drawing = _part(output, "xl/drawings/drawing1.xml")
        sizes = {int(value) for value in __import__("re").findall(r'sz="(\d+)"', drawing)}
        self.assertTrue(sizes and max(sizes) < 1100, sizes)
        comments = _part(output, "xl/comments/comment1.xml")
        self.assertIn("张三:", comments)
        self.assertIn("这里要核对\nÀ vérifier ici", comments)
        self.assertIn("对话批注原文\nCommentaire de discussion", _part(output, "xl/threadedComments/threadedComment1.xml"))
        # 兼容副本同步替换，开头的说明文字原样不动。
        self.assertIn("Comment:\n    对话批注原文\n    Commentaire de discussion", comments)
        self.assertIn("[Threaded comment]", comments)
        self.assertEqual(stats["shape_paragraphs"], 2)
        self.assertEqual(stats["comments"], 2)
        self.assertEqual(stats["shrunk_shapes"], 1)
        self.assertEqual(stats["mutated_cells"], 1)
        # 2 英寸 × 0.8 英寸的框缩到七成左右放得下，不进复核清单。
        self.assertTrue(min(sizes) >= 1100 * 0.6, sizes)
        self.assertEqual([item for item in review if item["category"] == REVIEW_CATEGORY_SHAPE_OVERFLOW], [])

    def test_roomy_box_keeps_font_size(self) -> None:
        build_workbook(self.source, box=(4572000, 1524000))  # 5 英寸 × 1.67 英寸，留了很多空
        stats: dict[str, object] = {}
        drawing = _part(self._write(stats=stats), "xl/drawings/drawing1.xml")
        self.assertEqual(set(__import__("re").findall(r'sz="(\d+)"', drawing)), {"1100"})
        self.assertEqual(stats["shrunk_shapes"], 0)

    def test_box_too_small_even_at_minimum_size_grows_down(self) -> None:
        build_workbook(self.source, box=(1097280, 457200))  # 1.2 英寸 × 0.5 英寸
        review: list[dict[str, str]] = []
        stats: dict = {}
        drawing = _part(self._write(review_positions=review, stats=stats), "xl/drawings/drawing1.xml")
        self.assertEqual(set(re.findall(r'sz="(\d+)"', drawing)), {"660"})
        # 缩到下限还放不下：锚点右下角和 xfrm 一起往下挪，不再列入复核。
        to_row, to_offset = re.search(
            r"<xdr:to>.*?<xdr:row>(\d+)</xdr:row><xdr:rowOff>(\d+)</xdr:rowOff>", drawing
        ).groups()
        grown_pt = (int(to_row) - 1) * 15 + int(to_offset) / 12700
        cy = int(re.search(r'<a:ext cx="\d+" cy="(\d+)"', drawing).group(1))
        self.assertGreater(grown_pt, 457200 / 12700)
        self.assertAlmostEqual(cy / 12700, grown_pt, delta=1)
        self.assertEqual(review, [])
        self.assertEqual(stats["grown_shapes"], 1)

    def test_box_grows_into_blank_rows_then_raises_the_rows_it_covers(self) -> None:
        # 框在 D 列第 2 行起，高 36 磅（落在第 4 行里）；D6 有表头。框先长到第 6 行顶，
        # 不盖住表头；还不够，把框覆盖的第 2～5 行平均加高，表头跟着下移。别的列有内容不算。
        build_workbook(self.source, box=(1097280, 457200), cells={"D6": "序号", "F5": "旁边一列"})
        review: list[dict[str, str]] = []
        stats: dict = {}
        output = self._write(review_positions=review, stats=stats)
        drawing = _part(output, "xl/drawings/drawing1.xml")
        self.assertIn("<xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>", drawing)
        self.assertEqual(review, [])
        self.assertEqual((stats["row_grown_shapes"], stats["overflow_shapes"]), (1, 0))
        heights = _row_heights(output)
        raised = [heights.get(row) for row in (2, 3, 4, 5)]
        self.assertTrue(all(height and height > 15 for height in raised), heights)
        self.assertEqual(len(set(raised)), 1, "四行加得一样多")
        self.assertNotIn(6, heights, "表头那一行不动")

        # 框下边那一行就有内容：不往下长，只加高框覆盖的第 2～4 行。
        build_workbook(self.source, box=(1097280, 457200), cells={"D4": "合计"})
        review = []
        output = self._write(review_positions=review)
        # 框原来就有 6 磅压在第 4 行里，加高后仍在第 4 行里、离行底的距离不变。
        self.assertRegex(_part(output, "xl/drawings/drawing1.xml"), r"<xdr:to>.*<xdr:row>3</xdr:row>")
        self.assertEqual(review, [])
        heights = _row_heights(output)
        self.assertEqual({row: heights.get(row) for row in (2, 3, 4)}, {2: 30.75, 3: 30.75, 4: 30.75})
        self.assertNotIn(5, heights)

    def test_raised_rows_keep_pictures_and_grown_boxes_in_shape(self) -> None:
        # 文本框为放下译文把第 2、3 行从 15 加到 30 磅、锚点按加高后的行写到第 4 行顶（高 60 磅）。
        # 单元格那一轮看到行变了：图片按原行高量是 30 磅，不被拉长；文本框按加高后的行量，保持 60 磅。
        from core.xlsx_patcher import _SheetGeometry, fix_drawing_anchors

        def sheet(height: float):
            return _SheetGeometry(
                etree.fromstring(
                    f'<worksheet xmlns="{NS_MAIN}"><sheetData><row r="2" ht="{height}" customHeight="1"/>'
                    f'<row r="3" ht="{height}" customHeight="1"/></sheetData></worksheet>'.encode()
                )
            )

        drawing = etree.fromstring(_drawing_xml(["图例说明"], anchor_box=(1097280, 0)))
        to = drawing.find(f".//{{{NS_XDR}}}to")
        to.find(f"{{{NS_XDR}}}row").text = "3"
        picture = etree.fromstring(
            f'<xdr:twoCellAnchor xmlns:xdr="{NS_XDR}"><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff>'
            "<xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff>"
            "<xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:pic/><xdr:clientData/></xdr:twoCellAnchor>"
        )
        drawing.append(picture)
        original, grown = sheet(15), sheet(30)
        fix_drawing_anchors(drawing, original, {2, 3}, new_geometry=grown, text_box_geometry=grown)
        self.assertEqual((to.find(f"{{{NS_XDR}}}row").text, to.find(f"{{{NS_XDR}}}rowOff").text), ("3", "0"))
        # 图片仍是两格锚定，右下角挪到第 3 行顶（加高后的第 2 行正好 30 磅），高度不变。
        self.assertEqual(
            (picture.findtext(f"{{{NS_XDR}}}to/{{{NS_XDR}}}row"), picture.findtext(f"{{{NS_XDR}}}to/{{{NS_XDR}}}rowOff")),
            ("2", "0"),
        )

    def test_locked_row_height_lists_the_box_instead_of_raising_rows(self) -> None:
        build_workbook(self.source, box=(1097280, 457200), cells={"D4": "合计"})
        review: list[dict[str, str]] = []
        output = self._write(review_positions=review, lock=True)
        self.assertIn("<xdr:row>1</xdr:row><xdr:rowOff>457200</xdr:rowOff></xdr:to>", _part(output, "xl/drawings/drawing1.xml"))
        # D4 的双语在锁死的默认行高里缩到最小字号也放不下，单元格那条另外列出
        self.assertEqual(
            [item["category"] for item in review], [REVIEW_CATEGORY_SHAPE_OVERFLOW, "locked_row_overflow"]
        )
        self.assertEqual(_row_heights(output), {})

    def test_raised_rows_stay_within_excel_limit(self) -> None:
        # 一行就到顶了也放不下：行高封顶 409.5 磅，框列入复核。
        from core.xlsx_embedded_text import _grow_covered_rows, _SheetGrid

        grid = _SheetGrid({}, {1: 400.0}, 48.0, 15.0)
        grown: dict[int, float] = {}
        added = _grow_covered_rows(grid, 1, 2, 0.0, 30.0, grown)
        self.assertEqual(grown, {1: 409.5})
        self.assertAlmostEqual(added, 9.5)

        # 两行，一行快到顶：多出来的摊给另一行。
        grid = _SheetGrid({}, {1: 405.0, 2: 15.0}, 48.0, 15.0)
        grown = {}
        added = _grow_covered_rows(grid, 1, 3, 0.0, 30.0, grown)
        self.assertEqual(grown, {1: 409.5, 2: 40.5})
        self.assertAlmostEqual(added, 30.0)

    def test_text_box_anchor_is_not_frozen_when_rows_change(self) -> None:
        # 行高一变，图片要冻成定尺寸锚点防变形；文本框不冻：写死的宽度在 Mac 上会比列窄。
        from core.xlsx_patcher import _SheetGeometry, fix_drawing_anchors

        sheet_root = etree.fromstring(f'<worksheet xmlns="{NS_MAIN}"><sheetData/></worksheet>'.encode())
        drawing = etree.fromstring(_drawing_xml(["图例说明"]))
        frozen = fix_drawing_anchors(drawing, _SheetGeometry(sheet_root), set(), freeze_all=True)
        self.assertEqual(frozen, 0)
        self.assertIsNotNone(drawing.find(f"{{{NS_XDR}}}twoCellAnchor"))

    def test_text_box_keeps_its_height_when_rows_shrink(self) -> None:
        # 第 2 行原来 40 磅高、写译文后被改成 15 磅：文本框仍是两格锚定（宽度跟着列走），
        # 右下角往下挪，高度保持 40 磅，不跟着行变矮把字裁掉。
        from core.xlsx_patcher import _SheetGeometry, fix_drawing_anchors

        def sheet(height: float):
            return etree.fromstring(
                f'<worksheet xmlns="{NS_MAIN}"><sheetData><row r="2" ht="{height}" customHeight="1"/></sheetData></worksheet>'.encode()
            )

        drawing = etree.fromstring(_drawing_xml(["图例说明"], anchor_box=(1097280, 40 * 12700)))
        to = drawing.find(f".//{{{NS_XDR}}}to")
        to.find(f"{{{NS_XDR}}}row").text = "2"
        to.find(f"{{{NS_XDR}}}rowOff").text = "0"
        changed = fix_drawing_anchors(
            drawing, _SheetGeometry(sheet(40)), {2}, new_geometry=_SheetGeometry(sheet(15))
        )
        self.assertEqual(changed, 1)
        self.assertIsNotNone(drawing.find(f"{{{NS_XDR}}}twoCellAnchor"))
        self.assertEqual(to.find(f"{{{NS_XDR}}}row").text, "3")
        self.assertEqual(to.find(f"{{{NS_XDR}}}rowOff").text, str(10 * 12700))

    def test_text_box_stops_above_content_when_rows_shrink(self) -> None:
        # 同上，但第 4 行（D4）有内容：框最多到第 4 行顶部，截短的记下来列入复核。
        from core.xlsx_patcher import _SheetGeometry, fix_drawing_anchors

        def sheet(height: float):
            return etree.fromstring(
                f'<worksheet xmlns="{NS_MAIN}"><sheetData><row r="2" ht="{height}" customHeight="1"/>'
                '<row r="4"><c r="D4" t="inlineStr"><is><t>表头</t></is></c></row></sheetData></worksheet>'.encode()
            )

        drawing = etree.fromstring(_drawing_xml(["图例说明"], anchor_box=(1097280, 40 * 12700)))
        to = drawing.find(f".//{{{NS_XDR}}}to")
        to.find(f"{{{NS_XDR}}}row").text = "2"
        to.find(f"{{{NS_XDR}}}rowOff").text = "0"
        shortened: list[str] = []
        fix_drawing_anchors(
            drawing, _SheetGeometry(sheet(40)), {2}, new_geometry=_SheetGeometry(sheet(15)), shortened=shortened
        )
        self.assertEqual((to.find(f"{{{NS_XDR}}}row").text, to.find(f"{{{NS_XDR}}}rowOff").text), ("3", "0"))
        self.assertEqual(shortened, ["TextBox 1"])

    def test_rotated_box_that_cannot_grow_is_listed_for_review(self) -> None:
        build_workbook(self.source, box=(1097280, 457200))
        with zipfile.ZipFile(self.source) as source:
            items = {name: source.read(name) for name in source.namelist()}
        items["xl/drawings/drawing1.xml"] = items["xl/drawings/drawing1.xml"].replace(
            b"<a:xfrm>", b'<a:xfrm rot="5400000">'
        )
        with zipfile.ZipFile(self.source, "w") as target:
            for name, data in items.items():
                target.writestr(name, data)
        review: list[dict[str, str]] = []
        drawing = _part(self._write(review_positions=review), "xl/drawings/drawing1.xml")
        self.assertIn("<xdr:row>1</xdr:row><xdr:rowOff>457200</xdr:rowOff></xdr:to>", drawing)
        self.assertIn(
            {"worksheet": "表1", "cell": "TextBox 1", "category": REVIEW_CATEGORY_SHAPE_OVERFLOW, "action": "listed_only"},
            review,
        )

    def test_two_cell_anchor_size_wins_over_recorded_shape_size(self) -> None:
        # xfrm 里写着 5 英寸宽，锚点实际只框住 1.2 英寸 × 0.5 英寸：按锚点算，缩到下限再加高。
        build_workbook(self.source, box=(4572000, 1524000))
        with zipfile.ZipFile(self.source) as source:
            items = {name: source.read(name) for name in source.namelist()}
        items["xl/drawings/drawing1.xml"] = _drawing_xml(
            ["图例说明", "红色表示待施工区域"], box=(4572000, 1524000), anchor_box=(1097280, 457200)
        )
        with zipfile.ZipFile(self.source, "w") as target:
            for name, data in items.items():
                target.writestr(name, data)
        stats: dict = {}
        drawing = _part(self._write(stats=stats), "xl/drawings/drawing1.xml")
        self.assertEqual(set(re.findall(r'sz="(\d+)"', drawing)), {"660"})
        self.assertEqual(stats["grown_shapes"], 1)

    def test_line_spacing_counts_when_judging_room(self) -> None:
        # 5 英寸 × 1.67 英寸的框单倍行距放得下（见上一条），两倍行距就要缩，
        # 固定行距的磅值也跟着字号一起缩。
        build_workbook(
            self.source,
            box=(4572000, 1524000),
            list_style='<a:lvl1pPr><a:lnSpc><a:spcPct val="200000"/></a:lnSpc></a:lvl1pPr>',
        )
        drawing = _part(self._write(), "xl/drawings/drawing1.xml")
        sizes = {int(value) for value in __import__("re").findall(r'sz="(\d+)"', drawing)}
        self.assertTrue(sizes and max(sizes) < 1100, sizes)

        build_workbook(
            self.source,
            box=(4572000, 1524000),
            list_style='<a:lvl1pPr><a:lnSpc><a:spcPts val="3300"/></a:lnSpc></a:lvl1pPr>',
        )
        drawing = _part(self._write(), "xl/drawings/drawing1.xml")
        line = int(__import__("re").search(r'<a:spcPts val="(\d+)"', drawing).group(1))
        self.assertLess(line, 3300)

    def test_comment_box_grows_to_show_the_translation(self) -> None:
        vml_name = next(
            name for name in zipfile.ZipFile(self.source).namelist() if name.endswith(".vml")
        )
        with zipfile.ZipFile(self.source) as source:
            items = {name: source.read(name) for name in source.namelist()}
        # Excel 自己写的批注框还带锚点：右下角落在第 6 行。
        vml = items[vml_name].decode()
        vml = vml.replace(
            "<ns2:Row>1</ns2:Row>",
            "<ns2:Anchor>2, 15, 0, 2, 4, 15, 5, 16</ns2:Anchor><ns2:Row>1</ns2:Row>",
            1,
        )
        items[vml_name] = vml.encode()
        with zipfile.ZipFile(self.source, "w") as target:
            for name, data in items.items():
                target.writestr(name, data)

        output = self.root / "grown.xlsx"
        shutil.copy2(self.source, output)
        long_translation = (
            "À vérifier ici avec le maître d'œuvre avant la réception définitive des travaux de charpente"
        )
        write_bilingual_workbook(
            output,
            translations={},
            target_lang="fr",
            keep_original_sheets=False,
            embedded_text=EmbeddedTextOptions(
                shapes=False,
                comments=True,
                translations={**TRANSLATIONS, "这里要核对": long_translation},
            ),
        )
        vml = _part(output, vml_name)
        heights = [int(value) for value in re.findall(r"height:(\d+)px", vml)]
        # B2 的批注接了一段长译文，框加高；C3 对话式批注的兼容副本按整段（原文 + 译文）加高。
        self.assertGreater(heights[0], 79)
        self.assertGreater(heights[1], 79)
        # 锚点的下边跟着挪：从上边量下来（默认行高 15 磅）不比新的框矮。
        anchor = [int(v) for v in re.search(r"Anchor>([^<]*)<", vml).group(1).split(",")]
        self.assertEqual(anchor[:6], [2, 15, 0, 2, 4, 15])
        anchored_pt = (anchor[6] - anchor[2]) * 15 - anchor[3] * 0.75 + anchor[7] * 0.75
        self.assertGreaterEqual(anchored_pt, heights[0] * 0.75 - 0.75)

    def test_comment_anchor_follows_actual_row_heights(self) -> None:
        from core.xlsx_embedded_text import _extend_anchor_bottom, _SheetGrid

        grid = _SheetGrid(columns={}, rows={1: 30.0, 2: 0.0}, default_column=48.0, default_row=15.0)
        values = ["2", "15", "0", "0", "4", "15", "1", "0"]
        # 要 60 磅：第 0 行 15 + 第 1 行 30（第 2 行隐藏不占高）+ 第 3 行再走 15 磅 = 第 4 行顶。
        _extend_anchor_bottom(values, 60.0, grid)
        self.assertEqual(values[6:], ["4", "0"])
        # 已经够高的不动。
        values = ["2", "15", "0", "0", "4", "15", "6", "0"]
        _extend_anchor_bottom(values, 60.0, grid)
        self.assertEqual(values[6:], ["6", "0"])

    def test_translation_only_replaces_text_and_keeps_author(self) -> None:
        output = self._write(translation_only=True)
        self.assertEqual(
            _shape_paragraphs(output),
            ["Légende", "Le rouge indique les zones où les travaux restent à réaliser"],
        )
        comments = _part(output, "xl/comments/comment1.xml")
        self.assertIn("张三:", comments)
        self.assertIn("À vérifier ici", comments)
        self.assertNotIn("这里要核对", comments)
        self.assertIn("<text>Commentaire de discussion</text>", _part(output, "xl/threadedComments/threadedComment1.xml"))

    def test_keep_original_sheet_clone_keeps_source_text(self) -> None:
        output = self._write(keep=True)
        with zipfile.ZipFile(output) as archive:
            names = archive.namelist()
            clone_rels = archive.read("xl/worksheets/_rels/sheet2.xml.rels").decode()
        self.assertNotIn(REL_THREADED, clone_rels)
        clone_drawing = next(
            name for name in names if name.startswith("xl/drawings/drawing") and name != "xl/drawings/drawing1.xml"
        )
        self.assertEqual(_shape_paragraphs(output, clone_drawing), ["图例说明", "红色表示待施工区域"])
        clone_comments = next(
            name for name in names if name.startswith("xl/comments") and name != "xl/comments/comment1.xml"
        )
        self.assertIn("这里要核对", _part(output, clone_comments))
        self.assertNotIn("À vérifier ici", _part(output, clone_comments))

    def test_untranslated_only_skips_what_already_has_a_translation(self) -> None:
        output = self._write()
        units = collect_embedded_units(output, shapes=True, comments=True, coverage=("zh", "fr"))
        self.assertEqual(units, [])
        only = self._write(translation_only=True)
        self.assertEqual(collect_embedded_units(only, shapes=True, comments=True, coverage=("zh", "fr")), [])
        # 原文件在补译模式下照样全部列出。
        self.assertEqual(
            len(collect_embedded_units(self.source, shapes=True, comments=True, coverage=("zh", "fr"))),
            4,
        )

    def test_allowed_keys_limit_what_is_written(self) -> None:
        units = collect_embedded_units(self.source, shapes=True, comments=False)
        output = self._write(allowed=frozenset({units[0].key}))
        self.assertEqual(_shape_paragraphs(output), ["图例说明", "Légende", "红色表示待施工区域"])
        self.assertNotIn("À vérifier ici", _part(output, "xl/comments/comment1.xml"))

    def test_without_options_drawing_and_comments_stay_byte_identical(self) -> None:
        output = self.root / "plain.xlsx"
        shutil.copy2(self.source, output)
        write_bilingual_workbook(output, translations={"单元格原文": "Texte"}, target_lang="fr")
        for part in ("xl/drawings/drawing1.xml", "xl/comments/comment1.xml", "xl/threadedComments/threadedComment1.xml"):
            self.assertEqual(_part(output, part), _part(self.source, part))


if __name__ == "__main__":
    unittest.main()


class _MappingEngine:
    engine_name = "embedded-text-test/mock"

    def __init__(self) -> None:
        self.requested: list[str] = []

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        self.requested.extend(texts)
        return {text: TRANSLATIONS.get(text, f"FR:{text}") for text in texts}

    def chat(self, system: str, user: str) -> str:
        raise NotImplementedError


class RunnerTranslatesEmbeddedTextTests(unittest.TestCase):
    """真跑 TaskRunner._run()：文本框和批注跟单元格一起送翻、写进输出文件。"""

    def setUp(self) -> None:
        import os
        from unittest.mock import patch

        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name) / "docs"
        self.root.mkdir()
        env = patch.dict(os.environ, {"TRANSLATOR_APP_DATA_DIR": str(Path(self._tmp.name) / "appdata")})
        env.start()
        self.addCleanup(env.stop)

    def _run(self, *, comments: bool, untranslated_only: bool = False, source: Path | None = None):
        from contextlib import ExitStack
        from unittest.mock import MagicMock, patch

        from core.api_config_check import ApiConfigCheckResult
        from core.file_scanner import FileItem
        from core.task_runner import DoneMsg, TaskRunner
        from settings import AppSettings, EngineSettings

        if source is None:
            source = self.root / "图例.xlsx"
            build_workbook(source)
        settings = AppSettings(
            engine=EngineSettings(
                mode="cloud",
                cloud_provider="custom_openai",
                cloud_model="embedded-test",
                cloud_base_url="https://example.invalid/v1",
                concurrency=1,
                batch_size=20,
            ),
            source_lang="zh",
            target_lang="fr",
            excel_domain_preset="无",
            excel_smart_context=False,
        )
        settings.excel_output.enable_excel_autofit = False
        settings.excel_output.translate_sheet_names = False
        settings.excel_output.keep_original_sheets = False
        settings.excel_output.translate_comments = comments
        runner = TaskRunner(
            [FileItem(path=source, name=source.stem, size_kb=1.0)],
            settings,
            source_root=self.root,
            untranslated_only=untranslated_only,
        )
        engine = _MappingEngine()
        with ExitStack() as stack:
            stack.enter_context(patch("core.task_runner.TaskLogger", return_value=MagicMock(task_id="embedded")))
            stack.enter_context(
                patch("core.task_runner.check_translation_api_config", return_value=ApiConfigCheckResult(ok=True))
            )
            stack.enter_context(patch("core.task_runner.build_engine", return_value=engine))
            stack.enter_context(patch("core.task_runner.get_system_prompt", return_value="system"))
            stack.enter_context(patch("core.task_runner.resolve_effective_model_config", return_value=object()))
            stack.enter_context(
                patch(
                    "core.task_runner.resolve_surface_rate",
                    side_effect=pinned_surface_rate(batch_items=20, concurrency=1),
                )
            )
            stack.enter_context(patch("core.task_runner.tm_manager.lookup_batch", return_value={}))
            stack.enter_context(patch("core.task_runner.tm_manager.insert_batch", return_value=0))
            runner._run()
        done = [m for m in list(runner._queue.queue) if isinstance(m, DoneMsg)]
        self.assertEqual(len(done), 1, list(runner._queue.queue))
        output = Path(done[0].file_results[0]["output_path"])
        return engine, output, done[0]

    def test_shapes_and_comments_are_translated_when_switched_on(self) -> None:
        engine, output, _done = self._run(comments=True)
        self.assertIn("图例说明", engine.requested)
        self.assertIn("这里要核对", engine.requested)
        self.assertIn("对话批注原文", engine.requested)
        self.assertNotIn("张三:\n这里要核对", engine.requested)
        self.assertIn("Légende", _shape_paragraphs(output))
        self.assertIn("À vérifier ici", _part(output, "xl/comments/comment1.xml"))

    def test_default_switches_translate_shapes_but_not_comments(self) -> None:
        from settings import ExcelOutputSettings

        defaults = ExcelOutputSettings()
        engine, output, _done = self._run(comments=defaults.translate_comments)
        self.assertIn("图例说明", engine.requested)
        self.assertNotIn("这里要核对", engine.requested)
        self.assertIn("Légende", _shape_paragraphs(output))
        self.assertEqual(
            _part(output, "xl/comments/comment1.xml"), _part(self.root / "图例.xlsx", "xl/comments/comment1.xml")
        )

    def test_shapes_have_no_switch_and_old_off_setting_is_ignored(self) -> None:
        # 文本框和形状里的字始终翻译。旧版 settings.json 里关掉的 translate_shapes
        # 读取时忽略，不能再让它把文本框跳过去。
        from settings import ExcelOutputSettings

        self.assertNotIn("translate_shapes", ExcelOutputSettings.model_fields)
        legacy = ExcelOutputSettings.model_validate({"translate_shapes": False, "translate_comments": False})
        self.assertNotIn("translate_shapes", legacy.model_dump())
        engine, output, _done = self._run(comments=False)
        self.assertIn("图例说明", engine.requested)
        self.assertIn("Légende", _shape_paragraphs(output))
        self.assertEqual(
            _part(output, "xl/comments/comment1.xml"), _part(self.root / "图例.xlsx", "xl/comments/comment1.xml")
        )

    def test_untranslated_only_fills_in_missing_shape_paragraphs(self) -> None:
        # 上一轮只翻了第一段文本框，这一轮补译只该送第二段。
        source = self.root / "半成品.xlsx"
        build_workbook(source, threaded=False)
        write_bilingual_workbook(
            source,
            translations={"单元格原文": TRANSLATIONS["单元格原文"]},
            target_lang="fr",
            embedded_text=EmbeddedTextOptions(
                shapes=True, translations={"图例说明": TRANSLATIONS["图例说明"]}
            ),
        )
        engine, output, _done = self._run(comments=False, untranslated_only=True, source=source)
        self.assertNotIn("图例说明", engine.requested)
        self.assertIn("红色表示待施工区域", engine.requested)
        self.assertEqual(
            _shape_paragraphs(output),
            [
                "图例说明",
                "Légende",
                "红色表示待施工区域",
                "Le rouge indique les zones où les travaux restent à réaliser",
            ],
        )
