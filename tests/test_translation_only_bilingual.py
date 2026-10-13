"""仅输出译文时，「中文＋外文」的格子也要翻成目标语言（反馈 FB-20261013-UVW95F）。

模型会把「6.1.1 土石方工程合计  Total of Earthworks」判成「已有双语」——中译法时
那半句英文并不是法文。原文＋译文模式照旧原样保留；仅译文模式下原样留着，用户看到
的就是没翻。
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from contextlib import ExitStack

from docx import Document
from openpyxl import load_workbook

from core.file_scanner import FileItem
from core.mixed_language import MIXED_ACTION_EXISTING_BILINGUAL, MixedLanguageResult
from core.task_runner import LogMsg, TaskRunner
from core.word_document import WordFileItem
from core.word_task_runner import WordTaskRunner
from settings import AppSettings
from tests import test_excel_resume as excel_helpers
from tests import test_word_resume as word_helpers

SUBTOTAL = "6.1.1 土石方工程合计  Total of Earthworks"


def _kept_as_bilingual(texts, **kwargs):
    stats = kwargs.get("stats")
    for text in texts:
        result = MixedLanguageResult(source=text, action=MIXED_ACTION_EXISTING_BILINGUAL)
        if stats is not None:
            stats.input_count += 1
            stats.record_result(result)
    return {
        text: MixedLanguageResult(source=text, action=MIXED_ACTION_EXISTING_BILINGUAL)
        for text in texts
    }


class TranslationOnlyBilingualTests(unittest.TestCase):
    def _run(self, *, translation_only: bool):
        calls: list[list[str]] = []

        def translate(texts, *args, **kwargs):
            calls.append(list(texts))
            return {text: f"FR {index}" for index, text in enumerate(texts)}

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = excel_helpers._make_xlsx(root / "boq.xlsx", {"A1": SUBTOTAL, "A2": "混凝土垫层"})
            settings = excel_helpers._settings(source_lang="zh", target_lang="fr")
            settings.excel_output.output_translation_only = translation_only
            settings.excel_output.keep_original_sheets = False
            with excel_helpers._pipeline_patches(translate_side_effect=translate), patch(
                "core.task_runner.translate_mixed_language_texts", side_effect=_kept_as_bilingual,
            ):
                runner = TaskRunner([FileItem(path=path, name="boq", size_kb=1.0)], settings, source_root=root)
                runner._run()
            outputs = [p for p in root.rglob("*.xlsx") if p != path]
            self.assertEqual(len(outputs), 1, outputs)
            workbook = load_workbook(outputs[0])
            values = [
                cell.value
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.value
            ]
            logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
        return calls, values, logs

    def test_translation_only_translates_cells_the_model_kept_as_bilingual(self) -> None:
        calls, values, logs = self._run(translation_only=True)
        self.assertIn([SUBTOTAL], calls)
        self.assertNotIn(SUBTOTAL, values)
        self.assertTrue(any("1 处已含两种语言的内容改为整体翻成目标语言" in line for line in logs), logs)

    def test_bilingual_output_keeps_them_as_they_are(self) -> None:
        calls, values, _logs = self._run(translation_only=False)
        self.assertNotIn([SUBTOTAL], calls)
        self.assertIn(SUBTOTAL, values)



class WordTranslationOnlyBilingualTests(unittest.TestCase):
    def _run(self, *, translation_only: bool):
        calls: list[list[str]] = []

        def translate_word(texts, *args, **kwargs):
            calls.append(list(texts))
            return {text: "Total des terrassements" for text in texts}

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "boq.docx"
            word_helpers.WordResumeTests._make_docx(path, [SUBTOTAL, "混凝土垫层"])
            settings = AppSettings(source_lang="zh", target_lang="fr")
            settings.word_output.output_translation_only = translation_only
            runner = WordTaskRunner(
                [WordFileItem(path=path, name=path.name, size_kb=1.0)], settings, source_root=root,
            )
            with ExitStack() as stack:
                word_helpers.WordResumeTests._runner_patches(self, stack, root=root)
                stack.enter_context(patch(
                    "core.word_task_runner.translate_mixed_language_texts",
                    side_effect=_kept_as_bilingual,
                ))
                stack.enter_context(patch(
                    "core.word_task_runner.translate_word_texts", side_effect=translate_word,
                ))
                runner._run()
            outputs = list((root / "out").rglob("*.docx"))
            self.assertEqual(len(outputs), 1, outputs)
            paragraphs = [p.text for p in Document(outputs[0]).paragraphs if p.text.strip()]
            logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
        return calls, paragraphs, logs

    def test_translation_only_translates_paragraphs_the_model_kept_as_bilingual(self) -> None:
        calls, paragraphs, logs = self._run(translation_only=True)
        self.assertIn([SUBTOTAL], calls)
        self.assertNotIn(SUBTOTAL, paragraphs)
        self.assertIn("Total des terrassements", paragraphs)
        self.assertTrue(any("1 处已含两种语言的内容改为整体翻成目标语言" in line for line in logs), logs)

    def test_bilingual_output_keeps_them_as_they_are(self) -> None:
        calls, paragraphs, _logs = self._run(translation_only=False)
        self.assertNotIn([SUBTOTAL], calls)
        self.assertIn(SUBTOTAL, paragraphs)

if __name__ == "__main__":
    unittest.main(verbosity=2)
