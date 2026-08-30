"""审计 2026-08-29 · Word 文档写入侧（中-25 / 中-26 / 低-Word×2）的回归测试。

四条都在 core/word_document.py 的写入路径上，共同点是"python-docx 看着正常、落到
XML 上不正常"：域被当成普通 run、修订删除不算文字、只读一个属性就凭空建出 part、
pPr 子元素顺序不合 schema。所以这里的 docx 全部按真实 XML 结构构造（w:fldChar 三件
套、w:del/w:delText、真实的 sectPr reference），断言也落在 XML 上，不看 python-docx
的便利视图——那层视图正是这批缺陷藏身的地方。
"""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.translation_protocol import REPLACE_TRANSLATION_PREFIX
from core.word_document import (
    _PPR_CHILD_ORDER,
    _paragraph_literal_text,
    apply_header_footer_translations,
    write_bilingual_docx,
)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _replace(text: str) -> str:
    return f"{REPLACE_TRANSLATION_PREFIX}{text}"


def _append_field(paragraph, instruction: str, *, result: str) -> None:
    """在段落末尾接一个真正的域：begin / instrText / separate / 缓存结果 / end。"""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    code_run = paragraph.add_run()
    code_run._r.append(begin)
    code_run._r.append(instr)
    code_run._r.append(separate)
    paragraph.add_run(result)
    end_run = paragraph.add_run()
    end_run._r.append(end)


def _add_hyperlink(paragraph, text: str, url: str):
    """在段落末尾追加一个真正的 w:hyperlink（带外部关系），而不是伪造的 run。"""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_element = OxmlElement("w:t")
    text_element.text = text
    text_element.set(qn("xml:space"), "preserve")
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def _text_around_field(paragraph) -> tuple[str, str]:
    """作者文字被域切成的前后两半（只认 fldChar 三件套的域，域结果本身不算）。

    断言必须落在「域夹在哪两段文字中间」上：``paragraph.text`` 把域结果和作者文字拼成
    一条连续的串，域插错位置在文本视图里完全看不出来——这正是"域被塞进年份里"能一路
    活到用户按 F9 才发作的原因。
    """
    before: list[str] = []
    after: list[str] = []
    depth = 0
    seen_field = False
    for node in paragraph._p.iter():
        if node.tag == qn("w:fldChar"):
            char_type = node.get(qn("w:fldCharType"))
            if char_type == "begin":
                depth += 1
                seen_field = True
            elif char_type == "end":
                depth = max(depth - 1, 0)
        elif node.tag == qn("w:t") and depth == 0:
            (after if seen_field else before).append(node.text or "")
    return "".join(before), "".join(after)


def _ppr_tags(paragraph) -> list[str]:
    p_pr = getattr(paragraph._p, "pPr", None)
    if p_pr is None:
        return []
    return [child.tag.replace(_W, "w:") for child in p_pr]


def _is_schema_ordered(tags: list[str]) -> bool:
    order = {name: index for index, name in enumerate(_PPR_CHILD_ORDER)}
    indices = [order.get(tag, -1) for tag in tags]
    return all(left <= right for left, right in zip(indices, indices[1:]))


class HeaderFieldReplacementTests(unittest.TestCase):
    """中-25：整行替换遇到含域的页眉段落时，不许退化成"原文 + 译文"追加。"""

    def _build_header_with_field(self, path: Path) -> str:
        doc = Document()
        doc.add_paragraph("施工内容")
        header = doc.sections[0].header.paragraphs[0]
        header.text = "施工组织设计 第 "
        _append_field(header, " STYLEREF 1 \\s ", result="3")
        header.add_run(" 章")
        doc.save(str(path))
        return Document(str(path)).sections[0].header.paragraphs[0].text

    def test_field_bearing_header_is_replaced_not_appended(self) -> None:
        """域在原位继续生效，作者写的那部分文字被译文顶掉，整行不重复。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "styleref.docx"
            line = self._build_header_with_field(source_path)
            self.assertEqual(line, "施工组织设计 第 3 章")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation Chapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertEqual(header.text, "Plan d'organisation Chapitre 3")
            # 原文一个字都不许留在页眉里——页眉高度是节边距定死的。
            self.assertNotIn("施工组织设计", header.text)
            self.assertNotIn("章", header.text)
            self.assertNotIn(" / ", header.text)
            # 域本体（指令码 + begin/end）必须原样还在，否则章号变成一次性死文本。
            xml = header._p.xml
            self.assertIn("STYLEREF", xml)
            self.assertIn('w:fldCharType="begin"', xml)
            self.assertIn('w:fldCharType="end"', xml)

    def test_translation_without_the_field_result_keeps_the_field(self) -> None:
        """译文里找不到域结果时不猜位置：整句写进第一处作者文字，域仍旧保留。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "styleref.docx"
            line = self._build_header_with_field(source_path)

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation du chantier")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertIn("Plan d'organisation du chantier", header.text)
            self.assertNotIn("施工组织设计", header.text)
            self.assertNotIn(" / ", header.text)
            self.assertIn("STYLEREF", header._p.xml)

    def test_unbalanced_field_falls_back_to_appending(self) -> None:
        """域的 begin/end 不配平时读不准边界，退回追加——宁可版式难看也不搅坏结构。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "broken.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计 第 "
            begin = OxmlElement("w:fldChar")
            begin.set(qn("w:fldCharType"), "begin")
            instr = OxmlElement("w:instrText")
            instr.text = " STYLEREF 1 \\s "
            broken = header.add_run()
            broken._r.append(begin)
            broken._r.append(instr)
            header.add_run("3 章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation Chapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertIn("Plan d'organisation Chapitre 3", header.text)
            self.assertIn("施工组织设计", header.text)
            self.assertIn("STYLEREF", header._p.xml)

    def test_plain_header_without_field_still_replaced_in_place(self) -> None:
        """没有域的页眉走原来的整段替换，行为不变。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "plain_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            doc.sections[0].header.paragraphs[0].text = "某某工程 抢工方案"
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={"某某工程 抢工方案": _replace("Plan de rattrapage")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]
            self.assertEqual(header.text, "Plan de rattrapage")


class FieldSplitAmbiguityTests(unittest.TestCase):
    """域结果往译文里对位时，认错一个数字就等于把域种进了别人的数字里。

    章号、题注号常常只有一两位数字，而工程文档的页眉里几乎一定还有个年份。按"第一处
    出现"去切，域就会落在年份中间：刚翻完渲染出来一模一样，用户按 F9、打印预览、或者
    改了章号触发域刷新之后，年份被章号改写，真正的章号反倒成了死文本。失效是静默的，
    所以断言落在「域夹在哪两段作者文字之间」，不看 ``paragraph.text``。
    """

    def test_field_is_not_planted_inside_a_year(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "year_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计 2023 第 "
            _append_field(header, " STYLEREF 1 \\s ", result="3")
            header.add_run(" 章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text
            self.assertEqual(line, "施工组织设计 2023 第 3 章")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan 2023 Chapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            before, after = _text_around_field(header)
            self.assertFalse(
                before.endswith("Plan 202"),
                f"域被塞进了年份 2023 的末位：域前的作者文字是 {before!r}",
            )
            # 章号变动、域刷新成 4 之后，年份必须还是 2023。
            self.assertIn("Plan 2023", f"{before}4{after}")
            self.assertNotIn("施工组织设计", header.text)
            self.assertIn("STYLEREF", header._p.xml)

    def test_unambiguous_field_result_still_splits_in_place(self) -> None:
        """域结果在译文里唯一时照旧原地切分——这条修的是认位，不是把功能关掉。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "plain_field_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计 第 "
            _append_field(header, " STYLEREF 1 \\s ", result="3")
            header.add_run(" 章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan Chapitre 3 fin")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            before, after = _text_around_field(header)
            self.assertEqual(before, "Plan Chapitre ")
            self.assertEqual(after, " fin")
            self.assertEqual(header.text, "Plan Chapitre 3 fin")

    def test_hyperlink_wrapping_a_field_falls_back_to_appending(self) -> None:
        """超链接里同时包着作者文字和域时，读不准就不改——退回追加，别把词序搅乱。

        ``_iter_paragraph_content_nodes`` 会把超链接内的 run 摊平成顶层内容节点，可锚点
        是插在整个 w:hyperlink **之前**的：域之后那一组文字于是跳到了域前面，产物读成
        「Prefixe  queue3」。与域 begin/end 不配平一样，这里也交回追加路径。
        """
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "hyperlink_around_field.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            self.assertEqual(len(header.runs), 0)
            hyperlink = _add_hyperlink(header, "前缀 ", "https://example.com/toc")
            _append_field(header, " STYLEREF 1 \\s ", result="3")
            tail = header.add_run(" 尾巴")
            # 域三件套和尾巴那段文字先落在段落级，再一起搬进超链接内部。
            for run in list(header._p.iterchildren(qn("w:r"))):
                hyperlink.append(run)
            self.assertIs(tail._r.getparent(), hyperlink)
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text
            self.assertEqual(line, "前缀 3 尾巴")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Prefixe 3 queue")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            # 走的是追加回退：原文原样留着、词序没乱，域和译文都在。
            self.assertTrue(
                header.text.startswith("前缀 3 尾巴"),
                f"超链接内的词序被打乱了：{header.text!r}",
            )
            self.assertIn("Prefixe 3 queue", header.text)
            self.assertIn("STYLEREF", header._p.xml)


class HeaderReplacementSideEffectTests(unittest.TestCase):
    """整行替换不许顺手毁掉段落里别的东西：超链接锚点、制表位、软换行。

    这两条是"含域页眉原地替换"（中-25）带出来的：新路径把一批原本走追加、因此毫发
    无损的页眉拉了进来。追加时超链接锚文本还是原文、制表位还在；改成原地替换后，
    锚点选错就把译文塞进了指向原 URL 的可点击链接，``Run.text`` setter 又会把
    w:tab / w:br 一并扫空——Word 默认的"左标题 <tab> 右章号"三段式页眉就此塌成
    左对齐一整行。所以断言全部落在 XML 节点上，不看 python-docx 的文本视图。
    """

    @staticmethod
    def _tab_run(paragraph):
        run = paragraph.add_run()
        run._r.append(OxmlElement("w:tab"))
        return run

    def test_translation_never_lands_inside_a_hyperlink(self) -> None:
        """段落以超链接开头时，译文不许写进超链接内部——整段译文会变成可点击链接。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "hyperlink_field_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            # 不能借 header.text = "" 清场：那会留下一个段落级空 run 顶在超链接前面，
            # 正好当上锚点，把要复现的坑掩盖掉。段落本来就是空的，直接往里挂。
            self.assertEqual(len(header.runs), 0)
            _add_hyperlink(header, "施工组织设计", "https://example.com/plan")
            header.add_run(" 第 ")
            _append_field(header, " STYLEREF 1 \\s ", result="3")
            header.add_run(" 章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text
            self.assertEqual(line, "施工组织设计 第 3 章")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan Chapitre 3 fin")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertEqual(header.text, "Plan Chapitre 3 fin")
            self.assertIn("STYLEREF", header._p.xml)
            linked_text = "".join(
                node.text or ""
                for hyperlink in header._p.findall(qn("w:hyperlink"))
                for node in hyperlink.iter(qn("w:t"))
            )
            # 超链接要么被清掉、要么里面一个字都不剩，绝不能裹着译文。
            self.assertEqual(linked_text, "")

    def test_field_header_keeps_its_tab_stop(self) -> None:
        """三段式页眉：译文没带回制表符时，原有的 w:tab 必须留在原位。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "tabbed_field_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计"
            self._tab_run(header)
            header.add_run("第 ")
            _append_field(header, " STYLEREF 1 \\s ", result="3")
            header.add_run(" 章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text
            self.assertEqual(line, "施工组织设计\t第 3 章")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation Chapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertNotIn("施工组织设计", header.text)
            self.assertEqual(len(header._p.findall(f".//{_W}tab")), 1)
            self.assertIn("STYLEREF", header._p.xml)

    def test_plain_header_keeps_its_tab_stop(self) -> None:
        """无域页眉走的是另一条替换路径，制表位同样不许丢。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "tabbed_plain_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计"
            self._tab_run(header)
            header.add_run("第三章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text
            self.assertEqual(line, "施工组织设计\t第三章")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation Chapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertEqual(header.text.replace("\t", ""), "Plan d'organisation Chapitre 3")
            self.assertEqual(len(header._p.findall(f".//{_W}tab")), 1)

    def test_header_keeps_its_soft_line_break(self) -> None:
        """两行页眉：译文只有一行时，原有的 w:br 必须留着，页眉高度才不变。

        制表位和软换行走的是同一条兜底（``_layout_tags_to_keep``），但换行丢掉的后果
        更重：两行页眉塌成一行，正文版心跟着上移一整行。所以单独钉一条。
        """
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "two_line_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计"
            break_run = header.add_run()
            break_run._r.append(OxmlElement("w:br"))
            header.add_run("第三章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text
            self.assertEqual(line, "施工组织设计\n第三章")

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan Chapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertNotIn("施工组织设计", header.text)
            self.assertIn("Plan Chapitre 3", header.text)
            self.assertEqual(len(header._p.findall(f".//{_W}br")), 1)

    def test_translation_carrying_its_own_tab_does_not_double_it(self) -> None:
        """译文自己带回了制表符时，原节点不再补挂一遍——两个制表位比一个更跑版。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "tabbed_roundtrip.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            header = doc.sections[0].header.paragraphs[0]
            header.text = "施工组织设计"
            self._tab_run(header)
            header.add_run("第三章")
            doc.save(str(source_path))
            line = Document(str(source_path)).sections[0].header.paragraphs[0].text

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation\tChapitre 3")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertEqual(header.text, "Plan d'organisation\tChapitre 3")
            self.assertEqual(len(header._p.findall(f".//{_W}tab")), 1)


class SelfClosingFieldParagraphTests(unittest.TestCase):
    """带 w:fldSimple 的段落：作者写的字必须一个不漏地被认出来。

    ``_paragraph_literal_text`` 原先用 ``id(text_node)`` 记住"哪些 w:t 属于 fldSimple"。
    lxml 的元素代理用完即弃，后建的代理会复用同一个内存地址——域外的一段正文因此被
    当成域结果扣掉，整段被判成"只有域、没正文"跳过不译。这条钉的是判定结果本身。
    """

    @staticmethod
    def _build(path: Path) -> str:
        doc = Document()
        doc.add_paragraph("施工内容")
        header = doc.sections[0].header.paragraphs[0]
        header.text = "施工组织设计 第 "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), " STYLEREF 1 \\s ")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "3"
        run.append(text)
        field.append(run)
        header._p.append(field)
        header.add_run(" 章")
        doc.save(str(path))
        return Document(str(path)).sections[0].header.paragraphs[0].text

    def test_literal_text_keeps_every_authored_run(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source_path = Path(temp_dir) / "fldsimple.docx"
            self._build(source_path)
            paragraph = Document(str(source_path)).sections[0].header.paragraphs[0]

            literal = _paragraph_literal_text(paragraph)
            self.assertIn("施工组织设计", literal)
            self.assertIn("章", literal)
            # 域结果不算作者写的字。
            self.assertNotIn("3", literal)

    def test_self_closing_field_header_is_replaced_in_place(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "fldsimple.docx"
            line = self._build(source_path)

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={line: _replace("Plan d'organisation Chapitre")},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            header = Document(str(out_path)).sections[0].header.paragraphs[0]

            self.assertNotIn("施工组织设计", header.text)
            self.assertNotIn(" / ", header.text)
            self.assertIn("Plan d'organisation Chapitre", header.text)
            self.assertIn("fldSimple", header._p.xml)
            self.assertIn("STYLEREF", header._p.xml)


class TrackedDeletionParagraphTests(unittest.TestCase):
    """中-26：文末"只剩一条修订删除"的段落不是空段，不许当垃圾清掉。"""

    @staticmethod
    def _append_tracked_deletion(paragraph, text: str) -> None:
        del_element = OxmlElement("w:del")
        del_element.set(qn("w:id"), "101")
        del_element.set(qn("w:author"), "审校")
        del_element.set(qn("w:date"), "2026-08-29T00:00:00Z")
        run = OxmlElement("w:r")
        del_text = OxmlElement("w:delText")
        del_text.set(qn("xml:space"), "preserve")
        del_text.text = text
        run.append(del_text)
        del_element.append(run)
        paragraph._p.append(del_element)

    def test_trailing_tracked_deletion_survives(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "tracked.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            self._append_tracked_deletion(
                doc.add_paragraph(),
                "原验收标准按国标执行。",
            )
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={"施工内容": "Contenu des travaux"},
                target_lang="fr",
                source_lang="zh",
            )
            with zipfile.ZipFile(str(out_path)) as archive:
                xml = archive.read("word/document.xml").decode("utf-8")

            # 修订记录还在：用户在 Word 里「拒绝所有修订」仍能拿回这段原文。
            self.assertIn("原验收标准按国标执行。", xml)
            self.assertIn("delText", xml)
            self.assertIn("<w:del ", xml)

    def test_real_empty_trailing_paragraphs_are_still_trimmed(self) -> None:
        """真正的空段照旧清掉——这条修的是判空口径，不是取消清理。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "blank_tail.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            doc.add_paragraph()
            doc.add_paragraph()
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={"施工内容": "Contenu des travaux"},
                target_lang="fr",
                source_lang="zh",
            )
            texts = [p.text for p in Document(str(out_path)).paragraphs]
            self.assertEqual(texts[-1], "Contenu des travaux")


class HeaderFooterPartInjectionTests(unittest.TestCase):
    """低-Word：没有页眉页脚的文档，翻译一趟不许凭空长出 6 个空 part。"""

    def test_document_without_headers_gets_no_new_parts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "plain.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={"施工内容": "Contenu des travaux"},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            with zipfile.ZipFile(str(out_path)) as archive:
                parts = [
                    name
                    for name in archive.namelist()
                    if "header" in name or "footer" in name
                ]
                document_xml = archive.read("word/document.xml").decode("utf-8")

            self.assertEqual(parts, [])
            self.assertNotIn("headerReference", document_xml)
            self.assertNotIn("footerReference", document_xml)

    def test_existing_header_is_still_translated(self) -> None:
        """真有页眉的文档照旧翻译，跳过的只是"根本不存在"的槽位。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "with_header.docx"
            doc = Document()
            doc.add_paragraph("施工内容")
            doc.sections[0].header.paragraphs[0].text = "某某工程 抢工方案"
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={"某某工程 抢工方案": "Plan de rattrapage"},
                target_lang="fr",
                source_lang="zh",
                translate_headers_footers=True,
            )
            out_doc = Document(str(out_path))
            self.assertEqual(
                out_doc.sections[0].header.paragraphs[0].text,
                "某某工程 抢工方案 / Plan de rattrapage",
            )
            with zipfile.ZipFile(str(out_path)) as archive:
                parts = sorted(
                    name
                    for name in archive.namelist()
                    if "header" in name or "footer" in name
                )
            self.assertEqual(parts, ["word/header1.xml"])

    def test_second_section_header_is_created_only_by_the_document(self) -> None:
        """第二节自己有页眉、第一节没有：只碰真实存在的那一个。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "second.docx"
            doc = Document()
            doc.add_paragraph("第一节")
            doc.add_section()
            second = doc.sections[1]
            second.header.is_linked_to_previous = False
            second.header.paragraphs[0].text = "某某工程 抢工方案"
            doc.save(str(source_path))

            out_doc = Document(str(source_path))
            written = apply_header_footer_translations(
                out_doc,
                {"某某工程 抢工方案": "Plan de rattrapage"},
                target_lang="fr",
                source_lang="zh",
            )
            self.assertEqual(written, 1)
            self.assertEqual(
                out_doc.sections[1].header.paragraphs[0].text,
                "某某工程 抢工方案 / Plan de rattrapage",
            )
            body_xml = out_doc.element.body.xml
            self.assertEqual(body_xml.count("footerReference"), 0)


class ParagraphPropertiesOrderTests(unittest.TestCase):
    """低-Word：合成出来的 w:pPr 子元素必须按 CT_PPr 的 schema 次序排列。"""

    def test_translation_paragraph_properties_follow_schema_order(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "ppr.docx"
            doc = Document()
            # List Paragraph 样式自带 w:ind + w:contextualSpacing，段落自己带
            # w:pStyle + w:jc：合并时"样式在前、段落在后"，pStyle 就被挤到了 ind
            # 后面；抑制自动编号补的 numPr 更是永远落在末尾。
            paragraph = doc.add_paragraph("施工内容", style="List Paragraph")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={"施工内容": "Contenu des travaux"},
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))
            translation = next(
                para
                for para in out_doc.paragraphs
                if para.text == "Contenu des travaux"
            )
            tags = _ppr_tags(translation)

            self.assertIn("w:pStyle", tags)
            self.assertIn("w:numPr", tags)
            self.assertTrue(
                _is_schema_ordered(tags),
                f"pPr 子元素次序不合 CT_PPr：{tags}",
            )
            # 顺序之外的东西一样不许丢：格式合并的结果原封不动。
            self.assertIn("w:ind", tags)
            self.assertIn("w:jc", tags)

    def test_all_generated_paragraphs_are_schema_ordered(self) -> None:
        """整份产物扫一遍，任何段落的 pPr 都不许出现逆序。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "mixed.docx"
            doc = Document()
            doc.add_paragraph("一级标题", style="Heading 1")
            doc.add_paragraph("项目要点", style="List Bullet")
            body = doc.add_paragraph("施工内容", style="List Paragraph")
            body.alignment = WD_ALIGN_PARAGRAPH.CENTER
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).paragraphs[0].text = "表内说明"
            doc.save(str(source_path))

            out_path = write_bilingual_docx(
                source_path=source_path,
                output_dir=temp_path / "out",
                translations={
                    "一级标题": "Titre",
                    "项目要点": "Points clés",
                    "施工内容": "Contenu des travaux",
                    "表内说明": "Note du tableau",
                },
                target_lang="fr",
                source_lang="zh",
            )
            out_doc = Document(str(out_path))
            paragraphs = list(out_doc.paragraphs)
            for cell_table in out_doc.tables:
                for row in cell_table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)

            offenders = [
                (para.text[:16], _ppr_tags(para))
                for para in paragraphs
                if not _is_schema_ordered(_ppr_tags(para))
            ]
            self.assertEqual(offenders, [])


if __name__ == "__main__":
    unittest.main()
