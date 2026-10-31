"""只输出译文时，表格一格的译文行数和原文段数对不上，先单独重译这一格。

本机实测（法译中，PDF 转 Word）：转换把两条「•」挤进了一段，模型翻译时顺手拆成两行，
这一格的译文比原文多一行。写回只能保留法文、把中文追加在下面，整份只译文的文档里
这两格成了双语。现在写盘前把这几格单独重译、要求一行对一段；重译仍对不上才保底。
"""
from __future__ import annotations

import json
import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from core.task_runner import LogMsg
from core.word_document import WordFileItem, cell_line_groups
from core.word_task_runner import WordTaskRunner, _cell_line_mismatches
from engines.base_engine import TranslationEngine
from settings import AppSettings
from tests import test_word_resume as word_helpers
from tests.app_data_isolation import IsolatedAppDataTestCase

HEADER = "Jour 2"
CELL = "Réunion de lancement, programme validé.• Enregistrement des relevés\nVisite du site"
SPLIT = "启动会议，确认工作计划。\n• 记录测量数据\n现场踏勘"
ALIGNED = "启动会议，确认工作计划。• 记录测量数据\n现场踏勘"


class SplitsBulletsEngine(TranslationEngine):
    """First answer splits the glued bullets onto their own line; a retry told the line count keeps them."""

    def __init__(self, *, obeys_line_rule: bool) -> None:
        self.obeys_line_rule = obeys_line_rule
        self.cell_prompts: list[str] = []
        self.judged: list[dict] = []

    @property
    def engine_name(self) -> str:
        return "fake/splits-bullets"

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        results = {}
        for text in texts:
            if text == CELL:
                self.cell_prompts.append(system_prompt)
                told = "原文共 2 行" in system_prompt
                results[text] = ALIGNED if told and self.obeys_line_rule else SPLIT
            else:
                results[text] = "第 2 天"
        return results


class JudgingEngine(SplitsBulletsEngine):
    """Also answers the meaning check after the retries."""

    def __init__(self, *, verdict: str) -> None:
        super().__init__(obeys_line_rule=False)
        self.verdict = verdict

    def chat(self, system: str, user: str) -> str:
        self.judged.append(json.loads(user))
        return json.dumps({"verdict": self.verdict, "reason": "test"})


class CellLineGroupTests(unittest.TestCase):
    def test_one_line_per_paragraph(self) -> None:
        self.assertEqual(cell_line_groups(["A", "B"], ["a", "b"]), ["A", "B"])

    def test_soft_breaks_inside_a_paragraph_take_their_share_of_lines(self) -> None:
        self.assertEqual(cell_line_groups(["A1", "A2", "B"], ["a1\na2", "b"]), ["A1\nA2", "B"])

    def test_blank_lines_from_the_model_do_not_count(self) -> None:
        self.assertEqual(cell_line_groups(["A", "", "B"], ["a", "b"]), ["A", "B"])

    def test_extra_line_does_not_fit(self) -> None:
        self.assertIsNone(cell_line_groups(["A", "B", "C"], ["a", "b"]))

    def test_only_table_cells_with_misaligned_lines_are_picked(self) -> None:
        locations = {
            CELL: [{"kind": "table_cell"}],
            "Ligne un\nLigne deux": [{"kind": "paragraph"}],
            "Case\nalignée": [{"kind": "table_cell"}],
        }
        translations = {CELL: SPLIT, "Ligne un\nLigne deux": "一\n二\n三", "Case\nalignée": "格\n对齐"}
        self.assertEqual(_cell_line_mismatches(translations, locations), [CELL])


class WordCellLineRealignTests(IsolatedAppDataTestCase):
    def _run(self, *, obeys_line_rule: bool = False, translation_only: bool = True, engine=None):
        engine = engine or SplitsBulletsEngine(obeys_line_rule=obeys_line_rule)
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "planning.docx"
            doc = Document()
            doc.add_paragraph(HEADER)
            cell = doc.add_table(rows=1, cols=1).cell(0, 0)
            first, second = CELL.split("\n")
            cell.paragraphs[0].text = first
            cell.add_paragraph(second)
            doc.save(path)
            settings = AppSettings(source_lang="fr", target_lang="zh")
            settings.word_output.output_translation_only = translation_only
            runner = WordTaskRunner(
                [WordFileItem(path=path, name=path.name, size_kb=1.0)], settings, source_root=root, source_lang="fr"
            )
            with ExitStack() as stack:
                word_helpers.WordResumeTests._runner_patches(self, stack, root=root)
                stack.enter_context(patch("core.word_task_runner.build_engine", return_value=engine))
                stack.enter_context(patch("core.word_task_runner.tm_manager.lookup_batch", return_value={}))
                runner._run()
            outputs = list((root / "out").rglob("*.docx"))
            self.assertEqual(len(outputs), 1, outputs)
            out_cell = Document(outputs[0]).tables[0].cell(0, 0)
            self.highlighted = any(
                run._r.find(".//" + qn("w:highlight")) is not None or run._r.find(".//" + qn("w:shd")) is not None
                for paragraph in out_cell.paragraphs
                for run in paragraph.runs
            )
            logs = [m.message for m in runner._queue.queue if isinstance(m, LogMsg)]
            return engine, out_cell.text, logs

    def test_retry_with_the_line_rule_writes_the_cell_in_chinese_only(self) -> None:
        engine, cell_text, logs = self._run(obeys_line_rule=True)
        self.assertEqual(cell_text, ALIGNED)
        self.assertNotIn("Enregistrement", cell_text)
        self.assertEqual(len(engine.cell_prompts), 2, "首次翻译一次，按行数重译一次")
        self.assertIn("原文共 2 行", engine.cell_prompts[-1])
        self.assertTrue([line for line in logs if "行数与原文一致" in line], logs)

    def test_same_meaning_after_retries_writes_the_translation_only(self) -> None:
        engine, cell_text, logs = self._run(engine=JudgingEngine(verdict="equivalent"))
        self.assertEqual(cell_text, SPLIT, "意思一致：整格只放译文，多出来的一行放在同一格里")
        self.assertNotIn("Enregistrement", cell_text)
        # 首次 + 「单段重试次数」那么多轮（默认 3），然后复查一次
        self.assertEqual(len(engine.cell_prompts), 4)
        self.assertEqual(len(engine.judged), 1)
        self.assertEqual(engine.judged[0]["items"][0]["source"], CELL)
        self.assertEqual(engine.judged[0]["items"][0]["translation"], SPLIT)
        self.assertFalse(self.highlighted)
        self.assertTrue([line for line in logs if "意思一致" in line], logs)

    def test_different_meaning_keeps_source_plus_translation_and_marks_it(self) -> None:
        engine, cell_text, logs = self._run(engine=JudgingEngine(verdict="not_equivalent"))
        self.assertIn("Enregistrement", cell_text)
        self.assertIn("记录测量数据", cell_text)
        self.assertEqual(len(engine.judged), 1)
        self.assertTrue(self.highlighted, "保留原文的格子要标底色")
        self.assertTrue([line for line in logs if "列入报告待复核" in line], logs)

    def test_engine_without_a_meaning_check_keeps_source_plus_translation(self) -> None:
        engine, cell_text, _logs = self._run(obeys_line_rule=False)
        self.assertIn("Enregistrement", cell_text)
        self.assertEqual(len(engine.cell_prompts), 4)
        self.assertTrue(self.highlighted)

    def test_bilingual_output_does_not_retry(self) -> None:
        engine, _cell_text, _logs = self._run(obeys_line_rule=True, translation_only=False)
        self.assertEqual(len(engine.cell_prompts), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)


class TitleFooterEngine(TranslationEngine):
    TITLE = "MÉTHODOLOGIE D’INTERVENTION"
    FOOTER = "Méthodologie d’intervention – Expertise structurelle"

    def __init__(self) -> None:
        self.footer_prompts: list[str] = []

    @property
    def engine_name(self) -> str:
        return "fake/title-footer"

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        results = {}
        for text in texts:
            if text == self.FOOTER:
                self.footer_prompts.append(system_prompt)
                told = "「MÉTHODOLOGIE D’INTERVENTION」已译作「实施方法」" in system_prompt
                results[text] = "实施方法——结构检测与鉴定" if told else "干预方法——结构检测与鉴定"
            else:
                results[text] = "实施方法"
        return results


class WordShortTextTermTests(IsolatedAppDataTestCase):
    def test_footer_is_retranslated_with_the_title_translation(self) -> None:
        engine = TitleFooterEngine()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            path = root / "methodologie.docx"
            doc = Document()
            doc.add_paragraph(engine.TITLE)
            doc.add_table(rows=1, cols=1).cell(0, 0).text = engine.FOOTER
            doc.save(path)
            settings = AppSettings(source_lang="fr", target_lang="zh")
            settings.word_output.output_translation_only = True
            runner = WordTaskRunner(
                [WordFileItem(path=path, name=path.name, size_kb=1.0)], settings, source_root=root, source_lang="fr"
            )
            with ExitStack() as stack:
                word_helpers.WordResumeTests._runner_patches(self, stack, root=root)
                stack.enter_context(patch("core.word_task_runner.build_engine", return_value=engine))
                stack.enter_context(patch("core.word_task_runner.tm_manager.lookup_batch", return_value={}))
                runner._run()
            output = next((root / "out").rglob("*.docx"))
            self.assertEqual(Document(output).tables[0].cell(0, 0).text, "实施方法——结构检测与鉴定")
        self.assertEqual(len(engine.footer_prompts), 2)
