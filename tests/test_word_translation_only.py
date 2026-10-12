"""Word translation-only output keeps the replace-only contract."""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from core.word_coverage import build_word_coverage_plan, write_untranslated_docx
from core.word_document import write_bilingual_docx
from settings import WordOutputSettings


class WordTranslationOnlyTests(unittest.TestCase):
    def _source(self, root: Path) -> Path:
        doc = Document()
        doc.add_paragraph("正文原文")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格原文"
        section = doc.sections[0]
        section.header.paragraphs[0].text = "页眉原文"
        section.footer.paragraphs[0].text = "页脚原文"
        path = root / "source.docx"
        doc.save(path)
        return path

    def test_full_write_replaces_body_table_header_footer_and_default_is_bilingual(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            source = self._source(root)
            translations = {
                "正文原文": "Body translation",
                "表格原文": "Table translation",
                "页眉原文": "Header translation",
                "页脚原文": "Footer translation",
            }
            only = write_bilingual_docx(
                source_path=source,
                output_dir=root / "only",
                translations=translations,
                target_lang="en",
                translate_headers_footers=True,
                output_translation_only=True,
            )
            bilingual = write_bilingual_docx(
                source_path=source,
                output_dir=root / "bilingual",
                translations=translations,
                target_lang="en",
                translate_headers_footers=True,
            )
            self.assertEqual(Document(only).paragraphs[0].text, "Body translation")
            self.assertEqual(Document(only).tables[0].cell(0, 0).text, "Table translation")
            self.assertEqual(Document(only).sections[0].header.paragraphs[0].text, "Header translation")
            self.assertEqual(Document(only).sections[0].footer.paragraphs[0].text, "Footer translation")
            self.assertIn("正文原文", Document(bilingual).paragraphs[0].text)
            self.assertIn("Body translation", Document(bilingual).paragraphs[1].text)

    def test_coverage_write_replaces_source_only_positions(self) -> None:
        with TemporaryDirectory() as temp:
            root = Path(temp)
            source = self._source(root)
            plan = build_word_coverage_plan(source, target_lang="en", source_lang="zh")
            output = write_untranslated_docx(
                source_path=source,
                output_dir=root / "coverage",
                plan=plan,
                translations={
                    "正文原文": "Body translation",
                    "表格原文": "Table translation",
                },
                target_lang="en",
                output_translation_only=True,
            )
            doc = Document(output)
            self.assertEqual(doc.paragraphs[0].text, "Body translation")
            self.assertEqual(doc.tables[0].cell(0, 0).text, "Table translation")

    def test_setting_defaults_to_bilingual(self) -> None:
        self.assertFalse(WordOutputSettings().output_translation_only)


class TranslationOnlyKeepsEmbeddedContentTests(unittest.TestCase):
    """「仅显示译文」改写整段时，段落里挂着的文本框和脚注编号不能被一起删掉。"""

    def test_textbox_and_footnote_reference_survive_replacement(self) -> None:
        import zipfile

        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls, qn

        mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
        wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
        box = "<w:txbxContent><w:p><w:r><w:t>文本框里的字</w:t></w:r></w:p></w:txbxContent>"
        with TemporaryDirectory() as temp:
            root = Path(temp)
            doc = Document()
            end = doc.element.body.find(qn("w:sectPr"))
            end.addprevious(parse_xml(
                f'<w:p {nsdecls("w")} xmlns:mc="{mc}" xmlns:wps="{wps}">'
                "<w:r><w:t>这一段带文本框</w:t></w:r><w:r><mc:AlternateContent>"
                f'<mc:Choice Requires="wps"><wps:txbx>{box}</wps:txbx></mc:Choice>'
                f"<mc:Fallback><w:pict>{box}</w:pict></mc:Fallback>"
                "</mc:AlternateContent></w:r></w:p>"
            ))
            end.addprevious(parse_xml(
                f'<w:p {nsdecls("w")}><w:r><w:t>这一段带脚注</w:t></w:r>'
                '<w:r><w:footnoteReference w:id="1"/></w:r></w:p>'
            ))
            source = root / "source.docx"
            doc.save(source)

            output = write_bilingual_docx(
                source_path=source,
                output_dir=root / "only",
                translations={
                    "这一段带文本框": "Paragraph with a text box",
                    "这一段带脚注": "Paragraph with a footnote",
                },
                target_lang="en",
                output_translation_only=True,
            )
            xml = zipfile.ZipFile(output).read("word/document.xml").decode("utf-8")
            self.assertEqual(xml.count("<w:txbxContent"), 2)
            self.assertEqual(xml.count("w:footnoteReference"), 1)
            texts = [paragraph.text for paragraph in Document(output).paragraphs]
            self.assertIn("Paragraph with a text box", texts)
            self.assertIn("Paragraph with a footnote", texts)
            self.assertNotIn("这一段带文本框", xml)


if __name__ == "__main__":
    unittest.main(verbosity=2)
