"""补译模式下公式格跟随「公式显示值回填」开关：全流程回归测试。

产品约定（设计初衷）：翻译产物里被译文命中的公式格以「值」覆盖——整格换成静态
双语文本、``<f>`` 删除；要看原公式去「_原文」分表或原始文件。全量模式一直如此，
补译模式必须同一条规则，否则同一份文件「全量翻会翻、补译却静默留中文」。

开关的两支：
  * ``formula_display_value_backfill=True``（默认）：公式格按缓存显示值进补译
    候选，写回时整格覆盖成静态双语文本（与全量一致）；
  * ``formula_display_value_backfill=False``：保护公式——覆盖率层把公式格判
    ignored（公式源码不送翻，不白花 API 调用），写入层的独立兜底保证即使坐标
    算错也绝不删 ``<f>``。

历史注记：审计高-6 的第一版修法是「补译公式格一律跳过」，与设计初衷相反，已
改成按开关拆分；这份文件同时守住两支的行为。
"""
from __future__ import annotations

import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from lxml import etree
from openpyxl import Workbook, load_workbook

from core.excel_coverage import build_excel_coverage_plan, write_untranslated_excel_file
from core.translation_coverage import COVERAGE_IGNORED, COVERAGE_SOURCE_ONLY
from core.xlsx_patcher import NS_MAIN, write_bilingual_workbook


def _m(tag: str) -> str:
    return f"{{{NS_MAIN}}}{tag}"


def _build_fixture(root: Path) -> Path:
    """一张真实 xlsx：A1 是待补译的普通源文本，A2 是公式格。

    openpyxl 保存公式格时只写 ``<f>``，不带缓存值——真实 Excel 保存过的文件
    才会带 ``<v>``。这里存完后用 lxml 补一个缓存值上去，模拟“Excel 算过一次”
    的公式格：显示值是 "施工内容合计"（含源语言文本，进补译候选的正是它）。
    """
    base = root / "base.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "报价"
    sheet["A1"] = "施工内容"
    sheet["A2"] = '=A1&"合计"'
    workbook.save(base)
    workbook.close()

    with zipfile.ZipFile(base) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    sheet_xml = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
    formula_cell = sheet_xml.find(f".//{_m('c')}[@r='A2']")
    formula_cell.set("t", "str")
    v_el = formula_cell.find(_m("v"))
    if v_el is None:
        v_el = etree.SubElement(formula_cell, _m("v"))
    v_el.text = "施工内容合计"
    parts["xl/worksheets/sheet1.xml"] = etree.tostring(
        sheet_xml, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    target = root / "fixture.xlsx"
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return target


def _formula_text(path: Path, sheet: str, coordinate: str) -> str | None:
    """直接读原始 XML 的 ``<f>``，不经过 openpyxl 的 data_only 折叠。"""
    with zipfile.ZipFile(path) as archive:
        wb = load_workbook(path, read_only=True)
        try:
            index = wb.sheetnames.index(sheet) + 1
        finally:
            wb.close()
        xml = archive.read(f"xl/worksheets/sheet{index}.xml")
    root = etree.fromstring(xml)
    cell = root.find(f".//{_m('c')}[@r='{coordinate}']")
    if cell is None:
        return None
    f_el = cell.find(_m("f"))
    return f_el.text if f_el is not None else None


class ExcelBackfillFormulaTests(unittest.TestCase):
    def setUp(self) -> None:
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.fixture = _build_fixture(self.root)

    def tearDown(self) -> None:
        self._temp.cleanup()

    def _run_full_chain(self, *, formula_display_value_backfill: bool) -> Path:
        plan = build_excel_coverage_plan(
            self.fixture,
            target_lang="en",
            source_lang="zh",
            formula_display_value_backfill=formula_display_value_backfill,
        )

        formula_unit = next(
            unit for unit in plan.units if unit.data.get("coordinate") == "A2"
        )
        if formula_display_value_backfill:
            # 开关开着：显示值按普通文本分类，进补译候选——与全量同一条规则。
            self.assertEqual(formula_unit.status, COVERAGE_SOURCE_ONLY)
            self.assertIn("施工内容合计", plan.source_texts)
        else:
            # 开关关着：公式格 ignored，公式源码绝不送翻。
            self.assertEqual(formula_unit.status, COVERAGE_IGNORED)
            self.assertNotIn(formula_unit.source_text, plan.source_texts)
            self.assertTrue(
                all(u.data.get("coordinate") != "A2" for u in plan.source_units)
            )

        # 非公式格两支都照常判定为待补译。
        plain_unit = next(
            unit for unit in plan.units if unit.data.get("coordinate") == "A1"
        )
        self.assertEqual(plain_unit.status, COVERAGE_SOURCE_ONLY)
        self.assertIn("施工内容", plan.source_texts)

        translations = {text: f"[EN]{text}" for text in plan.source_texts}
        # 关回填那支：万一实现有疏漏、公式的源文本真被送去"翻译"了，也让写入端
        # 有值可用，这样断言公式被保留就不会因为 KeyError 而意外通过。
        translations.setdefault(formula_unit.source_text, "[EN]" + formula_unit.source_text)

        out_path = write_untranslated_excel_file(
            source_path=self.fixture,
            output_dir=self.root / f"out_{formula_display_value_backfill}",
            plan=plan,
            translations=translations,
            target_lang="en",
            source_lang="zh",
            keep_original_sheets=False,
            formula_display_value_backfill=formula_display_value_backfill,
        )
        return out_path

    def test_formula_overwritten_with_backfill_on(self) -> None:
        """默认开关：补译把公式格覆盖成静态双语文本，<f> 删除——设计初衷。"""
        out_path = self._run_full_chain(formula_display_value_backfill=True)

        formula_text = _formula_text(out_path, "报价", "A2")
        self.assertIsNone(formula_text, "回填开着时公式格该被译文以值覆盖，<f> 却还在")

        wb = load_workbook(out_path)
        try:
            self.assertEqual(wb["报价"]["A2"].value, "施工内容合计\n[EN]施工内容合计")
            self.assertEqual(wb["报价"]["A1"].value, "施工内容\n[EN]施工内容")
        finally:
            wb.close()

    def test_formula_preserved_with_backfill_off(self) -> None:
        out_path = self._run_full_chain(formula_display_value_backfill=False)

        formula_text = _formula_text(out_path, "报价", "A2")
        self.assertIsNotNone(formula_text, "A2 的 <f> 元素被删掉了，公式丢失")
        self.assertEqual(formula_text, 'A1&"合计"')

        wb = load_workbook(out_path)
        try:
            self.assertEqual(wb["报价"]["A1"].value, "施工内容\n[EN]施工内容")
        finally:
            wb.close()

    # ── 绕过覆盖率层，直接考 xlsx_patcher 自己那层公式守卫 ──────────────────

    def test_second_formula_guard_blocks_when_backfill_off(self) -> None:
        """覆盖率层（core/excel_coverage.py）在回填关闭时本该把公式格挡在补译
        坐标集合外。core/xlsx_patcher.py 的 write_bilingual_workbook 自己还留了
        一道独立兜底：回填关闭 + 坐标限定（补译）时，公式格永远不可写。这条
        测试直接调 write_bilingual_workbook，故意把公式格的坐标塞进
        allowed_positions（模拟坐标计算在别处出错、越过了覆盖率层的场景），
        验证这道独立守卫真的兜住了，不依赖上游算对。
        """
        out_path = self.root / "bypass_out.xlsx"
        shutil.copy(self.fixture, out_path)

        write_bilingual_workbook(
            out_path,
            translations={
                "施工内容": "[EN]施工内容",
                # 公式源码和显示值都给出「译文」：如果守卫没拦住，mutation 一定
                # 会发生（不是恰好没命中翻译表才侥幸没改）。
                '=A1&"合计"': "[EN]施工内容合计",
                "施工内容合计": "[EN]施工内容合计",
            },
            target_lang="en",
            source_lang="zh",
            keep_original_sheets=False,
            formula_display_value_backfill=False,
            mark_review_items=False,
            # 故意把公式格 A2 也放进允许坐标集合——回填关闭时覆盖率层永远不会
            # 这样传，这里是直接白盒测试 xlsx_patcher 自己的独立兜底。
            allowed_positions={"报价": {"A1", "A2"}},
        )

        formula_text = _formula_text(out_path, "报价", "A2")
        self.assertIsNotNone(formula_text, "A2 的 <f> 元素被删掉了，公式丢失")
        self.assertEqual(formula_text, 'A1&"合计"')

        wb = load_workbook(out_path)
        try:
            # A2 仍然是公式，没有被坐标限定牵着改写成静态双语文本。
            self.assertEqual(wb["报价"]["A2"].value, '=A1&"合计"')
            # A1（非公式格）该翻的照常翻了——证明公式格是被专门挡下来的，
            # 不是 allowed_positions 机制整体失效导致两格都没写。
            self.assertEqual(wb["报价"]["A1"].value, "施工内容\n[EN]施工内容")
        finally:
            wb.close()

    def test_second_formula_guard_stands_down_when_backfill_on(self) -> None:
        """回填开着时这道兜底必须让路：补译坐标集合点名的公式格照常按显示值
        覆盖成静态双语文本，与全量模式行为一致。"""
        out_path = self.root / "allow_out.xlsx"
        shutil.copy(self.fixture, out_path)

        write_bilingual_workbook(
            out_path,
            translations={"施工内容合计": "[EN]施工内容合计"},
            target_lang="en",
            source_lang="zh",
            keep_original_sheets=False,
            formula_display_value_backfill=True,
            mark_review_items=False,
            allowed_positions={"报价": {"A2"}},
        )

        self.assertIsNone(
            _formula_text(out_path, "报价", "A2"),
            "回填开着、坐标点名的公式格该被覆盖，<f> 却还在",
        )
        wb = load_workbook(out_path)
        try:
            self.assertEqual(wb["报价"]["A2"].value, "施工内容合计\n[EN]施工内容合计")
            # A1 不在坐标集合里，一个字都不许动。
            self.assertEqual(wb["报价"]["A1"].value, "施工内容")
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
