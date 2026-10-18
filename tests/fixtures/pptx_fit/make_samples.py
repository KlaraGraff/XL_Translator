"""生成 tests/fixtures/pptx_fit/samples_zh.pptx：「中文原文刚好放满」的几种框。

要翻的框在 cNvPr descr 里写 "fit"，法文译文写进 translations_fr.json（形状名 → 各段译文，
表格按行列顺序）。中文按逐字 1 个字号宽、行高 1.2 个字号算框高，刚好放满。

需要 python-pptx（不进安装包，也不在 requirements 里）：
  python -m venv /tmp/pptxenv && /tmp/pptxenv/bin/pip install python-pptx
  /tmp/pptxenv/bin/python tests/fixtures/pptx_fit/make_samples.py tests/fixtures/pptx_fit
"""
import json
import math
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
INSET_W = 0.2  # in, 左右内边距合计
INSET_H = 0.1

translations: dict[str, list[str]] = {}


def set_fonts(run, size=18, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "zh-CN")
    rpr.set("altLang", "en-US")
    for tag in ("latin", "ea"):
        el = rpr.find(A + tag)
        if el is None:
            el = etree.SubElement(rpr, A + tag)
        el.set("typeface", "Calibri" if tag == "latin" else "微软雅黑")


def autofit(tf, kind, wrap=True):
    body = tf._txBody.find(A + "bodyPr")
    for child in list(body):
        body.remove(child)
    body.set("wrap", "square" if wrap else "none")
    if kind:
        etree.SubElement(body, A + kind)


def mark(shape, name, french):
    shape.name = name
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", "fit")
    translations[name] = french


def fill_height(texts, size, width_in, line=1.2):
    """中文逐字 1 个字号宽、行高 1.2 个字号：刚好放下这些段落的框高（英寸）。"""
    inner = (width_in - INSET_W) * 72
    per_line = max(1, int(inner // size))
    lines = sum(max(1, math.ceil(len(t) / per_line)) for t in texts)
    return lines * size * line / 72 + INSET_H + 0.02


def text_box(slide, name, x, y, w, texts, french, size=18, kind=None, wrap=True, h=None):
    h = h or fill_height(texts, size, w)
    shapes = getattr(slide, "shapes", slide)
    box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    for i, text in enumerate(texts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        set_fonts(run, size)
    autofit(tf, kind, wrap)
    box.line.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    box.line.width = Pt(0.75)
    mark(box, name, french)
    return box


def label(slide, text, y=0.15):
    box = slide.shapes.add_textbox(Inches(0.3), Inches(y), Inches(9.4), Inches(0.4))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    return box


def obstacle(slide, x, y, w, h, text="（下方已有内容）"):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0xDD, 0xE6, 0xF0)
    rect.line.color.rgb = RGBColor(0x88, 0x99, 0xAA)
    run = rect.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x33, 0x44, 0x55)
    return rect


def main(out_dir: Path) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]

    # 1. 不自动调整的文本框：下面空着 / 下面紧挨着别的形状
    s = prs.slides.add_slide(blank)
    label(s, "1 不自动调整（noAutofit）：左边下面空着，右边下面紧挨着形状")
    text_box(
        s, "noauto-free", 0.4, 0.8, 4.4,
        ["本项目位于城市新区核心地段，总用地面积约十四公顷，规划建设研发办公、配套商业和人才公寓三大功能板块。"],
        [("Le projet est situé au cœur de la nouvelle zone urbaine, sur une superficie totale d'environ quatorze hectares, "
         "et comprend trois volets : bureaux de recherche et développement, commerces de proximité et logements pour talents.")],
        kind="noAutofit",
    )
    box = text_box(
        s, "noauto-blocked", 5.2, 0.8, 4.4,
        ["施工期间请严格遵守安全规范，进入现场必须佩戴安全帽并系好安全带。"],
        [("Pendant les travaux, respectez strictement les règles de sécurité : le port du casque et du harnais est "
         "obligatoire pour accéder au chantier, sans exception, y compris pour les visiteurs accompagnés.")],
        kind="noAutofit",
    )
    obstacle(s, 5.2, (box.top + box.height) / 914400 + 0.02, 4.4, 0.6)

    # 2. 标题和内容占位符：继承母版的「溢出时缩排文字」
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "项目总体进度与下阶段计划"
    mark(s.shapes.title, "title", ["Avancement global du projet et planning de la phase suivante"])
    body = s.placeholders[1]
    items = [
        "主体结构已完成封顶，较原计划提前两周",
        "机电安装进入高峰期，现场人员约三百人",
        "幕墙材料第二批已到场，正在进行验收",
        "室外管网施工受雨季影响，需调整工序",
        "下月重点：样板间交付与消防联合调试",
        "风险提示：部分进口设备到货时间待确认",
        "质量安全：本月无事故，隐患整改率百分之百",
    ]
    body.text_frame.text = items[0]
    for item in items[1:]:
        body.text_frame.add_paragraph().text = item
    for p in body.text_frame.paragraphs:
        for run in p.runs:
            rpr = run._r.get_or_add_rPr()
            rpr.set("lang", "zh-CN")
    mark(body, "body", [
        "Le gros œuvre est terminé, avec deux semaines d'avance sur le planning initial",
        "Les installations électromécaniques sont en pleine phase, avec environ trois cents personnes sur site",
        "Le deuxième lot de matériaux de façade est arrivé et sa réception est en cours",
        "Les réseaux extérieurs sont retardés par la saison des pluies ; l'ordre des travaux doit être ajusté",
        "Mois prochain : livraison de l'appartement témoin et essais conjoints du système incendie",
        "Risque : la date de livraison de certains équipements importés reste à confirmer",
        "Qualité et sécurité : aucun accident ce mois-ci, 100 % des anomalies corrigées",
    ])

    # 3. 根据文字调整形状大小：下面空着 / 不换行的标签 / 下面不远处有形状
    s = prs.slides.add_slide(blank)
    label(s, "3 根据文字调整形状（spAutoFit）：加高 / 加宽 / 长到形状前再缩字")
    text_box(
        s, "spauto-free", 0.4, 0.8, 4.4,
        ["会议纪要须在会后两个工作日内发送全体参会人员确认。"],
        [("Le compte rendu de réunion doit être envoyé à tous les participants pour validation "
         "dans les deux jours ouvrables suivant la réunion.")],
        kind="spAutoFit",
    )
    text_box(
        s, "spauto-label", 0.4, 3.0, 1.6, ["设计负责人"], ["Responsable de la conception"],
        kind="spAutoFit", wrap=False, h=0.4,
    )
    box = text_box(
        s, "spauto-near", 5.2, 0.8, 4.4,
        ["监理单位应每周组织一次质量巡查，并形成书面记录。"],
        [("L'entreprise de supervision organise chaque semaine une inspection de la qualité sur l'ensemble du "
         "chantier et en établit un compte rendu écrit, transmis au maître d'ouvrage et à l'entreprise générale.")],
        kind="spAutoFit",
    )
    obstacle(s, 5.2, (box.top + box.height) / 914400 + 0.35, 4.4, 0.6)

    # 4. 表格
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "主要工程量"
    rows = [
        ["分项工程", "单位", "工程量"],
        ["钢筋混凝土", "立方米", "一万二千"],
        ["砌体墙", "平方米", "八千五百"],
        ["外墙保温", "平方米", "六千三百"],
    ]
    french_rows = [
        ["Ouvrage", "Unité", "Quantité"],
        ["Béton armé", "Mètre cube", "Douze mille"],
        ["Maçonnerie de remplissage", "Mètre carré", "Huit mille cinq cents"],
        ["Isolation thermique des façades", "Mètre carré", "Six mille trois cents"],
    ]
    frame = s.shapes.add_table(4, 3, Inches(1.0), Inches(1.8), Inches(6.0), Inches(4 * 0.42))
    table = frame.table
    for c, w in enumerate((2.2, 1.9, 1.9)):
        table.columns[c].width = Inches(w)
    for r in range(4):
        table.rows[r].height = Inches(0.42)
        for c in range(3):
            cell = table.cell(r, c)
            cell.text = rows[r][c]
            for run in cell.text_frame.paragraphs[0].runs:
                set_fonts(run, 18, bold=r == 0)
    frame.name = "table"
    frame._element.xpath("./*[1]/p:cNvPr")[0].set("descr", "fit")
    translations["table"] = [t for row in french_rows for t in row]

    # 5. 组合（组合放大了 1.25 倍）
    s = prs.slides.add_slide(blank)
    label(s, "5 组合里的文本框（组合整体放大 1.25 倍）")
    group = s.shapes.add_group_shape()
    text_box(group.shapes, "group-a", 0.4, 1.0, 3.2, ["请于周五前提交月度报表。"],
             ["Merci de remettre le rapport mensuel avant vendredi."], kind="noAutofit")
    text_box(group.shapes, "group-b", 0.4, 2.0, 3.2, ["图纸变更需经设计院书面确认。"],
             ["Toute modification des plans doit être confirmée par écrit par l'institut de conception."],
             kind="noAutofit")
    xfrm = group._element.find(".//" + A + "xfrm")
    ext = xfrm.find(A + "ext")
    ext.set("cx", str(int(int(ext.get("cx")) * 1.25)))
    ext.set("cy", str(int(int(ext.get("cy")) * 1.25)))

    # 6. 竖排（不动）与旋转 90° 的文本框（只缩字、不加高）
    s = prs.slides.add_slide(blank)
    label(s, "6 竖排框不动；旋转的框只缩字不加高")
    box = text_box(s, "vertical", 0.6, 1.0, 0.8, ["安全第一预防为主"], ["La sécurité avant tout"],
                   kind="noAutofit", h=3.0)
    box.text_frame._txBody.find(A + "bodyPr").set("vert", "eaVert")
    box = text_box(s, "rotated", 3.0, 2.5, 4.0, ["出入口请保持畅通。"],
                   ["Veuillez maintenir les entrées et sorties dégagées en permanence."], kind="noAutofit")
    box.rotation = 90

    out_dir.mkdir(parents=True, exist_ok=True)
    prs.save(out_dir / "samples_zh.pptx")
    (out_dir / "translations_fr.json").write_text(json.dumps(translations, ensure_ascii=False, indent=1), "utf-8")
    print("saved", out_dir / "samples_zh.pptx", Emu(prs.slide_width).inches, Emu(prs.slide_height).inches)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
