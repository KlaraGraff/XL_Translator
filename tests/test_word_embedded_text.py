"""Word 文本框和脚注/尾注的翻译（core/word_embedded_text.py）。"""

from __future__ import annotations

import re
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.word_coverage import build_word_coverage_plan, write_untranslated_docx
from core.word_document import (
    detect_hidden_word_content,
    extract_word_segments,
    extract_word_story_segments,
    write_bilingual_docx,
)

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
V = "urn:schemas-microsoft-com:vml"

BOX_TEXTS = ["图例说明", "红色表示待施工区域"]
TRANSLATIONS = {
    "正文原文": "Texte du corps",
    "图例说明": "Légende",
    "红色表示待施工区域": "Le rouge indique les zones où les travaux restent à réaliser",
    "脚注正文": "Texte de la note de bas de page",
    "尾注正文": "Texte de la note de fin",
}


def _box_xml(texts: list[str], ppr: str = "") -> str:
    paragraphs = "".join(
        f'<w:p>{ppr}<w:r><w:rPr><w:sz w:val="22"/></w:rPr><w:t>{text}</w:t></w:r></w:p>' for text in texts
    )
    return f"<w:txbxContent>{paragraphs}</w:txbxContent>"


def _textbox_paragraph(
    texts: list[str], *, cx: int, cy: int, autofit: bool = False, ppr: str = "", wrap: str = "wrapSquare"
) -> str:
    box = _box_xml(texts, ppr)
    fit = "<a:spAutoFit/>" if autofit else ""
    return (
        f'<w:p {nsdecls("w")} xmlns:mc="{MC}" xmlns:wps="{WPS}" xmlns:wp="{WP}" xmlns:a="{A}" xmlns:v="{V}">'
        "<w:r><w:t>正文原文</w:t></w:r>"
        "<w:r><mc:AlternateContent>"
        '<mc:Choice Requires="wps"><w:drawing>'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" '
        'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:{wrap}/><wp:docPr id="1" name="文本框 1"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
        f"<wps:txbx>{box}</wps:txbx>"
        f'<wps:bodyPr wrap="square">{fit}</wps:bodyPr>'
        "</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        f'<mc:Fallback><w:pict><v:shape style="width:144pt;height:60pt"><v:textbox>{box}</v:textbox></v:shape></w:pict></mc:Fallback>'
        "</mc:AlternateContent></w:r></w:p>"
    )


def _notes_part(doc, *, name: str, tag: str, ref_tag: str, text: str, rel_type: str, content_type: str) -> None:
    xml = (
        f'<w:{tag}s {nsdecls("w")}>'
        f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>'
        f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{tag}>'
        f'<w:{tag} w:id="1"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
        f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:{ref_tag}/></w:r>'
        f'<w:r><w:t xml:space="preserve"> {text}</w:t></w:r></w:p></w:{tag}>'
        f"</w:{tag}s>"
    ).encode()
    part = Part(PackURI(f"/word/{name}.xml"), content_type, xml, doc.part.package)
    doc.part.relate_to(part, rel_type)


def build_document(
    path: Path, *, cx: int = 2286000, cy: int = 914400, autofit: bool = False, ppr: str = "", wrap: str = "wrapSquare"
) -> None:
    """``ppr``：文本框里每段的 ``<w:pPr>``，用来给段落设行距。"""
    doc = Document()
    end = doc.element.body.find(qn("w:sectPr"))
    end.addprevious(parse_xml(_textbox_paragraph(BOX_TEXTS, cx=cx, cy=cy, autofit=autofit, ppr=ppr, wrap=wrap)))
    end.addprevious(
        parse_xml(
            f'<w:p {nsdecls("w")}><w:r><w:t>这一段带脚注</w:t></w:r>'
            '<w:r><w:footnoteReference w:id="1"/></w:r><w:r><w:endnoteReference w:id="1"/></w:r></w:p>'
        )
    )
    _notes_part(
        doc,
        name="footnotes",
        tag="footnote",
        ref_tag="footnoteRef",
        text="脚注正文",
        rel_type=RT.FOOTNOTES,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    )
    _notes_part(
        doc,
        name="endnotes",
        tag="endnote",
        ref_tag="endnoteRef",
        text="尾注正文",
        rel_type=RT.ENDNOTES,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
    )
    doc.save(path)


def _xml(path: Path, part: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode("utf-8")


def _box_texts(xml: str) -> list[list[str]]:
    """每个 w:txbxContent 里各段的文字（正本在前、VML 备份在后）。"""
    boxes = re.findall(r"<w:txbxContent>(.*?)</w:txbxContent>", xml, re.S)
    return [
        ["".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", p)) for p in re.findall(r"<w:p[ >].*?</w:p>", box, re.S)]
        for box in boxes
    ]


def _note_texts(xml: str, tag: str) -> list[str]:
    note = re.search(rf'<w:{tag} w:id="1">(.*?)</w:{tag}>', xml, re.S).group(1)
    return ["".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", p)).strip() for p in re.findall(r"<w:p[ >].*?</w:p>", note, re.S)]


class WordStoryTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.source = self.root / "source.docx"
        build_document(self.source)

    def _write(self, source: Path | None = None, *, only: bool = False, issues: list | None = None, name: str = "out") -> Path:
        return write_bilingual_docx(
            source_path=source or self.source,
            output_dir=self.root / name,
            translations=dict(TRANSLATIONS),
            target_lang="fr",
            output_translation_only=only,
            issue_callback=issues.append if issues is not None else None,
        )

    def test_segments_include_textbox_and_note_paragraphs_once(self) -> None:
        segments = extract_word_segments(self.source, target_lang="fr")
        by_kind: dict[str, list[str]] = {}
        for segment in segments:
            by_kind.setdefault(segment.kind, []).append(segment.source)
        # VML 备份里的同一个文本框不重复列出；分隔线条目不算。
        self.assertEqual(by_kind["textbox"], BOX_TEXTS)
        self.assertEqual(by_kind["footnote"], ["脚注正文"])
        self.assertEqual(by_kind["endnote"], ["尾注正文"])
        locations = {segment.source: segment.location for segment in segments}
        self.assertEqual(locations["红色表示待施工区域"], "textbox[0].paragraph[1]")
        self.assertEqual(locations["脚注正文"], "footnote[1].paragraph[0]")

    def test_bilingual_adds_translation_after_each_paragraph_and_syncs_fallback(self) -> None:
        output = self._write()
        boxes = _box_texts(_xml(output, "word/document.xml"))
        expected = ["图例说明", "Légende", "红色表示待施工区域", TRANSLATIONS["红色表示待施工区域"]]
        self.assertEqual(boxes, [expected, expected])
        self.assertEqual(_note_texts(_xml(output, "word/footnotes.xml"), "footnote"), ["脚注正文", TRANSLATIONS["脚注正文"]])
        self.assertEqual(_note_texts(_xml(output, "word/endnotes.xml"), "endnote"), ["尾注正文", TRANSLATIONS["尾注正文"]])
        # 文件仍能被 python-docx 打开，正文照常翻译。
        self.assertIn("Texte du corps", [p.text for p in Document(output).paragraphs])

    def test_translation_only_replaces_text_and_keeps_note_number(self) -> None:
        output = self._write(only=True)
        boxes = _box_texts(_xml(output, "word/document.xml"))
        expected = ["Légende", TRANSLATIONS["红色表示待施工区域"]]
        self.assertEqual(boxes, [expected, expected])
        footnotes = _xml(output, "word/footnotes.xml")
        self.assertEqual(_note_texts(footnotes, "footnote"), [TRANSLATIONS["脚注正文"]])
        note = re.search(r'<w:footnote w:id="1">(.*?)</w:footnote>', footnotes, re.S).group(1)
        # 编号标记还在第一个上标 run 里，译文写进了后面那个普通 run。
        runs = re.findall(r"<w:r>.*?</w:r>", note, re.S)
        self.assertIn("<w:footnoteRef/>", runs[0])
        self.assertNotIn("<w:t", runs[0])
        self.assertIn(TRANSLATIONS["脚注正文"], runs[1])
        self.assertNotIn("superscript", runs[1])

    def test_small_box_shrinks_font_then_grows_when_still_too_small(self) -> None:
        issues: list[dict] = []
        roomy = self._write(name="roomy", issues=issues)
        sizes = set(re.findall(r'<w:sz w:val="(\d+)"/>', _xml(roomy, "word/document.xml")))
        # 2.5 英寸 × 1 英寸：四段放得下（译文会折行），缩一点。
        self.assertTrue(sizes and all(12 <= int(size) < 22 for size in sizes), sizes)
        self.assertEqual(issues, [])

        # 1 英寸 × 0.25 英寸：缩到 60% 还放不下，框往下加高（正本、VML 备份一起），正文绕着框排、
        # 跟着往下让，不报问题。
        tiny = self.root / "tiny.docx"
        build_document(tiny, cx=914400, cy=228600)
        output = self._write(tiny, name="tiny", issues=issues)
        xml = _xml(output, "word/document.xml")
        self.assertIn("13", set(re.findall(r'<w:sz w:val="(\d+)"/>', xml)))  # 22 × 0.6
        extent_cy = int(re.search(r'<wp:extent cx="914400" cy="(\d+)"/>', xml).group(1))
        self.assertGreater(extent_cy, 228600)
        vml_height = float(re.search(r'style="width:144pt;height:([\d.]+)pt"', xml).group(1))
        self.assertAlmostEqual(vml_height, 60 + (extent_cy - 228600) / 12700, delta=0.1)
        self.assertEqual(issues, [])

        # 浮在文字上方的框（不环绕）加高会盖住正文：保持原样，报进质量报告。
        floating = self.root / "floating.docx"
        build_document(floating, cx=914400, cy=228600, wrap="wrapNone")
        output = self._write(floating, name="floating", issues=issues)
        self.assertIn('<wp:extent cx="914400" cy="228600"/>', _xml(output, "word/document.xml"))
        self.assertEqual([issue["kind"] for issue in issues], ["textbox_overflow"])
        issues.clear()

        # 旋转过的框加高会错位：保持原样，报进质量报告。
        rotated = self.root / "rotated.docx"
        build_document(rotated, cx=914400, cy=228600)
        with zipfile.ZipFile(rotated) as archive:
            items = {name: archive.read(name) for name in archive.namelist()}
        items["word/document.xml"] = items["word/document.xml"].replace(
            b"<wps:spPr>", b'<wps:spPr><a:xfrm rot="5400000"><a:off x="0" y="0"/><a:ext cx="914400" cy="228600"/></a:xfrm>'
        )
        with zipfile.ZipFile(rotated, "w") as archive:
            for name, data in items.items():
                archive.writestr(name, data)
        output = self._write(rotated, name="rotated", issues=issues)
        self.assertIn('<wp:extent cx="914400" cy="228600"/>', _xml(output, "word/document.xml"))
        self.assertEqual([issue["kind"] for issue in issues], ["textbox_overflow"])
        self.assertEqual(issues[0]["location"], "textbox[0].paragraph[0]")

    def test_big_or_autofit_box_keeps_font_size(self) -> None:
        for name, kwargs in (("big", {"cx": 5486400, "cy": 2743200}), ("autofit", {"cy": 228600, "autofit": True})):
            source = self.root / f"{name}.docx"
            build_document(source, **kwargs)
            output = self._write(source, name=name)
            sizes = set(re.findall(r'<w:sz w:val="(\d+)"/>', _xml(output, "word/document.xml")))
            self.assertEqual(sizes, {"22"}, name)

    def test_line_spacing_counts_when_judging_room(self) -> None:
        sizes = {}
        # 段距和行距都写在段落上，盖过默认模板 docDefaults 里的段后 10 磅、1.15 倍。
        for name, ppr in (
            ("single", '<w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'),
            ("wide", '<w:pPr><w:spacing w:after="0" w:line="360" w:lineRule="auto"/></w:pPr>'),
        ):
            source = self.root / f"{name}.docx"
            build_document(source, cx=180 * 12700, cy=100 * 12700, ppr=ppr)
            output = self._write(source, name=name)
            sizes[name] = set(re.findall(r'<w:sz w:val="(\d+)"/>', _xml(output, "word/document.xml")))
        # 同样大的框：单倍行距放得下，1.5 倍行距要缩字。
        self.assertEqual(sizes["single"], {"22"})
        self.assertTrue(all(int(size) < 22 for size in sizes["wide"]), sizes["wide"])

    def test_document_default_spacing_counts_when_judging_room(self) -> None:
        # Word 默认模板把段后 10 磅、1.15 倍行距写在 docDefaults 里，不在样式上：同样大的框，
        # 没有默认段距时放得下，有了就得缩字。
        sizes = {}
        for name, spacing in (("plain", ""), ("word_default", '<w:spacing w:after="200" w:line="276" w:lineRule="auto"/>')):
            source = self.root / f"{name}.docx"
            build_document(source, cx=180 * 12700, cy=100 * 12700)
            with zipfile.ZipFile(source) as archive:
                items = {item: archive.read(item) for item in archive.namelist()}
            styles = items["word/styles.xml"].decode()
            styles = re.sub(r"<w:pPrDefault>.*?</w:pPrDefault>|<w:pPrDefault/>", "", styles, flags=re.S)
            styles = styles.replace("</w:rPrDefault>", f"</w:rPrDefault><w:pPrDefault><w:pPr>{spacing}</w:pPr></w:pPrDefault>", 1)
            items["word/styles.xml"] = styles.encode()
            with zipfile.ZipFile(source, "w") as archive:
                for item, data in items.items():
                    archive.writestr(item, data)
            output = self._write(source, name=name)
            sizes[name] = set(re.findall(r'<w:sz w:val="(\d+)"/>', _xml(output, "word/document.xml")))
        self.assertEqual(sizes["plain"], {"22"})
        self.assertTrue(all(int(size) < 22 for size in sizes["word_default"]), sizes["word_default"])

    def test_exact_line_spacing_shrinks_with_the_font(self) -> None:
        source = self.root / "exact.docx"
        build_document(source, cx=180 * 12700, cy=80 * 12700, ppr='<w:pPr><w:spacing w:line="400" w:lineRule="exact"/></w:pPr>')
        output = self._write(source, name="exact")
        xml = _xml(output, "word/document.xml")
        size = {int(value) for value in re.findall(r'<w:sz w:val="(\d+)"/>', xml)}
        line = {int(value) for value in re.findall(r'w:line="(\d+)"', xml)}
        self.assertEqual(len(size), 1)
        self.assertEqual(len(line), 1)
        # 字缩到几成，固定行距也缩到几成；规则仍是固定值。
        self.assertAlmostEqual(line.pop() / 400, size.pop() / 22, delta=0.05)
        self.assertNotIn('w:lineRule="auto"', xml)

    def test_untranslated_only_fills_in_what_is_missing_without_repeating(self) -> None:
        first = self._write()
        # 已经译过的文档再补译一遍：什么也不加。
        self.assertEqual(extract_word_story_segments(first, target_lang="fr"), [])
        again = write_untranslated_docx(
            source_path=first,
            output_dir=self.root / "again",
            plan=build_word_coverage_plan(first, target_lang="fr"),
            translations=dict(TRANSLATIONS),
            target_lang="fr",
        )
        self.assertEqual(_box_texts(_xml(again, "word/document.xml"))[0], _box_texts(_xml(first, "word/document.xml"))[0])
        self.assertEqual(_note_texts(_xml(again, "word/footnotes.xml"), "footnote"), ["脚注正文", TRANSLATIONS["脚注正文"]])

        # 原文还没翻的文档：补译把文本框和脚注都补上。
        pending = extract_word_story_segments(self.source, target_lang="fr")
        self.assertEqual([segment.source for segment in pending], [*BOX_TEXTS, "脚注正文", "尾注正文"])
        filled = write_untranslated_docx(
            source_path=self.source,
            output_dir=self.root / "filled",
            plan=build_word_coverage_plan(self.source, target_lang="fr"),
            translations=dict(TRANSLATIONS),
            target_lang="fr",
        )
        self.assertIn("Légende", _box_texts(_xml(filled, "word/document.xml"))[0])
        self.assertEqual(_note_texts(_xml(filled, "word/endnotes.xml"), "endnote"), ["尾注正文", TRANSLATIONS["尾注正文"]])

    def test_textboxes_and_notes_no_longer_reported_as_untranslated(self) -> None:
        report = detect_hidden_word_content(self.source)
        self.assertEqual(report.textbox_count, 1)
        self.assertEqual(report.footnote_count, 2)
        self.assertFalse(report.found)
        self.assertEqual(report.describe(), "")

    def test_document_without_notes_or_boxes_is_untouched(self) -> None:
        plain = self.root / "plain.docx"
        doc = Document()
        doc.add_paragraph("正文原文")
        doc.save(plain)
        output = self._write(plain, name="plain")
        with zipfile.ZipFile(output) as archive:
            names = set(archive.namelist())
        self.assertNotIn("word/footnotes.xml", names)
        self.assertEqual([p.text for p in Document(output).paragraphs], ["正文原文", "Texte du corps"])


def _part(doc, name: str, content_type: str, payload: bytes, rel_type: str) -> str:
    """挂一个关系指向的部件，返回正文里引用的关系 id。"""
    part = Part(PackURI(name), content_type, payload, doc.part.package)
    return doc.part.relate_to(part, rel_type)


def _graphic_paragraph(graphic_data_xml: str, namespace: str, extra_ns: str = "") -> str:
    """正文里挂一个 ``wp:inline`` 图形；``extra_ns`` 补图形内容自己的命名空间声明。"""
    return (
        f'<w:p {nsdecls("w")} xmlns:wp="{WP}" xmlns:a="{A}" xmlns:r='
        '"http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f"{extra_ns}>"
        '<w:r><w:drawing><wp:inline><wp:extent cx="3000000" cy="2000000"/>'
        f'<a:graphic><a:graphicData uri="{namespace}">{graphic_data_xml}'
        "</a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
    )


class WordCommentTests(unittest.TestCase):
    """批注（comments.xml）写回时不能被破坏。

    审计 D2 线记的是「批注至今没有翻译路径也不进统计」；那是产品范围问题，靠测试
    补不上。能补、也值得钉住的是另一条下限：翻正文的时候不能顺手把批注弄坏。
    ``comments+xml`` 现在列在 ``_WORD_TEXT_PART_TYPES`` 的「允许改动」部件表里，
    照抄由「没改的部件必须逐字节一致」那条保存后核对兜着；批注的引用标记则落在
    正文 run 里，重整段落时最容易连带扫掉。这两点都没有现成用例。
    （若将来真的开始翻译批注，这里的逐字节断言要连同部件表一起改。）
    """

    def setUp(self) -> None:
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.source = self.root / "commented.docx"
        doc = Document()
        paragraph = doc.add_paragraph("正文原文")
        doc.add_comment(
            paragraph.runs, text="这段要请法务复核。", author="张三", initials="ZS"
        )
        doc.save(self.source)

    def _comments_xml(self, path: Path) -> bytes:
        with zipfile.ZipFile(path) as archive:
            return archive.read("word/comments.xml")

    def test_comment_part_and_its_anchor_markers_survive_a_whole_paragraph_rewrite(
        self,
    ) -> None:
        original = self._comments_xml(self.source)
        self.assertIn("这段要请法务复核。", original.decode("utf-8"))
        with zipfile.ZipFile(self.source) as archive:
            body = archive.read("word/document.xml").decode("utf-8")
        anchors = re.findall(r"<w:comment(?:RangeStart|RangeEnd|Reference)\b[^>]*>", body)
        self.assertEqual(len(anchors), 3, anchors)

        # 「仅显示译文」会整段改写正文，是扫掉引用标记风险最大的一条路。
        output = write_bilingual_docx(
            source_path=self.source,
            output_dir=self.root / "out",
            translations={"正文原文": "Texte du corps"},
            target_lang="fr",
            output_translation_only=True,
        )

        with zipfile.ZipFile(output) as archive:
            rewritten = archive.read("word/document.xml").decode("utf-8")
        self.assertEqual(
            re.findall(r"<w:comment(?:RangeStart|RangeEnd|Reference)\b[^>]*>", rewritten),
            anchors,
            "批注的引用标记在改写段落时被扫掉了，Word 里那条批注会挂空",
        )
        self.assertEqual(self._comments_xml(output), original)
        self.assertEqual([c.text for c in Document(output).comments], ["这段要请法务复核。"])

    def test_body_text_next_to_a_comment_is_still_extracted(self) -> None:
        # 批注挂在正文 run 上时，正文那一段不能被连带跳过。
        self.assertEqual(
            [s.source for s in extract_word_segments(self.source, target_lang="fr")],
            ["正文原文"],
        )
        self.assertEqual(
            [s.source for s in extract_word_story_segments(self.source, target_lang="fr")],
            [],
        )


class WordChartAndSmartArtTests(unittest.TestCase):
    """Word 里的图表和 SmartArt：文字住在独立的 chartN.xml / dgm dataN.xml 里。

    审计 D2 线把「Word 内 SmartArt/图表」列为未实测项。它们不在 python-docx 的
    正文遍历范围内，所以「翻不翻」是产品范围问题；这里补的是实测下限：翻正文的
    时候图形部件必须一个字节都不动（``_check_untouched_parts`` 会把改动报成部件名
    并让整次写出失败），图形在正文里的引用也不能丢。
    """

    CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    DIAGRAM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

    def setUp(self) -> None:
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.root = Path(self._tmp.name)
        self.source = self.root / "charts.docx"
        doc = Document()
        doc.add_paragraph("正文原文")
        self.chart_payload = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<c:chartSpace xmlns:c="{self.CHART_NS}" xmlns:a="{A}"><c:chart><c:plotArea>'
            '<c:barChart><c:ser><c:tx><c:strRef><c:strCache><c:pt idx="0">'
            "<c:v>销售额</c:v></c:pt></c:strCache></c:strRef></c:tx></c:ser></c:barChart>"
            "</c:plotArea></c:chart></c:chartSpace>"
        ).encode()
        self.diagram_payload = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<dgm:dataModel xmlns:dgm="{self.DIAGRAM_NS}" xmlns:a="{A}"><dgm:ptLst>'
            '<dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>进度计划</a:t></a:r></a:p>'
            "</dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>"
        ).encode()
        chart_rid = _part(
            doc,
            "/word/charts/chart1.xml",
            "application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
            self.chart_payload,
            RT.CHART,
        )
        diagram_rid = _part(
            doc,
            "/word/diagrams/data1.xml",
            "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
            self.diagram_payload,
            RT.DIAGRAM_DATA,
        )
        end = doc.element.body.find(qn("w:sectPr"))
        end.addprevious(
            parse_xml(
                _graphic_paragraph(
                    f'<c:chart r:id="{chart_rid}"/>',
                    self.CHART_NS,
                    f'xmlns:c="{self.CHART_NS}"',
                )
            )
        )
        end.addprevious(
            parse_xml(
                _graphic_paragraph(
                    f'<dgm:relIds r:dm="{diagram_rid}"/>',
                    self.DIAGRAM_NS,
                    f'xmlns:dgm="{self.DIAGRAM_NS}"',
                )
            )
        )
        doc.save(self.source)

    def test_chart_and_diagram_labels_stay_out_of_the_translation_scope(self) -> None:
        sources = [
            segment.source
            for segment in extract_word_segments(self.source, target_lang="fr")
        ]
        # 图表系列名「销售额」和 SmartArt 的「进度计划」都不在正文遍历范围内。
        self.assertEqual(sources, ["正文原文"])

    def test_chart_and_diagram_parts_are_copied_byte_for_byte(self) -> None:
        output = write_bilingual_docx(
            source_path=self.source,
            output_dir=self.root / "out",
            translations={"正文原文": "Texte du corps"},
            target_lang="fr",
        )
        with zipfile.ZipFile(output) as archive:
            self.assertEqual(archive.read("word/charts/chart1.xml"), self.chart_payload)
            self.assertEqual(archive.read("word/diagrams/data1.xml"), self.diagram_payload)
            self.assertIn(b"<w:drawing>", archive.read("word/document.xml"))
        # 图形段落本身没有文字，只看有字的段落。
        self.assertEqual(
            [p.text for p in Document(output).paragraphs if p.text],
            ["正文原文", "Texte du corps"],
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
