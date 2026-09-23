"""Cross-surface name settings and Excel's actual file-write integration."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from core.file_scanner import FileItem
from core.output_name_translation import (
    avoid_bilingual_name_collision,
    translate_names,
    translate_output_stem,
)
from core.pdf_image_translation import (
    PdfFileItem,
    PdfFileRecord,
    PdfImageTranslationRunner,
)
from core.task_runner import TaskRunner
from settings import AppSettings
from tests.test_excel_resume import (
    _fake_translate,
    _make_xlsx,
    _pipeline_patches,
    _run_and_get_done,
    _settings,
)


class _NameEngine:
    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        return {text: {"报告": "Report", "数据": "Data"}.get(text, text) for text in texts}


class OutputNameTranslationTests(unittest.TestCase):
    def test_defaults_and_translation_helper(self):
        settings = AppSettings()
        self.assertFalse(settings.excel_output.translate_output_filename)
        self.assertTrue(settings.excel_output.translate_sheet_names)
        self.assertFalse(settings.word_output.translate_output_filename)
        self.assertFalse(settings.pdf_output.translate_output_filename)
        self.assertEqual(translate_output_stem(_NameEngine(), "报告", "en", "zh"), "Report")
        self.assertEqual(translate_names(_NameEngine(), ["数据"], "en", "zh"), {"数据": "Data"})

    def test_excel_task_writes_translated_file_and_sheet_names(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = _make_xlsx(root / "报告.xlsx", {"A1": "你好"}, sheet_name="数据")
            settings = _settings()
            settings.excel_output.translate_output_filename = True
            settings.excel_output.translate_sheet_names = True
            with _pipeline_patches(translate_side_effect=_fake_translate([])):
                with patch("core.task_runner.translate_output_stem", return_value="Report"), patch(
                    "core.task_runner.translate_names", return_value={"数据": "Data"}
                ):
                    runner = TaskRunner(
                        [FileItem(path=source, name="报告", size_kb=1.0)],
                        settings,
                        source_root=root,
                    )
                    done = _run_and_get_done(runner)
            result = done.file_results[0]
            self.assertTrue(result.get("success"), result)
            output = Path(result["output_path"])
            self.assertEqual(output.name, "Report_英文_双语.xlsx")
            book = load_workbook(output)
            try:
                self.assertIn("Data", book.sheetnames)
                self.assertIn("数据_原文", book.sheetnames)
                self.assertEqual(book["Data"]["A1"].value, "你好\nT:你好")
            finally:
                book.close()

    def test_pdf_output_paths_use_translated_stem_only_when_enabled(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "报告.pdf"
            source.write_bytes(b"%PDF-1.4\n")
            item = PdfFileItem(path=source, name=source.name, size_kb=1.0)
            settings = AppSettings()
            settings.pdf_output.translate_output_filename = True
            settings.pdf.target_lang = "en"
            runner = PdfImageTranslationRunner([item], settings, source_root=root)
            record = PdfFileRecord(
                name=item.name,
                source_path=str(source),
                relative_path=item.name,
                source_copy_path=str(root / "output" / item.name),
            )
            with patch("core.engine_dispatcher.build_engine", return_value=_NameEngine()):
                prepared = runner._prepared_file_shell(
                    item,
                    relative_pdf=Path(item.name),
                    record=record,
                    output_dir=root / "output",
                    app_managed=False,
                )
            self.assertEqual(prepared.translated_pdf_path.name, "Report_英文_高清.pdf")
            self.assertEqual(prepared.compressed_pdf_path.name, "Report_英文_压缩.pdf")
            self.assertEqual(record.source_copy_path, str(root / "output" / "报告.pdf"))

    def test_identical_translated_file_names_get_distinct_outputs(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            (output / "Report_英文_双语.xlsx").touch()
            self.assertEqual(
                avoid_bilingual_name_collision(output, "Report.xlsx", "en"),
                "Report_2.xlsx",
            )

    def test_excel_dynamic_reference_keeps_entire_original_sheet_set(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = _make_xlsx(root / "报告.xlsx", {"A1": "你好"}, sheet_name="数据")
            book = load_workbook(source)
            book.create_sheet("Summary")["A1"] = '=INDIRECT("数据!A1")'
            book.save(source)
            book.close()
            settings = _settings()
            settings.excel_output.translate_sheet_names = True
            settings.excel_output.formula_display_value_backfill = False
            with _pipeline_patches(translate_side_effect=_fake_translate([])):
                with patch(
                    "core.task_runner.translate_names",
                    return_value={"数据": "Data", "Summary": "Summary"},
                ):
                    runner = TaskRunner(
                        [FileItem(path=source, name="报告", size_kb=1.0)],
                        settings,
                        source_root=root,
                    )
                    done = _run_and_get_done(runner)
            result = done.file_results[0]
            self.assertTrue(result["success"], result)
            self.assertEqual(result["sheet_name_status"], "preserved")
            self.assertTrue(any(issue.get("type") == "sheet_rename_preserved" for issue in done.issues))
            output = load_workbook(result["output_path"])
            try:
                self.assertIn("数据", output.sheetnames)
                self.assertIn("Summary", output.sheetnames)
                self.assertNotIn("Data", output.sheetnames)
            finally:
                output.close()


if __name__ == "__main__":
    unittest.main()
