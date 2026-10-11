"""自动识别源语言时，模型报「本来就是目标语言」的段落原样保留，不重试也不标需复核。"""
from __future__ import annotations

import tempfile
import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from core.language_preflight import TranslationLanguageResult
from core.task_runner import DoneMsg
from core.word_document import WordFileItem, WordSegment
from core.word_quality_report import _append_post_write_coverage_issues
from core.word_task_runner import WordTaskRunner
from tests.app_data_isolation import IsolatedAppDataTestCase
from tests import test_phase5_word_contracts as _contracts

ENGLISH = "Already English paragraph about the pump room."


class PostWriteCoverageTests(unittest.TestCase):
    def test_paragraphs_reported_as_target_language_are_not_flagged(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "out.docx"
            document = Document()
            document.add_paragraph(ENGLISH)
            document.save(path)

            flagged: list[dict] = []
            _append_post_write_coverage_issues(
                issues=flagged, file_name="out.docx", output_path=path,
                target_lang="fr", source_lang="en",
            )
            self.assertTrue(flagged, "前提：不排除时这一段会被报成未译源文")

            issues: list[dict] = []
            _append_post_write_coverage_issues(
                issues=issues, file_name="out.docx", output_path=path,
                target_lang="fr", source_lang="en",
                already_target_sources={ENGLISH},
            )
            self.assertEqual(issues, [])


class RunnerTests(IsolatedAppDataTestCase):
    def test_identity_in_target_language_skips_retry_and_reaches_post_write_check(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.docx"
            Document().save(source)
            prepared = SimpleNamespace(
                path=source, method="编号预处理：Python 兜底", temp_paths=(),
                fallback_messages=(), labels_seen=0, labels_prepended=0,
                conversion_method="not_required", conversion_fidelity="not_required",
                numbering_method="python_conservative", numbering_fallback_messages=(),
            )
            settings = _contracts.WordTaskResultContractTests._settings()
            runner = WordTaskRunner(
                [WordFileItem(path=source, name=source.name, size_kb=1.0)],
                settings, source_root=root, source_lang="auto",
            )

            class _PreflightEngine:
                engine_name = "mock"

                def chat(self, system: str, user: str) -> str:
                    return '{"source_langs":["en"]}'

            def translate(texts, _engine, target_lang, *_args, **kwargs):
                for text in texts:
                    kwargs["source_result_callback"](
                        text,
                        TranslationLanguageResult(
                            source_text=text, translation=text,
                            source_lang=target_lang, target_lang=target_lang,
                        ),
                    )
                    kwargs["candidate_callback"](text, text)
                return {text: text for text in texts}

            with ExitStack() as stack:
                _contracts.WordTaskResultContractTests._runner_patches(self, stack, root=root, prepared_by_path={source: prepared})
                coverage = stack.enter_context(
                    patch("core.word_task_runner._append_post_write_coverage_issues", return_value=0)
                )
                retries = stack.enter_context(
                    patch("core.word_task_runner._WordRecoveryPool.add_candidate")
                )
                stack.enter_context(patch("core.word_task_runner.build_engine", return_value=_PreflightEngine()))
                stack.enter_context(patch(
                    "core.word_task_runner.extract_word_segments",
                    return_value=[WordSegment(source=ENGLISH, kind="paragraph", location="正文第 1 段")],
                ))
                stack.enter_context(patch(
                    "core.word_task_runner.tm_manager.lookup_batch",
                    return_value={ENGLISH: None},
                ))
                stack.enter_context(patch("core.word_task_runner.translate_word_texts", side_effect=translate))
                runner._run()

            _contracts.WordTaskResultContractTests._terminal_message(runner, DoneMsg)
            retries.assert_not_called()
            self.assertEqual(coverage.call_args.kwargs["already_target_sources"], {ENGLISH})


if __name__ == "__main__":
    unittest.main(verbosity=2)
