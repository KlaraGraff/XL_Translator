"""PDF 转 Word 后的两处整理（本机实测，法文施工方法说明，3 页全走转成 Word）。

1. 两条「•」在 PDF 里各占一行，pdf2docx 却合进了同一段（各是一个 run、中间没有空格），
   Word 里排成「…programme validé.• Enregistrement…」。只输出译文时模型把它拆回两行，
   写回行数对不上，这一格就成了法文 + 中文。现在拆回两段。
2. 列表项的折行被写成制表符：「…si elle<tab>est prévue au programme.」，后半截被译成
   了俄文词「предусмотр」。现在换成空格，并用悬挂缩进让折行和正文对齐。
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from core.pdf_word_route import merge_wrapped_lines, split_glued_bullets

GLUED = (
    '<w:p {ns}><w:pPr><w:spacing w:line="259" w:lineRule="auto" w:before="128" w:after="40"/>'
    '<w:ind w:left="212" w:right="1584" w:firstLine="0"/><w:jc w:val="left"/>{sect}</w:pPr>'
    '<w:r><w:rPr><w:rFonts w:ascii="DejaVu Sans"/><w:sz w:val="16"/></w:rPr>'
    '<w:t>• Réalisation des essais au scléromètre selon un maillage défini.</w:t></w:r>'
    '<w:r><w:rPr><w:rFonts w:ascii="DejaVu Sans"/><w:sz w:val="16"/></w:rPr>'
    '<w:t>• Repérage et identification des points d’essai.</w:t></w:r></w:p>'
)
WRAPPED = (
    '<w:p {ns}><w:pPr><w:tabs><w:tab w:val="left" w:pos="352"/></w:tabs>'
    '<w:ind w:left="212" w:firstLine="0"/></w:pPr>'
    '<w:r><w:t xml:space="preserve">• Planification des essais, notamment la résistance du béton si elle </w:t></w:r>'
    '<w:r><w:tab/></w:r><w:r><w:t>est prévue au programme.</w:t></w:r></w:p>'
)


class PdfWordBulletTests(unittest.TestCase):
    def _docx(self, *xml_paragraphs: str, in_cell: bool = False) -> Path:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "converted.docx"
        doc = Document()
        container = doc.add_table(rows=1, cols=1).cell(0, 0) if in_cell else None
        for xml in xml_paragraphs:
            element = parse_xml(xml.format(ns=nsdecls("w"), sect=""))
            if container is not None:
                container._tc.append(element)
            else:
                doc.element.body.insert(len(doc.element.body) - 1, element)
        doc.save(str(path))
        return path

    def test_glued_bullets_in_a_table_cell_become_two_paragraphs(self) -> None:
        path = self._docx(GLUED, in_cell=True)
        self.assertEqual(split_glued_bullets(path), 1)
        cell = Document(str(path)).tables[0].cell(0, 0)
        texts = [p.text for p in cell.paragraphs if p.text.strip()]
        self.assertEqual(
            texts,
            [
                "• Réalisation des essais au scléromètre selon un maillage défini.",
                "• Repérage et identification des points d’essai.",
            ],
        )
        first, second = [p for p in cell.paragraphs if p.text.strip()]
        first_spacing = first._p.pPr.find(qn("w:spacing"))
        second_spacing = second._p.pPr.find(qn("w:spacing"))
        self.assertEqual(first_spacing.get(qn("w:before")), "128")
        self.assertEqual(first_spacing.get(qn("w:after")), "0")
        self.assertEqual(second_spacing.get(qn("w:before")), "0")
        self.assertEqual(second_spacing.get(qn("w:after")), "40")
        self.assertEqual(second.runs[0].font.size.pt, 8)

    def test_bullets_glued_inside_one_run_are_split_too(self) -> None:
        xml = GLUED.replace("</w:t></w:r><w:r><w:rPr><w:rFonts w:ascii=\"DejaVu Sans\"/><w:sz w:val=\"16\"/></w:rPr><w:t>", "")
        path = self._docx(xml)
        self.assertEqual(split_glued_bullets(path), 1)
        texts = [p.text for p in Document(str(path)).paragraphs if p.text.strip()]
        self.assertEqual(len(texts), 2, texts)
        self.assertTrue(texts[0].endswith("défini."))
        self.assertTrue(texts[1].startswith("• Repérage"))

    def test_section_break_stays_on_the_last_piece(self) -> None:
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        path = Path(tmp.name) / "sections.docx"
        doc = Document()
        sect = '<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>'
        doc.element.body.insert(0, parse_xml(GLUED.format(ns=nsdecls("w"), sect=sect)))
        doc.save(str(path))
        split_glued_bullets(path)
        paragraphs = [p for p in Document(str(path)).paragraphs if p.text.strip()]
        self.assertIsNone(paragraphs[0]._p.pPr.find(qn("w:sectPr")))
        self.assertIsNotNone(paragraphs[1]._p.pPr.find(qn("w:sectPr")))

    def test_separator_bullets_inside_a_sentence_are_left_alone(self) -> None:
        xml = (
            '<w:p {ns}><w:r><w:t xml:space="preserve">Rabat </w:t></w:r>'
            '<w:r><w:t>• Casablanca • Tanger</w:t></w:r></w:p>'
        )
        path = self._docx(xml)
        self.assertEqual(split_glued_bullets(path), 0)

    def test_tab_inside_a_wrapped_sentence_becomes_a_space_with_a_hanging_indent(self) -> None:
        path = self._docx(WRAPPED)
        merge_wrapped_lines(path)
        paragraph = next(p for p in Document(str(path)).paragraphs if p.text.strip())
        self.assertEqual(
            paragraph.text,
            "• Planification des essais, notamment la résistance du béton si elle est prévue au programme.",
        )
        ind = paragraph._p.pPr.find(qn("w:ind"))
        self.assertEqual(ind.get(qn("w:left")), "352")
        self.assertEqual(ind.get(qn("w:hanging")), "140")

    def test_alignment_tabs_are_kept(self) -> None:
        xml = '<w:p {ns}><w:r><w:t>Date</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>12/10/2026</w:t></w:r></w:p>'
        path = self._docx(xml)
        merge_wrapped_lines(path)
        paragraph = next(p for p in Document(str(path)).paragraphs if p.text.strip())
        self.assertEqual(paragraph.text, "Date\t12/10/2026")


if __name__ == "__main__":
    unittest.main(verbosity=2)
