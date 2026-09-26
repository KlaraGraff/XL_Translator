"""A failed Word write must not expose an original or a partial DOCX as output."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.document import Document as DocumentClass

from core.word_coverage import (
    apply_coverage_review_marks,
    build_word_coverage_plan,
    write_untranslated_docx,
)
from core.word_document import write_bilingual_docx


class WordAtomicOutputTests(unittest.TestCase):
    def _paths(self, root: Path) -> tuple[Path, Path]:
        source = root / "source.docx"
        document = Document()
        document.add_paragraph("项目名称")
        document.save(source)
        output_dir = root / "out"
        output_dir.mkdir()
        return source, output_dir

    def _write(self, kind: str, source: Path, output_dir: Path) -> Path:
        options = {
            "source_path": source,
            "output_dir": output_dir,
            "translations": {"项目名称": "Project name"},
            "target_lang": "en",
        }
        if kind == "bilingual":
            return write_bilingual_docx(**options)
        plan = build_word_coverage_plan(source, target_lang="en")
        return write_untranslated_docx(**options, plan=plan)

    def test_partial_save_preserves_existing_output_and_cleans_staging(self) -> None:
        for kind in ("bilingual", "untranslated"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as tmp:
                source, output_dir = self._paths(Path(tmp))
                final = self._write(kind, source, output_dir)
                previous = final.read_bytes()

                def partial_save(_document, path):
                    Path(path).write_bytes(b"partial Word package")
                    raise OSError("disk full during save")

                with patch.object(DocumentClass, "save", partial_save):
                    with self.assertRaisesRegex(OSError, "disk full"):
                        self._write(kind, source, output_dir)

                self.assertEqual(final.read_bytes(), previous)
                self.assertEqual(list(output_dir.iterdir()), [final])

    def test_failed_write_leaves_no_apparent_output(self) -> None:
        for kind in ("bilingual", "untranslated"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as tmp:
                source, output_dir = self._paths(Path(tmp))
                with patch.object(DocumentClass, "save", side_effect=OSError("save failed")):
                    with self.assertRaisesRegex(OSError, "save failed"):
                        self._write(kind, source, output_dir)
                self.assertEqual(list(output_dir.iterdir()), [])

    def test_unopenable_saved_package_does_not_replace_existing_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source, output_dir = self._paths(Path(tmp))
            final = self._write("bilingual", source, output_dir)
            previous = final.read_bytes()

            def corrupt_save(_document, path):
                Path(path).write_bytes(b"not a DOCX")

            with patch.object(DocumentClass, "save", corrupt_save):
                with self.assertRaises(Exception):
                    self._write("bilingual", source, output_dir)

            self.assertEqual(final.read_bytes(), previous)
            self.assertEqual(list(output_dir.iterdir()), [final])

    def test_successful_write_replaces_output_and_removes_staging(self) -> None:
        for kind in ("bilingual", "untranslated"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as tmp:
                source, output_dir = self._paths(Path(tmp))
                final = self._write(kind, source, output_dir)
                self.assertIn("Project name", [p.text for p in Document(final).paragraphs])
                self.assertEqual(list(output_dir.iterdir()), [final])

    def test_logging_callback_failure_does_not_turn_committed_file_into_failure(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source, output_dir = self._paths(Path(tmp))
            final = write_bilingual_docx(
                source_path=source,
                output_dir=output_dir,
                translations={"项目名称": "Project name"},
                target_lang="en",
                log_callback=lambda _message: (_ for _ in ()).throw(RuntimeError("log unavailable")),
            )
            self.assertIn("Project name", [p.text for p in Document(final).paragraphs])

    def test_review_mark_save_failure_preserves_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            source, output_dir = self._paths(Path(tmp))
            plan = build_word_coverage_plan(source, target_lang="en")
            final = output_dir / "review.docx"
            final.write_bytes(source.read_bytes())
            previous = final.read_bytes()

            def partial_save(_document, path):
                Path(path).write_bytes(b"partial review package")
                raise OSError("review save failed")

            with patch.object(DocumentClass, "save", partial_save):
                with self.assertRaisesRegex(OSError, "review save failed"):
                    apply_coverage_review_marks(final, plan=plan)

            self.assertEqual(final.read_bytes(), previous)
            self.assertEqual(list(output_dir.iterdir()), [final])


if __name__ == "__main__":
    unittest.main()
