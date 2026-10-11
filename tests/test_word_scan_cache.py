"""同一份 Word 文档在点「开始」前后会被连扫几遍，统计结果要复用，文件一改就重新算。"""

from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from core import word_document


class WordScanCountCacheTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "说明.docx"
        doc = Document()
        doc.add_paragraph("第一段")
        doc.add_paragraph("第二段")
        doc.save(str(self.path))

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_repeat_scan_reuses_counts_until_the_file_changes(self) -> None:
        real = word_document.extract_word_segments
        with patch.object(word_document, "extract_word_segments", side_effect=real) as extract:
            first = word_document.scan_word_path(self.path)[0]
            second = word_document.scan_word_path(self.path)[0]
            self.assertEqual(extract.call_count, 1)
            self.assertEqual(first.translatable_count, second.translatable_count)

            doc = Document(str(self.path))
            doc.add_paragraph("第三段")
            doc.save(str(self.path))
            stat = self.path.stat()
            os.utime(self.path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000))
            third = word_document.scan_word_path(self.path)[0]
            self.assertEqual(extract.call_count, 2)
            self.assertEqual(third.paragraph_count, first.paragraph_count + 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
