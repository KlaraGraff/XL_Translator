"""审计 高-12：「公式显示值回填」关闭时，公式格不能再送去翻译。"""

from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from core.task_runner import TaskRunner


def _write_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "项目名称"
    ws["B1"] = '=IF(1,"公式结果","否")'
    wb.save(path)
    # openpyxl 不算公式，存出来的 <v> 是空的；补上 Excel 会缓存的显示值，
    # 才能模拟真实文件里「公式格读出来是一段中文」的情况。
    with zipfile.ZipFile(path) as src:
        parts = {name: src.read(name) for name in src.namelist()}
    sheet = "xl/worksheets/sheet1.xml"
    parts[sheet] = parts[sheet].replace(
        b"</f><v></v></c>", '</f><v>公式结果</v></c>'.encode()
    ).replace(b'<c r="B1">', b'<c r="B1" t="str">')
    with zipfile.ZipFile(path, "w") as dst:
        for name, data in parts.items():
            dst.writestr(name, data)


class CollectTextsFormulaBackfillTests(unittest.TestCase):
    def test_formula_cells_are_skipped_when_backfill_is_off(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            _write_workbook(path)

            texts, sheet_count = TaskRunner._collect_texts(
                path, path.name, target_lang="en", formula_display_value_backfill=False
            )

        self.assertEqual(sheet_count, 1)
        self.assertEqual(texts, ["项目名称"])
        self.assertFalse(any(t.startswith("=") for t in texts))

    def test_plain_text_still_collected_when_backfill_is_on(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "book.xlsx"
            _write_workbook(path)

            texts, _ = TaskRunner._collect_texts(path, path.name, target_lang="en")

        self.assertCountEqual(texts, ["项目名称", "公式结果"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
