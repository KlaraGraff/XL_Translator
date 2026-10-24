"""Word 段内局部格式：用标记随原文发给模型，译完按标记还原。

整段翻译时译文写进一个 run，段内的局部格式（「注意：」加粗、句中某个词斜体、带链接
的文字、m² 的上标）就合成了一种。这里沿用专业翻译软件和 DeepL、BabelDOC 的做法：
段落里格式和主体不同的那几截用 ``<g1>…</g1>`` 包住随原文发出，模型把标记放到对应
的译文上，写回时每截按原来的格式（和链接）各写一个 run。

段里的制表符、软换行、行内图片、脚注编号、公式这类不能翻译的东西，写成自闭合的占位符
``<x1/>`` 随原文发出，模型把它放到译文里对应的位置，写回时那里原样放回原来的节点
（「1.2⇥**施工范围**见附图」的制表符和加粗都保得住；「强度按 [公式] 计算」的公式也留在
句中原处，双语版式的译文段里同样有一份）。段首段尾的占位符不发给模型，
写回时固定留在段首段尾。占位符去掉标记后还原成纯文本里对应的字符（制表符 ``\t``、
换行 ``\n``，图片、脚注编号和公式为空），下游看到的纯文本和以前一样。

标记的写法和校验与 PPT、Excel 共用：成对的 ``<gN>…</gN>`` 管格式，``<xN/>`` 管占位，
两套编号各自从 1 起。

保底：只有带局部格式、或段中间有占位符，且结构简单（只有文字 run、超链接和上面这些
占位内容，没有域、修订）的段落才加标记；译文里的标记少了、多了、乱了，就去掉标记、
按现在的整段写法写回，不报错。去掉标记后的纯文本照常走校验、记忆库和复核，和以前
完全一样。
"""
from __future__ import annotations

import re
import threading
from copy import deepcopy
from dataclasses import dataclass, field, replace

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.prompt_cache import append_batch_block

# 一段里最多加几对标记；再多说明格式很碎，模型放错的概率也高，按整段写法处理
MAX_GROUPS_PER_UNIT = 12
# 一段里最多几个占位符（相邻的已并成一个）；用制表符排成几列的段落常有七八个，学术论文
# 一段里的行内公式常有二三十个（反馈 #19）。模型弄丢几个时由 repair_marks 补回
MAX_PLACEHOLDERS_PER_UNIT = 40

_TAG_RE = re.compile(r"<\s*(/?)\s*g\s*(\d+)\s*>", re.IGNORECASE)
_SLOT_RE = re.compile(r"<\s*x\s*(\d+)\s*/\s*>", re.IGNORECASE)
_ANY_MARK_RE = re.compile(r"<\s*(/?)\s*g\s*(\d+)\s*>|<\s*x\s*(\d+)\s*/\s*>", re.IGNORECASE)

MARK_PROMPT_RULE = (
    "部分词条带有 <g1>…</g1> 这样成对的格式标记，标出原文中加粗、斜体、链接等局部格式。"
    "译文必须保留每一对标记，各包住与原文标记内容对应的译文；标记位置可以随译文语序调整，"
    "但不得增加、删除、合并、嵌套或改动标记。标记两侧的空格按译文语言的习惯书写，"
    "如「<g1>Note :</g1> les cotes…」冒号后照常留空格。"
    "部分词条还带有 <x1/> 这样的自闭合占位符，代表制表符、换行、图片、脚注编号、公式等不翻译的内容："
    "每个占位符在译文中必须原样出现恰好一次，放在译文里对应的位置，可以在标记里面，"
    "不得删除、复制或改写。"
)

# Word 公式（OMML）：段落里和 w:r 平级的节点，python-docx 的 paragraph.text 读不到它
_MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
MATH_TAGS = frozenset((f"{_MATH_NS}oMath", f"{_MATH_NS}oMathPara"))
_PARAGRAPH_CHILDREN = frozenset(
    qn(tag) for tag in ("w:pPr", "w:r", "w:hyperlink", "w:bookmarkStart", "w:bookmarkEnd", "w:proofErr",
                        "w:permStart", "w:permEnd")
) | MATH_TAGS
_HYPERLINK_CHILDREN = frozenset((qn("w:r"), qn("w:proofErr")))
_RUN_CHILDREN = frozenset((qn("w:rPr"), qn("w:t"), qn("w:lastRenderedPageBreak")))
# run 里排在文字中间、但不是文字的东西，作为占位符随原文发出、写回时照抄原节点。
# 版式类（制表符、软换行、不断行连字符、符号）在双语版式新插的译文段里照样复制一份；
_LAYOUT_SLOT_TAGS = frozenset(
    qn(tag) for tag in ("w:tab", "w:br", "w:cr", "w:ptab", "w:noBreakHyphen", "w:softHyphen", "w:sym")
)
# 对象类（图片、嵌入对象、文本框、脚注/尾注/批注的引用）只能有一份，双语的译文段里不复制
_OBJECT_SLOT_TAGS = frozenset(
    (
        *(qn(tag) for tag in ("w:drawing", "w:pict", "w:object", "w:footnoteReference", "w:endnoteReference",
                              "w:footnoteRef", "w:endnoteRef", "w:commentReference", "w:annotationRef")),
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent",
    )
)
_SLOT_TAGS = _LAYOUT_SLOT_TAGS | _OBJECT_SLOT_TAGS
# python-docx 的 Run.text 把这几种读成字符（见 docx.oxml.text.run.CT_R.text），其余读成空
_TEXT_EQUIVALENT_TAGS = frozenset(qn(tag) for tag in ("w:tab", "w:br", "w:cr", "w:ptab", "w:noBreakHyphen"))
# 不影响「看起来是不是同一种格式」的属性：字体名（中西文混排时 run 常按文种拆开）、
# 语言、拼写检查、复杂文种的对应项、字距、Word 2010 的文字效果
_IGNORED_RPR = frozenset(
    qn(tag) for tag in ("w:rFonts", "w:lang", "w:noProof", "w:szCs", "w:bCs", "w:iCs", "w:kern", "w:spacing",
                        "w:rtl", "w:cs", "w:eastAsianLayout", "w:specVanish", "w:webHidden", "w:snapToGrid",
                        "w:fitText", "w:w")
)
_ON_OFF = frozenset(
    qn(tag) for tag in ("w:b", "w:i", "w:caps", "w:smallCaps", "w:strike", "w:dstrike", "w:vanish", "w:outline",
                        "w:shadow", "w:emboss", "w:imprint")
)
_W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


@dataclass(frozen=True)
class MarkGroup:
    rpr: object | None
    hyperlink: object | None = None


@dataclass(frozen=True)
class Slot:
    """译文里的一个占位符 ``<xN/>``（parse_marked 拆出来的）。"""

    id: int


@dataclass(frozen=True)
class Placeholder:
    """原文里一处占位内容：一个或几个相邻的节点，各带所在 run 的 rPr，写回时原样放回。"""

    nodes: tuple  # ((节点, rPr 或 None), ...)
    text: str  # 在纯文本里对应的字符

    @property
    def copyable(self) -> bool:
        """只有制表符、软换行这类版式（不含分页、分栏）和公式，双语的译文段里可以再复制一份。"""
        return all(
            node.tag in MATH_TAGS
            or (
                node.tag in _LAYOUT_SLOT_TAGS
                and not (node.tag == qn("w:br") and (node.get(qn("w:type")) or "textWrapping") != "textWrapping")
            )
            for node, _rpr in self.nodes
        )


@dataclass
class ParagraphMarks:
    """一段（或单元格里的一段）带标记的原文，以及每对标记对应的格式、每个占位符对应的节点。"""

    marked_text: str
    plain_text: str
    groups: dict[int, MarkGroup] = field(default_factory=dict)
    base_rpr: object | None = None
    # 段中间的占位符（随原文发出）
    placeholders: dict[int, Placeholder] = field(default_factory=dict)
    # 段首、段尾不发给模型、写回时固定放回的内容：[(文字或 Placeholder, None)]
    lead: list = field(default_factory=list)
    trail: list = field(default_factory=list)

    @property
    def is_marked(self) -> bool:
        return bool(self.groups or self.placeholders)

    @property
    def placeholder_texts(self) -> dict[int, str]:
        return {key: item.text for key, item in self.placeholders.items()}

    @property
    def slot_hints(self) -> dict[int, str]:
        """补标记时告诉模型每个占位符是什么（公式写成 σ=M/W 这样的线性文字）。"""
        return {key: hint for key, item in self.placeholders.items() if (hint := placeholder_hint(item))}

    def pieces_for(self, translation: str):
        """带标记的译文拆成可写回的片段（占位符换成 Placeholder，前后接上段首段尾）；
        标记对不上返回 None。补标记时实在放不回的格式标记会去掉（那几截用主体格式），
        所以译文里的格式标记可以比原文少，占位符必须一个不少。"""
        present = {int(group) for closing, group in _TAG_RE.findall(translation) if not closing}
        pieces = None
        if present <= set(self.groups):
            pieces = parse_marked(translation, present, self.placeholders)
        if pieces is None:
            return None
        resolved = [(self.placeholders[text.id] if isinstance(text, Slot) else text, group) for text, group in pieces]
        return [*self.lead, *resolved, *self.trail]


@dataclass
class CellMarks:
    marked_text: str
    paragraphs: list[ParagraphMarks]

    @property
    def groups(self) -> dict[int, MarkGroup]:
        merged: dict[int, MarkGroup] = {}
        for item in self.paragraphs:
            merged.update(item.groups)
        return merged

    @property
    def is_marked(self) -> bool:
        return any(item.is_marked for item in self.paragraphs)

    @property
    def plain_text(self) -> str:
        return "\n".join(item.plain_text for item in self.paragraphs)

    @property
    def placeholder_texts(self) -> dict[int, str]:
        merged: dict[int, str] = {}
        for item in self.paragraphs:
            merged.update(item.placeholder_texts)
        return merged

    @property
    def slot_hints(self) -> dict[int, str]:
        merged: dict[int, str] = {}
        for item in self.paragraphs:
            merged.update(item.slot_hints)
        return merged


def _on_off_value(element) -> bool:
    value = str(element.get(qn("w:val")) or "").lower()
    return value not in ("0", "false", "off", "none")


def _look_key(rpr) -> tuple:
    if rpr is None:
        return ()
    items = []
    for child in rpr:
        tag = child.tag
        if not isinstance(tag, str) or tag in _IGNORED_RPR or tag.startswith(_W14):
            continue
        if tag in _ON_OFF:
            if _on_off_value(child):
                items.append((tag,))
            continue
        if tag == qn("w:u") and str(child.get(qn("w:val")) or "").lower() == "none":
            continue
        items.append((tag, tuple(sorted(child.attrib.items()))))
    return tuple(sorted(items))


@dataclass
class _Item:
    """段落里按顺序的一截：文字，或（node_pairs 非空时）一个占位符。"""

    text: str
    rpr: object | None
    hyperlink: object | None
    node_pairs: list = field(default_factory=list)

    @property
    def is_slot(self) -> bool:
        return bool(self.node_pairs)


def _content_runs(paragraph_element) -> list[tuple[object, object | None]] | None:
    """[(w:r 或公式, 所在的 w:hyperlink 或 None)]；结构不简单时返回 None。"""
    runs: list[tuple[object, object | None]] = []
    for child in paragraph_element:
        tag = child.tag
        if not isinstance(tag, str) or tag not in _PARAGRAPH_CHILDREN:
            return None
        if tag == qn("w:r") or tag in MATH_TAGS:
            runs.append((child, None))
        elif tag == qn("w:hyperlink"):
            for nested in child:
                if not isinstance(nested.tag, str) or nested.tag not in _HYPERLINK_CHILDREN:
                    return None
                if nested.tag == qn("w:r"):
                    runs.append((nested, child))
    for run, _hyperlink in runs:
        if run.tag in MATH_TAGS:
            continue
        if any(not isinstance(node.tag, str) or node.tag not in _RUN_CHILDREN | _SLOT_TAGS for node in run):
            return None
    return runs


def _content_items(runs) -> list[_Item]:
    """文字一截一项；相邻的占位内容（中间没有文字）并成一项。"""
    items: list[_Item] = []
    for run, hyperlink in runs:
        if run.tag in MATH_TAGS:
            # 整个公式是一个占位：不在纯文本里，写回时不套 w:r，直接放回段落
            if items and items[-1].is_slot:
                items[-1].node_pairs.append((run, None))
            else:
                items.append(_Item("", None, hyperlink, [(run, None)]))
            continue
        rpr = run.find(qn("w:rPr"))
        for node in run:
            if node.tag == qn("w:t"):
                if node.text:
                    items.append(_Item(node.text, rpr, hyperlink))
            elif node.tag in _SLOT_TAGS:
                text = str(node) if node.tag in _TEXT_EQUIVALENT_TAGS else ""
                if items and items[-1].is_slot:
                    items[-1].text += text
                    items[-1].node_pairs.append((node, rpr))
                else:
                    items.append(_Item(text, rpr, hyperlink, [(node, rpr)]))
    return items


def paragraph_marks(paragraph, *, start_id: int = 1, start_slot: int = 1) -> ParagraphMarks | None:
    """这一段能不能加标记：能就返回（is_marked 为 False 表示格式统一、中间也没有占位内容，
    不需要标记），不能返回 None。"""
    runs = _content_runs(paragraph._p)
    if runs is None:
        return None
    items = _content_items(runs)
    plain_full = "".join(item.text for item in items)
    plain = plain_full.strip()
    if (
        not plain
        or plain != (paragraph.text or "").strip()
        or _TAG_RE.search(plain)
        or _SLOT_RE.search(plain)
    ):
        return None

    looks = [None if item.is_slot else _look_key(item.rpr) for item in items]
    weights: dict[tuple, int] = {}
    for item, look in zip(items, looks):
        if look is not None:
            weights[look] = weights.get(look, 0) + len(item.text.strip())
    first_index = {}
    for index, look in enumerate(looks):
        if look is not None:
            first_index.setdefault(look, index)
    base_look = max(weights, key=lambda look: (weights[look], -first_index[look]))
    base_rpr = items[first_index[base_look]].rpr

    # 每一截归到哪一组：和主体格式一样、又不在链接里的不加标记。只有空白的那截和占位符
    # 夹在同一组中间就并进去，否则归主体
    keys: list[tuple | None] = []
    soft: list[bool] = []
    for item, look in zip(items, looks):
        is_soft = item.is_slot or (not item.text.strip() and item.hyperlink is None)
        soft.append(is_soft)
        if is_soft or (item.hyperlink is None and look == base_look):
            keys.append(None)
        else:
            keys.append((look, id(item.hyperlink) if item.hyperlink is not None else 0))
    for index, is_soft in enumerate(soft):
        if not is_soft:
            continue
        before = next((keys[i] for i in range(index - 1, -1, -1) if not soft[i]), None)
        after = next((keys[i] for i in range(index + 1, len(items)) if not soft[i]), None)
        keys[index] = before if before is not None and before == after else None

    # 去掉段首段尾的空白（和 paragraph.text.strip() 对齐）。段首段尾的空白和占位内容
    # 不发给模型，写回时固定放回段首段尾（段尾的空白文字照旧不留）
    lead_len = len(plain_full) - len(plain_full.lstrip())
    body_end = len(plain_full) - (len(plain_full) - len(plain_full.rstrip()))
    lead: list = []
    trail: list = []
    body: list[tuple[_Item, str, tuple | None]] = []
    offset = 0
    for item, key in zip(items, keys):
        start, end = offset, offset + len(item.text)
        offset = end
        if item.is_slot:
            if end <= lead_len:
                lead.append(item)
            elif start >= body_end:
                trail.append(item)
            elif start >= lead_len and end <= body_end:
                body.append((item, item.text, key))
            else:
                return None  # 占位内容跨在段首段尾的空白上，读不准，按整段写法
            continue
        if start < lead_len:
            lead.append(_Item(plain_full[start:min(end, lead_len)], item.rpr, None))
        lo, hi = max(start, lead_len), min(end, body_end)
        if lo < hi:
            body.append((item, plain_full[lo:hi], key))

    groups: dict[int, MarkGroup] = {}
    placeholders: dict[int, Placeholder] = {}
    parts: list[str] = []
    next_id, next_slot = start_id, start_slot
    current_key: tuple | None = None
    for item, text, key in body:
        if key != current_key:
            if current_key is not None:
                parts.append(f"</g{next_id - 1}>")
            if key is not None:
                rpr = item.node_pairs[0][1] if item.is_slot else item.rpr
                groups[next_id] = MarkGroup(rpr=rpr, hyperlink=item.hyperlink)
                parts.append(f"<g{next_id}>")
                next_id += 1
            current_key = key
        if item.is_slot:
            placeholders[next_slot] = Placeholder(nodes=tuple(item.node_pairs), text=item.text)
            parts.append(f"<x{next_slot}/>")
            next_slot += 1
        else:
            parts.append(text)
    if current_key is not None:
        parts.append(f"</g{next_id - 1}>")
    if len(groups) > MAX_GROUPS_PER_UNIT or len(placeholders) > MAX_PLACEHOLDERS_PER_UNIT:
        return None
    marked = "".join(parts)
    texts = {key: item.text for key, item in placeholders.items()}
    if strip_marks(marked, texts) != plain or placeholder_texts(marked, plain) != texts:
        return None
    return ParagraphMarks(
        marked_text=marked if groups or placeholders else plain,
        plain_text=plain,
        groups=groups,
        base_rpr=base_rpr,
        placeholders=placeholders,
        lead=[_fixed_piece(item) for item in lead],
        trail=[_fixed_piece(item) for item in trail if item.is_slot],
    )


def _fixed_piece(item: _Item):
    if item.is_slot:
        return (Placeholder(nodes=tuple(item.node_pairs), text=item.text), None)
    return (item.text, None)


_HINT_LIMIT = 80
_SLOT_HINT_NAMES = {
    **{qn(tag): "制表符" for tag in ("w:tab", "w:ptab")},
    **{qn(tag): "换行" for tag in ("w:br", "w:cr")},
    **{qn(tag): "图片" for tag in ("w:drawing", "w:pict")},
    qn("w:object"): "嵌入对象",
    **{qn(tag): "脚注编号" for tag in ("w:footnoteReference", "w:footnoteRef")},
    **{qn(tag): "尾注编号" for tag in ("w:endnoteReference", "w:endnoteRef")},
}


def placeholder_hint(placeholder: Placeholder) -> str:
    """一处占位内容给模型看的说明：公式写成线性文字（「σ=M/W」「x_i^2」），其余写类别。"""
    parts = []
    for node, _rpr in placeholder.nodes:
        if node.tag in MATH_TAGS:
            text = _math_linear(node).strip()
            parts.append(f"公式 {text}" if text else "公式")
        elif node.tag in _SLOT_HINT_NAMES:
            parts.append(_SLOT_HINT_NAMES[node.tag])
    hint = "、".join(dict.fromkeys(parts))
    return hint if len(hint) <= _HINT_LIMIT else hint[: _HINT_LIMIT - 1] + "…"


def _math_wrap(text: str) -> str:
    return text if len(text) <= 1 else f"({text})"


def _math_linear(node) -> str:
    """OMML 公式 → 线性文字，只求让模型看懂指的是哪个式子，不求能还原。"""
    if node is None or not isinstance(node.tag, str):
        return ""
    if not node.tag.startswith(_MATH_NS):
        return "".join(_math_linear(child) for child in node)
    local = node.tag[len(_MATH_NS):]
    if local == "t":
        return node.text or ""
    if local.endswith("Pr"):
        return ""
    kids = {child.tag[len(_MATH_NS):]: child for child in node if isinstance(child.tag, str)
            and child.tag.startswith(_MATH_NS)}

    def part(name: str) -> str:
        return _math_linear(kids.get(name)).strip()

    if local == "f":
        return f"{_math_wrap(part('num'))}/{_math_wrap(part('den'))}"
    if local == "sSup":
        return f"{part('e')}^{_math_wrap(part('sup'))}"
    if local == "sSub":
        return f"{part('e')}_{_math_wrap(part('sub'))}"
    if local == "sSubSup":
        return f"{part('e')}_{_math_wrap(part('sub'))}^{_math_wrap(part('sup'))}"
    if local == "rad":
        degree = part("deg")
        return f"{degree}√{_math_wrap(part('e'))}"
    if local == "nary":
        props = kids.get("naryPr")
        char = props.find(f"{_MATH_NS}chr") if props is not None else None
        symbol = (char.get(f"{_MATH_NS}val") if char is not None else None) or "∫"
        lower, upper = part("sub"), part("sup")
        bounds = (f"_{_math_wrap(lower)}" if lower else "") + (f"^{_math_wrap(upper)}" if upper else "")
        return f"{symbol}{bounds} {part('e')}"
    if local == "d":
        props = kids.get("dPr")

        def prop(name: str, default: str) -> str:
            found = props.find(f"{_MATH_NS}{name}") if props is not None else None
            value = found.get(f"{_MATH_NS}val") if found is not None else None
            return default if value is None else value

        inner = ",".join(_math_linear(child).strip() for child in node if child.tag == f"{_MATH_NS}e")
        return f"{prop('begChr', '(')}{inner}{prop('endChr', ')')}"
    if local == "eqArr":
        return "; ".join(_math_linear(child).strip() for child in node if child.tag == f"{_MATH_NS}e")
    return "".join(_math_linear(child) for child in node)


def cell_marks(paragraphs: list) -> CellMarks | None:
    """单元格：各段原文用换行拼起来（同 word_document._cell_source_text），标记和占位符的
    编号跨段连续。"""
    items: list[ParagraphMarks] = []
    next_id = next_slot = 1
    for paragraph in paragraphs:
        marks = paragraph_marks(paragraph, start_id=next_id, start_slot=next_slot)
        if marks is None:
            return None
        items.append(marks)
        next_id += len(marks.groups)
        next_slot += len(marks.placeholders)
    if (
        not items
        or not any(item.is_marked for item in items)
        or next_id - 1 > MAX_GROUPS_PER_UNIT
        or next_slot - 1 > MAX_PLACEHOLDERS_PER_UNIT
    ):
        return None
    return CellMarks(marked_text="\n".join(item.marked_text for item in items), paragraphs=items)


def strip_marks(text: str, placeholders: dict[int, str] | None = None) -> str:
    """去掉格式标记；占位符换回它在纯文本里对应的字符（placeholders 里没有的换成空）。"""
    texts = placeholders or {}
    without_tags = _TAG_RE.sub("", str(text or ""))
    return _SLOT_RE.sub(lambda match: texts.get(int(match.group(1)), ""), without_tags)


def placeholder_texts(marked_source: str, plain_source: str) -> dict[int, str] | None:
    """由带标记原文和纯文本原文反推每个占位符对应的字符；对不上返回 None。

    只凭两段字符串就能还原，扩展通道（只传字符串）也能用。生成时相邻的占位内容已并成
    一个、占位符之间至少隔着一个字，再加上 paragraph_marks 生成时核对过，不会有歧义。"""
    text = _TAG_RE.sub("", str(marked_source or ""))
    parts = _SLOT_RE.split(text)
    if len(parts) == 1:
        return {}
    ids = [int(value) for value in parts[1::2]]
    if len(set(ids)) != len(ids):
        return None
    pattern = "".join(re.escape(part) if index % 2 == 0 else "(.*?)" for index, part in enumerate(parts))
    match = re.fullmatch(pattern, str(plain_source or ""), re.DOTALL)
    if match is None:
        return None
    return dict(zip(ids, match.groups()))


def unmark(translation: str, marked_source: str, plain_source: str) -> str:
    """带标记的译文 → 纯文本译文（占位符按原文换回制表符、换行等）。"""
    return strip_marks(translation, placeholder_texts(marked_source, plain_source))


# 标点紧贴标记、标记后紧跟字母：「<g1>Note :</g1>les…」。原文是中文时冒号后本来没有空格，
# 模型在标记边界常照抄，去掉标记就成了「Note :les」。只补这一种明显缺的空格：
# 冒号、分号、逗号、问号、叹号、右括号后跟字母，或句号后跟大写字母；空格放在格式外面
# （闭标记之后、开标记之前）。上标「m<g1>2</g1>」、半个词加粗这类不碰。
_BOUNDARY_GAP_RE = re.compile(
    r"(?:(?<=[:;,!?)\]»])|(?<=\.)(?=(?:<[^>]*>)+[A-ZÀ-ÖØ-Þ]))"
    r"((?:<\s*/\s*[gG]\s*\d+\s*>)*)((?:<\s*[gG]\s*\d+\s*>)*)(?=[^\W\d_])"
)


# 脚注编号、图片这类不占字符的占位符夹在两个词中间：「béton<x1/>doit」。原文是中文时
# 本来没有空格，模型常照抄，写回成「béton¹doit」。在占位符后面补一个空格（放在格式外面）
_OBJECT_GAP_RE = re.compile(
    r"(?<=[^\W\d_]|[.,;:!?)\]»])(<\s*x\s*(\d+)\s*/\s*>)(?=(?:<\s*[gG]\s*\d+\s*>)*[^\W\d_])"
)


def add_boundary_spaces(
    translation: str, target_lang: str | None, placeholders: dict[int, str] | None = None
) -> str:
    """译文用空格分词时，补上标记边界处模型漏掉的空格（见 _BOUNDARY_GAP_RE）；placeholders
    给出占位符对应的字符时，不占字符的占位符（脚注编号、图片）两边也补（见 _OBJECT_GAP_RE）。"""
    if _no_word_spaces(target_lang):
        return translation

    def fill(match: re.Match) -> str:
        closing, opening = match.group(1), match.group(2)
        return f"{closing} {opening}" if closing or opening else match.group(0)

    text = _BOUNDARY_GAP_RE.sub(fill, str(translation or ""))
    if placeholders:
        text = _OBJECT_GAP_RE.sub(
            lambda match: match.group(1) + " " if placeholders.get(int(match.group(2))) == "" else match.group(0),
            text,
        )
    return text


_WHITESPACE_SPLIT_RE = re.compile(r"(\s+)")


def same_text(candidate: str, settled: str) -> bool:
    """去掉标记后的译文 candidate 和定稿译文 settled 算不算同一段：文字一字不差；定稿里
    只有空格的地方，candidate 可以是任何空白。记忆库里的定稿常把原文的制表符、软换行存成
    了空格，补标记时占位符顶替那个空格，写回的仍是这段译文；定稿里本来就有的制表符、换行
    不能挪到别处。"""
    left = _WHITESPACE_SPLIT_RE.split(str(candidate or "").strip())
    right = _WHITESPACE_SPLIT_RE.split(str(settled or "").strip())
    if len(left) != len(right):
        return False
    for index, (mine, theirs) in enumerate(zip(left, right)):
        if mine == theirs:
            continue
        if index % 2 == 0 or theirs.strip(" ") or not mine.isspace():
            return False
    return True


def parse_marked(text: str, allowed_ids, placeholder_ids=()) -> list[tuple[str | Slot, int | None]] | None:
    """把带标记的译文拆成 [(文字或 Slot, 组号或 None)]。每个组号必须恰好一开一闭、不嵌套，
    每个占位符恰好出现一次，组号和占位符的集合都和原文完全一致；否则返回 None。"""
    allowed = set(allowed_ids)
    slots_allowed = set(placeholder_ids)
    pieces: list[tuple[str | Slot, int | None]] = []
    seen: set[int] = set()
    seen_slots: set[int] = set()
    current: int | None = None
    position = 0
    for match in _ANY_MARK_RE.finditer(text):
        chunk = text[position:match.start()]
        if chunk:
            pieces.append((chunk, current))
        position = match.end()
        if match.group(3) is not None:
            slot = int(match.group(3))
            if slot not in slots_allowed or slot in seen_slots:
                return None
            seen_slots.add(slot)
            pieces.append((Slot(slot), current))
            continue
        closing, group = match.group(1) == "/", int(match.group(2))
        if group not in allowed:
            return None
        if closing:
            if current != group:
                return None
            current = None
        else:
            if current is not None or group in seen:
                return None
            seen.add(group)
            current = group
    tail = text[position:]
    if tail:
        pieces.append((tail, current))
    if current is not None or seen != allowed or seen_slots != slots_allowed:
        return None
    return pieces


def _new_run(text: str, rpr):
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(deepcopy(rpr))
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    run.append(t)
    return run


def _placeholder_runs(placeholder: Placeholder, *, move: bool) -> list:
    """占位内容按原来的 run 格式各放进一个 w:r（公式本来就和 w:r 平级，原样放回）。
    move=True 时挪用原节点（原段落随后整段换掉，挪用比复制稳妥：文本框里的文字、图片的
    编号都还是原来那一份）。"""
    runs = []
    for node, rpr in placeholder.nodes:
        if node.tag in MATH_TAGS:
            runs.append(node if move else deepcopy(node))
            continue
        run = OxmlElement("w:r")
        if rpr is not None:
            run.append(deepcopy(rpr))
        run.append(node if move else deepcopy(node))
        runs.append(run)
    return runs


def _without_objects(pieces) -> list:
    """去掉图片、脚注编号这类占位内容；它两边的空格并成一个，后面紧跟标点的去掉。"""
    out: list = []
    dropped = False
    for text, group in pieces:
        if isinstance(text, Placeholder) and not text.copyable:
            dropped = True
            continue
        if dropped and isinstance(text, str) and out and isinstance(out[-1][0], str) and out[-1][0].endswith(" "):
            if text.startswith(" ") or text[:1] in ",.;)]»":
                out[-1] = (out[-1][0][:-1], out[-1][1])
        dropped = False
        out.append((text, group))
    return out


def build_run_nodes(pieces, groups: dict[int, MarkGroup], base_rpr, *, keep_objects: bool = True) -> list:
    """按拆好的译文生成 w:r（链接那截包在 w:hyperlink 里，相邻同一个链接的并在一起），
    返回要插进段落的节点。片段里的 Placeholder 照抄原节点；keep_objects=False（双语版式
    新插的译文段）时只复制制表符、换行这类版式和公式，图片、脚注编号留在原文段里。"""
    nodes = []
    last_link: tuple[object, object] | None = None  # (原链接, 新建的 w:hyperlink)
    if not keep_objects:
        pieces = _without_objects(pieces)
    for text, group in pieces:
        if isinstance(text, Placeholder):
            if not keep_objects and not text.copyable:
                continue
            runs = _placeholder_runs(text, move=keep_objects)
        elif text:
            runs = [_new_run(text, base_rpr if group is None else groups[group].rpr)]
        else:
            continue
        source_link = groups[group].hyperlink if group is not None else None
        if source_link is None:
            nodes.extend(runs)
            last_link = None
            continue
        if last_link is None or last_link[0] is not source_link or nodes[-1] is not last_link[1]:
            link = OxmlElement("w:hyperlink")
            for key, value in source_link.attrib.items():
                link.set(key, value)
            nodes.append(link)
            last_link = (source_link, link)
        for run in runs:
            last_link[1].append(run)
    return nodes


def replace_paragraph_content(paragraph, nodes) -> None:
    """删掉段落里原有的文字 run 和超链接，把新节点放在第一个被删节点的位置。新节点里挪用的
    公式也先从原位摘下，免得插回去时位置算错。"""
    element = paragraph._p
    old = [
        child for child in element
        if child.tag in (qn("w:r"), qn("w:hyperlink")) or any(node is child for node in nodes)
    ]
    if old:
        anchor_index = list(element).index(old[0])
        for child in old:
            element.remove(child)
    else:
        anchor_index = len(element)
    for offset, node in enumerate(nodes):
        element.insert(anchor_index + offset, node)


def iter_new_runs(nodes):
    """新写的文字 run（占位内容的 run 照抄原样，不在其中）。"""
    for node in nodes:
        runs = node.iterchildren(qn("w:r")) if node.tag == qn("w:hyperlink") else (node,)
        for run in runs:
            if run.find(qn("w:t")) is not None:
                yield run


class InlineMarkEngine:
    """包住翻译引擎：带局部格式的段落换成带标记的原文发出，回来的译文校验标记、记下，
    再去掉标记交回，下游（校验、记忆库、复核）看到的仍是纯文本。

    标记没完整回来的段落不重翻：整段译文照用，回来的标记里对得上的留着
    （partial_translations），定稿后由 place_marks 只补缺的那几个（反馈 #19）。"""

    # engines.base_engine.engine_supports_chat 据此看里面那个引擎
    _delegates_capabilities = True

    def __init__(self, inner, marked_sources: dict[str, str]):
        self.inner = inner
        self._marked_sources = dict(marked_sources)
        self._lock = threading.Lock()
        self.marked_translations: dict[str, tuple[str, str]] = {}
        # 标记没完整回来的：原文 → (带标记原文, 只留下对得上的标记的译文)
        self.partial_translations: dict[str, tuple[str, str]] = {}
        # 排查用（只进 app.log，不含原文）：每段最后一次的结果 ok / partial
        self.outcomes: dict[str, str] = {}
        # 超长段落按句子切成的几块：原文 → [(这块的纯文本, 带标记的这块)]；见 split_long_sources
        self.parts: dict[str, list[tuple[str, str]]] = {}
        self._part_texts: set[str] = set()
        self._part_results: dict[str, str] = {}  # 超长段落各块最后交回的译文（去掉标记的）

    def __getattr__(self, name):
        return getattr(self.inner, name)

    def _outgoing(self, texts: list[str], system_prompt: str) -> tuple[list[str], str]:
        sent = [self._marked_sources.get(text, text) for text in texts]
        if sent == list(texts):
            return sent, system_prompt
        return sent, append_batch_block(system_prompt, MARK_PROMPT_RULE)

    def split_long_sources(self, split) -> None:
        """word_batching 把超长段落按句子切成几块分别发（split 和它切法相同）。这几块原样发出
        就没有标记，公式全靠定稿后补标记，一段二三十个公式时常补不上、全挤到段末（反馈
        #19）。这里把带标记的原文按同样的位置切开，每块也换成带标记的发出，定稿后再拼回
        （assemble_parts）。切不准的段落照旧。"""
        for plain, marked in list(self._marked_sources.items()):
            if plain in self.parts or plain in self._part_texts:
                continue
            pieces = split(plain)
            if len(pieces) <= 1:
                continue
            marked_pieces = split_marked_parts(marked, plain, pieces)
            if marked_pieces is None:
                continue
            self.parts[plain] = list(zip(pieces, marked_pieces))
            for piece, marked_piece in zip(pieces, marked_pieces):
                self._part_texts.add(piece)
                if marked_piece != piece:
                    self._marked_sources.setdefault(piece, marked_piece)

    def assemble_parts(self, source: str, join, flags=None) -> tuple[str, str] | None:
        """超长段落各块带标记的译文按 join（和 word_batching 拼译文的方式相同）拼回整段：
        (带标记原文, 带标记译文)；有一块缺了返回 None。给了 flags（place_marks 记下的
        要标出来核对的范围，块原文 → [(0, 开始, 结束)]）时，各块的范围换算成整段的，
        记到整段原文名下。"""
        pieces = self.parts.get(source)
        marked = self._marked_sources.get(source)
        if not pieces or marked is None:
            return None
        translated = []
        for piece, marked_piece in pieces:
            if marked_piece == piece:
                text = self._part_results.get(piece)
            else:
                entry = self.marked_translations.get(piece)
                text = entry[1] if entry is not None and entry[0] == marked_piece else None
            if not text:
                return None
            translated.append(text.strip())
        # 切开时跨块的那对标记各块各包一截，拼回时接缝处合成一对
        joined = _SEAM_RE.sub(r"\2", join(translated))
        present = {int(group) for closing, group in _TAG_RE.findall(joined) if not closing}
        if parse_marked(joined, present, _line_slots(marked)) is None or not present <= set(_line_ids(marked)):
            return None
        if flags is not None:
            gap = 0 if len(join(translated)) == sum(map(len, translated)) else 1
            offset, spans = 0, []
            for (piece, _marked_piece), text in zip(pieces, translated):
                spans += [(0, offset + begin, offset + end) for _line, begin, end in flags.get(piece, ())]
                offset += len(_zero_width(text)) + gap
            if spans:
                flags[source] = spans
        return marked, joined

    def parts_to_place(self, source: str, translation: str, join) -> list[tuple[str, str, str]] | None:
        """拼不回整段时，哪几块缺带标记的译文：[(块原文, 带标记块原文, 块译文)]，交给
        place_marks 一块一块补。各块译文拼起来和定稿对不上（定稿后又改过）、或有块没交回
        译文的返回 None。"""
        pieces = self.parts.get(source)
        if not pieces:
            return None
        results = [self._part_results.get(piece) for piece, _marked_piece in pieces]
        if any(not result for result in results) or not same_text(join([r.strip() for r in results]), translation):
            return None
        missing = []
        for (piece, marked_piece), result in zip(pieces, results):
            if marked_piece == piece:
                continue
            entry = self.marked_translations.get(piece)
            if entry is None or entry[0] != marked_piece or not same_text(unmark(entry[1], marked_piece, piece), result):
                missing.append((piece, marked_piece, result.strip()))
        return missing

    def _receive(self, plain: str, sent: str, translation: str, target_lang: str) -> str:
        result = self._receive_one(plain, sent, translation, target_lang)
        if plain in self._part_texts:
            with self._lock:
                self._part_results[plain] = result
        return result

    def _receive_one(self, plain: str, sent: str, translation: str, target_lang: str) -> str:
        if sent is plain or sent == plain:
            return translation
        texts = placeholder_texts(sent, plain)
        ok = marks_match(sent, translation)
        if not ok:
            # 不重翻：译文照用，对得上的标记留着，缺的定稿后由 place_marks 补
            translation = clean_marks(sent, translation)
        translation = add_boundary_spaces(translation, target_lang, texts)
        with self._lock:
            self.outcomes[plain] = "ok" if ok else "partial"
            if ok:
                self.marked_translations[plain] = (sent, translation)
                self.partial_translations.pop(plain, None)
            else:
                self.marked_translations.pop(plain, None)
                self.partial_translations[plain] = (sent, translation)
        return strip_marks(translation, texts)

    def translate_batch(self, texts, target_lang, system_prompt, source_lang="zh"):
        sent, prompt = self._outgoing(list(texts), system_prompt)
        if prompt is system_prompt:
            return self.inner.translate_batch(texts, target_lang, system_prompt, source_lang=source_lang)
        raw = dict(self.inner.translate_batch(sent, target_lang, prompt, source_lang=source_lang))
        result: dict[str, str] = {}
        for plain, outgoing in zip(texts, sent):
            if outgoing in raw:
                result[plain] = self._receive(plain, outgoing, raw[outgoing], target_lang)
        return result

    def translate_batch_with_sources(self, texts, target_lang, system_prompt, source_lang="auto"):
        sent, prompt = self._outgoing(list(texts), system_prompt)
        if prompt is system_prompt:
            return self.inner.translate_batch_with_sources(texts, target_lang, system_prompt, source_lang=source_lang)
        reported = self.inner.translate_batch_with_sources(sent, target_lang, prompt, source_lang=source_lang)
        plain_by_sent = dict(zip(sent, texts))
        out = []
        for item in reported:
            plain = plain_by_sent.get(item.source_text, item.source_text)
            out.append(replace(item, source_text=plain,
                               translation=self._receive(plain, item.source_text, item.translation, target_lang)))
        return out


def _no_word_spaces(target_lang) -> bool:
    """译文不用空格分词（中日韩泰等）。"""
    lang = str(target_lang or "").strip().lower().replace("_", "-").split("-")[0]
    return lang in ("zh", "ja", "ko", "th", "lo", "km", "my")


# 句末：中文句号类直接断；西文句号类后面要有空格，且下一个字是大写字母、数字、汉字或标记，
# 免得在「e.g. the」「Fig. 3」这种地方断开（「Fig. 3」仍可能断，只影响拆在哪儿，不影响结果）
_SENTENCE_END_RE = re.compile(r"[。！？；]\s*|[.!?;]\s+(?=[A-ZÀ-ÖØ-Þ0-9<\u4e00-\u9fff])")


_SEAM_RE = re.compile(r"<\s*/\s*g\s*(\d+)\s*>(\s*)<\s*g\s*\1\s*>", re.IGNORECASE)


_PART_SENTENCE_ENDS = frozenset("。！？；.!?;")


def split_marked_parts(marked: str, plain: str, parts: list[str]) -> list[str] | None:
    """带标记原文按纯文本的切块（word_batching 按句子切的，每块去掉首尾空白）切成对应的
    几块。块之间的空白不要；夹在块边界上的标记：闭标记归前一块，开标记归后一块，占位符
    紧贴前一块的末字时归前一块（前一块以句末标点结尾时归后一块），否则归后一块。一对格式标记跨两块时各块各包一截（同一个
    编号）。对不上返回 None。"""
    texts = placeholder_texts(marked, plain)
    if texts is None:
        return None
    tokens: list[tuple[str, str, int, int]] = []  # (种类 char/open/close/slot, 原串, 纯文本起点, 终点)
    position = cursor = 0
    for match in _ANY_MARK_RE.finditer(marked):
        for char in marked[cursor:match.start()]:
            tokens.append(("char", char, position, position + 1))
            position += 1
        cursor = match.end()
        if match.group(3) is not None:
            width = len(texts.get(int(match.group(3)), ""))
            tokens.append(("slot", match.group(0), position, position + width))
            position += width
        else:
            tokens.append(("close" if match.group(1) == "/" else "open", match.group(0), position, position))
    for char in marked[cursor:]:
        tokens.append(("char", char, position, position + 1))
        position += 1
    spans: list[tuple[int, int]] = []
    start = 0
    for part in parts:
        found = plain.find(part, start)
        if found < 0:
            return None
        spans.append((found, found + len(part)))
        start = found + len(part)

    def owner(kind: str, at: int, until: int) -> int | None:
        if until > at:  # 占字符的：整个落在一块里才行
            return next((index for index, (begin, end) in enumerate(spans) if begin <= at and until <= end), None)
        for index, (begin, end) in enumerate(spans):
            if begin < at < end:
                return index
            nxt = spans[index + 1][0] if index + 1 < len(spans) else None
            if at == begin:
                return index - 1 if kind == "close" and index > 0 else index
            if at == end:
                # 句末之后紧跟的占位符属于下一句（中文句子之间没有空白）
                ends_sentence = kind == "slot" and plain[end - 1:end] in _PART_SENTENCE_ENDS
                return index + 1 if kind in ("open", "slot") and nxt is not None and (kind == "open" or ends_sentence) else index
            if nxt is not None and end < at < nxt:
                return index if kind == "close" else index + 1
        return 0 if at < spans[0][0] else len(spans) - 1

    chunks: list[list[str]] = [[] for _ in parts]
    for kind, raw, at, until in tokens:
        index = owner(kind, at, until)
        if index is None:
            if kind == "char" and raw.isspace():
                # 块之间的空白不要，除非前面已经有归到下一块的占位符（"<x2/> is" 的空格）
                following = next((i for i, (begin, _end) in enumerate(spans) if begin > at), None)
                if following is not None and chunks[following]:
                    chunks[following].append(raw)
                continue
            return None
        chunks[index].append(raw)
    out: list[str] = []
    carried: str | None = None  # 上一块结尾还开着的那个格式标记
    for part, chunk in zip(parts, chunks):
        text = (carried or "") + "".join(chunk)
        opened = None
        for match in _TAG_RE.finditer(text):
            opened = None if match.group(1) == "/" else match.group(2)
        carried = f"<g{opened}>" if opened is not None else None
        if opened is not None:
            text += f"</g{opened}>"
        if strip_marks(text, texts).strip() != part:
            return None
        out.append(text)
    return out


def _line_ids(text: str) -> list[int]:
    return [int(group) for closing, group in _TAG_RE.findall(text) if not closing]


def _line_slots(text: str) -> list[int]:
    return [int(slot) for slot in _SLOT_RE.findall(text)]


def marks_match(marked_source: str, translation: str) -> bool:
    """译文里的标记和占位符与原文一一对应。单元格的原文按段用换行拼接，译文也必须一行
    对一段，标记和占位符不能跨行。"""
    sent_lines = marked_source.split("\n")
    lines = str(translation or "").split("\n")
    if len(lines) != len(sent_lines) or not (_line_ids(marked_source) or _line_slots(marked_source)):
        return False
    return all(
        parse_marked(line, _line_ids(sent_line), _line_slots(sent_line)) is not None
        for line, sent_line in zip(lines, sent_lines)
    )


# 模型写走样的占位符：「<x1>」「</x1>」「<x 1 />」、转义成「&lt;x1/&gt;」、全角尖括号
_LOOSE_SLOT_RE = re.compile(r"(?:<|&lt;|＜)\s*/?\s*x\s*(\d+)\s*/?\s*(?:>|&gt;|＞)", re.IGNORECASE)


def _clean_line(sent_line: str, line: str) -> str:
    slots = set(_line_slots(sent_line))
    groups = set(_line_ids(sent_line))
    seen: set[int] = set()
    parts: list[str] = []
    cursor = 0
    for match in _LOOSE_SLOT_RE.finditer(line):
        parts.append(line[cursor:match.start()])
        cursor = match.end()
        slot = int(match.group(1))
        if slot in slots and slot not in seen:  # 重复的、原文没有的去掉
            seen.add(slot)
            parts.append(f"<x{slot}/>")
    parts.append(line[cursor:])
    fixed = "".join(parts)
    present = {int(group) for closing, group in _TAG_RE.findall(fixed) if not closing}
    if not (present <= groups and parse_marked(fixed, present, seen) is not None):
        fixed = _TAG_RE.sub("", fixed)  # 格式标记乱了：整套去掉，缺的再补
    return fixed


def clean_marks(marked_source: str, translation: str) -> str:
    """模型带回的标记不全时，只留下对得上的：写走样的占位符改正，重复的、原文没有的去掉；
    格式标记乱了（不成对、嵌套、原文没有）就整套去掉。缺的不猜位置，留给 place_marks
    补。行数和原文对不上时去掉全部标记。"""
    sent_lines = marked_source.split("\n")
    lines = str(translation or "").split("\n")
    if len(lines) != len(sent_lines):
        return _TAG_RE.sub("", _LOOSE_SLOT_RE.sub("", str(translation or "")))
    return "\n".join(_clean_line(sent_line, line) for sent_line, line in zip(sent_lines, lines))


# 补标记：每组最多几个标记（一个分句里的公式、格式标记，一对格式标记算一个）、一次请求
# 最多几个标记（几组打包一起问），见 place_marks
PLACE_MARKS_GROUP = 5
PLACE_MARKS_REQUEST = 20
PLACE_MARKS_BATCH = 20  # 一次请求最多几组
PLACE_MARKS_PROMPT = (
    "你是文档排版助手。下面 JSON 数组的每一项给出带标记的「原文」、它已经定稿的「译文」，以及这次要插入"
    "译文的标记「插入」。<g1>…</g1> 是成对的格式标记，标出加粗、斜体、链接等局部格式；<x1/> 是自闭合占位符，"
    "代表公式、图片、脚注编号、制表符等不翻译的内容，「说明」写出部分占位符代表的内容（公式写成线性文字），"
    "便于判断它在译文里对应的位置。请只把「插入」里列出的标记插进译文：成对标记包住与原文标记内容对应的译文，"
    "占位符放在译文里与原文位置对应的地方；每个标记恰好出现一次。译文里已有的标记保持原位，"
    "不要插入没列出的标记，译文的文字一个都不能改。"
    "按原顺序返回 JSON 字符串数组，长度与输入相同，每项是插好标记的译文；不要输出解释或 markdown 代码块。"
)
# 一组放不进去时：第 1 次重试原样再问，第 2 次拆成两半，之后拆开的再问
_SPLIT_AT_ATTEMPT = 2
_PLACE_MIN_SIMILARITY = 0.85
_CLAUSE_CHARS = frozenset("，,；;。.:：!?！？、")


def _seed_slots(marked: str, translation: str, texts: dict[int, str]) -> str:
    """定稿译文里制表符、换行这类有字符的占位内容，按原文顺序先换成占位符，模型只需补
    图片、脚注编号、公式和格式标记。按顺序找不齐就原样交给模型。"""
    seeded, cursor = [], 0
    for slot in _line_slots(marked):
        text = texts.get(slot, "")
        if not text:
            continue
        found = translation.find(text, cursor)
        if found < 0:
            return translation
        seeded.append(translation[cursor:found])
        seeded.append(f"<x{slot}/>")
        cursor = found + len(text)
    seeded.append(translation[cursor:])
    return "".join(seeded)


def _zero_width(text: str) -> str:
    """去掉所有标记（占位符也不换回字符）：补标记时比对位置用。"""
    return _ANY_MARK_RE.sub("", text)


def _mark_tokens(line: str) -> list[tuple[str, int, str]]:
    """[(标记的名字 x3 / g2 / /g2, 去掉标记后的位置, 原串)]，按出现顺序。"""
    tokens = []
    position = cursor = 0
    for match in _ANY_MARK_RE.finditer(line):
        position += match.start() - cursor
        cursor = match.end()
        if match.group(3) is not None:
            tokens.append((f"x{int(match.group(3))}", position, match.group(0)))
        else:
            tokens.append((f"{'/' if match.group(1) == '/' else ''}g{int(match.group(2))}", position, match.group(0)))
    return tokens


def _units(marked_line: str) -> list[tuple[str, int]]:
    """原文这一行要放的标记单元（占位符一个算一个，一对格式标记算一个）和位置，按原文顺序。"""
    return [(name, position) for name, position, _raw in _mark_tokens(marked_line) if not name.startswith("/")]


def _present_units(line: str) -> set[str]:
    names = [name for name, _position, _raw in _mark_tokens(line)]
    return {name for name in names if name.startswith("x") or (name.startswith("g") and f"/{name}" in names)}


def _group_units(marked_line: str, units: list[str], size: int) -> list[list[str]]:
    """缺的标记按原文顺序分组，每组最多 size 个，尽量在分句（逗号、分号、句号）处断开。"""
    positions = dict(_units(marked_line))
    plain = _zero_width(marked_line)

    def clause_between(left: str, right: str) -> bool:
        return any(char in _CLAUSE_CHARS for char in plain[positions[left]:positions[right]])

    groups: list[list[str]] = []
    rest = list(units)
    while rest:
        if len(rest) <= size:
            groups.append(rest)
            break
        cut = size  # 找不到分句处就按个数断
        for index in range(size, max(size // 2, 1) - 1, -1):
            if clause_between(rest[index - 1], rest[index]):
                cut = index
                break
        groups.append(rest[:cut])
        rest = rest[cut:]
    return groups


def _map_position(blocks, position: int, answer_plain: str = "") -> int | None:
    """答案里（去掉标记后）的位置换算成定稿译文里的位置。落在模型多写的空白里的，归到
    前面那段对得上的文字后面；落在改过的文字中间的返回 None。"""
    previous = None
    for a, b, size in blocks:
        if a <= position < a + size or (position == a + size and size):
            return b + (position - a)
        if size == 0 and position == a:
            return b
        if previous is not None and previous[0] < position < a:
            if not answer_plain[previous[0]:a].strip():
                return previous[1]
            return None
        previous = (a + size, b + size)
    return None


def _answer_positions(
    answer: str, base_plain: str, units: list[str], texts: dict[int, str] | None = None,
    hints: dict[int, str] | None = None,
) -> dict[str, tuple[int, int | None, bool]] | None:
    """从模型的回答里取出这一组标记在定稿译文里的位置：占位符 → (位置, None, 是否顶替那里的
    空白)，格式标记 → (开始, 结束, False)。缺了、重了、落在改过的文字里的标记不在结果里；
    回答的文字和定稿差太多时返回 None。

    空白：制表符、换行这类有字符的占位符，回答里紧贴后面的字、定稿那里却是空白（记忆库把
    制表符存成了空格）时顶替这个空白；脚注编号、图片这类不占字符的，落在空白和下一个词之间
    （「béton <x1/>doit」）时挪到空白前面，贴着前一个词。"""
    texts = texts or {}
    hints = hints or {}
    from difflib import SequenceMatcher

    answer = _LOOSE_SLOT_RE.sub(lambda match: f"<x{int(match.group(1))}/>", str(answer or ""))
    if "\n" in answer.strip():
        return None
    tokens = _mark_tokens(answer)
    plain = _zero_width(answer)
    matcher = SequenceMatcher(None, plain, base_plain, autojunk=False)
    if base_plain and matcher.ratio() < _PLACE_MIN_SIMILARITY:
        return None
    blocks = matcher.get_matching_blocks()
    found: dict[str, tuple[int, int | None, bool]] = {}
    for unit in units:
        if unit.startswith("x"):
            hits = [position for name, position, _raw in tokens if name == unit]
            if len(hits) != 1:
                continue
            at = _map_position(blocks, hits[0], plain)
            if at is None:
                continue
            slot = int(unit[1:])
            eat = False
            if texts.get(slot):
                eat = base_plain[at:at + 1].isspace() and not plain[hits[0]:hits[0] + 1].isspace()
            elif not hints.get(slot, "").startswith("公式"):
                while at > 0 and base_plain[at - 1] == " " and base_plain[at:at + 1].strip():
                    at -= 1
            found[unit] = (at, None, eat)
        else:
            opens = [index for index, (name, _p, _r) in enumerate(tokens) if name == unit]
            closes = [index for index, (name, _p, _r) in enumerate(tokens) if name == f"/{unit}"]
            if len(opens) != 1 or len(closes) != 1 or closes[0] < opens[0]:
                continue
            start, end = (_map_position(blocks, tokens[i][1], plain) for i in (opens[0], closes[0]))
            if start is None or end is None or end <= start:
                continue
            found[unit] = (start, end, False)
    return found


def _rank(name: str, order: dict[str, int]) -> tuple[int, int]:
    """同一位置上几个标记的先后：闭标记在前，其余按原文里的顺序。"""
    return (0, order.get(name[1:], 0)) if name.startswith("/") else (1, order.get(name, 0))


def _insert_mark(line: str, position: int, raw: str, name: str, order: dict[str, int]) -> str:
    """在去掉标记后的 position 处插入一个标记，那里已有别的标记时按 _rank 排进去。"""
    rank = _rank(name, order)
    zero = cursor = 0
    for match, (existing, _p, _r) in zip(_ANY_MARK_RE.finditer(line), _mark_tokens(line)):
        here = zero + (match.start() - cursor)
        if here > position or (here == position and _rank(existing, order) > rank):
            break
        zero, cursor = here, match.end()
    index = cursor + (position - zero)
    return line[:index] + raw + line[index:]


@dataclass
class _PlaceLine:
    key: tuple[str, int]  # (原文, 第几行)
    marked: str  # 这一行带标记的原文
    current: str  # 定稿译文，已经放好的标记都在里面
    order: dict[str, int]
    hints: dict[int, str]
    texts: dict[int, str]


@dataclass
class _PlaceGroup:
    line: _PlaceLine
    units: list[str]
    attempts: int = 0
    split: bool = False


def _group_request(group: _PlaceGroup) -> dict:
    line = group.line
    marks = [f"<{unit}/>" if unit.startswith("x") else f"<{unit}>…</{unit}>" for unit in group.units]
    item = {"原文": line.marked, "译文": line.current, "插入": "、".join(marks)}
    notes = {f"<{unit}/>": line.hints[int(unit[1:])] for unit in group.units
             if unit.startswith("x") and int(unit[1:]) in line.hints}
    if notes:
        item["说明"] = notes
    return item


def _delete_char(line: str, position: int) -> str:
    """删掉去掉标记后第 position 个字。"""
    zero = cursor = 0
    for match in _ANY_MARK_RE.finditer(line):
        if zero + (match.start() - cursor) > position:
            break
        zero += match.start() - cursor
        cursor = match.end()
    index = cursor + (position - zero)
    return line[:index] + line[index + 1:]


def _apply_group(line: _PlaceLine, found: dict[str, tuple[int, int | None, bool]]) -> bool:
    """把一组标记按位置插进这一行；插完格式标记交叉、嵌套的不要。"""
    updated = line.current
    for unit, (start, end, eat) in sorted(found.items(), key=lambda item: -(item[1][1] or item[1][0])):
        if unit.startswith("x"):
            if eat:
                updated = _delete_char(updated, start)
            updated = _insert_mark(updated, start, f"<{unit}/>", unit, line.order)
        else:
            updated = _insert_mark(updated, end, f"</{unit}>", f"/{unit}", line.order)
            updated = _insert_mark(updated, start, f"<{unit}>", unit, line.order)
    present = _present_units(updated)
    groups = {int(name[1:]) for name in present if name.startswith("g")}
    slots = {int(name[1:]) for name in present if name.startswith("x")}
    if parse_marked(updated, groups, slots) is None:
        return False
    line.current = updated
    return True


def _sentence_spans(plain: str) -> list[tuple[int, int, int]]:
    """[(句子开始, 句末标点的位置（没有标点时是句尾）, 句子结束)]，去掉标记后的位置。"""
    spans = []
    start = 0
    for match in _SENTENCE_END_RE.finditer(plain):
        spans.append((start, match.start(), match.start() + 1))
        start = match.end()
    end = len(plain.rstrip())
    if plain[start:end].strip():
        # 最后一句：句号后面没有下一句，_SENTENCE_END_RE 不算它
        before = end - 1 if plain[end - 1] in _PART_SENTENCE_ENDS else end
        spans.append((start, before, end))
    return spans or [(0, end, end)]


def _settle_leftovers(line: _PlaceLine, units: list[str]) -> list[tuple[int, int]]:
    """试满次数仍放不进去的：占位符放到它在原文所在那一句的译文句末（句末标点前），译文
    句数和原文对不上时放到这一行末尾；格式标记不要了（那截用主体格式）。返回要标出来
    核对的范围（去掉标记后的位置）。"""
    source_spans = _sentence_spans(_zero_width(line.marked))
    positions = dict(_units(line.marked))
    flagged: list[tuple[int, int]] = []
    for unit in units:
        if not unit.startswith("x"):
            continue
        spans = _sentence_spans(_zero_width(line.current))
        at = positions[unit]
        sentence = next((index for index, (begin, _end, close) in enumerate(source_spans) if at < close), None)
        if sentence is None:
            sentence = len(source_spans) - 1
        if len(spans) == len(source_spans):
            begin, before, close = spans[sentence]
        else:
            begin, before, close = 0, spans[-1][1], spans[-1][2]
        line.current = _insert_mark(line.current, before, f"<{unit}/>", unit, line.order)
        flagged.append((begin, close))
    return sorted(set(flagged))


def place_marks(
    engine,
    items,
    *,
    should_stop=None,
    log=None,
    reasons=None,
    seeds=None,
    hints=None,
    retries: int = 0,
    group_size: int = PLACE_MARKS_GROUP,
    request_marks: int = PLACE_MARKS_REQUEST,
    concurrency: int = 1,
    flags=None,
    stats=None,
) -> dict[str, tuple[str, str]]:
    """给已经定稿、但标记不全的译文补上标记（记忆库直接套用的、夹杂多种语言那条路翻的、
    模型弄丢标记的、复查改过的）。译文文字不动，只让模型说标记放在哪：缺的标记按分句分组，
    每组最多 group_size 个，几组打包成一次请求（一次最多 request_marks 个标记）；回答里
    取出这组标记的位置，换算到定稿译文上插进去，所以模型顺手改了几个字也不要紧。

    items: [(原文, 带标记原文, 定稿译文)]。seeds: 原文 → 已经放好部分标记的译文（翻译时
    对得上的那些，见 InlineMarkEngine.partial_translations），去掉标记要和定稿一致才用。
    hints: 原文 → {占位符编号: 说明}（公式的线性文字）。一组放不进去时最多重试 retries 次：
    第 1 次原样再问，第 2 次拆成两半，之后拆开的再问。

    给了 flags（dict）时，试满次数仍放不进去的占位符放到所在那一句的句末，原文 → [(第几行,
    开始, 结束)] 记下要标出来核对的范围；格式标记放不进去的去掉。不给 flags 时这样的段落
    不算补上（调用方按整段写法写回）。reasons（dict）记下没补上的原因：no_answer 没拿到
    回答、broken 标记仍对不上、text 改了译文文字。stats（dict）记下请求次数等（排查用）。"""
    import json
    from concurrent.futures import ThreadPoolExecutor

    from engines.base_engine import engine_supports_chat, strip_markdown_json

    reasons = {} if reasons is None else reasons
    # first / retried / split：首次、重试、拆开后放好的标记个数
    counts = {"requests": 0, "groups": 0, "marks": 0, "first": 0, "retried": 0, "split": 0, "settled": 0,
              "dropped": 0}
    placed: dict[str, tuple[str, str]] = {}
    lines: dict[str, list[_PlaceLine]] = {}
    pending: list[_PlaceGroup] = []
    for source, marked, translation in items:
        texts = placeholder_texts(marked, source)
        if texts is None:
            continue
        sent_lines = marked.split("\n")
        target_lines = str(translation).split("\n")
        if len(sent_lines) != len(target_lines):
            reasons[source] = "text"
            continue
        seed = (seeds or {}).get(source)
        seed_lines = None
        if seed and same_text(strip_marks(seed, texts), translation):
            seed_lines = seed.split("\n")
            if len(seed_lines) != len(sent_lines):
                seed_lines = None
        work = []
        for index, (sent_line, target_line) in enumerate(zip(sent_lines, target_lines)):
            current = _clean_line(sent_line, seed_lines[index]) if seed_lines else _seed_slots(
                sent_line, target_line, texts
            )
            order = {name: position for position, (name, _p, _r) in enumerate(_mark_tokens(sent_line))
                     if not name.startswith("/")}
            line = _PlaceLine((source, index), sent_line, current, order, dict((hints or {}).get(source, {})), texts)
            work.append(line)
            missing = [name for name, _position in _units(sent_line) if name not in _present_units(current)]
            for units in _group_units(sent_line, missing, max(1, group_size)) if missing else []:
                pending.append(_PlaceGroup(line, units))
        lines[source] = work
    counts["groups"] = len(pending)
    counts["marks"] = sum(len(group.units) for group in pending)
    chat = engine_supports_chat(engine)
    if pending and not chat:
        reasons.update((group.line.key[0], "no_answer") for group in pending)

    def ask(batch: list[_PlaceGroup]):
        payload = json.dumps([_group_request(group) for group in batch], ensure_ascii=False)
        try:
            answers = json.loads(strip_markdown_json(engine.chat(PLACE_MARKS_PROMPT, payload)))
        except Exception as exc:  # noqa: BLE001 - 补标记失败只是重试或退回
            if log is not None:
                log(f"段内格式补标记失败：{exc!r}")
            return batch, None
        return batch, answers if isinstance(answers, list) and len(answers) == len(batch) else None

    gave_up: list[_PlaceGroup] = []
    while pending and chat:
        if should_stop is not None and should_stop():
            break
        round_groups: list[_PlaceGroup] = []
        for group in pending:
            if group.attempts == _SPLIT_AT_ATTEMPT and len(group.units) > 1 and not group.split:
                half = len(group.units) // 2
                round_groups += [_PlaceGroup(group.line, group.units[:half], group.attempts, True),
                                 _PlaceGroup(group.line, group.units[half:], group.attempts, True)]
            else:
                round_groups.append(group)
        batches: list[list[_PlaceGroup]] = []
        for group in round_groups:
            if (batches and len(batches[-1]) < PLACE_MARKS_BATCH
                    and sum(len(item.units) for item in batches[-1]) + len(group.units) <= max(1, request_marks)):
                batches[-1].append(group)
            else:
                batches.append([group])
        counts["requests"] += len(batches)
        with ThreadPoolExecutor(max_workers=max(1, min(concurrency, len(batches)))) as pool:
            results = list(pool.map(ask, batches))
        pending = []
        for batch, answers in results:
            for index, group in enumerate(batch):
                key = group.line.key[0]
                answer = answers[index] if answers is not None else None
                base = _zero_width(group.line.current)
                found = (_answer_positions(answer, base, group.units, group.line.texts, group.line.hints)
                         if isinstance(answer, str) else None) or {}
                if found and not _apply_group(group.line, found):
                    found = {}
                if found:
                    counts["first" if group.attempts == 0 else "split" if group.split else "retried"] += len(found)
                rest = [unit for unit in group.units if unit not in found]
                if not rest:
                    continue
                reasons[key] = "no_answer" if answer is None else "broken"
                group.units = rest
                group.attempts += 1
                (pending if group.attempts <= retries else gave_up).append(group)
    gave_up += pending  # 停止时还没轮到的
    leftovers: dict[tuple[str, int], list[str]] = {}
    for group in gave_up:
        leftovers.setdefault(group.line.key, []).extend(group.units)

    for source, marked, translation in items:
        work = lines.get(source)
        if work is None:
            continue
        texts = placeholder_texts(marked, source) or {}
        flagged: list[tuple[int, int, int]] = []
        failed = False
        for line in work:
            missing = leftovers.get(line.key)
            if not missing:
                continue
            if flags is None:
                failed = True
                break
            counts["settled"] += sum(1 for unit in missing if unit.startswith("x"))
            counts["dropped"] += sum(1 for unit in missing if unit.startswith("g"))
            flagged += [(line.key[1], begin, end) for begin, end in _settle_leftovers(line, missing)]
        if failed:
            continue
        final = "\n".join(line.current for line in work)
        if not same_text(strip_marks(final, texts), translation):
            reasons[source] = "text"
            continue
        sent_lines = marked.split("\n")
        if not all(parse_marked(line.current, {int(n[1:]) for n in _present_units(line.current) if n[0] == "g"},
                                _line_slots(sent_line)) is not None for line, sent_line in zip(work, sent_lines)):
            reasons[source] = "broken"
            continue
        placed[source] = (marked, final)
        reasons.pop(source, None)
        if flagged and flags is not None:
            flags[source] = flagged
    if stats is not None:
        for key, value in counts.items():
            stats[key] = stats.get(key, 0) + value
    return placed
