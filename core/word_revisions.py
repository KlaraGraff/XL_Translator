"""Word 内容控件（w:sdt）与未接受的修订（w:ins / w:del）。

python-docx 有两处看不见的正文：

- 内容控件里的段落不在 ``doc.paragraphs`` 里——控件是 body 的直接子节点，段落挂在
  它下面的 ``w:sdtContent`` 里，读不到、也写不进去；
- 未接受的修订不在段落的可见文本里——``w:ins`` 包着的 ``w:t`` 读不到，删除的文字在
  ``w:del`` 里的 ``w:delText``。这些字以前一律漏译，只有任务跑起来之后一句「存在无法
  读取的内容」告警（V10.5 之前）。

现在的约定：

- 内容控件（自动目录除外）里的段落当成普通正文段落，参与抽取和写回，不由用户选择；
- 修订在**扫描时**就数出来（``tracked_change_count``），用户在点「开始翻译」时决定：
  ``skip`` 保持原文不动（默认），``translate`` 连修订里的文字一起翻译、修订标记
  （作者、日期、id）原样保留。

本模块只负责这两件事的读写，不在模块顶层 import ``core.word_document``（那边要 import
这里），避免循环导入；自动目录的判定复用那边同一套实现，函数内延迟导入。
"""
from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from difflib import SequenceMatcher

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# 用户对未接受的修订的两种处理方式。缺省（旧客户端不传）按 skip。
REVISION_MODE_SKIP = "skip"
REVISION_MODE_TRANSLATE = "translate"
REVISION_MODES = (REVISION_MODE_SKIP, REVISION_MODE_TRANSLATE)


def normalize_revision_mode(value: object) -> str:
    """把请求里的 revision_mode 归一成两个合法值之一；认不出的按 skip。

    旧客户端根本不传这个字段。绝不能因为缺一个字段就改变已有行为（把修订里的原文
    翻掉），所以默认值是不会改动任何修订内容的 skip。
    """
    cleaned = str(value or "").strip().lower()
    return cleaned if cleaned in REVISION_MODES else REVISION_MODE_SKIP


@dataclass(frozen=True)
class WordRevisionReport:
    """扫描期数到的未接受修订：``w:ins`` + ``w:del``，只数有可见文字的。"""

    tracked_change_count: int = 0

    @property
    def found(self) -> bool:
        return self.tracked_change_count > 0

    def as_dict(self) -> dict[str, int]:
        return {"tracked_change_count": self.tracked_change_count}


@dataclass(frozen=True)
class RevisionWriteStats:
    """一次写回里修订的处理结果（逐文件记录与日志用）。"""

    mode: str = REVISION_MODE_SKIP
    tracked_change_count: int = 0
    translated_count: int = 0

    def describe(self) -> str:
        if self.mode == REVISION_MODE_TRANSLATE:
            return f"翻译了 {self.translated_count} 处修订（保留修订标记）"
        return f"跳过了 {self.tracked_change_count} 处修订（保持原文）"

    def as_dict(self) -> dict[str, object]:
        # 键名不叫 "text"：任务中心的脱敏会把 text 当成正文内容整键丢掉，那句「跳过了
        # N 处修订（保持原文）」就到不了界面上。
        return {
            "mode": self.mode,
            "count": (
                self.translated_count
                if self.mode == REVISION_MODE_TRANSLATE
                else self.tracked_change_count
            ),
            "label": self.describe(),
        }


# ── 结构遍历 ─────────────────────────────────────────────────────────


def _has_ancestor_tag(element, tag: str, root) -> bool:
    qualified = qn(tag)
    parent = element.getparent()
    while parent is not None and parent is not root:
        if parent.tag == qualified:
            return True
        parent = parent.getparent()
    return False


def _iter_outermost_elements(root, tag: str):
    """root 下同名嵌套里最外层的那一个，避免嵌套结构被重复计数。"""
    qualified = qn(tag)
    for element in root.iter(qualified):
        if not _has_ancestor_tag(element, tag, root):
            yield element


def _is_toc_control(element) -> bool:
    from core.word_document import _is_toc_content_control

    return _is_toc_content_control(element)


def toc_controls(body) -> list:
    """正文里的自动目录控件（``w:sdt``）。目录是域，翻译它没有意义。"""
    return [
        control
        for control in _iter_outermost_elements(body, "w:sdt")
        if _is_toc_control(control)
    ]


def _inside(element, containers) -> bool:
    return any(element in list(container.iter()) for container in containers)


def _visible_text(element, tags: tuple[str, ...]) -> str:
    return "".join(
        node.text or "" for node in element.iter(*(qn(tag) for tag in tags))
    ).strip()


# ── 未接受的修订 ─────────────────────────────────────────────────────

_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _has_ancestor_clark(element, clark_tag: str, root) -> bool:
    parent = element.getparent()
    while parent is not None and parent is not root:
        if parent.tag == clark_tag:
            return True
        parent = parent.getparent()
    return False


def _in_embedded_story(element, root) -> bool:
    """在文本框里，或在新版图形的旧版备份（``mc:Fallback``）里。

    文本框的段落由 core/word_embedded_text.py 单独翻译；``mc:Fallback`` 是同一个文本框
    给旧版 Word 的第二份拷贝，数进去就成了一处修订算两处。
    """
    return _has_ancestor_tag(element, "w:txbxContent", root) or _has_ancestor_clark(
        element, _MC_FALLBACK, root
    )


def iter_tracked_revisions(body, *, excluded=()) -> list:
    """正文里未接受的修订元素（最外层 ``w:ins`` / ``w:del``），只留有可见文字的。

    ``excluded`` 是不参与处理的容器（调用方传自动目录控件）。

    ``w:ins`` 也用来标记「这个段落标记是新插入的」（``w:pPr/w:rPr/w:ins``），那种
    节点一个 ``w:t`` 都没有，不算一处修订；``w:del`` 里的文字在 ``w:delText`` 里。
    文本框在 ``mc:Fallback`` 里的旧版备份不算，否则同一处修订数两遍。
    """
    found: list = []
    for tag, text_tags in (
        ("w:ins", ("w:t",)),
        ("w:del", ("w:delText", "w:t")),
    ):
        for element in _iter_outermost_elements(body, tag):
            if not _visible_text(element, text_tags):
                continue
            if _has_ancestor_clark(element, _MC_FALLBACK, body):
                continue
            if excluded and _inside(element, excluded):
                continue
            found.append(element)
    return found


def count_tracked_changes(body, *, excluded=()) -> int:
    """正文里未接受的修订处数（扫描期的 ``tracked_change_count``）。"""
    return len(iter_tracked_revisions(body, excluded=excluded))


def paragraph_of(element):
    """element 所在的 ``w:p``。"""
    parent = element.getparent()
    while parent is not None and parent.tag != qn("w:p"):
        parent = parent.getparent()
    return parent


_REVISION_TAGS = (qn("w:ins"), qn("w:del"))


def _revision_kind(run, paragraph) -> str | None:
    """run 落在哪种修订里（"ins" / "del"），不在修订里返回 None。"""
    parent = run.getparent()
    while parent is not None and parent is not paragraph:
        if parent.tag == qn("w:ins"):
            return "ins"
        if parent.tag == qn("w:del"):
            return "del"
        parent = parent.getparent()
    return None


def _run_text(run) -> str:
    pieces = []
    for child in run.iterchildren():
        if child.tag in (qn("w:t"), qn("w:delText")):
            pieces.append(child.text or "")
        elif child.tag in (qn("w:tab"), qn("w:ptab")):
            pieces.append("\t")
        elif child.tag in (qn("w:br"), qn("w:cr")):
            pieces.append("\n")
    return "".join(pieces)


def _paragraph_runs(paragraph) -> list:
    """段落自己的 run（含修订、超链接、行内控件里的），不进文本框和旧版备份。"""
    return [
        run
        for run in paragraph.iter(qn("w:r"))
        if not _in_embedded_story(run, paragraph) and paragraph_of(run) is paragraph
    ]


def paragraph_versions(paragraph) -> tuple[str, str]:
    """(修改前, 修改后)：拒绝全部修订、接受全部修订后这一段各是什么字。

    python-docx 的 ``paragraph.text`` 两边都不算——插入和删除的字都读不到，「合同金额为
    <删>100</删><增>120</增>万元」读出来是「合同金额为万元」。整句翻译要的是这两个版本。
    """
    before: list[str] = []
    after: list[str] = []
    for run in _paragraph_runs(paragraph):
        kind = _revision_kind(run, paragraph)
        text = _run_text(run)
        if kind != "ins":
            before.append(text)
        if kind != "del":
            after.append(text)
    return "".join(before).strip(), "".join(after).strip()


def paragraph_has_revisions(paragraph) -> bool:
    return any(
        _visible_text(element, ("w:t", "w:delText"))
        for element in paragraph.iter(*_REVISION_TAGS)
        if not _in_embedded_story(element, paragraph)
    )


def paragraph_coverage_text(paragraph) -> str:
    """覆盖率判定用的段落文字：普通段落就是 ``paragraph.text``；整段都在修订里的段落
    （「修订一起翻译」写出的译文段常常是这样）读接受修订后的版本，否则它看起来是空段，
    前面的原文就被误判成「未翻译」。"""
    text = (paragraph.text or "").strip()
    if text or not paragraph_has_revisions(paragraph._p):
        return text
    return paragraph_versions(paragraph._p)[1]


@dataclass
class RevisionUnit:
    """「修订一起翻译」时单独处理的一段。

    含修订的段落整段由这里接手（正文、表格、内容控件三处的常规写回都跳过它）：修改前
    和修改后各当一句完整的原文送翻，写回时对比两句译文，不同的词标成删除 / 新增。
    同一个单元格里没有修订的段落也归这里——那一格的常规写法是整格一起翻，有一段被
    接走，剩下的就得逐段跟着走，不然这一格的译文会缺一截。
    """

    paragraph: object  # w:p
    has_revisions: bool
    before: str
    after: str
    plain: str
    revision_count: int = 0
    protected: bool = False


@dataclass
class RevisionLayout:
    units: list[RevisionUnit]
    # 常规写回要跳过的段落（w:p）和单元格（w:tc）
    paragraphs: set
    cells: set

    def skips_paragraph(self, element) -> bool:
        return element in self.paragraphs

    def skips_cell(self, element) -> bool:
        return element in self.cells


EMPTY_REVISION_LAYOUT = RevisionLayout(units=[], paragraphs=set(), cells=set())


def _top_level_child(element, body):
    while element is not None and element.getparent() is not body:
        element = element.getparent()
    return element


def front_matter_start(doc, front_matter):
    """「保护封面和目录」的正文起点在 body 里的那个直接子节点；不保护或没找到返回 None。"""
    if front_matter is None or not getattr(front_matter, "found", False):
        return None
    paragraphs = doc.paragraphs
    index = front_matter.body_start_index
    if index is None or not 0 <= index < len(paragraphs):
        return None
    return paragraphs[index]._p


def is_front_matter(element, body, start) -> bool:
    """element 排在正文起点之前（封面页内容控件、封面表格都算）。"""
    if start is None:
        return False
    child = _top_level_child(element, body)
    if child is None:
        return False
    return body.index(child) < body.index(start)


def _container_of(paragraph, body):
    """段落最近的容器："content_control" / ("cell", w:tc) / "body"。"""
    parent = paragraph.getparent()
    while parent is not None and parent is not body:
        if parent.tag == qn("w:sdtContent"):
            return "content_control", None
        if parent.tag == qn("w:tc"):
            return "cell", parent
        parent = parent.getparent()
    return "body", None


def revision_layout(doc, *, front_matter=None) -> RevisionLayout:
    """找出「修订一起翻译」要单独处理的段落。没有修订时返回空布局。"""
    body = doc.element.body
    excluded = toc_controls(body)
    elements = iter_tracked_revisions(body, excluded=excluded)
    if not elements:
        return EMPTY_REVISION_LAYOUT
    counts: dict = {}
    for element in elements:
        paragraph = paragraph_of(element)
        if paragraph is None or _in_embedded_story(paragraph, body):
            continue
        counts[paragraph] = counts.get(paragraph, 0) + 1
    if not counts:
        return EMPTY_REVISION_LAYOUT

    start = front_matter_start(doc, front_matter)
    paragraphs: set = set()
    cells: set = set()
    for paragraph in counts:
        kind, cell = _container_of(paragraph, body)
        if kind == "cell":
            cells.add(cell)
            paragraphs.update(cell.iterchildren(qn("w:p")))
        else:
            paragraphs.add(paragraph)

    units: list[RevisionUnit] = []
    for paragraph in body.iter(qn("w:p")):
        if paragraph not in paragraphs:
            continue
        before, after = paragraph_versions(paragraph)
        plain = "".join(
            _run_text(run)
            for run in _paragraph_runs(paragraph)
            if _revision_kind(run, paragraph) is None
        ).strip()
        units.append(
            RevisionUnit(
                paragraph=paragraph,
                has_revisions=paragraph in counts,
                before=before,
                after=after,
                plain=plain,
                revision_count=counts.get(paragraph, 0),
                protected=is_front_matter(paragraph, body, start),
            )
        )
    return RevisionLayout(units=units, paragraphs=paragraphs, cells=cells)


def _looks_translated(text: str, *, target_lang: str, source_lang: str) -> bool:
    from core.translation_coverage import looks_like_source_text, looks_like_target_text

    return bool(
        text
        and looks_like_target_text(text, source_lang=source_lang, target_lang=target_lang)
        and not looks_like_source_text(text, source_lang=source_lang, target_lang=target_lang)
    )


def _unit_already_translated(unit: RevisionUnit, *, target_lang: str, source_lang: str) -> bool:
    """补译：这一段自己已是译文（只输出译文的旧产物），或后面紧跟着一段译文。"""
    if _looks_translated(unit.after, target_lang=target_lang, source_lang=source_lang):
        return True
    following = unit.paragraph.getnext()
    if following is None or following.tag != qn("w:p"):
        return False
    return _looks_translated(
        paragraph_coverage_text(Paragraph(following, None)),
        target_lang=target_lang,
        source_lang=source_lang,
    )


def _unit_sources(unit: RevisionUnit) -> list[str]:
    return [unit.before, unit.after] if unit.has_revisions else [unit.plain]


def extract_revision_segments(
    doc,
    *,
    target_lang: str,
    source_lang: str,
    untranslated_only: bool = False,
    front_matter=None,
) -> list:
    """「修订一起翻译」时含修订段落的待译原文：修改前、修改后各一句整句。"""
    from core.word_document import WordSegment, _is_translatable_source

    segments = []
    for index, unit in enumerate(revision_layout(doc, front_matter=front_matter).units):
        if unit.protected:
            continue
        if untranslated_only and _unit_already_translated(
            unit, target_lang=target_lang, source_lang=source_lang
        ):
            continue
        for source in _unit_sources(unit):
            if not _is_translatable_source(source, target_lang=target_lang, source_lang=source_lang):
                continue
            segments.append(
                WordSegment(
                    source=source,
                    kind="tracked_change",
                    location=f"revision[{index}]",
                    section_path="修订",
                )
            )
    return segments


_TOKEN_RE = re.compile(r"\s+|\w+|[^\w\s]")
# 两句译文差得太多时，逐词对比标出来是一片碎片，不如整句删、整句加清楚。
_MIN_DIFF_RATIO = 0.5


def revision_pieces(before: str, after: str, *, target_lang: str) -> list[tuple[str, str]]:
    """对比修改前、修改后两句译文：[("equal" | "delete" | "insert", 文字), ...]。"""
    from core.word_inline_marks import _no_word_spaces

    if before == after:
        return [("equal", after)] if after else []
    if not before:
        return [("insert", after)]
    if not after:
        return [("delete", before)]
    if _no_word_spaces(target_lang):
        old, new = list(before), list(after)
    else:
        old, new = _TOKEN_RE.findall(before), _TOKEN_RE.findall(after)
    matcher = SequenceMatcher(None, old, new, autojunk=False)
    if matcher.ratio() < _MIN_DIFF_RATIO:
        return [("delete", before), ("insert", after)]
    pieces: list[tuple[str, str]] = []

    def push(op: str, text: str) -> None:
        if not text:
            return
        if pieces and pieces[-1][0] == op:
            pieces[-1] = (op, pieces[-1][1] + text)
        else:
            pieces.append((op, text))

    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        if op == "equal":
            push("equal", "".join(new[j1:j2]))
            continue
        push("delete", "".join(old[i1:i2]))
        push("insert", "".join(new[j1:j2]))
    return _absorb_whitespace(pieces)


def _absorb_whitespace(pieces: list[tuple[str, str]]) -> list[tuple[str, str]]:
    """一处改动连成一截：夹在两处改动之间、只剩空格的「未改动」并进去，连着的删除并成
    一个、新增并成一个，删除在前。不然标记会碎成一词一截，读起来像打了补丁。"""
    spread: list[tuple[str, str]] = []
    for index, (op, text) in enumerate(pieces):
        if (
            op == "equal"
            and not text.strip()
            and 0 < index < len(pieces) - 1
            and pieces[index - 1][0] != "equal"
            and pieces[index + 1][0] != "equal"
        ):
            spread.extend((("delete", text), ("insert", text)))
        else:
            spread.append((op, text))
    result: list[tuple[str, str]] = []
    pending = {"delete": "", "insert": ""}

    def flush() -> None:
        for op in ("delete", "insert"):
            if pending[op]:
                result.append((op, pending[op]))
                pending[op] = ""

    for op, text in spread:
        if op == "equal":
            flush()
            result.append((op, text))
        else:
            pending[op] += text
    flush()
    return result


class _RevisionIds:
    """新写出来的 w:ins / w:del 要有自己的 w:id：从文档里现有的最大编号往上数。"""

    def __init__(self, body) -> None:
        highest = 0
        for element in body.iter():
            value = element.get(qn("w:id"))
            if value is not None and value.lstrip("-").isdigit():
                highest = max(highest, int(value))
        self._next = highest + 1

    def take(self) -> str:
        value = self._next
        self._next += 1
        return str(value)


def _wrapper_template(paragraph, tag: str):
    """新修订标记沿用原段落里同类修订的作者和日期；没有同类的就借另一类的。"""
    for wanted in (qn(tag), *_REVISION_TAGS):
        for element in paragraph.iter(wanted):
            if not _in_embedded_story(element, paragraph) and _visible_text(
                element, ("w:t", "w:delText")
            ):
                return element
    return None


def _new_wrapper(tag: str, template, ids: _RevisionIds):
    wrapper = OxmlElement(tag)
    if template is not None:
        for name in ("w:author", "w:date"):
            value = template.get(qn(name))
            if value is not None:
                wrapper.set(qn(name), value)
    wrapper.set(qn("w:id"), ids.take())
    return wrapper


def _append_pieces(target, pieces, *, rpr, templates: dict, ids: _RevisionIds) -> None:
    for op, text in pieces:
        run = OxmlElement("w:r")
        if rpr is not None:
            run.append(deepcopy(rpr))
        run.text = text
        if op == "equal":
            target.append(run)
            continue
        if op == "delete":
            for node in run.iter(qn("w:t")):
                node.tag = qn("w:delText")
        tag = "w:del" if op == "delete" else "w:ins"
        wrapper = _new_wrapper(tag, templates.get(tag), ids)
        wrapper.append(run)
        target.append(wrapper)


def _translation_run_shape(paragraph: Paragraph, *, target_lang: str):
    """译文 run 的格式：借双语版式插译文段的那套格式复制规则，拿到 rPr 再撤掉临时段。"""
    from core.word_document import _insert_translation_paragraph_after

    probe = _insert_translation_paragraph_after(paragraph, "x", target_lang=target_lang)
    run = probe._p.find(qn("w:r"))
    rpr = run.find(qn("w:rPr")) if run is not None else None
    rpr = deepcopy(rpr) if rpr is not None else None
    probe._p.getparent().remove(probe._p)
    return rpr


def _clear_paragraph_text(paragraph) -> None:
    """只输出译文：段落里的文字连同修订一起清掉；图片、书签、段落格式（含段落标记上的
    修订）留着，译文接在后面。"""
    from core.word_document import _drop_emptied_hyperlinks

    for element in [
        node
        for node in paragraph.iter(*_REVISION_TAGS)
        if not _in_embedded_story(node, paragraph)
        and not _has_ancestor_tag(node, "w:pPr", paragraph)
        and not any(_has_ancestor_tag(node, tag, paragraph) for tag in ("w:ins", "w:del"))
    ]:
        element.getparent().remove(element)
    for run in _paragraph_runs(paragraph):
        for node in list(run.iterchildren()):
            if node.tag in (qn("w:t"), qn("w:delText"), qn("w:tab"), qn("w:br"), qn("w:cr")):
                run.remove(node)
        if all(child.tag == qn("w:rPr") for child in run.iterchildren()):
            run.getparent().remove(run)
    _drop_emptied_hyperlinks(Paragraph(paragraph, None))


def apply_revision_translations(
    doc,
    translations: dict[str, str],
    *,
    target_lang: str,
    source_lang: str,
    mode: str = REVISION_MODE_SKIP,
    untranslated_only: bool = False,
    output_translation_only: bool = False,
    front_matter=None,
    layout: RevisionLayout | None = None,
) -> RevisionWriteStats:
    """按用户选择的模式处理未接受的修订（调用方负责 save）。

    ``skip``：一个字都不动，只数出有几处（逐文件记录写「跳过了 N 处修订」）。

    ``translate``：含修订的段落修改前、修改后各译一句，对比两句译文，相同的部分写成
    普通文字，不同的部分包进 ``w:del`` / ``w:ins``（作者、日期沿用原修订）。在 Word 里
    「接受全部」得到修改后的译文，「拒绝全部」得到修改前的译文。

    - 双语版式：原文段落一个字不动（修订标记照旧），修订的译文写进紧跟其后的译文段；
    - 只输出译文：原段落的文字原地换成上面那样带标记的译文。

    ``layout`` 传常规写回前取的那一份（常规写回会在段落后插译文段，事后再取一遍就会把
    插进去的段落也算进来）；不传就现取。
    """
    from core.word_document import (
        _insert_translation_paragraph_after,
        _is_translatable_source,
        _paragraph_has_field,
        _replace_paragraph_text,
        _resolve_translation,
    )
    from core.word_layout_fixes import loosen_row_height

    body = doc.element.body
    elements = iter_tracked_revisions(body, excluded=toc_controls(body))
    stats = RevisionWriteStats(
        mode=normalize_revision_mode(mode), tracked_change_count=len(elements)
    )
    if stats.mode != REVISION_MODE_TRANSLATE:
        return stats
    if layout is None:
        layout = revision_layout(doc, front_matter=front_matter)

    def translated(source: str) -> str | None:
        if not source:
            return ""
        if not _is_translatable_source(source, target_lang=target_lang, source_lang=source_lang):
            return source
        resolved = _resolve_translation(source, translations)
        if resolved is None:
            # 译文和原文一样（编号、代号）时 _resolve_translation 也返回 None。
            raw = str(translations.get(source.strip()) or "").strip()
            return source if raw and raw.casefold() == source.casefold() else None
        return resolved.text.strip()

    ids = _RevisionIds(body)
    count = 0
    for unit in layout.units:
        if unit.protected:
            continue
        if untranslated_only and _unit_already_translated(
            unit, target_lang=target_lang, source_lang=source_lang
        ):
            continue
        paragraph = Paragraph(unit.paragraph, doc._body)
        if unit.has_revisions:
            before, after = translated(unit.before), translated(unit.after)
            if before is None or after is None or not (before or after):
                continue
            pieces = revision_pieces(before, after, target_lang=target_lang)
        else:
            text = translated(unit.plain)
            if not text or text == unit.plain:
                continue
            pieces = [("equal", text)]
        in_place = output_translation_only and not _paragraph_has_field(paragraph)
        if not unit.has_revisions:
            if in_place:
                _replace_paragraph_text(paragraph, pieces[0][1], target_lang=target_lang)
            else:
                _insert_translation_paragraph_after(paragraph, pieces[0][1], target_lang=target_lang)
        else:
            rpr = _translation_run_shape(paragraph, target_lang=target_lang)
            # 只输出译文时原段落马上要清空，作者和日期先取下来。
            templates = {tag: _wrapper_template(unit.paragraph, tag) for tag in ("w:ins", "w:del")}
            if in_place:
                _clear_paragraph_text(unit.paragraph)
                target = unit.paragraph
            else:
                target = _insert_translation_paragraph_after(paragraph, "", target_lang=target_lang)._p
                for run in list(target.iterchildren(qn("w:r"))):
                    target.remove(run)
            _append_pieces(target, pieces, rpr=rpr, templates=templates, ids=ids)
            count += unit.revision_count
        cell = unit.paragraph.getparent()
        if cell is not None and cell.tag == qn("w:tc"):
            loosen_row_height(_CellShim(cell))
    return RevisionWriteStats(
        mode=stats.mode,
        tracked_change_count=stats.tracked_change_count,
        translated_count=count,
    )


class _CellShim:
    """loosen_row_height 只用 ``cell._tc``。"""

    def __init__(self, tc) -> None:
        self._tc = tc


# ── 内容控件 ─────────────────────────────────────────────────────────


def _owning_content_control(element, body):
    """这个节点所在的最内层内容控件（``w:sdtContent`` 的父节点）。"""
    parent = element.getparent()
    while parent is not None and parent is not body:
        if parent.tag == qn("w:sdtContent"):
            holder = parent.getparent()
            return holder if holder is not None and holder.tag == qn("w:sdt") else None
        parent = parent.getparent()
    return None


def content_control_paragraph_entries(
    body, *, excluded=(), parent=None, front_matter_start=None
) -> list[tuple[Paragraph, int, int]]:
    """内容控件里的段落：[(Paragraph, 控件序号, 段落在控件内的序号), ...]。

    嵌套控件里的段落只归最内层那个控件；``excluded``（自动目录）整块跳过。文本框里的
    段落不算（core/word_embedded_text.py 那一趟翻它们，这里再算一次就译两遍）。

    ``front_matter_start``：「保护封面和目录」的正文起点（见 ``front_matter_start()``）。
    排在它前面的控件整个不算——Word 的封面页本身就是一个内容控件（docPartGallery
    「Cover Pages」），保护封面时它得原样留着。控件序号仍按全文计数，开不开保护同一
    个控件的定位不变。

    ``parent`` 传文档的 ``_Body``：段落包装对象要带上它才读得到样式（``paragraph.style``
    得经 ``part`` 找样式表），复制段内格式、判断标题都靠这一步。不传也改得了文字，
    只是样式相关的格式复制会退化成默认。
    """
    controls = [
        control
        for control in _iter_outermost_elements(body, "w:sdt")
        if not (excluded and _inside(control, excluded))
    ]
    entries: list[tuple[Paragraph, int, int]] = []
    for control_index, control in enumerate(controls):
        if is_front_matter(control, body, front_matter_start):
            continue
        own = [
            paragraph
            for paragraph in control.iter(qn("w:p"))
            if _owning_content_control(paragraph, body) is control
            and not _in_embedded_story(paragraph, control)
        ]
        for paragraph_index, paragraph in enumerate(own):
            entries.append((Paragraph(paragraph, parent), control_index, paragraph_index))
    return entries


def _already_translated(paragraphs: list[Paragraph], index: int, *, target_lang: str, source_lang: str) -> bool:
    """补译：这一段后面紧跟着一段目标语言的段落，说明上次已经译过。"""
    if index + 1 >= len(paragraphs):
        return False
    return _looks_translated(
        paragraph_coverage_text(paragraphs[index + 1]),
        target_lang=target_lang,
        source_lang=source_lang,
    )


def _control_entries(doc, *, front_matter=None):
    """内容控件段落清单，外加按控件归组的段落（补译判定要看「下一段是不是已有译文」）。"""
    body = doc.element.body
    entries = content_control_paragraph_entries(
        body,
        excluded=toc_controls(body),
        parent=doc._body,
        front_matter_start=front_matter_start(doc, front_matter),
    )
    control_paragraphs: dict[int, list[Paragraph]] = {}
    for paragraph, control_index, _paragraph_index in entries:
        control_paragraphs.setdefault(control_index, []).append(paragraph)
    return entries, control_paragraphs


def extract_content_control_segments(
    doc,
    *,
    target_lang: str,
    source_lang: str,
    untranslated_only: bool = False,
    front_matter=None,
    layout: RevisionLayout | None = None,
) -> list:
    """内容控件里的段落，作为 ``WordSegment`` 列表。

    ``front_matter``：「保护封面和目录」划出的边界，封面页控件不抽；``layout``：
    「修订一起翻译」时由修订那一趟接手的段落，这里跳过。

    控件里的段落和普通正文一样参与翻译（用户没有选择），所以定位用
    ``content_control[i].paragraph[j]``，写回时按同一个序号找回。自动目录整块不算。
    """
    from core.word_document import (
        WordSegment,
        _is_toc_or_field_paragraph,
        _is_translatable_source,
        _paragraph_source_text,
    )

    entries, control_paragraphs = _control_entries(doc, front_matter=front_matter)
    skip = layout or EMPTY_REVISION_LAYOUT
    segments = []
    for paragraph, control_index, paragraph_index in entries:
        if skip.skips_paragraph(paragraph._p):
            continue
        source = _paragraph_source_text(paragraph)
        if not _is_translatable_source(source, target_lang=target_lang, source_lang=source_lang):
            continue
        if _is_toc_or_field_paragraph(paragraph):
            continue
        if untranslated_only and _already_translated(
            control_paragraphs.get(control_index, []),
            paragraph_index,
            target_lang=target_lang,
            source_lang=source_lang,
        ):
            continue
        segments.append(
            WordSegment(
                source=source,
                kind="content_control",
                location=f"content_control[{control_index}].paragraph[{paragraph_index}]",
                section_path="内容控件",
            )
        )
    return segments


def apply_content_control_translations(
    doc,
    translations: dict[str, str],
    *,
    target_lang: str,
    source_lang: str,
    output_translation_only: bool = False,
    untranslated_only: bool = False,
    front_matter=None,
    layout: RevisionLayout | None = None,
) -> int:
    """把译文写进内容控件里的段落（调用方负责 save）。返回写入的段数。

    ``front_matter`` / ``layout`` 同 ``extract_content_control_segments``。

    控件里夹着域（日期占位之类）时走 ``_replace_paragraph_text_around_fields``，域
    留在原位；否则只输出译文就原地替换，双语版式就在控件里另起一段译文——控件本身
    不动，用户还在同一个表单格子里看到原文和译文。
    """
    from core.word_document import (
        _insert_translation_paragraph_after,
        _is_toc_or_field_paragraph,
        _is_translatable_source,
        _paragraph_has_field,
        _paragraph_source_text,
        _replace_paragraph_text,
        _replace_paragraph_text_around_fields,
        _resolve_translation,
    )

    entries, control_paragraphs = _control_entries(doc, front_matter=front_matter)
    skip = layout or EMPTY_REVISION_LAYOUT
    written = 0
    for paragraph, control_index, paragraph_index in entries:
        if skip.skips_paragraph(paragraph._p):
            continue
        source = _paragraph_source_text(paragraph)
        if not _is_translatable_source(source, target_lang=target_lang, source_lang=source_lang):
            continue
        if _is_toc_or_field_paragraph(paragraph):
            continue
        if untranslated_only and _already_translated(
            control_paragraphs.get(control_index, []),
            paragraph_index,
            target_lang=target_lang,
            source_lang=source_lang,
        ):
            continue
        resolved = _resolve_translation(source, translations)
        if resolved is None:
            continue
        if (resolved.replace_only or output_translation_only) and not _paragraph_has_field(paragraph):
            _replace_paragraph_text(paragraph, resolved.text, target_lang=target_lang)
        elif (resolved.replace_only or output_translation_only) and _replace_paragraph_text_around_fields(
            paragraph, resolved.text, target_lang=target_lang
        ):
            pass
        else:
            _insert_translation_paragraph_after(paragraph, resolved.text, target_lang=target_lang)
        written += 1
    return written


def inline_content_control_count(body, *, excluded=()) -> int:
    """行内内容控件（日期、勾选框之类）的个数。

    这类控件挂在 ``w:r`` 上、和正文混在同一段里，本版还没法把它们当成独立段落翻译。
    调用方据此保留「存在无法读取的内容」告警，而不是悄悄漏掉。
    """
    total = 0
    for control in _iter_outermost_elements(body, "w:sdt"):
        if excluded and _inside(control, excluded):
            continue
        if next(control.iter(qn("w:p")), None) is not None:
            continue
        if _visible_text(control, ("w:t",)):
            total += 1
    return total
