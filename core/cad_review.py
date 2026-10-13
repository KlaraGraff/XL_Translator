"""CAD 译文对照表：翻完以后查看、修改每句译文，改完只重新写回图纸，不再调用模型。

数据都来自任务自己留下的记录：批次清单（``manifest.json``，列出每张图纸的原图、
输出和逐张清单）和逐张清单里的 ``units``（每段原文和它的译文）。用户改过的译文存在
批次目录的 ``cad-review-edits.json`` 里，按逐张清单的路径分开记。

重新写回走的仍是 ``CadTranslationPipeline``：只是把「翻译」换成按对照表查表，所以
双语、防重叠、译文字体、表格单元格这些写法和第一次翻译完全一样。
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import shutil
import tempfile
from collections.abc import Callable, Iterable, Mapping
from dataclasses import asdict, fields
from pathlib import Path
from typing import Any

from core.cad_translation import (
    CadPipelineError,
    CadPipelineOptions,
    CadTranslationPipeline,
)

EDITS_FILENAME = "cad-review-edits.json"
# 对照表里的状态：translated 已翻译、kept 已是目标语言保留原文、unresolved 没译出来。
# review：译文已写入，但未采用记忆库「始终采用」词条的设定译法，需复核。
_STATUSES = frozenset({"translated", "kept", "unresolved", "review"})
# Excel 里的列。第一列是图纸文件名，导回时按「图纸 + 原文」找到对应的那一行。
EXPORT_HEADERS = ("图纸", "原文", "译文", "状态", "出现次数")
_STATUS_LABELS = {"translated": "已翻译", "kept": "保留原文", "unresolved": "未译出", "review": "需复核", "edited": "已修改"}


class CadReviewError(Exception):
    """对照表做不了这件事：消息直接给用户看。"""


def _read_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise CadReviewError("这次任务的译文记录已被清理或损坏，无法查看。请重新翻译这批图纸。") from exc
    if not isinstance(value, dict):
        raise CadReviewError("这次任务的译文记录已被清理或损坏，无法查看。请重新翻译这批图纸。")
    return value


def _edits_path(batch_manifest: Path) -> Path:
    return batch_manifest.parent / EDITS_FILENAME


def load_edits(batch_manifest: Path) -> dict[str, dict[str, str]]:
    path = _edits_path(batch_manifest)
    if not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    files = value.get("files") if isinstance(value, dict) else None
    if not isinstance(files, dict):
        return {}
    return {
        str(key): {str(source): str(target) for source, target in edits.items()}
        for key, edits in files.items()
        if isinstance(edits, dict)
    }


def _write_edits(batch_manifest: Path, edits: Mapping[str, Mapping[str, str]]) -> None:
    path = _edits_path(batch_manifest)
    payload = {"schema_version": 1, "files": {key: dict(value) for key, value in edits.items() if value}}
    staged = path.with_name(f".{path.name}.tmp")
    staged.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(staged, path)


def _file_entries(batch: Mapping[str, Any]) -> list[dict[str, Any]]:
    entries = []
    for item in batch.get("files") or []:
        if not isinstance(item, dict):
            continue
        manifest = str(item.get("manifest_path") or "").strip()
        if not manifest:
            continue  # 失败、停止的图纸没有译文记录
        entries.append(item)
    return entries


def _rows(file_manifest: Mapping[str, Any], edits: Mapping[str, str]) -> list[dict[str, Any]]:
    """同一段原文在图上出现多次时只列一行（写回时也是按原文整体替换）。"""
    crowded = {
        str(item.get("source") or "")
        for item in file_manifest.get("layout_overflow") or []
        if isinstance(item, dict)
    }
    rows: dict[str, dict[str, Any]] = {}
    for unit in file_manifest.get("units") or []:
        if not isinstance(unit, dict):
            continue
        source = str(unit.get("source") or "")
        if not source:
            continue
        row = rows.get(source)
        if row is not None:
            row["count"] += 1
            continue
        status = str(unit.get("status") or "")
        rows[source] = {
            "source": source,
            "original": str(unit.get("translation") or ""),
            "status": status if status in _STATUSES else "translated",
            "count": 1,
            "crowded": source in crowded,
            "reason": str(unit.get("reason") or "") if status == "review" else "",
        }
    out = []
    for row in rows.values():
        edited = edits.get(row["source"])
        row["translation"] = edited if edited is not None else row["original"]
        row["edited"] = edited is not None
        out.append(row)
    return out


def load_review(batch_manifest: str | Path) -> dict[str, Any]:
    """批次里每张图纸的对照表，以及能不能按它重新写回。"""
    batch_path = Path(batch_manifest)
    batch = _read_json(batch_path)
    edits = load_edits(batch_path)
    rewritable_batch = isinstance(batch.get("pipeline_options"), dict)
    files = []
    for entry in _file_entries(batch):
        key = str(entry["manifest_path"])
        path = Path(key)
        source = Path(str(entry.get("source_path") or ""))
        output = Path(str(entry.get("output_path") or ""))
        if not path.is_file():
            files.append({
                "key": key, "name": source.name or output.name, "rows": [], "available": False,
                "rewritable": False, "reason": "这张图纸的译文记录已被清理。",
            })
            continue
        manifest = _read_json(path)
        reason = ""
        if not rewritable_batch:
            reason = "这批图纸是旧版本翻译的，只能查看和导出，不能直接重新写回；重新翻译一次后即可使用。"
        elif not source.is_file():
            reason = "原图纸已被移动或删除，无法重新写回。"
        elif not output.parent.is_dir():
            reason = "译好的图纸所在文件夹已被移动或删除，无法重新写回。"
        files.append({
            "key": key,
            "name": source.name or output.name,
            "output_name": output.name,
            "rows": _rows(manifest, edits.get(key, {})),
            "available": True,
            "rewritable": not reason,
            "reason": reason,
            "edited": len(edits.get(key, {})),
        })
    return {
        "target_lang": str((batch.get("pipeline_options") or {}).get("target_lang") or ""),
        "files": files,
    }


def save_edits(batch_manifest: str | Path, key: str, changes: Mapping[str, str]) -> dict[str, str]:
    """记下一张图纸里改过的译文；改回和第一次译文一样（或清空）就当没改过。

    返回这次真正改动、需要存进记忆库的 ``{原文: 新译文}``。
    """
    batch_path = Path(batch_manifest)
    batch = _read_json(batch_path)
    entry = next((item for item in _file_entries(batch) if str(item["manifest_path"]) == key), None)
    if entry is None:
        raise CadReviewError("找不到这张图纸，请刷新后重试。")
    manifest = _read_json(Path(key))
    originals = {row["source"]: row["original"] for row in _rows(manifest, {})}
    all_edits = load_edits(batch_path)
    current = dict(all_edits.get(key, {}))
    saved: dict[str, str] = {}
    for source, raw in changes.items():
        if source not in originals:
            continue
        value = str(raw or "").strip()
        if not value or value == originals[source]:
            current.pop(source, None)
            continue
        if current.get(source) != value:
            saved[source] = value
        current[source] = value
    all_edits[key] = current
    _write_edits(batch_path, all_edits)
    return saved


# ── Excel 导出 / 导入 ─────────────────────────────────────────


def export_workbook(review: Mapping[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "译文对照"
    sheet.append(list(EXPORT_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    review_fill = PatternFill("solid", fgColor="FFF2CC")
    for file in review.get("files") or []:
        for row in file.get("rows") or []:
            status = "edited" if row.get("edited") else str(row.get("status") or "")
            sheet.append([file.get("name", ""), row["source"], row["translation"], _STATUS_LABELS.get(status, status), row["count"]])
            if status in {"unresolved", "review"} or row.get("crowded"):
                for cell in sheet[sheet.max_row]:
                    cell.fill = review_fill
    for column, width in zip("ABCDE", (28, 48, 48, 12, 10)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row[1:3]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "B2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def parse_workbook(review: Mapping[str, Any], payload: bytes) -> dict[str, dict[str, str]]:
    """导回的表里改过的行：``{逐张清单路径: {原文: 新译文}}``。

    和当前译文、第一次的译文都不一样才算改过。
    """
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise CadReviewError("读不出这个文件，请选择从这里导出、改好后保存的 Excel 文件（.xlsx）。") from exc
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(value or "").strip() for value in next(rows, ())]
    try:
        name_col, source_col, target_col = (header.index(name) for name in EXPORT_HEADERS[:3])
    except ValueError as exc:
        raise CadReviewError("表头不对：需要「图纸」「原文」「译文」三列，请用从这里导出的表格修改。") from exc
    by_name: dict[str, list[Mapping[str, Any]]] = {}
    for file in review.get("files") or []:
        by_name.setdefault(str(file.get("name") or ""), []).append(file)
    changes: dict[str, dict[str, str]] = {}
    for values in rows:
        if not values:
            continue
        cells = list(values) + [None] * 3

        def text(index: int) -> str:
            value = cells[index]
            return "" if value is None else str(value)

        source, target = text(source_col), text(target_col).strip()
        if not source or not target:
            continue
        for file in by_name.get(text(name_col).strip(), []):
            current = next((row for row in file.get("rows") or [] if row["source"] == source), None)
            # 表里还是第一次的译文，说明拿到表的人没改这一行：不能拿它盖掉导出之后在这里
            # 做的修改。
            if current is None or target in (current["translation"], current["original"]):
                continue
            changes.setdefault(str(file["key"]), {})[source] = target
    workbook.close()
    return changes


# ── 重新写回 ─────────────────────────────────────────────────


def _pipeline_options(snapshot: Mapping[str, Any]) -> CadPipelineOptions:
    known = {item.name for item in fields(CadPipelineOptions)}
    values = {key: value for key, value in snapshot.items() if key in known}
    for key in ("locked_blocks", "locked_layers"):
        if key in values:
            values[key] = tuple(values[key] or ())
    return CadPipelineOptions(**values)


def pipeline_options_snapshot(options: CadPipelineOptions) -> dict[str, Any]:
    """存进批次清单的写回设置，重新写回时原样还原。"""
    return asdict(options)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rewrite_files(
    batch_manifest: str | Path,
    keys: Iterable[str],
    *,
    converter: Any,
    dxf_cache: Any = None,
    progress: Callable[[str, int, int], None] | None = None,
    stop_requested: Callable[[], bool] | None = None,
) -> list[dict[str, Any]]:
    """按对照表（第一次的译文 + 用户的修改）重新生成这些图纸，覆盖上次译好的文件。

    不调用模型、不查记忆库：每段文字的译文都从对照表里取。原图纸改过（和翻译时的
    内容不一样）就不写，免得把旧的译文套到新图纸上。
    """
    batch_path = Path(batch_manifest)
    batch = _read_json(batch_path)
    snapshot = batch.get("pipeline_options")
    if not isinstance(snapshot, dict):
        raise CadReviewError("这批图纸是旧版本翻译的，不能直接重新写回；重新翻译一次后即可使用。")
    options = _pipeline_options(snapshot)
    edits = load_edits(batch_path)
    entries = {str(item["manifest_path"]): item for item in _file_entries(batch)}
    wanted = [key for key in keys if key in entries]
    results: list[dict[str, Any]] = []
    for index, key in enumerate(wanted):
        entry = entries[key]
        source = Path(str(entry.get("source_path") or ""))
        output = Path(str(entry.get("output_path") or ""))
        name = source.name or output.name
        if progress is not None:
            progress(name, index, len(wanted))
        if stop_requested is not None and stop_requested():
            break
        try:
            results.append({"key": key, "name": name, **_rewrite_one(
                Path(key), source, output, edits.get(key, {}), options,
                converter=converter, dxf_cache=dxf_cache, stop_requested=stop_requested,
            )})
        except (CadReviewError, CadPipelineError) as exc:
            results.append({"key": key, "name": name, "ok": False, "message": str(exc)})
    if progress is not None:
        progress("", len(wanted), len(wanted))
    return results


def _rewrite_one(
    manifest_path: Path,
    source: Path,
    output: Path,
    edits: Mapping[str, str],
    options: CadPipelineOptions,
    *,
    converter: Any,
    dxf_cache: Any,
    stop_requested: Callable[[], bool] | None,
) -> dict[str, Any]:
    manifest = _read_json(manifest_path)
    if not source.is_file():
        raise CadReviewError("原图纸已被移动或删除，无法重新写回。")
    expected = str(manifest.get("source_sha256") or "")
    if expected and _sha256(source) != expected:
        raise CadReviewError("原图纸在翻译之后改过，对照表已经对不上，请重新翻译这张图纸。")
    if source.suffix.lower() == ".dwg" and converter is None:
        raise CadReviewError("重新写回 DWG 图纸前，请先在「插件」中连接图纸转换工具。")
    rows = _rows(manifest, edits)
    mapping = {row["source"]: row["translation"] for row in rows if row["translation"]}
    kept = {row["source"] for row in rows if row["status"] == "kept" and not row["edited"]}
    target = str(options.target_lang or "").split("-", 1)[0].lower()

    def translator(texts: list[str], glossary: Mapping[str, str], **_: Any) -> dict[str, str]:
        return {text: mapping[text] for text in texts if text in mapping}

    def language_translator(texts: list[str], glossary: Mapping[str, str], **_: Any) -> dict[str, Any]:
        # 第一次由模型确认「已是目标语言」的文字照样保留原文；其余按对照表。
        out: dict[str, Any] = {}
        for text in texts:
            if text in mapping:
                out[text] = (mapping[text], "")
            elif text in kept:
                out[text] = (text, target)
        return out

    with tempfile.TemporaryDirectory(prefix="xl-cad-rewrite-") as temporary:
        staged = Path(temporary) / output.name
        result = CadTranslationPipeline(
            converter=converter,
            translator=translator,
            language_translator=language_translator,
            dxf_cache=dxf_cache,
        ).translate_file(source, staged, options=options, stop_requested=stop_requested)
        output.parent.mkdir(parents=True, exist_ok=True)
        publish = output.with_name(f".{output.name}.rewrite")
        shutil.copy2(result.output, publish)
        os.replace(publish, output)
        # 新的逐张清单换掉旧的：对照表里的「第一次译文」从此就是这次写进图纸的译文，
        # 已经写进去的修改从修改记录里拿掉。
        new_manifest = _read_json(Path(result.manifest))
        new_manifest.update({key: manifest[key] for key in manifest if key not in new_manifest})
        new_manifest["output"] = str(output)
        new_manifest["manifest"] = str(manifest_path)
        new_manifest["rewritten_from_review"] = True
        # 未采用设定译法、这次也没改的文字仍是需复核（写回不经过记忆库核对）。
        still_review = {
            row["source"]: row
            for row in rows
            if row["status"] == "review" and not row["edited"]
        }
        if still_review:
            reasons = {
                str(unit.get("source") or ""): str(unit.get("reason") or "")
                for unit in manifest.get("units") or []
                if isinstance(unit, dict) and unit.get("status") == "review"
            }
            for unit in new_manifest.get("units") or []:
                source_text = str(unit.get("source") or "") if isinstance(unit, dict) else ""
                if source_text in still_review and unit.get("status") == "translated":
                    unit["status"] = "review"
                    unit["reason"] = reasons.get(source_text, "")
        manifest_path.write_text(json.dumps(new_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        stats = result.stats
        return {
            "ok": True,
            "changed": stats.changed,
            "untranslated": stats.untranslated,
            "layout_unresolved": stats.layout_unresolved,
            "message": "",
        }


def clear_applied_edits(batch_manifest: str | Path, keys: Iterable[str]) -> None:
    """写回成功的图纸：修改已经进了新的逐张清单，修改记录清掉。"""
    batch_path = Path(batch_manifest)
    edits = load_edits(batch_path)
    for key in keys:
        edits.pop(key, None)
    _write_edits(batch_path, edits)
