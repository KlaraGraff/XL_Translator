"""宽容地读取用户自己整理的词表文件，交给记忆库导入预览。

用户手里的表很少正好是我们导出的格式：列名五花八门（「中文 / Tiếng Việt / 说明」），
用 GBK 或 UTF-16 保存，用分号或制表符分隔，中间夹着空行、序号列、合并标题行。
这里只做一件事：把文件读成一张二维表，并猜出哪列是原文、译文、备注。
真正写库仍走 ``tm_manager.import_entries``，预览里用户可以改列。

支持：.xlsx/.xlsm（openpyxl）、.xls（xlrd）、.csv/.tsv/.txt（自动识别编码和分隔符）、
.json（我们导出的 ``entries``、对象数组、{原文: 译文} 字典、二维数组）。
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from typing import Any

from config import SUPPORTED_LANGS
from core.language_registry import BUILTIN_LANGUAGE_ENGLISH_NAMES

MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_IMPORT_ROWS = 200_000
# 在前几行里找表头：有的表第一行是大标题，第二行才是列名。
HEADER_SCAN_ROWS = 5

ROLE_SOURCE = "source"
ROLE_TARGET = "target"
ROLE_NOTE = "note"
ROLE_WORD_TYPE = "word_type"
ROLE_PINNED = "pinned"

# 列名同义词（统一小写、去空白和标点后比较）。
_ROLE_ALIASES: dict[str, tuple[str, ...]] = {
    ROLE_SOURCE: (
        "source_text", "source", "src", "original", "originaltext", "sourcetext", "sourceterm",
        "term", "原文", "源文", "原词", "原文词条", "源语言", "原语言", "词条", "术语", "中文术语",
        "texte source", "source term", "original text", "terme", "terme source",
    ),
    ROLE_TARGET: (
        "target_text", "target", "tgt", "translation", "translated", "targettext", "targetterm",
        "译文", "翻译", "译词", "译名", "目标语言", "目标", "对应译文", "外文",
        "traduction", "texte cible", "cible", "terme cible", "target term",
    ),
    ROLE_NOTE: (
        "note", "notes", "备注", "说明", "注释", "注", "场景", "适用场景", "使用场景", "用法说明",
        "comment", "comments", "remark", "remarks", "context", "description", "desc",
        "remarque", "commentaire", "contexte",
    ),
    ROLE_WORD_TYPE: ("word_type", "type", "来源"),
    ROLE_PINNED: ("pinned", "固定", "置顶"),
}

# 常见语言的本族语叫法，补在 SUPPORTED_LANGS（中文名）和英文名之外。
_NATIVE_LANG_NAMES: dict[str, tuple[str, ...]] = {
    "zh": ("中文", "汉语", "汉字", "简体中文", "简体", "繁体中文", "繁體中文", "繁体", "中", "chinese", "zh-cn", "zh-tw", "chinois"),
    "en": ("英文", "英语", "英", "english", "anglais"),
    "fr": ("法文", "法语", "法", "français", "francais", "french"),
    "vi": ("越南语", "越南文", "tiếng việt", "tieng viet", "vietnamese"),
    "ar": ("阿拉伯语", "阿拉伯文", "العربية", "arabic", "arabe"),
    "es": ("español", "espanol", "西班牙文", "espagnol"),
    "pt": ("português", "portugues", "葡萄牙文"),
    "de": ("deutsch", "德语", "德文", "allemand"),
    "ru": ("русский", "俄语", "俄文", "russe"),
    "ja": ("日本語", "日语", "日文", "japonais"),
    "ko": ("한국어", "韩语", "韩文", "coréen"),
    "km": ("ខ្មែរ", "高棉语", "柬埔寨语", "khmer"),
    "th": ("ไทย", "泰语", "泰文"),
}

_PUNCT_RE = re.compile(r"[\s　_\-:：()（）\[\]【】/\\|.,，。*#]+")


def _norm_header(value: str) -> str:
    return _PUNCT_RE.sub("", str(value or "").strip().casefold())


def _lang_aliases(code: str) -> set[str]:
    names: set[str] = {code}
    for zh_name, lang_code in SUPPORTED_LANGS.items():
        if lang_code == code:
            names.add(zh_name)
            # 「柬埔寨语（高棉语）」→ 也认「柬埔寨语」
            names.add(re.sub(r"[（(].*?[）)]", "", zh_name))
    names.update(BUILTIN_LANGUAGE_ENGLISH_NAMES.get(code, set()))
    names.update(_NATIVE_LANG_NAMES.get(code, ()))
    return {_norm_header(name) for name in names if name}


_ROLE_NORMALIZED = {role: {_norm_header(alias) for alias in aliases} for role, aliases in _ROLE_ALIASES.items()}


@dataclass
class ParsedTable:
    file_format: str
    rows: list[list[str]]
    encoding: str = ""
    delimiter: str = ""
    sheet: str = ""
    sheets_skipped: int = 0
    warnings: list[str] = field(default_factory=list)


# ── 读文件 ────────────────────────────────────────────────

_ENCODINGS = ("utf-8-sig", "gb18030", "big5", "shift_jis", "cp1252", "latin-1")


def _decode_text(data: bytes) -> tuple[str, str]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "UTF-16"
    # 没有 BOM 的 UTF-16（Excel「Unicode 文本」偶尔这样）：一半字节是 0。
    sample = data[:4000]
    if sample and sample.count(b"\x00") > len(sample) // 4:
        even_zero = sample[1::2].count(b"\x00") > sample[0::2].count(b"\x00")
        enc = "utf-16-le" if even_zero else "utf-16-be"
        return data.decode(enc, errors="replace"), "UTF-16"
    for enc in _ENCODINGS:
        try:
            text = data.decode(enc)
        except UnicodeDecodeError:
            continue
        label = {"utf-8-sig": "UTF-8", "gb18030": "GBK", "big5": "Big5", "shift_jis": "Shift-JIS"}.get(enc, enc.upper())
        return text, label
    return data.decode("utf-8", errors="replace"), "UTF-8"


def _sniff_delimiter(text: str) -> str:
    lines = [line for line in text.splitlines() if line.strip()][:50]
    if not lines:
        return ","
    best = ","
    best_score = -1.0
    for delim in ("\t", ",", ";", "|", "，"):
        counts = [line.count(delim) for line in lines]
        if not any(counts):
            continue
        # 分数：多数行都有、且列数稳定的分隔符胜出。
        hit = sum(1 for count in counts if count) / len(counts)
        common = max(set(counts), key=counts.count)
        stable = counts.count(common) / len(counts)
        score = hit * 2 + stable + (0.1 if delim == "\t" else 0)
        if common and score > best_score:
            best, best_score = delim, score
    return best


def _read_delimited(data: bytes, file_format: str) -> ParsedTable:
    text, encoding = _decode_text(data)
    delimiter = "\t" if file_format == "tsv" else _sniff_delimiter(text)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [[cell for cell in row] for row in reader]
    return ParsedTable(file_format=file_format, rows=rows, encoding=encoding, delimiter=delimiter)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _pick_sheet(sheets: list[tuple[str, list[list[str]]]]) -> tuple[str, list[list[str]], int]:
    # 取内容最多的那张表；其余表（说明页、封面）不导入，在预览里提一句。
    non_empty = [(name, rows) for name, rows in sheets if any(any(c.strip() for c in row) for row in rows)]
    if not non_empty:
        return (sheets[0][0] if sheets else ""), [], 0
    name, rows = max(non_empty, key=lambda item: sum(1 for row in item[1] if sum(1 for c in row if c.strip()) >= 2))
    return name, rows, len(non_empty) - 1


def _read_xlsx(data: bytes) -> ParsedTable:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets: list[tuple[str, list[list[str]]]] = []
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([_cell_text(value) for value in row])
                if len(rows) > MAX_IMPORT_ROWS:
                    break
            sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    name, rows, skipped = _pick_sheet(sheets)
    return ParsedTable(file_format="xlsx", rows=rows, sheet=name if len(sheets) > 1 else "", sheets_skipped=skipped)


def _read_xls(data: bytes) -> ParsedTable:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheets = []
    for sheet in book.sheets():
        rows = [[_cell_text(sheet.cell_value(r, c)) for c in range(sheet.ncols)] for r in range(min(sheet.nrows, MAX_IMPORT_ROWS))]
        sheets.append((sheet.name, rows))
    name, rows, skipped = _pick_sheet(sheets)
    return ParsedTable(file_format="xls", rows=rows, sheet=name if len(sheets) > 1 else "", sheets_skipped=skipped)


def _rows_from_objects(items: list[dict[str, Any]]) -> list[list[str]]:
    headers: list[str] = []
    for item in items:
        for key in item:
            if key not in headers:
                headers.append(str(key))
    rows = [headers]
    for item in items:
        rows.append([_cell_text(item.get(key)) if not isinstance(item.get(key), (dict, list)) else "" for key in headers])
    return rows


def _read_json(data: bytes) -> tuple[ParsedTable, str]:
    text, encoding = _decode_text(data)
    payload = json.loads(text)
    lang_pair = ""
    items: Any = payload
    if isinstance(payload, dict):
        lang_pair = str(payload.get("lang_pair") or "")
        for key in ("entries", "terms", "glossary", "items", "data", "rows"):
            if isinstance(payload.get(key), list):
                items = payload[key]
                break
        else:
            # {"原文": "译文", ...}
            pairs = [(k, v) for k, v in payload.items() if isinstance(v, (str, int, float)) and k != "lang_pair"]
            if pairs:
                return ParsedTable(file_format="json", rows=[[str(k), _cell_text(v)] for k, v in pairs], encoding=encoding), lang_pair
            items = []
    if not isinstance(items, list):
        items = []
    if items and all(isinstance(item, dict) for item in items):
        rows = _rows_from_objects(items)
    else:
        rows = [[_cell_text(cell) for cell in item] if isinstance(item, list) else [_cell_text(item)] for item in items]
    return ParsedTable(file_format="json", rows=rows, encoding=encoding), lang_pair


def read_table(file_name: str, data: bytes) -> tuple[ParsedTable, str]:
    """返回 (表格, 文件里自带的语言对)。认不出的格式抛 ValueError（中文提示）。"""
    if len(data) > MAX_IMPORT_BYTES:
        raise ValueError("文件超过 20 MB，无法导入记忆库。")
    lower = file_name.casefold()
    ext = lower.rsplit(".", 1)[-1] if "." in lower else ""
    if data[:4] == b"PK\x03\x04" or ext in {"xlsx", "xlsm"}:
        try:
            return _read_xlsx(data), ""
        except Exception as exc:  # noqa: BLE001 - 给用户一句能看懂的话
            raise ValueError("无法打开这个 Excel 文件，请确认文件没有损坏或加密。") from exc
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or ext == "xls":
        try:
            return _read_xls(data), ""
        except Exception as exc:  # noqa: BLE001
            raise ValueError("无法打开这个 Excel 文件，请确认文件没有损坏或加密。") from exc
    stripped = data.lstrip(b"\xef\xbb\xbf \t\r\n")
    if ext == "json" or stripped[:1] in {b"{", b"["}:
        try:
            return _read_json(data)
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            if ext == "json":
                raise ValueError("JSON 文件格式有误，无法读取。") from exc
    return _read_delimited(data, "tsv" if ext == "tsv" else ("txt" if ext == "txt" else "csv")), ""


# ── 认列 ──────────────────────────────────────────────────


def _header_role(header: str, source_lang: str, target_lang: str) -> str | None:
    norm = _norm_header(header)
    if not norm:
        return None
    for role in (ROLE_SOURCE, ROLE_TARGET, ROLE_NOTE, ROLE_WORD_TYPE, ROLE_PINNED):
        if norm in _ROLE_NORMALIZED[role]:
            return role
    if source_lang and norm in _lang_aliases(source_lang):
        return ROLE_SOURCE
    if target_lang and norm in _lang_aliases(target_lang):
        return ROLE_TARGET
    # 「中文名称」「法文译名」「Chinese term」这类：语言名开头/结尾。
    for lang, role in ((source_lang, ROLE_SOURCE), (target_lang, ROLE_TARGET)):
        if not lang:
            continue
        for alias in _lang_aliases(lang):
            if len(alias) >= 2 and (norm.startswith(alias) or norm.endswith(alias)):
                return role
    return None


def _match_header_row(row: list[str], source_lang: str, target_lang: str) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(row):
        role = _header_role(header, source_lang, target_lang)
        if role and role not in mapping:
            mapping[role] = index
    return mapping


_SERIAL_RE = re.compile(r"^\s*[#№]?\d{1,6}[.、)]?\s*$")


def _looks_like_serial_column(rows: list[list[str]], index: int) -> bool:
    values = [row[index].strip() for row in rows[:50] if index < len(row) and row[index].strip()]
    return bool(values) and sum(1 for v in values if _SERIAL_RE.match(v)) >= max(1, int(len(values) * 0.8))


def _guess_without_header(rows: list[list[str]]) -> dict[str, int]:
    # 没有可认的表头：取前两列有内容、且不是序号的列作原文、译文。
    width = max((len(row) for row in rows), default=0)
    candidates = []
    for index in range(width):
        filled = sum(1 for row in rows[:200] if index < len(row) and row[index].strip())
        if filled and not _looks_like_serial_column(rows, index):
            candidates.append(index)
    mapping: dict[str, int] = {}
    if candidates:
        mapping[ROLE_SOURCE] = candidates[0]
    if len(candidates) > 1:
        mapping[ROLE_TARGET] = candidates[1]
    return mapping


def detect_layout(rows: list[list[str]], lang_pair: str) -> tuple[int, dict[str, int]]:
    """返回 (表头所在行号，没有表头为 -1；{角色: 列号})。"""
    source_lang, _, target_lang = (lang_pair or "").partition("-")
    best_row = -1
    best_mapping: dict[str, int] = {}
    seen = 0
    for index, row in enumerate(rows):
        if not any(cell.strip() for cell in row):
            continue
        seen += 1
        if seen > HEADER_SCAN_ROWS:
            break
        mapping = _match_header_row(row, source_lang, target_lang)
        has_core = ROLE_SOURCE in mapping or ROLE_TARGET in mapping
        if has_core and len(mapping) > len(best_mapping):
            best_row, best_mapping = index, mapping
    data_rows = rows[best_row + 1 :] if best_row >= 0 else rows
    if best_row < 0:
        return -1, _guess_without_header(rows)
    # 表头只认出一边：另一边取剩下第一列有内容、非序号、非备注的列。
    if ROLE_SOURCE not in best_mapping or ROLE_TARGET not in best_mapping:
        missing = ROLE_TARGET if ROLE_SOURCE in best_mapping else ROLE_SOURCE
        used = set(best_mapping.values())
        width = max((len(row) for row in data_rows), default=0)
        for index in range(width):
            if index in used or _looks_like_serial_column(data_rows, index):
                continue
            if any(index < len(row) and row[index].strip() for row in data_rows[:200]):
                best_mapping[missing] = index
                break
    return best_row, best_mapping


def _header_lang(header: str) -> str:
    norm = _norm_header(header)
    if not norm:
        return ""
    for code in sorted(set(SUPPORTED_LANGS.values())):
        if norm in _lang_aliases(code):
            return code
    return ""


def _suggest_lang_pair(headers: list[str], mapping: dict[str, int], lang_pair: str) -> str:
    """列名写明了语言（「中文 | Tiếng Việt」）而和当前语言对不同时，建议换成文件里的语言对。"""
    if ROLE_SOURCE not in mapping or ROLE_TARGET not in mapping:
        return ""
    src = _header_lang(headers[mapping[ROLE_SOURCE]]) if mapping[ROLE_SOURCE] < len(headers) else ""
    tgt = _header_lang(headers[mapping[ROLE_TARGET]]) if mapping[ROLE_TARGET] < len(headers) else ""
    current_src, _, current_tgt = (lang_pair or "").partition("-")
    suggested = f"{src or current_src}-{tgt or current_tgt}"
    if not (src or tgt) or suggested == lang_pair or not current_src or not current_tgt:
        return ""
    if (src or current_src) == (tgt or current_tgt):
        return ""
    return suggested


def parse_import_file(file_name: str, data: bytes, lang_pair: str) -> dict[str, Any]:
    """读文件并猜列，交给前端预览；前端按（可改的）列映射拼出词条。"""
    table, file_lang_pair = read_table(file_name, data)
    effective_pair = file_lang_pair or lang_pair
    rows = [[str(cell).replace("\r\n", "\n").strip() for cell in row] for row in table.rows[:MAX_IMPORT_ROWS + 1]]
    header_row, mapping = detect_layout(rows, effective_pair)
    width = max((len(row) for row in rows), default=0)
    if header_row >= 0:
        headers = [rows[header_row][i] if i < len(rows[header_row]) else "" for i in range(width)]
        data_rows = rows[header_row + 1 :]
        skipped_above = sum(1 for row in rows[:header_row] if any(row))
    else:
        headers = [""] * width
        data_rows = rows
        skipped_above = 0
    blank = sum(1 for row in data_rows if not any(row))
    data_rows = [row + [""] * (width - len(row)) for row in data_rows if any(row)]
    return {
        "suggested_lang_pair": _suggest_lang_pair(headers, mapping, effective_pair),
        "format": table.file_format,
        "encoding": table.encoding,
        "delimiter": table.delimiter,
        "sheet": table.sheet,
        "sheets_skipped": table.sheets_skipped,
        "lang_pair": file_lang_pair,
        "has_header": header_row >= 0,
        "headers": headers,
        "rows": data_rows,
        "mapping": mapping,
        "skipped_blank": blank,
        "skipped_above_header": skipped_above,
    }
