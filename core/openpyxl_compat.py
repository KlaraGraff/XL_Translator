"""openpyxl 读文件的兼容补丁。

项目里所有 ``load_workbook`` 都从这里取，好让下面的补丁在读第一份文件前装上。

空填充 ``<fill/>``：styles.xml 的填充表里出现一个没有子元素的 ``<fill/>``，Excel 照常
打开，当成「无填充」。openpyxl 的 ``Fill.from_tree`` 遇到它返回 None，样式表随后报
``TypeError("expected <class 'openpyxl.styles.fills.Fill'>")``，整份文件读不了——扫描
阶段直接把它当坏文件跳过，用户只看到「跳过 1」。实测是一份用户产值表，59 个填充里
有 3 个空的，跟内容毫无关系。这里让空 ``<fill/>`` 读成无填充，与 Excel 一致。

只改读：本项目写 xlsx 走 ``core.xlsx_patcher`` 的部件级补丁，不经 openpyxl 保存，
所以原文件里的 ``<fill/>`` 原样留在输出里，填充表下标也不会错位。
"""
from __future__ import annotations

_installed = False


def _install() -> None:
    global _installed
    if _installed:
        return
    from openpyxl.styles.fills import Fill, PatternFill

    original = Fill.from_tree.__func__

    def from_tree(cls, el):
        value = original(cls, el)
        if value is None:
            return PatternFill()
        return value

    Fill.from_tree = classmethod(from_tree)
    _installed = True


def load_workbook(*args, **kwargs):
    """``openpyxl.load_workbook``，先装好上面的兼容补丁。"""
    _install()
    from openpyxl import load_workbook as _load_workbook

    return _load_workbook(*args, **kwargs)
